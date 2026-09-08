from __future__ import annotations

import hashlib
import json
import sys
import zipfile
from pathlib import Path

from docx import Document


def iter_table_text(table):
    for row_index, row in enumerate(table.rows):
        yield {
            "row": row_index,
            "cells": ["\n".join(p.text for p in cell.paragraphs).strip() for cell in row.cells],
        }


def main() -> None:
    source = Path(sys.argv[1]).resolve()
    out_dir = Path(sys.argv[2]).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    doc = Document(source)
    report = {
        "source": str(source),
        "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "sections": len(doc.sections),
        "paragraphs": [
            {"index": i, "style": p.style.name if p.style else None, "text": p.text}
            for i, p in enumerate(doc.paragraphs)
            if p.text.strip()
        ],
        "tables": [list(iter_table_text(table)) for table in doc.tables],
        "inline_shapes": len(doc.inline_shapes),
        "package_media": [],
    }

    media_dir = out_dir / "media"
    media_dir.mkdir(exist_ok=True)
    with zipfile.ZipFile(source) as package:
        for name in sorted(package.namelist()):
            if not name.startswith("word/media/") or name.endswith("/"):
                continue
            data = package.read(name)
            target = media_dir / Path(name).name
            target.write_bytes(data)
            report["package_media"].append(
                {
                    "part": name,
                    "file": str(target),
                    "bytes": len(data),
                    "sha256": hashlib.sha256(data).hexdigest(),
                }
            )

    (out_dir / "content.json").write_text(
        json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    print(json.dumps(report, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
