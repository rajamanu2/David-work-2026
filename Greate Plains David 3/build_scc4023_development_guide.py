from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"C:\Users\LIKKI\Documents\ChatGPT\Greate Plains David 3\SCC-4023-Development-Guide.docx")

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "5E6A78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_GREEN = "EAF4EA"
PALE_GOLD = "FFF4CE"
PALE_RED = "FDECEC"
WHITE = "FFFFFF"
BORDER = "AEB9C6"

CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_table_header(row):
    repeat_header(row)


def set_run_font(run, name="Calibri", size=11, color=INK, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_paragraph_border(paragraph, color=BORDER, size=8, side="bottom"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    edge = OxmlElement(f"w:{side}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), "1")
    edge.set(qn("w:color"), color)
    p_bdr.append(edge)


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    r = paragraph.add_run()._r
    r.append(fld_char1)
    r.append(instr_text)
    r.append(fld_char2)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    code = doc.styles.add_style("Code Block", 1)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(8.5)
    code.font.color.rgb = RGBColor.from_string("263238")
    code.paragraph_format.left_indent = Inches(0.14)
    code.paragraph_format.right_indent = Inches(0.14)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(7)
    code.paragraph_format.line_spacing = 1.0

    small = doc.styles.add_style("Small Note", 1)
    small.font.name = "Calibri"
    small._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    small._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    small.font.size = Pt(9)
    small.font.color.rgb = RGBColor.from_string(MUTED)
    small.paragraph_format.space_after = Pt(4)
    small.paragraph_format.line_spacing = 1.15


def add_numbering_definition(doc, num_id, bullet=False):
    numbering = doc.part.numbering_part.element
    abstract_num = OxmlElement("w:abstractNum")
    abstract_num.set(qn("w:abstractNumId"), str(num_id))
    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), f"000000{num_id:02d}")
    abstract_num.append(nsid)
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract_num.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if bullet else "%1.")
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    abstract_num.append(lvl)
    numbering.append(abstract_num)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(num_id))
    num.append(abstract_num_id)
    numbering.append(num)


def add_list_item(doc, text, num_id=41, bold_lead=None, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)
    p_pr.append(num_pr)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc, text, bold_lead=None):
    return add_list_item(doc, text, num_id=42, bold_lead=bold_lead)


def add_checkbox(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(f"☐ {text}")
    set_run_font(r)
    return p


def add_callout(doc, label, text, fill=LIGHT_BLUE, label_color=NAVY):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.2
    shade_paragraph(p, fill)
    r = p.add_run(f"{label}: ")
    set_run_font(r, bold=True, color=label_color)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_code(doc, text):
    p = doc.add_paragraph(style="Code Block")
    shade_paragraph(p, LIGHT_GRAY)
    for index, line in enumerate(text.splitlines()):
        if index:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, name="Consolas", size=8.5, color="263238")
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    repeat_header(hdr)
    for i, text in enumerate(headers):
        set_cell_shading(hdr.cells[i], LIGHT_BLUE)
        p = hdr.cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=9.5, bold=True, color=NAVY)
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            p = row.cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(str(value))
            set_run_font(r, size=9.2)
    set_table_geometry(table, widths)
    return table


def add_page_break(doc):
    doc.add_page_break()


def add_heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_source_link(doc, label, url):
    p = doc.add_paragraph(style="Small Note")
    r = p.add_run(f"{label}: {url}")
    set_run_font(r, size=9, color=MUTED)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("SCC-4023  |  Salesforce Development Guide")
    set_run_font(hr, size=9, color=MUTED, bold=True)
    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_number(fp)


def build():
    doc = Document()
    configure_styles(doc)
    configure_document(doc)
    add_numbering_definition(doc, 41, bullet=False)
    add_numbering_definition(doc, 42, bullet=True)
    doc.core_properties.title = "SCC-4023 Salesforce Development Guide"
    doc.core_properties.subject = "Fiber Installation Method global restricted picklist implementation"
    doc.core_properties.author = "Great Plains Communications Delivery Team"
    doc.core_properties.keywords = "SCC-4023, Salesforce, Field Service, picklist, metadata"

    # Opening page: memo masthead, adapted for an implementation guide.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("SALESFORCE DEVELOPMENT GUIDE")
    set_run_font(r, size=10, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("SCC-4023")
    set_run_font(r, size=28, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("Standardize Fiber Installation Method across Premise, Work Order, and Service Appointment")
    set_run_font(r, size=15, color=DARK_BLUE, bold=True)
    set_paragraph_border(p, color=BLUE, size=14, side="bottom")

    meta = [
        ("Document type", "Development implementation guide"),
        ("Prepared for", "Salesforce delivery and peer review"),
        ("Prepared on", "September 1, 2026"),
        ("Change type", "Metadata configuration; no Apex expected"),
        ("Source", "SCC-4023 ticket export"),
    ]
    add_table(doc, ["Item", "Detail"], meta, [2700, 6660])

    add_heading(doc, "Executive summary", 1)
    add_body(doc, "SCC-4023 should be implemented as one shared, restricted global picklist named Fiber Installation Method. The same value set must be referenced by the existing fields on Premise, Work Order, and Service Appointment.")
    add_callout(doc, "Target result", "Users and integrations can save only Aerial, Buried, Conduit, or Temporary. Fixed Wireless and Unknown must no longer be selectable. Existing records require an agreed data treatment before the restriction is enabled.", fill=PALE_GREEN)

    add_heading(doc, "Required outcome", 2)
    add_table(
        doc,
        ["Object", "Field", "Target configuration"],
        [
            ("Premise", "vlocity_cmt__Premises__c.Fiber_Installation_Method__c", "Restricted; shared global value set"),
            ("Work Order", "WorkOrder.Fiber_Installation_Method__c", "Restricted; shared global value set"),
            ("Service Appointment", "ServiceAppointment.Fiber_Installation_Method__c", "Restricted; shared global value set"),
        ],
        [1750, 4450, 3160],
    )
    add_callout(doc, "Implementation position", "Use metadata only unless the dependency review finds hard-coded legacy values in Flow, Apex, OmniStudio, or integration mappings. Do not recreate the fields; update the existing field metadata in place.")

    add_page_break(doc)

    add_heading(doc, "Ticket corrections and decisions", 1)
    add_body(doc, "The ticket describes the business outcome clearly, but the technical section should be corrected before development is considered complete.")
    add_table(
        doc,
        ["Finding", "Required correction"],
        [
            ("Work Order is missing from Components", "Add WorkOrder.Fiber_Installation_Method__c to the component list."),
            ("Global value set is missing", "Add GlobalValueSet:Fiber_Installation_Method as a new or updated component."),
            ("Objects/fields says None", "Replace with the three updated custom fields and the global value set."),
            ("Testing says Restrict picklist is unchecked", "Treat this as the current-state defect. The target state is restricted on all three fields."),
            ("Legacy data treatment is undefined", "Agree whether Fixed Wireless and Unknown will be mapped, cleared, or retained temporarily before restriction."),
        ],
        [3000, 6360],
    )

    add_heading(doc, "Decision required before build", 2)
    add_bullet(doc, "Data owner decision: define the replacement for records currently holding Fixed Wireless or Unknown.", bold_lead="Data owner decision:")
    add_bullet(doc, "Integration owner decision: confirm that Vetro will send only the four approved values after the Salesforce change is released.", bold_lead="Integration owner decision:")
    add_bullet(doc, "Release decision: Salesforce must be ready before Vetro begins sending the new values, or record creation may fail.", bold_lead="Release decision:")

    add_heading(doc, "Planned metadata inventory", 2)
    add_table(
        doc,
        ["Metadata type", "API name", "Action"],
        [
            ("GlobalValueSet", "Fiber_Installation_Method", "Create or update with four approved values"),
            ("CustomField", "vlocity_cmt__Premises__c.Fiber_Installation_Method__c", "Point to shared set; restricted = true"),
            ("CustomField", "WorkOrder.Fiber_Installation_Method__c", "Point to shared set; restricted = true"),
            ("CustomField", "ServiceAppointment.Fiber_Installation_Method__c", "Point to shared set; restricted = true"),
        ],
        [2050, 4670, 2640],
    )
    add_callout(doc, "Current-state verification", "Retrieve the latest metadata from the development org before editing. This guide is based on the ticket export and a workspace dependency snapshot; it does not replace a current-org retrieve.", fill=PALE_GOLD, label_color="7A5A00")

    add_heading(doc, "Development procedure", 1)
    add_body(doc, "Complete the following sequence in a development sandbox and a source-controlled branch. Do not develop directly in the Merge sandbox.")

    add_heading(doc, "Retrieve the current fields", 2)
    add_body(doc, "Start by retrieving the three existing fields. This protects any current labels, descriptions, required settings, formulas, or other field properties that are not part of SCC-4023.")
    add_code(doc, """sf project retrieve start `
  --metadata \"CustomField:ServiceAppointment.Fiber_Installation_Method__c\" `
  --metadata \"CustomField:WorkOrder.Fiber_Installation_Method__c\" `
  --metadata \"CustomField:vlocity_cmt__Premises__c.Fiber_Installation_Method__c\" `
  --target-org <development-alias>""")
    add_body(doc, "Check whether a suitable global value set already exists before creating a new one:")
    add_code(doc, """sf org list metadata `
  --metadata-type GlobalValueSet `
  --target-org <development-alias> `
  --json""")
    add_callout(doc, "Stop condition", "If an existing global value set has the same business meaning, confirm ownership and downstream use before reusing it. Do not introduce a second similarly named set without review.", fill=PALE_GOLD, label_color="7A5A00")

    add_heading(doc, "Audit existing data", 2)
    add_body(doc, "Run a distinct-value count for each object. Save the results in the work item or PR so reviewers can see whether legacy values exist and how many records are affected.")
    add_code(doc, """SELECT Fiber_Installation_Method__c, COUNT(Id)
FROM ServiceAppointment
WHERE Fiber_Installation_Method__c != NULL
GROUP BY Fiber_Installation_Method__c""")
    add_body(doc, "Repeat the same query for WorkOrder and vlocity_cmt__Premises__c. If Fixed Wireless, Unknown, or any unexpected value is present, do not enable the restriction until the data decision is approved.")

    add_page_break(doc)

    add_heading(doc, "Build the shared value set", 1)
    add_body(doc, "Create the following metadata file if no approved shared set already exists:")
    add_code(doc, "force-app/main/default/globalValueSets/Fiber_Installation_Method.globalValueSet-meta.xml")
    add_code(doc, """<?xml version=\"1.0\" encoding=\"UTF-8\"?>
<GlobalValueSet xmlns=\"http://soap.sforce.com/2006/04/metadata\">
    <customValue>
        <fullName>Aerial</fullName>
        <default>false</default>
        <label>Aerial</label>
    </customValue>
    <customValue>
        <fullName>Buried</fullName>
        <default>false</default>
        <label>Buried</label>
    </customValue>
    <customValue>
        <fullName>Conduit</fullName>
        <default>false</default>
        <label>Conduit</label>
    </customValue>
    <customValue>
        <fullName>Temporary</fullName>
        <default>false</default>
        <label>Temporary</label>
    </customValue>
    <masterLabel>Fiber Installation Method</masterLabel>
    <sorted>false</sorted>
</GlobalValueSet>""")
    add_callout(doc, "Value rule", "The set contains exactly four active values: Aerial, Buried, Conduit, and Temporary. Do not add Fixed Wireless or Unknown to the new set.")

    add_heading(doc, "Update the three existing fields", 2)
    add_body(doc, "In each existing field file, preserve all unrelated properties and replace only the picklist value-set definition with the shared value-set reference:")
    add_code(doc, """<type>Picklist</type>
<valueSet>
    <restricted>true</restricted>
    <valueSetName>Fiber_Installation_Method</valueSetName>
</valueSet>""")
    add_body(doc, "Apply the change to these exact field files:")
    add_bullet(doc, "objects/vlocity_cmt__Premises__c/fields/Fiber_Installation_Method__c.field-meta.xml")
    add_bullet(doc, "objects/WorkOrder/fields/Fiber_Installation_Method__c.field-meta.xml")
    add_bullet(doc, "objects/ServiceAppointment/fields/Fiber_Installation_Method__c.field-meta.xml")

    add_page_break(doc)

    add_heading(doc, "Dependency and integration review", 1)
    add_body(doc, "A restricted picklist can reject values that previously saved successfully. Review every producer, transformer, and consumer of Fiber Installation Method before merging.")

    add_heading(doc, "Workspace dependency snapshot", 2)
    add_body(doc, "The available workspace snapshot identifies the following references. Verify the current org and branch because the snapshot may not reflect the latest configuration.")
    add_table(
        doc,
        ["Category", "Components to verify"],
        [
            ("Field Set", "FSL_WorkOrder_Mobile"),
            ("Flows", "Create Ad hoc Work Order Screen Flow; Service Appointment - After Save - Create Update; ServiceAppointment_OrderLookup"),
            ("Layouts", "Field Service Dispatcher Service Appointment Layout; Field Service Dispatcher Work Order Layout; Field Service Technician Service Appointment Layout; Field Service Technician Work Order Layout"),
            ("Layouts", "Service Appointment Layout; Serviceability Request Layout; System Administrator Service Appointment Layout; System Administrator Work Order Layout; Work Order Layout"),
        ],
        [1800, 7560],
    )

    add_heading(doc, "What to inspect", 2)
    add_bullet(doc, "Flow decisions, assignments, formulas, and choices that use Fixed Wireless or Unknown.")
    add_bullet(doc, "Apex, OmniStudio, DataRaptors, Integration Procedures, rules, and formulas with hard-coded values.")
    add_bullet(doc, "FSL mobile experiences, layouts, and permission/profile entries that expose the fields.")
    add_bullet(doc, "Any automation that copies the value across Premise, Work Order, and Service Appointment.")

    add_heading(doc, "Vetro contract check", 2)
    add_table(
        doc,
        ["Contract item", "Required evidence"],
        [
            ("Outbound values", "Vetro sends only Aerial, Buried, Conduit, or Temporary after cutover."),
            ("Case and spelling", "Exact API values and capitalization match Salesforce."),
            ("Null behavior", "Blank values are either permitted by the field design or handled before insert/update."),
            ("Error handling", "Rejected records are logged, alerted, and replayable."),
            ("Sequence", "Salesforce metadata is released before Vetro starts sending the new values."),
        ],
        [2250, 7110],
    )
    add_callout(doc, "Merge blocker", "Do not merge until hard-coded legacy values are removed or deliberately handled, and the integration owner confirms the four-value contract.", fill=PALE_RED, label_color="9B1C1C")

    add_heading(doc, "PR and validation package", 1)
    add_heading(doc, "Minimum PR scope", 2)
    add_bullet(doc, "One GlobalValueSet metadata file: Fiber_Installation_Method.")
    add_bullet(doc, "Three CustomField metadata files, one for each object.")
    add_bullet(doc, "Only automations or integration metadata that genuinely require a change after dependency review.")
    add_bullet(doc, "Data audit results and the approved treatment for legacy values.")
    add_bullet(doc, "Validation output and focused functional-test evidence.")

    add_heading(doc, "Recommended Jira technical documentation", 2)
    add_table(
        doc,
        ["Section", "Recommended entry"],
        [
            ("New components", "GlobalValueSet - Fiber_Installation_Method"),
            ("Updated fields", "Premise, Work Order, and Service Appointment Fiber_Installation_Method__c"),
            ("Triggers or automations", "None, unless dependency review requires a specific Flow/Apex/OmniStudio update"),
            ("Deployment", "Deploy value set and all three field references in the same release unit"),
            ("Post-deployment", "Read-back metadata, run smoke tests, and monitor integration failures"),
        ],
        [2200, 7160],
    )

    add_heading(doc, "Check-only deployment", 2)
    add_body(doc, "Run validation against the intended validation org. Use the actual project paths if the repository structure differs.")
    add_code(doc, """sf project deploy start --dry-run `
  --manifest manifest/scc-4023.xml `
  --target-org <validation-alias>""")
    add_body(doc, "The manifest must contain the one global value set and the three CustomField members listed in the metadata inventory.")
    add_callout(doc, "Validation gate", "A successful dry run proves metadata deployability only. It does not prove runtime propagation, FSL mobile behavior, profile visibility, or Vetro compatibility.")

    add_page_break(doc)

    add_heading(doc, "Functional verification", 1)
    add_body(doc, "Use focused tests that directly prove the acceptance criteria and the highest-risk integration behavior.")
    add_table(
        doc,
        ["Test", "Action", "Expected result"],
        [
            ("Approved values - Premise", "Save each of the four approved values on a Premise.", "Every value saves; no extra values are selectable."),
            ("Approved values - Work Order", "Save each approved value on a Work Order.", "Every value saves; no extra values are selectable."),
            ("Approved values - Service Appointment", "Save each approved value on a Service Appointment.", "Every value saves; no extra values are selectable."),
            ("Invalid value", "Attempt to load Fixed Wireless, Unknown, and a random string through an integration-style update.", "Salesforce rejects the invalid values and the failure is logged."),
            ("Propagation", "Create or update the source Premise and progress through Work Order and Service Appointment creation.", "The selected value is copied without transformation or loss."),
            ("FSL mobile", "Open and update the affected record in the technician mobile experience.", "The field is visible as designed and only four values appear."),
            ("Legacy records", "Open and update a record included in the data audit.", "Behavior matches the approved mapping or cleanup decision."),
            ("Security", "Test with representative dispatcher, technician, and administrator access.", "Existing field-level security and layout visibility remain correct."),
        ],
        [1850, 3940, 3570],
    )

    add_heading(doc, "Release sequence", 1)
    add_table(
        doc,
        ["Phase", "Required actions"],
        [
            ("Pre-deployment", "Freeze the integration cutover window; rerun the data audit; complete approved legacy cleanup; confirm Vetro payload values and rollback contacts."),
            ("Deployment", "Deploy the global value set and all three field references as one release unit. Include automation changes only when dependency review proves they are necessary."),
            ("Post-deployment", "Read back all four metadata components, perform the end-to-end smoke test, monitor failed API records, and confirm Vetro cutover."),
            ("Rollback", "Pause inbound writes that send unsupported values, restore the prior metadata version if required, and replay failed records after correction."),
        ],
        [1900, 7460],
    )

    add_page_break(doc)

    add_heading(doc, "Completion checklist", 1)
    for item in (
        "Current field metadata retrieved from the development org.",
        "Existing values counted on Premise, Work Order, and Service Appointment.",
        "Legacy-value treatment approved by the business/data owner.",
        "Global value set contains exactly Aerial, Buried, Conduit, and Temporary.",
        "All three fields reference the same global value set and are restricted.",
        "Dependencies reviewed for hard-coded Fixed Wireless and Unknown values.",
        "Vetro contract and release order confirmed by the integration owner.",
        "PR contains one global value set and three field changes at minimum.",
        "Check-only deployment succeeds in the validation org.",
        "Functional tests pass across the three objects and FSL mobile.",
        "Jira technical documentation lists the complete component scope.",
        "Post-deployment monitoring and rollback ownership are assigned.",
    ):
        add_checkbox(doc, item)

    add_heading(doc, "Definition of done", 2)
    add_callout(doc, "Done", "SCC-4023 is complete when the three fields share one restricted global value set with exactly four approved values, legacy records have an approved treatment, dependent automation and integrations are verified, and the PR includes successful validation and functional evidence.", fill=PALE_GREEN)

    add_heading(doc, "Sources and evidence", 2)
    add_source_link(doc, "Primary requirement", r"SCC-4023.docx - provided Jira ticket export")
    add_source_link(doc, "Salesforce Metadata API coverage", "https://developer.salesforce.com/docs/success/metadata-coverage-report/references/metadata-types/v67.0")
    add_source_link(doc, "Dependency evidence", "Workspace custom-field dependency snapshot; verify against the current org before implementation")

    add_body(doc, "No Salesforce metadata was changed or deployed while preparing this guide.", bold_lead="No Salesforce metadata was changed or deployed")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(str(OUT))


if __name__ == "__main__":
    build()
