from pathlib import Path
import json, re, zipfile
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image

OUT=Path(__file__).resolve().parent
R=OUT.parent/'SALDEV-1403-dry-run-2026-09-07'
run=R/'recheck-232030'
results=json.loads((run/'deployment/verification-results.json').read_text())
assert results['passed']==results['total']==236
items=json.loads((run/'deployment/after-mapper-items.json').read_text())
spec=json.loads((R/'candidate/mapper-change-spec.json').read_text())
d=Document();sec=d.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
for st in d.styles:
 if st.type==1:
  st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.space_after=Pt(8)
d.styles['Normal'].paragraph_format.line_spacing=1.05
d.styles['Title'].font.size=Pt(26)
d.styles['Title'].paragraph_format.space_after=Pt(12)
d.styles['Heading 1'].font.size=Pt(18)
d.styles['Heading 2'].font.size=Pt(13)
d.styles['Caption'].font.size=Pt(10)
sec.header.paragraphs[0].text='SALDEV-1403  |  Quote document changes and testing'
sec.header.paragraphs[0].style='Caption'
f=sec.footer.paragraphs[0];f.text='FlywirePartial  •  7 September 2026     |     Page '
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
d.core_properties.author='Salesforce QA'
d.core_properties.title='SALDEV 1403 Component Changes and Test Guide'
d.core_properties.subject='Quote proposal component inventory and reproduction steps'
def p(t,style=None):return d.add_paragraph(t,style)
def sub(t):d.add_heading(t,2)
def page(t):d.add_page_break();d.add_heading(t,1)
def steps(values,start=1):
 for i,v in enumerate(values,start):p(f'{i}. {v}')
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 for values in rows:
  for c,v in zip(t.add_row().cells,values):c.text=str(v)
 for j,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for i,c in enumerate(row.cells):
   c.width=Inches(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
   cp.append(b);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF3' if j==0 else 'FFFFFF');cp.append(sh)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(5);para.paragraph_format.space_before=Pt(5)
    for rr in para.runs:rr.font.size=Pt(10);rr.bold=j==0
 p('').paragraph_format.space_after=Pt(0)
 return t
def link(label,url):
 pp=p('');rel=pp.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel);rr=OxmlElement('w:r');tt=OxmlElement('w:t');tt.text=label;rr.append(tt);h.append(rr);pp._p.append(h)
def shot(path,box,width,caption):
 iw,ih=Image.open(path).size;x1,y1,x2,y2=box
 sh=p('').add_run().add_picture(str(path),width=Inches(width),height=Inches(width*(y2-y1)/(x2-x1)))
 crop=OxmlElement('a:srcRect')
 for k,v in dict(l=x1/iw,t=y1/ih,r=(iw-x2)/iw,b=(ih-y2)/ih).items():crop.set(k,str(round(v*100000)))
 sh._inline.graphic.graphicData.pic.blipFill.insert(1,crop)
 sh._inline.docPr.set('descr',caption);p(caption,'Caption')

d.add_heading('SALDEV 1403 Component Changes and Test Guide',0)
p('CPQ Quote Proposal and OmniStudio','Subtitle')
p('Environment: FlywirePartial  |  Evidence date: 7 September 2026')
p('This document brings together the SALDEV-1403 component changes, expected quote-document behavior, dry-run results, test records, and steps to reproduce. It is intended for the story reviewer and business tester.')
p('The latest read-only backend dry run passed 236 of 236 checks across eight quotes. The remaining acceptance work is to inspect fresh generated documents, especially sequential Term numbering, final table layout, and the cases with incomplete test data.')
sub('What the story delivers')
for text in ['Groups appear in start-date order, with the longest subscription term first only when start dates match. Remaining ties follow the Quote Line Editor sequence.',
'Grouped headings show Term 1, Term 2, and so on with each group’s subscription term in Months. Service-period text is removed from both grouped branches.',
'Products retain quote-line order within their group or ungrouped table. Bundle options stay beneath the correct parent in the tested data.',
'Rate Details are appended to Description with appropriate spacing and blank-value handling. Applicable Cancellation, Quantity Increase, and Quantity Reduction labels remain.',
'Ungrouped output retains its original table layout, with Ship-To visibility controlled by Multiple Ship-To Accounts.']:
 p(text,'List Bullet')
sub('Dry run outcome')
table(['Check','Recorded result'],[['Sandbox identity','FlywirePartial, org 00DhG0000000jOXUAY'],['Native Extract and Transform runs','8 completed without mapper errors'],['Backend assertions','236 passed out of 236'],['Eligible product rows checked','153 across the eight quotes'],['Scoped configuration and template','Match the previously verified state'],['Fresh generated document appearance','Pending business verification']],[3.2,3.9])
p('SALDEV-1420 material is reference only. This guide covers the SALDEV-1403 quote proposal; it does not assign approval-email testing to this story.')

page('Changed components and exact items')
sub('Extract Data Mapper')
p('Name: GetQuoteProposalDataSteppedUpPricing\nRecord ID: 0jIhG0000000GsnUAE')
newrows=[]
labels=['Group order','Grouped line order','Ungrouped line order','Group term mapping']
details=['Start Date ASC NULLS LAST; Subscription Term DESC NULLS LAST; group Number ASC.','Quote-line Number ASC for lines inside groups.','Quote-line Number ASC for ungrouped lines.','Map group SBQQ__SubscriptionTerm__c to QuoteLineGroup:SubscriptionTerm.']
for c,label,detail in zip(spec['creates'],labels,details):
 actual=next(i for i in items if i['GlobalKey']==c['GlobalKey'])
 newrows.append([label,actual['Id'],detail])
table(['Added item','Record ID','Behavior'],newrows,[1.3,1.7,4.1])
sub('Transform Data Mapper')
p('Name: CPQQuoteProposalDocumentSteppedUpPricing\nRecord ID: 0jIhG0000000HX7UAM')
table(['Updated item','Record ID','Formula result'],[['Grouped Description','0kdhG0000002nUTQAY','QuoteLineGroup:QuoteLine:LineDescriptionWithType'],['Ungrouped Description','0kdhG0000002m5NQAQ','UngroupedLine:LineDescriptionWIthType']],[1.5,1.7,3.9])
p('Both formulas concatenate nonblank Description and Rate Details, inserting a space only when needed. They retain the relevant line-type suffix. Example: “Service description Rate text (Quantity Increase)”. The existing result-path spelling is preserved.')
sub('Proposal document template')
p('Component: CPQ Quote Proposal\nTemplate ID: 2dtPb0000000BwLIAU\nContent document ID: 069hG000003DZfnQAG\nVerified content version: 068hG000003PryiQAC')
p('Both grouped branches use a Word-numbered Term heading with {{SubscriptionTerm}} Months. Both Service Period paragraphs are removed. Product tables, signature content, grouped/ungrouped conditions, and Ship-To conditions are retained.')
p('Change scope: four new Extract items, two Transform formula updates, and one proposal DOCX. No Apex class, custom field, or OmniScript component change is part of this SALDEV-1403 fix.')

page('Steps to reproduce on the main quote')
p('Start with Q-37912, the reported ordering case. Confirm the browser is connected to flywire--partial.sandbox before testing. Current inputs should be rechecked because other testers can edit shared records.')
link('Open Q-37912 in FlywirePartial','https://flywire--partial.sandbox.lightning.force.com/lightning/r/SBQQ__Quote__c/a2NhG000003lAjmUAE/view')
steps(['Open the quote and Edit Lines. Capture the group dates, subscription terms, group sequence, product sequence, and Multiple Ship-To Accounts value. Do not change the quote merely to match an expected screenshot.',
'Return to the quote. Open the action menu and choose Generate Document or Generate Quote Document using the existing stepped-up-pricing action.',
'If template selection appears, choose CPQ Quote Proposal. Complete the existing generation screen and select Preview or Generate according to the established test workflow.',
'Open the fresh output. Do not use a previously attached PDF as evidence of this execution.',
'Compare the group headings against the expected sequence below. Verify that service dates are absent from the group headings.'])
table(['Heading','Source group','Expected term'],[['Term 1','Group4','36 Months'],['Term 2','Group1','12 Months'],['Term 3','Group2','12 Months'],['Term 4','Group3','12 Months']],[1.8,2.7,2.6])
steps(['Check all 78 eligible product rows against the source quote. Products must remain in the correct group, follow quote-line order, and appear once. Keep bundle children and nested options with their parent.',
'Confirm Ship-To Account is hidden for Q-37912 because Multiple Ship-To Accounts is No.',
'Generate another fresh preview and confirm Term numbering restarts at 1, with no skipped or repeated numbers. Check page breaks, table widths, and readable descriptions.',
'Record the quote ID, output filename/time, expected versus actual result, screenshot, and Pass/Fail/Blocked status.'],start=6)
p('Expected backend group sequence was confirmed in the latest run. Whether Word numbering survives native document repetition is still a visual acceptance check.')

page('Regression records and expected output')
p('Close and reopen the generation screen for every quote so the correct record context is used. Repeat the page 3 preview procedure and compare each fresh document with the saved quote inputs.')
table(['Quote','Case','Expected output'],[
['Q-37640','Grouped; Ship-To No','3 groups with 12-month terms; 12 rows; Ship-To hidden.'],
['Q-38009','Ungrouped new; Ship-To Yes','9 rows; Ship-To visible; no Term/Group heading; 2 populated Rate Details lines.'],
['Q-38023','Ungrouped new; Ship-To No','2 rows; Ship-To hidden; no Term/Group heading.'],
['Q-38021','Ungrouped renewal','20 rows; Ship-To hidden; no group heading; 2 populated Rate Details lines.'],
['Q-37723','Grouped renewal; Ship-To Yes','5 groups; 14 rows; Ship-To visible. Group terms missing; structure checks only.'],
['Q-37780','Grouped; Ship-To No','5 groups; 18 rows; Ship-To hidden; 4 populated Rate Details lines. Terms missing; structure checks only.'],
['Q-38029','Amendment exclusion','2 source Amendment lines excluded; 0 eligible output rows. Exclusion check only.']],[1.0,2.0,4.1])
sub('Description and Rate Details checks')
steps(['Use Q-38009, Q-37780, and Q-38021 for populated Rate Details. Match the output text to the source External Description and Rate Details.',
'Confirm Rate Details appear in Description, not a separate column. Check readable spacing and no literal NULL or blank-value placeholders.',
'Check description-only, rates-only, and both-blank cases using suitable existing test data. Blank values must not add dangling separators.',
'For an eligible line with Cancellation, Quantity Increase, or Quantity Reduction, verify the appropriate suffix remains. Use a suitable positive amendment record; Q-38029 cannot demonstrate positive output.'])
p('Backend status: all eight quotes passed the captured checks, including the main Q-37912 case. Visual status: Pending until fresh generated output is reviewed. Q-37640 is grouped and must not be used as the ungrouped test case.')

page('Proposal template reference screenshot')
p('The excerpts below show the actual local proposal template used for the verified content comparison. Merge tokens and conditional branches are visible because this is a template view, not a Salesforce-generated customer document.')
shot(R/'qa-render/page-1.png',(64,783,1140,1045),7.05,'Figure 1. Ungrouped template branches with and without the Ship-To Account column. Only the applicable branch should appear in a generated document.')
shot(R/'qa-render/page-1.png',(64,1108,1140,1262),7.05,'Figure 2. Grouped template heading and product table. The template uses the group subscription-term token and a numbered Term heading.')
sub('Capture fresh output evidence')
p('For Q-37912, capture each Term heading and enough product rows to establish group and bundle order. Record the generated filename and time. For ungrouped quotes, capture the product table and the absence of a group heading. Include source QLE screenshots so the expected sequence can be reproduced.')
p('These reference images do not prove repeated Term numbering, conditional section removal, or final pagination. Those checks must use fresh output from the normal quote-document action.')
sub('Technical source files')
p('The local change record is mapper-change-spec.json. The proposal source file is CPQ Quote Proposal - SALDEV-1403 - dry-run.docx. They describe the six mapper-item changes and the template changes listed on page 2. The recorded native output and source-line captures support the backend results on page 1.')

page('Acceptance checklist and result record')
sub('What remains before acceptance')
for text in ['Confirm sequential Term numbering begins at 1, repeats correctly for each group, and restarts for each generated document.',
'Confirm service dates are absent from headings and the final document has readable tables, descriptions, and page breaks.',
'Complete grouped and ungrouped Ship-To Yes/No checks with fresh document screenshots.',
'Verify positive amendment output with eligible lines. The existing Q-38029 case only proves exclusion.',
'Use grouped records with populated group subscription terms for full heading acceptance. Ten groups across Q-37723 and Q-37780 lack terms; blank Months headings cannot be marked passed.',
'Review output as an intended business user and obtain business-owner acceptance.']:
 p(text,'List Bullet')
sub('Record each business test')
table(['Field','Entry'],[['Quote number and record ID','________________________________'],['Scenario and tester','________________________________'],['Test date and output filename','________________________________'],['Source screenshot reference','________________________________'],['Output screenshot reference','________________________________'],['Expected result','________________________________'],['Actual result and issue reference','________________________________'],['Result','Pending / Pass / Fail / Blocked']],[2.6,4.5])
sub('Evidence boundaries')
p('The 236 backend checks validate captured source order, product counts, description text, conditions, and scoped configuration. Product-name multiplicities do not independently establish quote-line identity because the output does not include quote-line record IDs. Backend success does not replace review of the generated document.')
p('Business reviewer: ____________________    Review date: ______________')
p('Final acceptance: _____________________    Open issues: ______________')

out=OUT/'SALDEV-1403-Component-Changes-and-Test-Guide.docx';d.save(out)
with zipfile.ZipFile(out) as z:
 xml=' '.join(z.read(n).decode('utf-8',errors='ignore') for n in z.namelist() if n.endswith('.xml'))
 assert not re.search(r'\b(?:AI|Codex|ChatGPT|OpenAI|GPT|deployment|deploy|deployed)\b',xml,re.I)
print(out)
