from __future__ import annotations

import shutil
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Pt

from build_architect_review import (
    AMBER,
    BLUE,
    INK,
    LIGHT_AMBER,
    LIGHT_BLUE,
    LIGHT_RED,
    LIGHT_TEAL,
    MUTED,
    NAVY,
    RED,
    TEAL,
    add_body,
    add_figure,
    add_heading,
    arrow,
    draw_wrapped,
    load_font,
    rounded,
    set_cell_text,
    set_run_font,
    set_table_geometry,
    shade_cell,
)
from build_scc3387_architect_review import (
    configure_document,
    normalize_table_margins,
    review_callout,
    review_table,
)


ROOT = Path(__file__).resolve().parent
REFERENCE = ROOT / "SCC-4179-architect-review.docx"
ASSET_DIR = ROOT / "doc-assets-scc4487"
OUTPUT = ROOT / "SCC-4487-PR-Review.docx"
ASSET_DIR.mkdir(exist_ok=True)


def clear_document_body(document: Document) -> None:
    body = document._body._element
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)


def create_segment_path_diagram(path: Path) -> None:
    image = Image.new("RGB", (1600, 1000), "#F7F9FC")
    draw = ImageDraw.Draw(image)
    title = load_font(42, bold=True)
    subtitle = load_font(22)
    heading = load_font(24, bold=True)
    body = load_font(18)
    small = load_font(16, bold=True)

    draw.text((70, 50), "SCC-4487 | Declarative Segment path", font=title, fill="#0B2545")
    draw.text(
        (70, 110),
        "DevA contains the complete declarative path: Account picklist, read-only Order formula, both UI placements, and field visibility.",
        font=subtitle,
        fill="#5B6777",
    )

    boxes = [
        (
            (55, 225, 350, 605),
            "Account source",
            ["Account.Segment__c exists", "Residential + MDU Tenant", "Blank and 3 other values"],
            "#E7F5F1",
            "#16836B",
        ),
        (
            (430, 225, 755, 605),
            "Order formula",
            ["Order.Segment__c present", "Text formula", "TEXT(Account.Segment__c)"],
            "#E7F5F1",
            "#16836B",
        ),
        (
            (835, 225, 1160, 605),
            "Inside Sales UI",
            ["Layout placement correct", "Dynamic Forms placement correct", "Read-only + Inside Sales only"],
            "#E7F5F1",
            "#16836B",
        ),
        (
            (1240, 225, 1545, 605),
            "Access + PR",
            ["3 profiles can read", "Runtime edit access false", "PR diff parity still required"],
            "#EAF2FA",
            "#2E74B5",
        ),
    ]

    for box, box_title, lines, fill, outline in boxes:
        rounded(draw, box, fill, outline=outline, radius=22, width=4)
        draw.text((box[0] + 24, box[1] + 24), box_title, font=heading, fill="#0B2545")
        y = box[1] + 98
        for line in lines:
            draw.ellipse((box[0] + 26, y + 7, box[0] + 38, y + 19), fill=outline)
            wrapped = draw_wrapped(
                draw,
                (box[0] + 52, y),
                line,
                body,
                "#1F2937",
                box[2] - box[0] - 78,
                spacing=6,
            )
            y += max(76, 27 * len(wrapped) + 34)

    for start, end, label, color in (
        ((350, 415), (430, 415), "MIRROR", "#16836B"),
        ((755, 415), (835, 415), "PLACE", "#16836B"),
        ((1160, 415), (1240, 415), "GRANT", "#2E74B5"),
    ):
        arrow(draw, start, end, color=color, width=6)
        draw.text((start[0] + 4, start[1] - 42), label, font=small, fill=color)

    rounded(draw, (70, 685, 1530, 905), "#EAF2FA", outline="#2E74B5", radius=22, width=3)
    draw.text((105, 720), "PR reviewer reading", font=heading, fill="#0B2545")
    decision = (
        "No SCC-4487 implementation defect was identified in GreatPlainsDevA. The formula, UI ordering, visibility, effective "
        "read-only access, blank behavior, reporting filters, and observed order values all align with the story. Final PR approval "
        "still requires the submitted Git diff to match this verified DevA snapshot."
    )
    draw_wrapped(draw, (105, 780), decision, body, "#1F2937", 1360, spacing=8)
    image.save(path)


def add_metadata_table(document: Document) -> None:
    metadata = [
        ("Review target", "SCC-4487 implementation in GreatPlainsDevA (org 00DEa00000Fc086MAB)"),
        ("Review date", "28 August 2026"),
        ("Evidence", "Focused retrieve | live schema | 533-order formula comparison | UI and profile review"),
        ("Decision", "Implementation Approved | PR diff parity pending"),
        ("Boundary", "PR review only | no deployment, activation, assignment, metadata, or record changes"),
    ]
    table = document.add_table(rows=len(metadata), cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [1900, 7460])
    normalize_table_margins(table)
    for row, (label, value) in zip(table.rows, metadata):
        shade_cell(row.cells[0], LIGHT_BLUE)
        set_cell_text(row.cells[0], label, size=9.5, color=NAVY, bold=True)
        set_cell_text(row.cells[1], value, size=9.5)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def build_document() -> None:
    diagram_path = ASSET_DIR / "scc4487-segment-path.png"
    create_segment_path_diagram(diagram_path)

    shutil.copy2(REFERENCE, OUTPUT)
    document = Document(OUTPUT)
    clear_document_body(document)
    configure_document(document)
    for section in document.sections:
        for footer in (section.footer, section.even_page_footer, section.first_page_footer):
            for paragraph in footer.paragraphs:
                for run in paragraph.runs:
                    if "GreatPlainsMerge" in run.text:
                        run.text = run.text.replace("GreatPlainsMerge", "GreatPlainsDevA")

    # Page 1: PR decision.
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(2)
    set_run_font(paragraph.add_run("PR REVIEW"), size=10, color=BLUE, bold=True)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(
        paragraph.add_run("SCC-4487 | Add Segment to Inside Sales Order"),
        size=25,
        color=NAVY,
        bold=True,
    )

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(16)
    set_run_font(
        paragraph.add_run("GreatPlainsDevA | Declarative metadata pull-request assessment"),
        size=13,
        color=MUTED,
    )

    add_metadata_table(document)
    review_callout(
        document,
        "PR decision",
        "Implementation Approved in DevA. No story-related code or metadata defect was identified. The formula, both Inside Sales placements, and effective read-only access satisfy the acceptance criteria. Because no PR URL, branch, commit, or patch was supplied, final Git PR approval is conditional on diff parity with this verified snapshot.",
        fill=LIGHT_TEAL,
        accent=TEAL,
    )

    add_heading(document, "PR review outcome", 1)
    outcome_rows = [
        ("1", "PR branch / diff", "Boundary", "No PR URL, branch, commit, or patch was supplied; approve the Git PR only after confirming its diff matches DevA."),
        ("2", "Order.Segment__c", "Pass", "Present as a Text formula: TEXT(Account.Segment__c); schema reports calculated=true, createable=false, and updateable=false."),
        ("3", "Inside Sales layout", "Pass", "Segment__c is directly after Customer_Type__c and is explicitly Readonly."),
        ("4", "Lightning record page", "Pass", "Record.Segment__c is directly after Customer Type, readonly, and visible only when record type name is Inside Sales."),
        ("5", "Named profiles", "Pass", "Admin, API-only Admin, and Standard User have effective read=true/edit=false field access."),
        ("6", "Formula results", "Pass", "All 533 sampled Inside Sales orders matched Account Segment: 158 Residential, 3 MDU Tenant, 372 blank, 0 mismatches."),
    ]
    review_table(
        document,
        ["#", "Scope", "Result", "Reviewer evidence"],
        outcome_rows,
        [500, 2150, 1450, 5260],
        status_column=2,
    )

    # Page 2: declarative design path.
    document.add_page_break()
    add_heading(document, "Declarative design path", 1)
    add_body(
        document,
        "The DevA implementation is metadata-only: a text formula on Order mirrors the Account picklist at read time, the field is placed in both Inside Sales presentation layers, and the named profiles receive effective read-only field visibility.",
    )
    add_figure(
        document,
        diagram_path,
        "Figure 1. SCC-4487 declarative Segment path",
        "Diagram showing Account Segment feeding the verified Order formula field, then the verified Inside Sales layout and Dynamic Forms placements, followed by effective read-only profile access and a remaining Git diff parity check.",
        width=6.45,
    )
    review_callout(
        document,
        "Design reading",
        "Green components are verified in DevA. Blue marks access and the remaining PR-evidence boundary. The formula calculates at read time, so SCC-4487 does not require Apex, Flow, trigger, or persisted copy logic.",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )

    # Page 3: component evidence.
    document.add_page_break()
    add_heading(document, "Component-by-component PR evidence", 1)
    component_rows = [
        ("CustomField", "Order.Segment__c", "Verified", "Text formula TEXT(Account.Segment__c); label, description, and help text are story-specific."),
        ("Layout", "Order-Inside Sales", "Verified", "Customer_Type__c -> Segment__c (Readonly) -> Additional_Information_Online_Form__c."),
        ("FlexiPage", "OSE_CPQOrderRecordPage", "Verified", "Record.Segment__c follows Customer Type, is readonly, and has an Inside Sales visibility rule."),
        ("Profile", "Admin", "Verified", "Field is readable; effective PermissionsEdit is false because the field is calculated."),
        ("Profile", "System Administrator - API Only", "Verified", "Field is readable; effective PermissionsEdit is false because the field is calculated."),
        ("Profile", "Standard", "Verified", "Field is readable and not editable; no unrelated record-type visibility change is present."),
    ]
    review_table(
        document,
        ["Type", "Component", "Observed DevA state", "Reviewer evidence"],
        component_rows,
        [1450, 2750, 2000, 3160],
    )

    add_heading(document, "Source-field and packaging evidence", 2)
    source_rows = [
        ("Account.Segment__c", "Picklist; editable; blank allowed", "Source field exists."),
        ("Active source values", "MDU Bulk; MDU Tenant; SFU; Residential; Other", "Direct mirroring correctly preserves every current Account value; the ticket's two values are examples."),
        ("Order schema", "Calculated; non-createable; non-updateable; nillable; filterable", "Read-only, blank-safe, and usable as a report/list filter."),
        ("Observed records", "533 Inside Sales orders; 0 mismatches", "Includes Residential, MDU Tenant, and blank results."),
        ("Expected package", "1 field + 1 layout + 1 FlexiPage + 3 profiles", "Confirm the submitted PR contains these six components without unrelated churn."),
    ]
    review_table(document, ["Evidence area", "Observed value", "Reviewer interpretation"], source_rows, [2200, 3600, 3560])
    review_callout(
        document,
        "Evidence boundary",
        "These findings are based on a focused read-only GreatPlainsDevA snapshot and live formula results. They validate the org implementation but cannot prove that an unsupplied Git branch or pull-request diff contains the same metadata.",
        fill=LIGHT_AMBER,
        accent=AMBER,
    )

    # Page 4: findings and modifications.
    heading = add_heading(document, "PR findings and reviewer notes", 1)
    heading.paragraph_format.page_break_before = True
    finding_rows = [
        ("Gate", "PR traceability", "The actual Git diff is unavailable.", "Attach the PR URL/branch/commit and compare the six metadata components with DevA."),
        ("Pass", "Core field", "The formula is implemented as designed.", "No modification required if the PR contains the verified field metadata."),
        ("Pass", "UI placement", "Both presentation layers place Segment below Customer Type.", "No modification required if the PR diff matches DevA."),
        ("Pass", "Field access", "All three named profiles have effective read-only access.", "Do not add edit access or unrelated permission changes."),
        ("Note", "Value semantics", "Account currently has five active values.", "Direct mirroring is consistent with 'same value'; no restrictive CASE mapping is needed."),
        ("Note", "Change scope", "The formula calculates at read time.", "Keep the PR declarative; do not add Apex, Flow, trigger, or record-copy logic."),
        ("Note", "Profile XML", "Admin source XML may show editable=true.", "The live FieldPermissions API reports edit=false and schema updateable=false; this is not a runtime defect."),
    ]
    review_table(
        document,
        ["Result", "Area", "Finding", "Reviewer action"],
        finding_rows,
        [1100, 1700, 2900, 3660],
    )

    add_heading(document, "Formula review standard", 2)
    formula_rows = [
        ("Return type", "Text formula", "The source is a picklist; Order must display its current label as read-only text."),
        ("Recommended formula", "TEXT(Account.Segment__c)", "Mirrors the current Account value at read time and naturally returns blank when the source is blank."),
        ("Security", "Effective read=true; edit=false", "Live schema and FieldPermissions confirm the formula cannot be edited."),
        ("Dependencies", "Account relationship only", "No trigger, Flow, scheduled process, or persisted synchronization is required."),
    ]
    review_table(document, ["Review item", "Required design", "Reason"], formula_rows, [1900, 2900, 4560])
    review_callout(
        document,
        "PR approval gate",
        "The DevA implementation is acceptable. Approve the submitted PR after confirming its six-component diff is identical in substance, contains no unrelated churn, and retains the metadata-only design.",
        fill=LIGHT_TEAL,
        accent=TEAL,
    )

    # Page 5: reviewer checklist and response.
    document.add_page_break()
    add_heading(document, "Reviewer approval checklist", 1)
    checklist_rows = [
        ("1", "PR identity", "PR URL, source branch, target branch, commit SHA, and Jira key are present."),
        ("2", "Package completeness", "All six components are in the diff; omissions and extra dependencies are explained."),
        ("3", "Formula metadata", "Order.Segment__c is Text Formula with TEXT(Account.Segment__c), a clear description, and blank-safe behavior."),
        ("4", "Layout order", "Segment is directly after Customer_Type__c on Order-Inside Sales."),
        ("5", "Dynamic Forms order", "Record.Segment__c directly follows Customer Type, is readonly, and is limited to Inside Sales."),
        ("6", "Profile access", "Admin, API-only Admin, and Standard User retain effective read=true/edit=false access."),
        ("7", "Scope discipline", "No Apex, Flow, trigger, record-update logic, record-type assignment, or unrelated metadata churn is introduced."),
        ("8", "Org parity", "The PR metadata matches the verified DevA snapshot summarized in this review."),
    ]
    review_table(document, ["#", "Reviewer check", "Approval evidence"], checklist_rows, [500, 2600, 6260])

    add_heading(document, "Ready-to-paste PR review comment", 2)
    response = (
        "PR Review completed for SCC-4487 - Implementation Approved in GreatPlainsDevA. Order.Segment__c is a Text formula "
        "using TEXT(Account.Segment__c), is calculated and non-editable, and is filterable for reports and list views. Segment "
        "appears directly below Customer Type on both Order-Inside Sales and OSE_CPQOrderRecordPage; the Dynamic Forms field is "
        "readonly and limited to the Inside Sales record type. Admin, System Administrator - API Only, and Standard User have "
        "effective read=true/edit=false access. A read-only comparison of 533 Inside Sales orders found 158 Residential, 3 MDU "
        "Tenant, and 372 blank values with zero Account-to-Order mismatches. No SCC-4487 code change is required. Final Git PR "
        "approval is conditional on the submitted six-component diff matching this DevA snapshot and containing no unrelated "
        "metadata churn. This review made no Salesforce changes."
    )
    review_callout(document, "PR comment", response, fill=LIGHT_BLUE, accent=BLUE)

    document.core_properties.title = "SCC-4487 Add Segment to Inside Sales Order PR Review"
    document.core_properties.subject = "Salesforce declarative metadata pull-request review"
    document.core_properties.author = "Solution Architecture Review"
    document.core_properties.keywords = (
        "Salesforce, Great Plains, SCC-4487, Segment, Order, formula, layout, FlexiPage, profile, PR review"
    )
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
