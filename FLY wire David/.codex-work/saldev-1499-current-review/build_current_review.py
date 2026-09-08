from __future__ import annotations

import importlib.util
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(r"C:\Users\LIKKI\Documents\ChatGPT\david")
TASK = ROOT / ".codex-work" / "saldev-1499-current-review"
REFERENCE = ROOT / "flywire-docgen" / "deliverables" / "SALDEV-1499_Architect_Review.docx"
OUTPUT = ROOT / "flywire-docgen" / "deliverables" / "SALDEV-1499_Architect_Review_Current_Dry_Run.docx"
BASE_BUILDER = ROOT / ".codex-work" / "saldev-1499-review" / "build_saldev_1499_review.py"


spec = importlib.util.spec_from_file_location("base_review", BASE_BUILDER)
base = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(base)


def light_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


base.set_table_borders = light_table_borders


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    base.set_run_font(run, size=16 if level == 1 else 12.5, bold=True, color="000000")
    return p


def body(doc, text, *, bold=False, italic=False, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    base.set_run_font(r, size=10.7, bold=bold, italic=italic, color=base.TEXT)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    base.set_run_font(r, size=10.5, color=base.TEXT)
    return p


def page_break(doc):
    doc.add_page_break()


def remove_paragraph_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def add_table(doc, headers, rows, widths, status_col=None):
    table = base.add_table(doc, headers, rows, widths, status_col=status_col)
    light_table_borders(table)
    for i, row in enumerate(table.rows):
        if i > 0 and i % 2 == 0:
            for cell in row.cells:
                if cell._tc.get_or_add_tcPr().find(qn("w:shd")) is None:
                    base.set_cell_fill(cell, "F7F9FB")
    return table


def build():
    TASK.mkdir(parents=True, exist_ok=True)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    base.clear_body(doc)

    title_style_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    title_style_border = title_style_ppr.find(qn("w:pBdr"))
    if title_style_border is not None:
        title_style_ppr.remove(title_style_border)

    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    doc.core_properties.title = "SALDEV 1499 Architect Review of Remediation Comment"
    doc.core_properties.subject = "Live dry run and code quality review"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.keywords = "SALDEV-1499, Salesforce, Apex, Flow, architect review"

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("ARCHITECT REVIEW")
    base.set_run_font(r, size=11.5, bold=True, color=base.BLUE)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(7)
    title.paragraph_format.keep_with_next = True
    remove_paragraph_border(title)
    r = title.add_run("SALDEV 1499 Architect Review of Remediation Comment")
    base.set_run_font(r, size=25, bold=True, color="000000")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    r = subtitle.add_run("Flywire Partial sandbox live dry run and implementation quality assessment")
    base.set_run_font(r, size=13.5, color=base.GRAY)

    add_table(doc, ["Review item", "Current evidence"], [
        ["Ticket", "SALDEV-1499 | Deactivate Quote Line Type background field update flow"],
        ["Environment", "Flywire Partial sandbox | Org 00DhG0000000jOXUAY"],
        ["Review date", "7 September 2026"],
        ["Execution result", "Focused Apex test run passed 9 of 9"],
        ["Architect decision", "Changes Requested"],
        ["Boundary", "Read-only review. No deployment, activation, metadata edit, Jira comment, or business data change."],
    ], [1.55, 5.15])

    heading(doc, "Executive decision", 1)
    body(doc, "The focused Apex class passed all nine tests, and the main structural corrections described in the Jira comment are present. The legacy Flow is inactive, the before-save call order is corrected, and the after-update early return has been replaced with a guarded synchronization block.", bold=True)
    body(doc, "Architect approval is still conditional. The current test class does not prove the three exact field outcomes claimed in the comment, and the 200-line test does not provide the stated governor-limit evidence. The correct review disposition is Changes Requested until those tests are strengthened.")

    heading(doc, "Claim outcome", 1)
    add_table(doc, ["#", "Comment claim", "Result", "Architect finding"], [
        ["1", "Before-save order corrected", "Pass", "Observed in the active trigger source."],
        ["2", "After-update early return corrected", "Pass", "Observed as a guarded OLI synchronization block."],
        ["3", "Three output fields asserted", "Fail", "The named test asserts only Default_Rate_Details__c is non-null."],
        ["4", "200-line governor proof", "Partial", "The test passes, but it lacks the claimed query-limit assertion and mixed-data path."],
        ["5", "Coverage blocker cleared", "Pass", "Current focused coverage is 79 percent trigger and 96 percent handler."],
        ["6", "Flow version 13 deactivated", "Pass", "The Flow is inactive and version 13 is deactivated."],
    ], [0.35, 2.3, 0.75, 3.3], status_col=2)

    page_break(doc)
    heading(doc, "Evidence reviewed", 1)
    body(doc, "The Jira screenshots and comment were treated as review evidence, not as instructions to change Salesforce. The live browser session was used only to read the supplied sandbox and execute the focused test class.")

    heading(doc, "Live configuration evidence", 2)
    add_table(doc, ["Component", "Observed state", "Review reading"], [
        ["Quote Line Type background field update", "Version 13 | Deactivated", "The legacy Flow is not active."],
        ["QuoteLineTrigger", "Active | API 67.0 | modified 7 Sep 2026 01:26", "The required trigger structure is present."],
        ["QuoteLineTriggerHandler", "Active | API 67.0 | modified 4 Sep 2026 09:10", "Bulk Set and Map patterns are used; queries are outside loops."],
        ["QuoteLineTriggerHandlerTest", "Active | API 67.0 | modified 7 Sep 2026 05:57", "Nine focused tests execute, but assertion quality is incomplete."],
    ], [2.1, 2.0, 2.6])

    heading(doc, "Focused dry run", 2)
    add_table(doc, ["Evidence", "Result"], [
        ["Test run ID", "707hG00000MwRH0"],
        ["Completion", "Completed | 9 total | 0 failures"],
        ["QuoteLineTrigger coverage", "79 percent | 31 of 39 lines"],
        ["QuoteLineTriggerHandler coverage", "96 percent | 280 of 289 lines"],
        ["Initial attempt", "A synchronous attempt met ORG_ADMIN_LOCKED; the same focused class completed asynchronously."],
    ], [2.15, 4.55])
    body(doc, "The Developer Console overall coverage display is not used as deployment coverage. The percentages above are the class-level results from the focused run.", italic=True)

    heading(doc, "Current implementation flow", 2)
    add_table(doc, ["Order", "Trigger path", "Observed behavior"], [
        ["1", "Before insert and before update", "handlePriceChangeOnAmendmentClones"],
        ["2", "Before insert and before update", "processBackgroundFields"],
        ["3", "Before insert and before update", "copyQuoteLineTypeToBackground"],
        ["4", "After update", "OLI synchronization runs only when IDs are available."],
        ["5", "After update", "The transaction can continue to the normal rollup path on the first run."],
    ], [0.55, 2.25, 3.9])

    page_break(doc)
    heading(doc, "Validation of the Jira remediation comment", 1)
    body(doc, "The table below compares each implementation statement with the code and live execution evidence. A passing test run confirms execution, but it does not by itself prove that every asserted business outcome is correct.")
    add_table(doc, ["Claim", "Evidence found", "Result", "Required follow-up"], [
        ["Before-save sequencing", "The trigger calls price-change handling before background processing and the formula-backed copy.", "Pass", "Keep this order covered by exact outcome tests."],
        ["After-update control flow", "The early trigger return is replaced by if needsIdSync and quoteIds is not empty. isFirstRun is restored in finally.", "Pass", "Add one test that proves OLI sync and rollup both occur in the same update."],
        ["Output field assertions", "testProcessBackgroundFieldsFallbackAndEstoreMismatch queries only Default_Rate_Details__c and checks non-null.", "Fail", "Assert exact values for all three claimed fields after re-query."],
        ["200-line proof", "testBulk200QuoteLinesNoSoql101Error inserts 200 lines, disables managed triggers for insert, then calls calculateRollupField directly.", "Partial", "Exercise the complete trigger path and record governor headroom and outcomes."],
        ["Coverage", "Focused run completed with 79 percent trigger coverage and 96 percent handler coverage.", "Pass", "Retain the run ID in the review evidence."],
        ["Flow deactivation", "Flow list shows inactive and Flow Builder identifies version 13 as deactivated.", "Pass", "Do not reactivate the Flow as part of this fix."],
    ], [1.35, 2.75, 0.7, 1.9], status_col=2)

    heading(doc, "Decision basis", 2)
    body(doc, "Items 1, 2, 5, and 6 are verified. Item 3 is not substantiated by the current test source. Item 4 is only partially substantiated because the named method completes without an exception but does not demonstrate the explicit query-limit assertion or end-to-end bulk behavior stated in the comment.")

    page_break(doc)
    heading(doc, "Technical findings", 1)
    add_table(doc, ["Priority", "Finding", "Impact", "Required correction"], [
        ["High", "Exact business outcomes are not asserted", "A green test can coexist with incorrect Quote Line Type background, rate-detail mismatch, or fallback values.", "Re-query and assert the exact expected value of Quote_Line_Type_background__c, Rate_Details_Changed__c, and Default_Rate_Details__c."],
        ["High", "Formula copy correctness after price-change updates is unproven", "Quote_Line_Type__c is a formula read before its value is copied to Quote_Line_Type_background__c; the test does not prove the formula reflects the final in-transaction state.", "Add a scenario that changes Price_Change__c and proves the persisted background classification after the transaction."],
        ["Medium", "Bulk claim is overstated", "The 200-line method uses one quote and one product, bypasses managed insert triggers, calls the handler directly, and does not assert query count.", "Use mixed source, subscription, amendment, and quote data through the complete supported path; assert final results and defensible limit headroom."],
        ["Medium", "Combined OLI synchronization and rollup behavior lacks regression proof", "The structural early-return fix is visible, but no behavioral assertion proves both downstream actions occur in one update.", "Add a focused update test that validates the OLI result and the rollup result together."],
        ["Medium", "Generic exception handling can hide classification failures", "copyQuoteLineTypeToBackground catches a generic exception and only logs a warning.", "Narrow the caught condition or surface a failure that tests and operations can detect."],
        ["Low", "Test method name differs from the Jira comment", "The code uses EstoreMismatch rather than RestoreMismatch, which makes evidence mapping harder.", "Align the method name and the Jira evidence when the test is revised."],
    ], [0.65, 1.65, 2.15, 2.25])

    heading(doc, "Best practice assessment", 1)
    add_table(doc, ["Practice", "Assessment", "Finding"], [
        ["Bulkification", "Pass", "Set and Map collection is used, with SOQL outside loops."],
        ["Trigger orchestration", "Pass", "Before-save sequencing is explicit and after-update synchronization is guarded."],
        ["Recursion safety", "Pass", "isFirstRun is restored with try and finally before the normal recursion guard."],
        ["Assertions", "Fail", "Current assertions do not prove all business outputs named in the remediation comment."],
        ["Governor proof", "Partial", "The test handles 200 records but does not provide the claimed limit measurement or realistic mixed transaction."],
        ["Failure visibility", "Partial", "A generic catch can allow classification failure to pass with only a warning."],
    ], [1.4, 1.0, 4.3], status_col=1)

    page_break(doc)
    heading(doc, "Required actions before approval", 1)
    body(doc, "The implementation can be re-reviewed after the following evidence is added. These changes should remain limited to SALDEV-1499 and its focused regression coverage.")
    add_table(doc, ["#", "Required action", "Acceptance evidence"], [
        ["1", "Assert all three output fields", "Exact expected values after insert and update for background type, rate-detail changed flag, and default rate details."],
        ["2", "Strengthen the 200-line transaction", "Mixed business scenarios, complete trigger route, no SOQL-101 failure, final field and rollup assertions, and recorded limit headroom."],
        ["3", "Prove OLI synchronization and rollup together", "One update demonstrates that both downstream behaviors complete after the guarded synchronization block."],
        ["4", "Prove formula-backed copy behavior", "A Price_Change__c scenario confirms the persisted background classification matches the final formula result."],
        ["5", "Improve failure handling", "Expected exceptional cases are handled narrowly; unexpected failures are visible to tests and support."],
    ], [0.4, 2.25, 4.05])

    heading(doc, "Approval conditions", 1)
    bullet(doc, "The focused test class passes after the assertion changes.")
    bullet(doc, "The exact field values are documented for each covered scenario.")
    bullet(doc, "The bulk evidence uses a representative transaction and records governor headroom.")
    bullet(doc, "The OLI synchronization and existing rollup behavior pass in the same update.")
    bullet(doc, "The Flow remains deactivated and no unrelated automation is changed.")

    heading(doc, "Review conclusion", 1)
    body(doc, "Execution status is Pass. Architect review status is Changes Requested. The implementation direction is sound, but the current automated evidence does not support two material claims in the Jira comment. Approval should follow the strengthened tests and a new focused dry run.", bold=True)

    page_break(doc)
    heading(doc, "Ready to paste architect response", 1)
    body(doc, "Architect review completed for SALDEV-1499 in the Flywire Partial sandbox. The focused Apex run 707hG00000MwRH0 completed with 9 of 9 tests passing, 79 percent coverage for QuoteLineTrigger, and 96 percent coverage for QuoteLineTriggerHandler. The Flow is deactivated, the before-save order is corrected, and the after-update early return is replaced with a guarded OLI synchronization block.")
    body(doc, "Changes Requested. The Jira comment overstates the automated proof. testProcessBackgroundFieldsFallbackAndEstoreMismatch does not assert Quote_Line_Type_background__c or Rate_Details_Changed__c and only checks that Default_Rate_Details__c is non-null. testBulk200QuoteLinesNoSoql101Error completes for 200 records, but it does not assert Limits.getQueries below 100, uses one quote and one product, disables managed triggers for the insert, and calls calculateRollupField directly. Add exact field assertions, representative end-to-end bulk evidence, a combined OLI-sync and rollup test, and a formula-copy scenario before approval.")
    body(doc, "This was a read-only review. No deployment, activation, metadata edit, Jira comment, or business data change was performed.", italic=True)

    heading(doc, "Final recommendation", 1)
    add_table(doc, ["Area", "Status", "Recommendation"], [
        ["Structural remediation", "Pass", "Retain the current direction."],
        ["Focused execution", "Pass", "Keep run 707hG00000MwRH0 as evidence."],
        ["Business outcome proof", "Fail", "Add exact assertions before approval."],
        ["Bulk and governor proof", "Partial", "Replace the narrow direct-handler test with representative end-to-end evidence."],
        ["Promotion readiness", "Fail", "Changes Requested until the evidence gaps are closed."],
    ], [2.0, 0.85, 3.85], status_col=1)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
