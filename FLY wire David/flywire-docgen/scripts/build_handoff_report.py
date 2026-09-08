from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('deliverables/Flywire_Stepped_Up_Pricing_DocGen_Handoff.docx')
NAVY = '17365D'; BLUE = '2E74B5'; LIGHT = 'E8EEF5'; GRAY = 'F2F4F7'; GREEN = 'E2F0D9'; AMBER = 'FFF2CC'; WHITE='FFFFFF'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc.get_or_add_tcPr(); tcMar = tc.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tc.append(tcMar)
    for name,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        x=tcMar.find(qn('w:'+name))
        if x is None: x=OxmlElement('w:'+name); tcMar.append(x)
        x.set(qn('w:w'),str(val)); x.set(qn('w:type'),'dxa')

def set_repeat(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    r=p.add_run(str(text)); r.bold=bold; r.font.name='Calibri'; r.font.size=Pt(size)
    if color: r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(cell)

def table(doc, headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,(h,w) in enumerate(zip(headers,widths)):
        t.columns[i].width=Inches(w); set_cell_text(t.rows[0].cells[i],h,True,WHITE,9); shade(t.rows[0].cells[i],NAVY)
    set_repeat(t.rows[0])
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells
        for i,(v,w) in enumerate(zip(row,widths)):
            cells[i].width=Inches(w); set_cell_text(cells[i],v,size=8.6)
            if ridx%2: shade(cells[i],GRAY)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def bullet(doc, text, level=0):
    p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.add_run(text); return p

def numbered(doc, text):
    p=doc.add_paragraph(style='List Number'); p.add_run(text); return p

def callout(doc, title, text, fill=LIGHT):
    t=doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    c=t.cell(0,0); shade(c,fill); margins(c,140,160,140,160); c.text=''
    p=c.add_paragraph(); p.paragraph_format.space_after=Pt(3); r=p.add_run(title); r.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY)
    p=c.add_paragraph(text); p.paragraph_format.space_after=Pt(0)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
styles=doc.styles
styles['Normal'].font.name='Calibri'; styles['Normal'].font.size=Pt(10.5); styles['Normal'].paragraph_format.space_after=Pt(6); styles['Normal'].paragraph_format.line_spacing=1.1
for name,size,before,after,color in [('Title',27,0,8,NAVY),('Subtitle',13,0,18,'666666'),('Heading 1',16,16,8,BLUE),('Heading 2',13,12,6,BLUE),('Heading 3',11.5,8,4,NAVY)]:
    s=styles[name]; s.font.name='Calibri'; s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
    if name!='Subtitle': s.font.bold=True

header=sec.header.paragraphs[0]; header.text='FLYWIRE PARTIAL SANDBOX  |  IMPLEMENTATION HANDOFF'; header.runs[0].font.size=Pt(8); header.runs[0].font.color.rgb=RGBColor.from_string('777777')
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT; footer.add_run('Stepped-Up Pricing Quote Document  |  August 2026').font.size=Pt(8)

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(36); r=p.add_run('IMPLEMENTATION HANDOFF'); r.bold=True; r.font.size=Pt(10); r.font.color.rgb=RGBColor.from_string(BLUE)
doc.add_paragraph('Stepped-Up Pricing Quote Document', style='Title')
doc.add_paragraph('User story, technical implementation, deployment inventory, validation evidence, and reproduction guide', style='Subtitle')
table(doc,['Field','Value'],[
    ('Target environment','Flywire Partial sandbox'),('Salesforce org alias','FlywirePartial'),('Primary document template','CPQ Quote Proposal'),('Active OmniScript','CPQ / QuoteDocument / English - Version 2'),('Prepared for','David and Flywire CPQ / OmniStudio stakeholders'),('Prepared on','13 August 2026')],[1.75,4.75])
callout(doc,'Outcome','Grouped and ungrouped quote scenarios now return a valid document-generation data structure. The active Word template renders grouped pricing with group names, service periods, and conditional Ship-To columns while preserving the ordinary quote layout.',GREEN)

doc.add_heading('1. Executive summary',1)
doc.add_paragraph('The request was to enhance the existing CPQ quote proposal so stepped-up pricing quotes display each quote-line group as a distinct section. The implementation was completed in the Partial sandbox by correcting the OmniStudio Data Mapper output structure and updating the Microsoft Word document template. The solution also preserves the pre-existing ungrouped quote presentation.')
for x in ['Grouped quotes generate one section per quote-line group.','Group name and Start Date / End Date are displayed as the service period.','Quote lines stay beneath their owning group and are not duplicated by the mapper.','The Ship-To Account column is conditional: shown for Multiple Ship-To Accounts = Yes and omitted for No.','Ungrouped quotes continue to use the original product table and do not display an empty Group 1 section.']: bullet(doc,x)

doc.add_heading('2. User story and requested behavior',1)
doc.add_heading('Business story',2)
doc.add_paragraph('As a CPQ user generating a quote proposal, I need stepped-up pricing products to be organised by quote-line group so the customer can understand which products and prices apply during each service period. The proposal must adapt its columns when multiple Ship-To accounts are used and must not alter the ordinary quote layout.')
doc.add_heading('Acceptance criteria',2)
table(doc,['Area','Required behavior'],[
('Grouped output','Every quote-line group appears separately.'),('Group header','Display the group name.'),('Service period','Display group Start Date and End Date.'),('Line ownership','Products appear beneath the correct group without duplication.'),('Conditional column','Show Ship-To Account only when Multiple Ship-To Accounts = Yes.'),('Line fields','Populate product, quantity, billing frequency, description, and price.'),('Regression','Ungrouped quotes retain the existing table with no group heading or empty Group 1.')],[1.45,5.05])

doc.add_heading('3. Solution design',1)
doc.add_paragraph('Runtime flow: the OmniScript receives the Quote record ID as Context ID, extracts quote and quote-line data, selects the active Document Template, runs the template Data Mapper, and passes the transformed JSON plus the DOCX template to the client-side document-generation component.')
numbered(doc,'GetQuoteProposalData executes GetQuoteProposalDataSteppedUpPricing with ContextId as Id.')
numbered(doc,'GetDocumentTemplate selects CPQ Quote Proposal (Microsoft Word, client-side generation).')
numbered(doc,'CPQQuoteProposalDocumentSteppedUpPricing converts the extract result into the document token model.')
numbered(doc,'The corrected model contains one root key-value object. Group is a list; each Group contains its QuoteLine list.')
numbered(doc,'The Word template loops over Group and then QuoteLine, using conditional sections for grouped/ungrouped and multiple/non-multiple Ship-To scenarios.')
callout(doc,'Critical defect corrected','The earlier transform fanned each group into a separate top-level record. Client-side DocGen rejected that list with “Input object is not a key-value pair.” The final declarative mapping outputs one root map with Group as the nested list.',AMBER)

doc.add_heading('4. Component and change inventory',1)
table(doc,['Component','Type','Created / changed','Final state'],[
('GetQuoteProposalDataSteppedUpPricing','Omni Data Mapper - Extract','Existing; used and validated','Active; no extract errors for all scenarios'),
('CPQQuoteProposalDocumentSteppedUpPricing','Omni Data Mapper - Transform','Changed output mappings and group-list structure','Active; JSON root Map with Group list'),
('CPQ Quote Proposal','DocumentTemplate','Existing record; content link updated','Active; ID 2dtPb0000000BwLIAU'),
('CPQ Quote Proposal - SteppedUpPricing-v5.docx','ContentVersion / DOCX','Created final template file','Active linked file; 068hG0000033ObzQAE'),
('Generate Quote Document - SteppedUpPricing v2','OmniScript','Existing active version used for testing','Active; ID 0jNhG0000000umzUAA'),
('CPQQuoteProposalOutputSerializer','Apex class','Created during investigation','Deployed but NOT connected; no runtime effect'),
('CPQQuoteProposalOutputSerializerTest','Apex test class','Created to test serializer experiment','3 focused tests passed; class not used by final design')],[1.55,1.25,2.15,1.55])

doc.add_heading('5. Detailed Data Mapper changes',1)
doc.add_paragraph('The final solution uses a whole-list declarative mapping. The existing QuoteLineGroup array is mapped directly to the output Group array. Child mappings that caused top-level fan-out were disabled. Existing quote-level fields and formulas remain available at the single root level.')
table(doc,['Change','Input','Output / behavior'],[
('Whole group list mapping','QuoteLineGroup','Group'),('Nested source lines','Group:QuoteLine','Template repeats QuoteLine within each Group'),('Group identity','QuoteLineGroup:Name','Group.Name'),('Service period','QuoteLineGroup:StartDate / EndDate','Group.StartDate / Group.EndDate'),('Ship-To conditional','MultipleShipToAccounts','IF_MultipleShipToAccounts / IF_NoMultipleShipToAccounts'),('Grouped conditional','QuoteLineGroupsPresent','IF_QuoteLineGroupsPresentTrue / False'),('Fan-out prevention','12 former child/group mappings','Disabled after whole-list mapping')],[1.55,2.25,2.7])
doc.add_paragraph('Earlier targeted corrections also included the grouped adjusted-rate path, grouped line-field paths, the false-group formula, and the LineDescriptionWithType field spelling. These were superseded where necessary by the final whole-list mapping; the active output was verified after the final change.')

doc.add_heading('6. Word template changes',1)
table(doc,['Template area','Implementation'],[
('Grouped section','Added Group repeat blocks for both Ship-To variants.'),('Group heading','Token: {{Name}}.'),('Service period','Tokens: {{StartDate}} and {{EndDate}}.'),('Grouped product rows','Nested repeat changed to {{#QuoteLine}} ... {{/QuoteLine}}.'),('Multiple Ship-To grouped table','Includes {{ShipToAccountName}}.'),('No Multiple Ship-To grouped table','Omits the Ship-To Account column.'),('Ungrouped branch','Retains existing {{#Line}} table and prior layout.'),('Template integrity','Removed duplicate nested Group block and repaired conditional start/end nesting.'),('Active version','CPQ Quote Proposal - SteppedUpPricing-v5.docx')],[1.85,4.65])

doc.add_heading('7. Files created in the local workspace',1)
table(doc,['Path','Purpose','Delivery relevance'],[
('templates/CPQ Quote Proposal - SteppedUpPricing-v5.docx','Final Word template uploaded to Salesforce','Final runtime artifact'),
('backup/CPQ Quote Proposal - v3-org-backup.docx','Original org template backup','Rollback reference'),
('scripts/build_stepped_template.py','Reproducibly builds the corrected DOCX','Engineering support'),
('scripts/validate_stepped_bundle.py','Checks tokens, nesting, table branches, and mapper bundle','Validation support'),
('scripts/apex/deploy_group_list_mapping.apex','Applies the final declarative Group list mapping','Deployment script'),
('scripts/apex/dry_run_three_quote_scenarios.apex','Validates the three acceptance scenarios','Test evidence'),
('force-app/.../CPQQuoteProposalOutputSerializer.cls','Investigative custom-output approach','Not connected / not used'),
('force-app/.../CPQQuoteProposalOutputSerializerTest.cls','Tests investigative serializer','Not part of runtime solution')],[2.9,2.3,1.3])

doc.add_heading('8. Deployment record and identifiers',1)
table(doc,['Item','Identifier / result'],[
('Target org','FlywirePartial - Partial sandbox'),('DocumentTemplate','2dtPb0000000BwLIAU'),('DocumentTemplateContentDoc link','2ddhG0000001bwnQAA'),('Final ContentDocument','069hG000003DZfnQAG'),('Final ContentVersion','068hG0000033ObzQAE'),('Active OmniScript version','0jNhG0000000umzUAA'),('Transform Data Mapper','0jIhG0000000HX7UAM'),('Final mapper update','1 whole-list mapping enabled; 12 fan-out mappings disabled'),('Apex focused deployment','Serializer and test classes deployed; 3 tests passed; serializer not attached')],[2.15,4.35])

doc.add_heading('9. Validation evidence',1)
doc.add_paragraph('The final read-only Extract -> Transform dry run returned the following results. These checks prove data shape and conditional routing; visual PDF review confirms final pagination and appearance.')
table(doc,['Quote / Context ID','Scenario','Dry-run evidence','Result'],[
('Q-37723\na2NhG000002ktrcUAA','5 groups; Multiple Ship-To = Yes','root Map; groups=5; IF grouped=true; IF multi=true','PASS'),
('Q-37780\na2NhG000003EetVUAS','5 groups; Multiple Ship-To = No','root Map; groups=5; IF grouped=true; IF no-multi=true','PASS'),
('Q-37640\na2NhG000002YPkPUAW','Ungrouped ordinary quote','root Map; groups=0; ungrouped lines=4; IF grouped=false','PASS')],[1.85,1.65,2.35,.65])

doc.add_heading('10. How to reproduce and test in the org',1)
doc.add_heading('Where the Context ID comes from',2)
doc.add_paragraph('The Context ID is the Salesforce Quote record ID beginning with a2N. It is taken from the quote record URL and pasted into the OmniScript Preview Context ID field.')
table(doc,['Quote','Record URL suffix / Context ID'],[('Q-37723','a2NhG000002ktrcUAA'),('Q-37780','a2NhG000003EetVUAS'),('Q-37640','a2NhG000002YPkPUAW')],[1.5,5.0])
doc.add_heading('Preview procedure',2)
for x in ['Open OmniScript Version 2 - Generate Quote Document - SteppedUpPricing.','Click Preview.','Paste the applicable Quote ID into Context ID.','Click the circular refresh button.','Wait for GetQuoteProposalData and GetDocumentTemplate to complete.','Generate or preview the proposal PDF.','Close and reopen Preview between tests if the browser retains a cached template.']: numbered(doc,x)

doc.add_heading('11. Visual acceptance checklist',1)
doc.add_heading('Q-37723 - grouped, Multiple Ship-To = Yes',2)
for x in ['Five group sections appear separately.','Group name and service-period dates appear.','Products are under the correct group with no duplicates.','Ship-To Account column is visible.','Product, quantity, billing frequency, description, and price are populated.']: bullet(doc,x)
doc.add_heading('Q-37780 - grouped, Multiple Ship-To = No',2)
for x in ['Five group sections appear separately.','Group name and service-period dates appear.','Products are under the correct group with no duplicates.','Ship-To Account column is not present.','All other required line values are populated.']: bullet(doc,x)
doc.add_heading('Q-37640 - ungrouped regression',2)
for x in ['No group heading or service-period section appears.','The normal product table appears.','The previous ungrouped layout is preserved.','No empty Group 1 appears.']: bullet(doc,x)

doc.add_heading('12. Troubleshooting and rollback',1)
table(doc,['Situation','Action'],[
('Old template appears','Exit Preview, close the OmniScript tab, reopen it, and refresh the Context ID.'),('Generic key-value error returns','Inspect the Transform output first. Root must be a Map, not a List.'),('Groups missing','Confirm Group is a list and QuoteLine is nested within each Group.'),('Wrong Ship-To layout','Check MultipleShipToAccounts and the two conditional flags.'),('Rollback required','Relink DocumentTemplateContentDoc 2ddhG0000001bwnQAA to a previous ContentDocument and restore the backed-up mapper item values.'),('Unused Apex classes','CPQQuoteProposalOutputSerializer and its test are not connected. They can be removed in a separate cleanup change after stakeholder approval.')],[1.65,4.85])

doc.add_heading('13. Final status',1)
callout(doc,'Ready for stakeholder handoff','The Partial sandbox implementation is complete. All three Extract/Transform scenarios passed with the expected root structure, group counts, line counts, and conditional flags. The active template is v5. Final business acceptance consists of confirming the generated PDFs visually using the checklist above.',GREEN)

OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT); print(OUT.resolve())
