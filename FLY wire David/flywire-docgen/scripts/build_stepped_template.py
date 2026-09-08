from copy import deepcopy
from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def insert_paragraph_before(anchor, text, *, bold=False, size=None, space_before=0, space_after=0):
    paragraph = anchor.insert_paragraph_before()
    run = paragraph.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(space_after)
    return paragraph


def copy_table_before(anchor, source_table):
    copied = deepcopy(source_table._tbl)
    anchor._p.addprevious(copied)
    return copied


def set_paragraph_text(paragraph, value):
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def build(source: Path, target: Path):
    doc = Document(source)
    group_open = next(p for p in doc.paragraphs if "{{#Group}}" in p.text)
    group_close = next(p for p in doc.paragraphs if "{{/Group}}" in p.text)
    multi_table = doc.tables[0]
    single_table = doc.tables[1]

    # Replace the original grouped section completely. Keeping its existing
    # nodes would nest a second {{#Group}} inside the new group loop, which the
    # client-side document generator rejects as a non key-value input object.
    node = group_open._p.getnext()
    while node is not None and node is not group_close._p:
        next_node = node.getnext()
        node.getparent().remove(node)
        node = next_node

    generated_date = next(p for p in doc.paragraphs if "Generated Date:" in p.text)
    set_paragraph_text(
        generated_date,
        "Generated Date: {{Date}}"
        "{{#IF_QuoteLineGroupsPresentFalse}}"
        "{{#IF_MultipleShipToAccounts}}",
    )

    set_paragraph_text(
        group_open,
        "{{/IF_NoMultipleShipToAccounts}}"
        "{{/IF_QuoteLineGroupsPresentFalse}}"
        "{{#IF_QuoteLineGroupsPresentTrue}}"
        "{{#IF_MultipleShipToAccounts}}"
        "{{#Group}}",
    )

    insert_paragraph_before(group_close, "{{Name}}", bold=True, size=12, space_before=8, space_after=2)
    insert_paragraph_before(
        group_close,
        "Service Period: {{StartDate}} to {{EndDate}}",
        size=10,
        space_after=4,
    )
    copy_table_before(group_close, multi_table)

    insert_paragraph_before(
        group_close,
        "{{/Group}}{{/IF_MultipleShipToAccounts}}"
        "{{#IF_NoMultipleShipToAccounts}}{{#Group}}",
    )
    insert_paragraph_before(group_close, "{{Name}}", bold=True, size=12, space_before=8, space_after=2)
    insert_paragraph_before(
        group_close,
        "Service Period: {{StartDate}} to {{EndDate}}",
        size=10,
        space_after=4,
    )
    copy_table_before(group_close, single_table)

    set_paragraph_text(
        group_close,
        "{{/Group}}{{/IF_NoMultipleShipToAccounts}}"
        "{{/IF_QuoteLineGroupsPresentTrue}}"
        "{{#IF_PrintSignatureBlock}}",
    )

    # The whole QuoteLineGroup list is mapped directly to Group, so grouped
    # tables iterate the source child list name and field names.
    for table in doc.tables[2:4]:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    value = (
                        paragraph.text
                        .replace("{{#Line}}", "{{#QuoteLine}}")
                        .replace("{{/Line}}", "{{/QuoteLine}}")
                        .replace("{{ShipToAccount}}", "{{ShipToAccountName}}")
                    )
                    if value != paragraph.text:
                        set_paragraph_text(paragraph, value)

    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)


if __name__ == "__main__":
    build(Path(sys.argv[1]), Path(sys.argv[2]))
