from __future__ import annotations

import shutil
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Pt

from build_architect_review import (
    AMBER,
    BLUE,
    LIGHT_AMBER,
    LIGHT_BLUE,
    LIGHT_RED,
    LIGHT_TEAL,
    MUTED,
    NAVY,
    RED,
    TEAL,
    add_body,
    add_figure,
    add_heading,
    arrow,
    draw_wrapped,
    load_font,
    rounded,
    set_run_font,
)
from build_scc3387_architect_review import (
    configure_document,
    normalize_table_margins,
    review_callout,
    review_table,
)


ROOT = Path(__file__).resolve().parent
REFERENCE = ROOT / "SCC-4179-architect-review.docx"
ASSET_DIR = ROOT / "doc-assets-scc4485"
OUTPUT = ROOT / "SCC-4485-PR-Review-DevA.docx"
ASSET_DIR.mkdir(exist_ok=True)


def clear_document_body(document: Document) -> None:
    body = document._body._element
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)


def set_deva_footer(document: Document) -> None:
    for section in document.sections:
        for footer in (section.footer, section.even_page_footer, section.first_page_footer):
            for paragraph in footer.paragraphs:
                for run in paragraph.runs:
                    if "GreatPlainsMerge" in run.text:
                        run.text = run.text.replace("GreatPlainsMerge", "GreatPlainsDevA")


def create_review_path_diagram(path: Path) -> None:
    image = Image.new("RGB", (1600, 1000), "#F7F9FC")
    draw = ImageDraw.Draw(image)
    title = load_font(42, bold=True)
    subtitle = load_font(22)
    heading = load_font(24, bold=True)
    body = load_font(18)
    small = load_font(16, bold=True)

    draw.text((70, 50), "SCC-4485 | DevA implementation and remaining PR gate", font=title, fill="#0B2545")
    draw.text(
        (70, 110),
        "The field logic aligns with the story; Data Mapper activation and the actual PR diff remain unresolved.",
        font=subtitle,
        fill="#5B6777",
    )

    boxes = [
        (
            (55, 225, 350, 605),
            "Buy-flow input",
            ["Place.buildingType", "Place.Type", "SQResponse or sqJson path"],
            "#EAF2FA",
            "#2E74B5",
        ),
        (
            (430, 225, 755, 605),
            "MDU decision",
            ["MDU / MDU Tenant", "Multi-Family / Multifamily", "Type = Unit"],
            "#E7F5F1",
            "#16836B",
        ),
        (
            (835, 225, 1160, 605),
            "Account values",
            ["LOB: Residential", "Customer Type: Consumer", "Segment: MDU Tenant or Residential"],
            "#E7F5F1",
            "#16836B",
        ),
        (
            (1240, 225, 1545, 605),
            "Remaining gate",
            ["AccountCreation inactive", "No PR diff supplied", "No new runtime execution"],
            "#FFF6DE",
            "#B7791F",
        ),
    ]

    for box, box_title, lines, fill, outline in boxes:
        rounded(draw, box, fill, outline=outline, radius=22, width=4)
        draw.text((box[0] + 24, box[1] + 24), box_title, font=heading, fill="#0B2545")
        y = box[1] + 98
        for line in lines:
            draw.ellipse((box[0] + 26, y + 7, box[0] + 38, y + 19), fill=outline)
            wrapped = draw_wrapped(draw, (box[0] + 52, y), line, body, "#1F2937", box[2] - box[0] - 78, spacing=6)
            y += max(76, 27 * len(wrapped) + 34)

    for start, end, label, color in (
        ((350, 415), (430, 415), "CLASSIFY", "#2E74B5"),
        ((755, 415), (835, 415), "MAP", "#16836B"),
        ((1160, 415), (1240, 415), "VERIFY", "#B7791F"),
    ):
        arrow(draw, start, end, color=color, width=6)
        draw.text((start[0] + 8, start[1] - 42), label, font=small, fill=color)

    rounded(draw, (70, 685, 1530, 905), "#EAF2FA", outline="#2E74B5", radius=22, width=3)
    draw.text((105, 720), "PR reviewer reading", font=heading, fill="#0B2545")
    decision = (
        "DevA now contains the required MDU classification and all three Account-field mappings. Do not merge "
        "solely from this org snapshot: activate or prove runtime resolution of AccountCreation, attach the exact "
        "PR/DataPack diff, and provide one post-change execution result."
    )
    draw_wrapped(draw, (105, 780), decision, body, "#1F2937", 1360, spacing=8)
    image.save(path)


def add_metadata_table(document: Document) -> None:
    metadata = [
        ("Environment", "GreatPlainsDevA | 00DEa00000Fc086MAB | Sandbox"),
        ("Review date", "28 August 2026"),
        ("Evidence", "Live OmniStudio configuration | Account schema | Aggregate Account read-back"),
        ("Decision", "Changes Requested - logic aligns; runtime/version gate remains"),
        ("Boundary", "PR review only; no deployment, activation, account creation, or org change"),
    ]
    table = document.add_table(rows=len(metadata), cols=2)
    table.style = "Table Grid"
    from build_architect_review import LIGHT_BLUE, set_cell_text, set_table_geometry, shade_cell

    set_table_geometry(table, [1900, 7460])
    normalize_table_margins(table)
    for row, (label, value) in zip(table.rows, metadata):
        shade_cell(row.cells[0], LIGHT_BLUE)
        set_cell_text(row.cells[0], label, size=9.5, color=NAVY, bold=True)
        set_cell_text(row.cells[1], value, size=9.5)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def build_document() -> None:
    diagram_path = ASSET_DIR / "scc4485-deva-review-path.png"
    create_review_path_diagram(diagram_path)

    shutil.copy2(REFERENCE, OUTPUT)
    document = Document(OUTPUT)
    clear_document_body(document)
    configure_document(document)
    set_deva_footer(document)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(2)
    set_run_font(paragraph.add_run("PR REVIEW"), size=10, color=BLUE, bold=True)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(paragraph.add_run("SCC-4485 | Consumer Account GL Classification"), size=25, color=NAVY, bold=True)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(16)
    set_run_font(paragraph.add_run("GreatPlainsDevA | OmniStudio implementation and merge-readiness review"), size=13, color=MUTED)

    add_metadata_table(document)
    review_callout(
        document,
        "PR decision",
        "Changes Requested. The acceptance-criteria logic is present in DevA, but AccountCreation v1 is inactive while active pipeline v5 invokes it. Resolve or prove that runtime/version state, attach the actual PR or DataPack diff, and provide one post-change execution result before merge approval.",
        fill=LIGHT_AMBER,
        accent=AMBER,
    )

    add_heading(document, "Finding summary", 1)
    summary_rows = [
        ("P0", "AccountCreation v1", "Blocker", "The only Data Mapper record is inactive; runtime resolution was not executed during this read-only review."),
        ("P1", "SQ/OrderPipeline v5", "Verified", "Active Consumer action supplies Residential, Consumer, and Residential fallback inputs."),
        ("P1", "MDU classification", "Verified", "Formula detects documented MDU building types or Type = Unit and returns MDU Tenant."),
        ("P1", "Account mappings", "Verified", "All three target Account fields have enabled mappings and valid picklist values."),
        ("P1", "PR evidence", "Blocked", "No PR URL, branch, commit, or exported DataPack diff was supplied for line-by-line review."),
    ]
    review_table(document, ["Severity", "Location", "Result", "Reviewer finding"], summary_rows, [1050, 2050, 1400, 4860], status_column=2)

    document.add_page_break()
    add_heading(document, "Implemented path and remaining gate", 1)
    add_body(
        document,
        "DevA now classifies the documented Place variants inside AccountCreation, maps the derived values to the Consumer Account, and preserves Residential fallback behavior. The outstanding gate is evidence that the inactive Data Mapper is the version executed by active SQ/OrderPipeline v5.",
    )
    add_figure(
        document,
        diagram_path,
        "Figure 1. SCC-4485 DevA implementation and remaining PR gate",
        "Diagram showing buy-flow Place inputs, MDU classification, the three Account values, and the remaining Data Mapper activation, PR-diff, and runtime-evidence gates.",
        width=6.45,
    )
    review_callout(
        document,
        "Reviewer interpretation",
        "The implementation correctly centralizes Segment classification in AccountCreation and keeps the Consumer action explicit. Salesforce guidance requires an active Data Mapper for runtime use when versioning is enforced; DevA exposes AccountCreation as inactive, so the PR must remove that ambiguity.",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )

    document.add_page_break()
    add_heading(document, "Detailed PR findings", 1)
    finding_rows = [
        (
            "PR-01",
            "P0",
            "AccountCreation v1 state",
            "The only AccountCreation Load record is IsActive = false, last modified 28 August 2026 by API Service. Active SQ/OrderPipeline v5 references it by bundle name.",
            "Activate the intended Data Mapper version, or attach runtime/versioning evidence proving this exact inactive record resolves and executes safely after deployment.",
        ),
        (
            "PR-02",
            "P1",
            "AccountCreation > SegmentValue",
            "The enabled formula recognizes MDU, MDU Tenant, Multi-Family, Multifamily, or Type = Unit across SQResponse and sqJson paths; otherwise it uses Segment or Residential.",
            "No acceptance-criteria change required. Preserve this formula and its SegmentValue -> Segment__c mapping in the PR diff.",
        ),
        (
            "PR-03",
            "P1",
            "AccountCreation > Customer Type",
            "An enabled formula returns Consumer for the Consumer record type and CustomerTypeValue maps to Customer_Type__c.",
            "No acceptance-criteria change required. Keep the Consumer-only guard and mapping together in the exported component.",
        ),
        (
            "PR-04",
            "P1",
            "SQ/OrderPipeline v5",
            "CreateConsumerAccount is active and passes LineofBusiness = Residential, CustomerType = Consumer, and Segment = Residential under the Consumer record-type branch.",
            "Preserve the Consumer-only condition and prove v5 is the version included and activated by the deployment mechanism.",
        ),
        (
            "PR-05",
            "P2",
            "Input contract maintainability",
            "CustomerType = Consumer is passed by the pipeline, but AccountCreation derives CustomerTypeValue from record type rather than consuming the CustomerType input.",
            "Document the intentional record-type derivation or remove the unused input in a controlled follow-up; do not change Consumer behavior in this story.",
        ),
        (
            "PR-06",
            "P1",
            "PR artifact",
            "No repository PR URL, source branch, commit, or exported DataPack diff was available.",
            "Attach the exact SQ_OrderPipeline/SQtoOrderPipeline and AccountCreation diff for formal approval.",
        ),
    ]
    review_table(document, ["ID", "Severity", "Component / location", "Evidence", "Reviewer action"], finding_rows, [650, 1050, 2100, 2880, 2680])
    review_callout(
        document,
        "Review conclusion",
        "The earlier Merge-based findings about missing Customer Type and missing MDU logic do not apply to DevA. PR-01 is the remaining configuration blocker; PR-06 prevents line-by-line PR approval.",
        fill=LIGHT_AMBER,
        accent=AMBER,
    )

    document.add_page_break()
    add_heading(document, "Acceptance and merge checklist", 1)
    acceptance_rows = [
        ("Line of Business", "Residential", "Verified", "Pipeline input maps to Line_of_Business__c."),
        ("Customer Type", "Consumer", "Verified", "Consumer formula maps CustomerTypeValue to Customer_Type__c."),
        ("MDU Segment", "MDU Tenant", "Verified", "Documented building types and Type = Unit are covered."),
        ("Standard Segment", "Residential", "Verified", "Non-MDU path falls back to input Segment or Residential."),
        ("Existing Accounts", "Unchanged", "Verified", "Account Id mapping is disabled; the Data Mapper remains create-only."),
        ("Service Account", "Unchanged", "Verified", "Service action does not pass the three Consumer classification inputs."),
    ]
    review_table(document, ["Criterion", "Expected", "Result", "Evidence"], acceptance_rows, [1800, 1800, 1500, 4260], status_column=2)

    add_heading(document, "Required before merge", 2)
    gate_rows = [
        ("1", "Resolve AccountCreation inactive status or prove runtime/version resolution", "Blocked"),
        ("2", "Attach exact PR or exported DataPack diff", "Blocked"),
        ("3", "Show one post-change MDU buy-flow execution result", "Pending"),
        ("4", "Show one post-change standard Residential execution result", "Pending"),
        ("5", "Confirm downstream package activates SQ/OrderPipeline v5 and intended Data Mapper", "Pending"),
    ]
    review_table(document, ["#", "Merge check", "Status"], gate_rows, [600, 6900, 1860], status_column=2)
    review_callout(
        document,
        "Scope note",
        "These are PR evidence and configuration gates. No deployment steps or execution instructions are included in this review.",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )

    document.add_page_break()
    add_heading(document, "Review evidence and limitations", 1)
    evidence_rows = [
        ("Target org", "GreatPlainsDevA | Great Plains Communications | 00DEa00000Fc086MAB | Sandbox"),
        ("Active pipeline", "SQ / OrderPipeline | SQtoOrderPipeline v5 | active | 7 active elements"),
        ("Data Mapper", "AccountCreation | Load v1 | inactive | 49 items | modified 28 August 2026"),
        ("Account fields", "Line_of_Business__c, Customer_Type__c, Segment__c; required picklist values exist"),
        ("30-day outcomes", "17 Consumer Accounts: 2 expected MDU, 4 expected Residential, 11 legacy blank Customer Type"),
        ("Today", "No Consumer Account was created today; no post-change runtime result exists in the reviewed data"),
        ("Version scan", "SQtoOrderPipeline v1-v5 reviewed; only v5 is active; v4-v5 carry SCC-4485 inputs"),
    ]
    review_table(document, ["Evidence source", "Confirmed result"], evidence_rows, [2200, 7160])

    review_callout(
        document,
        "Review limitation",
        "The two MDU outcome records predate today's component modification timestamps and cannot prove the currently deployed version executed. No source PR diff or new buy-flow transaction was inspected or run.",
        fill=LIGHT_AMBER,
        accent=AMBER,
    )

    add_heading(document, "Ready-to-paste PR review comment", 2)
    response = (
        "Changes Requested for SCC-4485. DevA now contains the required logic: active SQ/OrderPipeline v5 passes "
        "Residential / Consumer / Residential defaults for the Consumer branch, AccountCreation classifies the "
        "documented MDU address variants as MDU Tenant, and all three Account fields are mapped. The remaining "
        "blocker is version/runtime state: the only AccountCreation v1 record is inactive while the active pipeline "
        "invokes it by bundle name. Please activate the intended mapper or provide evidence that this exact version "
        "resolves and executes under the org's runtime configuration. Also attach the exact PR/DataPack diff and one "
        "post-change MDU plus standard Residential execution result. The prior findings that Customer Type and MDU "
        "logic were missing do not apply to DevA. No deployment, activation, account creation, or org write was "
        "performed during this review."
    )
    review_callout(document, "PR comment", response, fill=LIGHT_BLUE, accent=BLUE)

    document.core_properties.title = "SCC-4485 Consumer Account GL Classification PR Review - DevA"
    document.core_properties.subject = "DevA OmniStudio implementation and merge-readiness review"
    document.core_properties.author = "Solution Architecture Review"
    document.core_properties.keywords = "Salesforce, Great Plains, SCC-4485, PR review, OmniStudio, AccountCreation, OrderPipeline, MDU Tenant"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
