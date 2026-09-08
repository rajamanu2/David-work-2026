from copy import deepcopy
from pathlib import Path

from docx import Document


SOURCE = Path("templates/CPQ Quote Proposal - SteppedUpPricing-v5.docx")
OUTPUT = Path("templates/SALDEV1420 Approval Quote Document - v1.docx")


def replace_paragraph_text_preserving_first_run(paragraph, text):
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


doc = Document(SOURCE)

# Both grouped layouts use the same dynamic Term header. The SubscriptionTerm
# section suppresses the colon/month suffix when the source value is blank.
term_header = "Term {{Number}}{{#SubscriptionTerm}}: {{SubscriptionTerm}} Months{{/SubscriptionTerm}}"
replace_paragraph_text_preserving_first_run(doc.paragraphs[7], term_header)
replace_paragraph_text_preserving_first_run(doc.paragraphs[10], term_header)

# SALDEV-1420 is an approval document variant; retain all established quote
# proposal formatting and conditional table behavior from the validated v5.
doc.core_properties.title = "SALDEV-1420 Approval Quote Document"
doc.core_properties.subject = "OmniStudio-only grouped approval document"
doc.core_properties.comments = (
    "Inactive Partial sandbox draft. Dynamic Term headers; conditional Ship-To; "
    "grouped and ungrouped layouts."
)

doc.save(OUTPUT)
print(OUTPUT.resolve())
