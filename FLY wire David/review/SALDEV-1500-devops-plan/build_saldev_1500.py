from pathlib import Path
from textwrap import wrap
from datetime import date

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT.parent.parent / "output" / "documents" / "SALDEV-1500_GTM_DevOps_and_Jira_Operating_Model.docx"
ARCH_IMG = ROOT / "target_operating_model.png"
JIRA_IMG = ROOT / "jira_workflow.png"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
GRAY = "5B6777"
LIGHT_GRAY = "F2F4F7"
BLUE_GRAY = "E8EEF5"
PALE_BLUE = "EAF3F8"
GREEN = "16836B"
PALE_GREEN = "E8F5F1"
AMBER = "B7791F"
PALE_AMBER = "FFF4D6"
RED = "C53030"
PALE_RED = "FDECEC"
WHITE = "FFFFFF"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def set_run_font(run, name="Calibri", size=11, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="CBD5E1", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)
        tag.set(qn("w:space"), "0")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def set_table_geometry(table, widths_dxa, indent_dxa=TABLE_INDENT_DXA):
    assert sum(widths_dxa) == CONTENT_WIDTH_DXA, (widths_dxa, sum(widths_dxa))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        set_row_cant_split(row)
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_cell_text(cell, size=8.7, color=INK, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    for p in cell.paragraphs:
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.08
        for run in p.runs:
            set_run_font(run, size=size, color=color, bold=bold)


def add_table(doc, headers, rows, widths_dxa, header_fill=BLUE_GRAY, font_size=8.7,
              center_cols=None, status_col=None):
    center_cols = center_cols or set()
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_header(hdr)
    for i, text in enumerate(headers):
        hdr.cells[i].text = str(text)
        set_cell_shading(hdr.cells[i], header_fill)
        style_cell_text(hdr.cells[i], size=8.6, color=NAVY, bold=True,
                        align=WD_ALIGN_PARAGRAPH.CENTER if i in center_cols else WD_ALIGN_PARAGRAPH.LEFT)
    for row_values in rows:
        row = table.add_row()
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            cell.text = str(value)
            align = WD_ALIGN_PARAGRAPH.CENTER if i in center_cols else WD_ALIGN_PARAGRAPH.LEFT
            color = INK
            bold = False
            if status_col is not None and i == status_col:
                value_low = str(value).lower()
                if any(k in value_low for k in ("required", "blocked", "fail", "stop")):
                    color = RED
                elif any(k in value_low for k in ("approved", "pass", "complete", "ready")):
                    color = GREEN
                elif any(k in value_low for k in ("conditional", "pending", "tbd", "gate")):
                    color = AMBER
                bold = True
            style_cell_text(cell, size=font_size, color=color, bold=bold, align=align)
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def add_hyperlink(paragraph, text, url, color=BLUE):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(r_fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "20")
    r_pr.append(sz)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_keep_with_next(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if keep is None:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)
    keep.set(qn("w:val"), "1" if value else "0")


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)
    set_run_font(run, size=9, color=GRAY)


def add_numbering(doc):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    next_abs = max(existing_abs or [0]) + 1
    next_num = max(existing_num or [0]) + 1

    def make_abstract(abs_id, fmt, text0, font=None):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abs_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "multilevel")
        abstract.append(multi)
        for level in range(3):
            lvl = OxmlElement("w:lvl")
            lvl.set(qn("w:ilvl"), str(level))
            start = OxmlElement("w:start")
            start.set(qn("w:val"), "1")
            lvl.append(start)
            num_fmt = OxmlElement("w:numFmt")
            num_fmt.set(qn("w:val"), fmt)
            lvl.append(num_fmt)
            lvl_text = OxmlElement("w:lvlText")
            if fmt == "bullet":
                lvl_text.set(qn("w:val"), [text0, "o", "▪"][level])
            else:
                lvl_text.set(qn("w:val"), f"%{level + 1}.")
            lvl.append(lvl_text)
            suff = OxmlElement("w:suff")
            suff.set(qn("w:val"), "tab")
            lvl.append(suff)
            p_pr = OxmlElement("w:pPr")
            tabs = OxmlElement("w:tabs")
            tab = OxmlElement("w:tab")
            tab.set(qn("w:val"), "num")
            tab.set(qn("w:pos"), str(720 + level * 360))
            tabs.append(tab)
            p_pr.append(tabs)
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(720 + level * 360))
            ind.set(qn("w:hanging"), "360")
            p_pr.append(ind)
            spacing = OxmlElement("w:spacing")
            spacing.set(qn("w:after"), "160")
            spacing.set(qn("w:line"), "280")
            spacing.set(qn("w:lineRule"), "auto")
            p_pr.append(spacing)
            lvl.append(p_pr)
            if font:
                r_pr = OxmlElement("w:rPr")
                fonts = OxmlElement("w:rFonts")
                fonts.set(qn("w:ascii"), font)
                fonts.set(qn("w:hAnsi"), font)
                r_pr.append(fonts)
                lvl.append(r_pr)
            abstract.append(lvl)
        return abstract

    bullet_abstract = make_abstract(next_abs, "bullet", "•")
    decimal_abstract = make_abstract(next_abs + 1, "decimal", "%1.")
    numbering.append(bullet_abstract)
    numbering.append(decimal_abstract)

    def add_num(num_id, abstract_id):
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)
        return num_id

    bullet_id = add_num(next_num, next_abs)
    decimal_id = add_num(next_num + 1, next_abs + 1)
    return bullet_id, decimal_id, next_abs, next_abs + 1


def new_numbering_instance(doc, abstract_id):
    numbering = doc.part.numbering_part.element
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    num_id = max(existing_num or [0]) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    lvl_override = OxmlElement("w:lvlOverride")
    lvl_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    lvl_override.append(start_override)
    num.append(lvl_override)
    numbering.append(num)
    return num_id


def add_list_item(doc, text, num_id, level=0, bold_prefix=None):
    p = doc.add_paragraph(style="Normal")
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    n = OxmlElement("w:numId")
    n.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(n)
    p_pr.append(num_pr)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_callout(doc, label, text, tone="blue"):
    colors = {
        "blue": (PALE_BLUE, BLUE),
        "green": (PALE_GREEN, GREEN),
        "amber": (PALE_AMBER, AMBER),
        "red": (PALE_RED, RED),
    }
    fill, accent = colors[tone]
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.1
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    for edge in ("left", "top", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8" if edge == "left" else "4")
        border.set(qn("w:space"), "6")
        border.set(qn("w:color"), accent)
        p_bdr.append(border)
    p_pr.append(p_bdr)
    r = p.add_run(f"{label}: ")
    set_run_font(r, size=10.2, color=accent, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.2, color=NAVY, bold=True)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_keep_with_next(p)
    return p


def add_body(doc, text, bold_prefix=None, italic=False, align=None):
    p = doc.add_paragraph(style="Normal")
    if align is not None:
        p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, italic=italic)
    return p


def page_break(doc):
    # Major sections flow continuously to avoid sparse continuation pages.
    # Clean chapter breaks are inserted explicitly at selected landmarks.
    return None


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = rgb(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.keep_together = True

    if "Source Note" not in [s.name for s in doc.styles]:
        st = doc.styles.add_style("Source Note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        st = doc.styles["Source Note"]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    st.font.size = Pt(9)
    st.font.italic = True
    st.font.color.rgb = rgb(GRAY)
    st.paragraph_format.space_before = Pt(4)
    st.paragraph_format.space_after = Pt(4)


def set_headers_footers(doc):
    for section in doc.sections:
        header = section.header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("GTM SYSTEMS DEVOPS OPERATING MODEL")
        set_run_font(r, size=8.5, color=GRAY, bold=True)
        p_pr = p._p.get_or_add_pPr()
        p_bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "D7DBE2")
        p_bdr.append(bottom)
        p_pr.append(p_bdr)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = fp.add_run("SALDEV-1500 | Proposed operating standard | Page ")
        set_run_font(r, size=8.5, color=GRAY)
        add_field(fp, "PAGE")
        r = fp.add_run(" of ")
        set_run_font(r, size=8.5, color=GRAY)
        add_field(fp, "NUMPAGES")


def find_font(size, bold=False):
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for p in candidates:
        if p.exists():
            return ImageFont.truetype(str(p), size)
    return ImageFont.load_default()


def pil_color(value):
    if isinstance(value, str) and len(value) in (6, 8) and not value.startswith("#"):
        return f"#{value}"
    return value


def centered_text(draw, box, text, font, fill, max_chars=22, line_gap=6):
    lines = []
    for paragraph in text.split("\n"):
        lines.extend(wrap(paragraph, width=max_chars) or [""])
    heights = []
    widths = []
    for line in lines:
        b = draw.textbbox((0, 0), line, font=font)
        widths.append(b[2] - b[0])
        heights.append(b[3] - b[1])
    total_h = sum(heights) + line_gap * max(0, len(lines) - 1)
    y = box[1] + (box[3] - box[1] - total_h) / 2
    for line, w, h in zip(lines, widths, heights):
        x = box[0] + (box[2] - box[0] - w) / 2
        draw.text((x, y), line, font=font, fill=pil_color(fill))
        y += h + line_gap


def rounded_box(draw, box, fill, outline, radius=24, width=4):
    draw.rounded_rectangle(box, radius=radius, fill=pil_color(fill), outline=pil_color(outline), width=width)


def arrow(draw, p1, p2, color=BLUE, width=6):
    color = pil_color(color)
    draw.line((p1, p2), fill=color, width=width)
    x2, y2 = p2
    draw.polygon([(x2, y2), (x2 - 18, y2 - 11), (x2 - 18, y2 + 11)], fill=color)


def create_architecture_diagram(path):
    img = Image.new("RGB", (1900, 1050), pil_color(WHITE))
    d = ImageDraw.Draw(img)
    title = find_font(50, True)
    body = find_font(31, True)
    small = find_font(25, False)
    d.text((80, 50), "One governed path from demand to production", font=title, fill=pil_color(NAVY))
    xs = [80, 420, 760, 1100, 1440]
    labels = [
        ("Jira", "work, scope, approvals"),
        ("GitHub", "branch, PR, review"),
        ("Quality gates", "lint, tests, validate"),
        ("Release platform", "Gearset preferred\nCopado alternative"),
        ("Salesforce", "INT → QA → UAT → PROD"),
    ]
    fills = [PALE_AMBER, PALE_BLUE, LIGHT_GRAY, PALE_GREEN, PALE_BLUE]
    for i, (x, (name, detail)) in enumerate(zip(xs, labels)):
        box = (x, 260, x + 270, 520)
        rounded_box(d, box, fills[i], BLUE if i != 0 else AMBER)
        centered_text(d, (x + 15, 280, x + 255, 365), name, body, NAVY, 16)
        centered_text(d, (x + 18, 365, x + 252, 505), detail, small, INK, 22)
        if i < len(xs) - 1:
            arrow(d, (x + 275, 390), (xs[i + 1] - 15, 390))
    rounded_box(d, (150, 690, 1750, 915), LIGHT_GRAY, "CBD5E1", radius=20, width=3)
    d.text((210, 725), "Control plane", font=body, fill=pil_color(NAVY))
    controls = "CODEOWNERS • required reviews • protected branches • service accounts • environment approvals • evidence retention • drift monitoring"
    centered_text(d, (235, 790, 1680, 885), controls, small, INK, 88)
    img.save(path, dpi=(180, 180))


def create_jira_diagram(path):
    img = Image.new("RGB", (1900, 1050), pil_color(WHITE))
    d = ImageDraw.Draw(img)
    title = find_font(48, True)
    body = find_font(25, True)
    small = find_font(21, False)
    d.text((80, 45), "Jira workflow and release evidence", font=title, fill=pil_color(NAVY))
    statuses = [
        ("BACKLOG", "triaged"), ("READY", "DoR met"), ("IN DEV", "branch active"),
        ("CODE REVIEW", "PR open"), ("QA", "deployed + tests"), ("UAT", "business sign-off"),
        ("READY RELEASE", "release approved"), ("DEPLOYED", "prod job success"), ("DONE", "smoke + evidence"),
    ]
    positions = []
    for i in range(5):
        positions.append((80 + i * 350, 220))
    for i in range(4):
        positions.append((1480 - i * 350, 610))
    for i, ((name, detail), (x, y)) in enumerate(zip(statuses, positions)):
        box = (x, y, x + 270, y + 160)
        fill = PALE_GREEN if name in ("DEPLOYED", "DONE") else PALE_BLUE if name not in ("UAT", "READY RELEASE") else PALE_AMBER
        rounded_box(d, box, fill, BLUE if fill != PALE_AMBER else AMBER, radius=18, width=3)
        centered_text(d, (x + 10, y + 18, x + 260, y + 75), name, body, NAVY, 18)
        centered_text(d, (x + 10, y + 78, x + 260, y + 145), detail, small, INK, 24)
        if i < 4:
            arrow(d, (x + 275, y + 80), (positions[i + 1][0] - 15, y + 80))
        elif i == 4:
            d.line((x + 135, y + 165, x + 135, 500, 1615, 500, 1615, 595), fill=pil_color(BLUE), width=6)
            d.polygon([(1615, 610), (1604, 590), (1626, 590)], fill=pil_color(BLUE))
        elif i < len(statuses) - 1:
            arrow(d, (x - 5, y + 80), (positions[i + 1][0] + 285, y + 80))
    d.text((80, 930), "Rejected validation or approval returns the ticket to IN DEV with the failure evidence attached.", font=small, fill=pil_color(RED))
    img.save(path, dpi=(180, 180))


def set_image_alt(paragraph, description):
    for doc_pr in paragraph._p.xpath(".//wp:docPr"):
        doc_pr.set("descr", description)


def add_picture_with_caption(doc, path, caption, alt):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(6.45))
    set_image_alt(p, alt)
    cap = doc.add_paragraph(style="Source Note")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    set_run_font(r, size=9, color=GRAY, italic=True)


def add_source_note(doc, text):
    p = doc.add_paragraph(style="Source Note")
    r = p.add_run(text)
    set_run_font(r, size=9, color=GRAY, italic=True)
    return p


def add_story(doc, story_id, summary, story, criteria, owner, deps, estimate, wave):
    add_heading(doc, f"{story_id} — {summary}", 2)
    add_table(
        doc,
        ["Field", "Value"],
        [
            ("Jira issue type", "Story"),
            ("Suggested owner", owner),
            ("Dependency", deps),
            ("Indicative estimate", estimate),
            ("Delivery wave", wave),
        ],
        [1800, 7560],
        header_fill=LIGHT_GRAY,
        font_size=9,
    )
    add_body(doc, "User story", bold_prefix="User story")
    add_body(doc, story)
    add_body(doc, "Acceptance criteria", bold_prefix="Acceptance criteria")
    for criterion in criteria:
        add_list_item(doc, criterion, BULLET_ID)


doc = Document()
configure_styles(doc)
BULLET_ID, DECIMAL_ID, BULLET_ABSTRACT_ID, DECIMAL_ABSTRACT_ID = add_numbering(doc)
set_headers_footers(doc)
create_architecture_diagram(ARCH_IMG)
create_jira_diagram(JIRA_IMG)

# Cover / masthead
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(14)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("GTM SYSTEMS • PLATFORM & DEVOPS")
set_run_font(r, size=10, color=BLUE, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
r = p.add_run("DevOps, Code Repository and Jira Operating Model")
set_run_font(r, size=25, color=NAVY, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(16)
r = p.add_run("SALDEV-1500 — source control, CI/CD, deployment governance")
set_run_font(r, size=13, color=GRAY)
r.add_break()
r = p.add_run("and implementation backlog")
set_run_font(r, size=13, color=GRAY)

add_table(
    doc,
    ["Document control", "Value"],
    [
        ("Status", "Proposed operating standard — approval and implementation required"),
        ("Prepared for", "Flywire GTM Systems, Platform & DevOps"),
        ("Prepared on", "28 August 2026"),
        ("Decision requested", "Approve target toolchain, repository standard, Jira workflow and implementation backlog"),
        ("Scope", "Salesforce and adjacent GTM systems; configuration, code, APIs, integrations, tests and release evidence"),
        ("Ticket", "SALDEV-1500"),
    ],
    [2100, 7260],
    header_fill=BLUE_GRAY,
    font_size=9.2,
)

add_callout(
    doc,
    "Executive decision",
    "Make GitHub Enterprise Cloud the canonical source repository, keep Jira as the canonical work and approval record, and use a Salesforce-aware release platform on top. Prefer Gearset for a new implementation; use Copado when an existing enterprise investment makes it the lower-risk choice.",
    "blue",
)

add_heading(doc, "Approval requested", 1)
for text in [
    "Approve the target operating model and the principle that every production change starts with a Jira ticket and ends with evidence linked back to that ticket.",
    "Authorize procurement or confirmation of GitHub Enterprise Cloud and the selected Salesforce DevOps platform for the required users and service accounts.",
    "Create the implementation epic and stories in Section 17, then pilot the process with one low-risk Salesforce change before scaling.",
]:
    add_list_item(doc, text, BULLET_ID)

doc.add_page_break()

add_heading(doc, "Contents", 1)
contents = [
    "1. Executive summary and ticket outcomes",
    "2. Scope, exclusions and operating principles",
    "3. Tool recommendation and decision rationale",
    "4. Target operating architecture",
    "5. Repository portfolio and Salesforce folder structure",
    "6. Version-control and branching standard",
    "7. End-to-end delivery and deployment process",
    "8. Jira workflow, gates and automation",
    "9. CI/CD quality gates by stage",
    "10. Environment, deployment and release controls",
    "11. Security, access and secrets",
    "12. Rollback, hotfix and drift management",
    "13. Roles and accountability",
    "14. Definition of Ready and Definition of Done",
    "15. Confluence publication structure",
    "16. Implementation roadmap and success measures",
    "17. Jira-ready implementation stories",
    "18. Acceptance-criteria traceability and ready-to-paste response",
    "Appendix A. Repository tree",
    "Appendix B. Sources and assumptions",
]
for item in contents:
    add_list_item(doc, item, BULLET_ID)

add_heading(doc, "How to use this document", 1)
add_body(doc, "The opening sections establish the decision. Sections 5 through 14 are the operating standard. Sections 15 through 17 are the implementation package. The appendices contain a copy-ready repository tree and the source basis. This document is designed to be imported or copied into the GTM Systems Confluence space without changing its control logic.")

doc.add_page_break()

add_heading(doc, "1. Executive summary and ticket outcomes", 1)
add_body(doc, "Flywire’s parallel GTM roadmap creates a version-control problem and an operating-model problem. Code, declarative Salesforce metadata, CPQ/OmniStudio configuration, integration assets, tests and runbooks must move through one traceable process even when different teams work at the same time. SALDEV-1500 is therefore not only a repository setup task; it establishes the rules connecting Jira, Git, automated validation, environment promotion and production evidence.")

add_heading(doc, "Decision", 2)
for text in [
    "Repository: GitHub Enterprise Cloud, with private repositories, enterprise identity controls, protected default branches, CODEOWNERS, required pull-request reviews and required status checks.",
    "Work management: Jira remains the authoritative backlog, scope, acceptance, approval and release-evidence record. Git commits, branches, pull requests and deployment jobs must carry the Jira key.",
    "Salesforce delivery: Gearset is the preferred greenfield release layer because it provides Salesforce-aware comparisons, dependency analysis, source-control integration, CI jobs and Jira integration. Copado is the approved alternative when Flywire already has the licenses, platform ownership and skills to operate it effectively.",
    "Cross-GTM CI: GitHub Actions provides repository-level checks and reusable automation for integration code, scripts, documentation and non-Salesforce components. The Salesforce deployment platform owns metadata-aware promotion and validation.",
]:
    add_list_item(doc, text, BULLET_ID)
add_source_note(doc, "Evidence basis: GitHub documents protected branches and deployment environments with required checks/reviewers [S1-S3]; Gearset documents end-to-end CI and Git/Jira integrations [S4-S6].")

add_heading(doc, "SALDEV-1500 outcomes", 2)
add_table(
    doc,
    ["Ticket requirement", "Deliverable in this operating model", "Outcome"],
    [
        ("Recommend a repository/version-control tool", "Section 3 decision and decision matrix", "Addressed"),
        ("Define repository and folder structure", "Sections 5-6 and Appendix A", "Addressed"),
        ("Document deployment and CI/CD process", "Sections 7, 9 and 10", "Addressed"),
        ("Document Jira ticket movement", "Section 8 workflow and transition evidence", "Addressed"),
        ("Create implementation stories", "Section 17 copy-ready Jira stories", "Addressed"),
        ("Publish in GTM Systems Confluence", "Section 15 publication tree and owner actions", "Implementation story required"),
    ],
    [2550, 5010, 1800],
    status_col=2,
)

add_heading(doc, "Decision boundaries", 2)
add_callout(doc, "Not authorized by this document", "No repository, Jira workflow, connected app, Salesforce environment, CI/CD pipeline or production deployment is created or changed by this planning deliverable. Each implementation action requires its approved Jira story and named owner.", "red")

page_break(doc)

add_heading(doc, "2. Scope, exclusions and operating principles", 1)
add_heading(doc, "In scope", 2)
for text in [
    "Salesforce Apex, Lightning Web Components, Flows, objects, fields, layouts, permissions, OmniStudio assets, CPQ configuration represented in supported source formats, Experience Cloud assets and deployment manifests.",
    "Integration code and configuration for APIs, middleware, eventing and downstream GTM systems, provided secrets and environment values are externalized.",
    "Automated tests, static analysis rules, test data definitions, release scripts, runbooks, architecture decisions and deployment evidence.",
    "Business-as-usual changes, platform-unification work and new quote-to-cash capabilities, all using the same control path with risk-scaled gates.",
]:
    add_list_item(doc, text, BULLET_ID)

add_heading(doc, "Out of scope for the initial implementation", 2)
for text in [
    "Replacing Jira, Confluence or an enterprise identity provider.",
    "Placing production data, credentials, tokens, certificates, unmasked exports or customer documents in Git.",
    "Automating high-risk production changes before the pilot proves the validation, approval and rollback controls.",
    "Forcing all GTM systems into one repository when they have different release cadences, owners or access boundaries.",
]:
    add_list_item(doc, text, BULLET_ID)

add_heading(doc, "Non-negotiable principles", 2)
principles = [
    ("One traceable unit of change", "Jira key → branch → commit → pull request → validation → deployment → evidence."),
    ("Git is the source of truth", "No unmanaged production-intent metadata exists only in an org or on a workstation."),
    ("No direct production changes", "Emergency changes use the hotfix path and are reconciled to Git immediately."),
    ("Separation of duties", "The author cannot be the sole reviewer and production approver."),
    ("Small, reversible changes", "Prefer focused tickets and deployable increments; isolate destructive changes."),
    ("Evidence before status", "A Jira transition represents completed evidence, not an expectation."),
    ("Least privilege and no secrets in Git", "Use enterprise identity, service accounts, protected secrets and short-lived credentials."),
    ("Automation with a human production gate", "Automate repeatable checks; retain explicit approval for production."),
]
add_table(doc, ["Principle", "Operating rule"], principles, [2600, 6760], font_size=9.1)

page_break(doc)

add_heading(doc, "3. Tool recommendation and decision rationale", 1)
add_callout(
    doc,
    "Name confirmation",
    "The meeting transcript says “GitSense.” This document interprets that as Gearset, the Salesforce DevOps platform. If a different product was intended, confirm the product name and repeat the procurement comparison before commitment.",
    "amber",
)
add_callout(doc, "Recommended stack", "GitHub Enterprise Cloud + Jira + Gearset, with GitHub Actions for repository-native checks. If Flywire already operates Copado at enterprise scale, use GitHub Enterprise Cloud + Jira + Copado while preserving the same controls and evidence model.", "green")

add_heading(doc, "Why GitHub is the repository decision", 2)
add_body(doc, "The repository tool must serve Salesforce and adjacent GTM engineering, not only metadata deployment. GitHub Enterprise Cloud provides the shared Git system of record, review controls and automation surface. GitHub rulesets and protected branches can require pull requests and successful checks, while deployment environments can require reviewers and restrict deployment branches. These are foundational controls regardless of which Salesforce-aware release product is selected. [S1-S3]")

add_heading(doc, "Why the Salesforce release layer is separate", 2)
add_body(doc, "Salesforce changes have dependency, profile/permission, destructive-change and environment-drift behavior that generic Git automation does not fully model. A Salesforce-aware platform reduces manual packaging and gives admins a guided promotion path. Gearset documents CI jobs that compare version control with Salesforce orgs, build packages, analyze dependencies, run tests and retain deployment reports; it also documents GitHub and Jira integrations. [S4-S6]")

doc.add_page_break()
add_heading(doc, "Decision matrix", 2)
tool_rows = [
    ("GitHub Enterprise + Gearset", "Preferred greenfield", "Broad Git governance plus Salesforce-aware CI/CD, dependency analysis, Jira integration and admin-friendly operations.", "Confirm enterprise/security review, data residency, license tiers and service-account model."),
    ("GitHub Enterprise + Copado", "Approved alternative", "Strong Salesforce-native promotion model, user-story/branch associations and pipeline controls; good fit when Copado capability already exists.", "Avoid duplicate Jira/Copado story administration; confirm Source Format feature limitations and operating ownership. [S7-S8]"),
    ("GitHub Enterprise + Actions + sf CLI", "Engineering-led fallback", "Maximum code transparency and cross-platform reuse; no additional Salesforce deployment vendor required.", "Higher build/maintenance burden and a steeper path for declarative admins; team owns dependency handling, UX and audit reporting."),
    ("Salesforce DevOps Center", "Pilot/limited fit", "Native Salesforce option with GitHub connection and work-item concepts.", "Assess fit for complex GTM integrations, CPQ/OmniStudio scope, cross-system releases and enterprise reporting before selection."),
]
add_table(doc, ["Option", "Position", "Strength", "Primary watch-out"], tool_rows, [2000, 1550, 3100, 2710], font_size=8.2, status_col=1)

add_heading(doc, "Procurement decision gates", 2)
for text in [
    "Confirm whether Copado licenses, administrators, pipeline configuration and support ownership already exist at Flywire.",
    "Run the same pilot scenario in Gearset and Copado if there is no decisive existing investment: retrieve a representative Salesforce change, create a ticket-linked branch, validate a pull request, promote to a test org and export audit evidence.",
    "Score security architecture, SSO/SCIM, role segregation, audit retention, Salesforce/CPQ/OmniStudio support, Jira linkage, API/export access, service-account licensing, support model and total three-year cost.",
    "Do not choose a product solely because it can deploy metadata; the winning tool must support the operating controls in this document.",
]:
    add_list_item(doc, text, DECIMAL_ID)

page_break(doc)

add_heading(doc, "4. Target operating architecture", 1)
add_picture_with_caption(doc, ARCH_IMG, "Figure 1. Target control path. Git and Jira remain authoritative even when the Salesforce release platform performs promotion.", "Flow from Jira to GitHub, automated quality gates, Gearset or Copado, and Salesforce environments with governance controls underneath.")
add_heading(doc, "System-of-record boundaries", 2)
add_table(
    doc,
    ["System", "Authoritative for", "Must link to"],
    [
        ("Jira", "Business scope, acceptance criteria, ownership, risk, approvals, status and release evidence", "Branch, PR, validation and deployment records"),
        ("GitHub", "Versioned code/configuration, branch history, reviews, tags and repository automation", "Jira ticket and deployment platform"),
        ("Gearset or Copado", "Salesforce comparison, package construction, validation, promotion and deployment audit", "Git commit/PR and Jira ticket"),
        ("Confluence", "Published standards, runbooks, architecture decisions, training and release calendar", "Jira implementation stories and repository source"),
        ("Salesforce orgs", "Runtime state only; not the long-term version history", "Deployed commit/tag and release record"),
    ],
    [1500, 4860, 3000],
    font_size=8.8,
)

add_heading(doc, "Minimum integration contract", 2)
for text in [
    "Branch format begins with the Jira key; pull request title begins with the Jira key; commits include the Jira key.",
    "The Jira development panel or equivalent links the branch, PR and build/deployment events.",
    "A deployment record identifies source commit/tag, target environment, component scope, test level/results, approver, operator, timestamps and outcome.",
    "Production credentials are available only to the protected production job/environment, after approval, and never to pull-request jobs.",
    "The evidence link is written back to Jira before the ticket moves from Deployed to Done.",
]:
    add_list_item(doc, text, BULLET_ID)

page_break(doc)

add_heading(doc, "5. Repository portfolio and Salesforce folder structure", 1)
add_heading(doc, "Portfolio decision: domain-aligned repositories", 2)
add_body(doc, "Use a small set of repositories aligned to deployable systems and access boundaries. Do not create one repository per ticket and do not place every GTM technology in one monorepo by default. Start with the portfolio below and add a repository only when ownership, release cadence, security boundary or build tooling is materially different.")
repo_rows = [
    ("gtm-salesforce-core", "Unified Salesforce metadata, Apex/LWC, supported CPQ and OmniStudio source, manifests and tests", "GTM Salesforce engineering"),
    ("gtm-integrations", "API/middleware services, schemas, mappings, contract tests and deployment code", "Integration engineering"),
    ("gtm-devops-templates", "Reusable workflows, repository policies, scripts and quality configurations", "Platform & DevOps"),
    ("gtm-architecture-and-runbooks", "ADRs, diagrams, operating standards, release/rollback runbooks and Confluence source", "Architecture + release management"),
    ("gtm-data-config", "Non-secret, deployable configuration and synthetic seed definitions that cannot live naturally with system code", "System owners; create only when needed"),
]
add_table(doc, ["Repository", "Contents", "Primary owner"], repo_rows, [2200, 4900, 2260], font_size=8.6)

add_heading(doc, "Salesforce repository standard", 2)
sf_rows = [
    (".github/", "CODEOWNERS, PR template, issue forms, workflows and dependency/update policy."),
    ("force-app/main/default/", "Salesforce DX source format. Organize deployable metadata by native metadata folders."),
    ("packages/<domain>/", "Optional package directories for bounded capabilities such as quoting, pricing, contracting or shared platform; introduce only after dependency analysis."),
    ("manifest/", "Canonical manifests, environment-safe destructive manifests and retrieval/deployment scopes."),
    ("config/", "Scratch-org/project definitions and non-secret tool configuration."),
    ("scripts/ci/", "Deterministic validation, delta/package generation, analysis and test orchestration."),
    ("scripts/release/", "Release packaging, evidence capture, tag verification and approved operational helpers."),
    ("tests/", "Integration, UI, contract and test-data definitions not colocated with source."),
    ("data/seed/", "Synthetic or masked seed plans only; never production exports or personal data."),
    ("docs/", "Architecture decisions, component maps, runbooks, release notes and manual deployment steps."),
    ("Root files", "README.md, CONTRIBUTING.md, SECURITY.md, CODEOWNERS reference, sfdx-project.json, package.json and ignore files."),
]
add_table(doc, ["Path", "Purpose and rule"], sf_rows, [2550, 6810], font_size=8.6)

add_heading(doc, "Repository content rules", 2)
for text in [
    "Commit text-based source whenever the platform supports it. Use Git LFS only for necessary large binaries and after platform/security approval.",
    "Exclude credentials, tokens, certificates, auth URLs, customer data, production exports, generated build outputs, local CLI state and machine-specific files through .gitignore and secret scanning.",
    "Keep environment-specific IDs, endpoints and secrets outside deployable source where possible; use named configuration, protected variables or post-deployment tasks with documented ownership.",
    "Every manual step must live in a versioned runbook and be referenced by the release record; undocumented manual setup is a release blocker.",
]:
    add_list_item(doc, text, BULLET_ID)

page_break(doc)

add_heading(doc, "6. Version-control and branching standard", 1)
add_callout(doc, "Branch model", "Use short-lived ticket branches plus controlled environment branches for the Salesforce promotion path. Keep main aligned to production intent and tag every production release.", "blue")

branch_rows = [
    ("main", "Production intent and release history", "Protected; no direct push; release manager merge; tag after successful production deployment."),
    ("uat", "Approved UAT candidate", "Protected; promotion PR only; successful UAT validation required."),
    ("integration", "Shared integration baseline", "Protected; feature PRs target this branch after review and validation."),
    ("SALDEV-<key>-<slug>", "Short-lived feature/fix branch", "Create from the current approved base; delete after promotion; rebase/merge per repository policy."),
    ("release/<yyyy.mm>", "Optional coordinated release train", "Create only when a freeze or parallel release is required; not a permanent default."),
    ("hotfix/SALDEV-<key>-<slug>", "Emergency correction from main", "Requires incident/change reference, expedited review and mandatory back-promotion."),
]
add_table(doc, ["Branch", "Purpose", "Control"], branch_rows, [2200, 2550, 4610], font_size=8.5)

add_heading(doc, "Naming and traceability", 2)
for text in [
    "Branch: SALDEV-1523-price-ramp-validation",
    "Commit: SALDEV-1523 Add validation for ramped quote groups",
    "Pull request: SALDEV-1523 — Validate stepped pricing dates",
    "Production tag: salesforce-v2026.09.15.1",
    "Release record: includes the immutable commit SHA and production deployment job ID.",
]:
    add_list_item(doc, text, BULLET_ID)

add_heading(doc, "Default branch/ruleset controls", 2)
for text in [
    "Require pull requests, two approvals for production-path merges and at least one CODEOWNER approval for owned paths.",
    "Dismiss stale approvals when code changes; require conversation resolution and successful status checks.",
    "Block force pushes and branch deletion; restrict bypass to a very small audited break-glass group.",
    "Require signed commits if enterprise policy supports the user/tooling model; otherwise prioritize enforced reviews and immutable audit logs.",
    "Use merge commits or squash merges consistently. Recommended: squash feature PRs for one-ticket/one-change traceability, preserve promotion merges as configured by the release platform.",
]:
    add_list_item(doc, text, BULLET_ID)
add_source_note(doc, "GitHub documents protected branches and enterprise rulesets that can require PRs, reviews and status checks [S2-S3].")

page_break(doc)

doc.add_page_break()
add_heading(doc, "7. End-to-end delivery and deployment process", 1)
process_rows = [
    ("1. Refine", "Product owner + architect", "Scope, acceptance criteria, metadata/data/API impact, risk, dependencies, test personas, rollback outline", "Ready for Dev"),
    ("2. Branch", "Developer/admin", "Ticket-linked branch from approved base; environment synchronized; no unrelated components", "In Development"),
    ("3. Build", "Developer/admin", "Source captured in DX/source format; local checks; manual steps documented; no secrets", "In Development"),
    ("4. Pull request", "Author", "Focused diff, risk, test evidence, component list, rollback, screenshots where relevant", "Code Review"),
    ("5. Validate", "CI + reviewer", "Static analysis, unit tests, metadata validation/check-only and dependency review", "Ready for QA"),
    ("6. Integrate", "Release platform", "Promote approved commit to INT/QA; capture job ID and scope", "In QA"),
    ("7. Test", "QA", "Regression, permission/persona, integration and negative-path evidence", "Ready for UAT"),
    ("8. UAT", "Business owner", "Acceptance criteria executed; named approval recorded; defects linked", "Ready for Release"),
    ("9. Release", "Release manager", "Change window, production validation, approvals, runbook, backup/rollback readiness", "Scheduled"),
    ("10. Deploy", "Approved operator/service", "Protected production job; immutable source commit; no scope substitution", "Deployed"),
    ("11. Verify", "Release + business", "Smoke tests, monitoring, business validation, deployment evidence and known issues", "Done"),
]
add_table(doc, ["Stage", "Accountable role", "Required evidence", "Exit status"], process_rows, [1550, 1800, 4450, 1560], font_size=8.0)

add_heading(doc, "Change classes", 2)
class_rows = [
    ("Standard", "Repeatable, low-risk, no destructive/security/integration impact", "Normal PR + automated validation + QA; predefined release approval."),
    ("Normal", "Most functional/configuration/code changes", "Full process including UAT and release-manager approval."),
    ("High risk", "Permissions, authentication, data model deletion, critical integration, billing/quote logic or broad automation", "Architect/security review, expanded regression, rehearsed rollback and explicit business approval."),
    ("Emergency", "Active incident or urgent regulatory/security fix", "Hotfix path, expedited independent review, production approval and retrospective/back-promotion within one business day."),
]
add_table(doc, ["Class", "Definition", "Required path"], class_rows, [1600, 3650, 4110], font_size=8.6)

page_break(doc)

add_heading(doc, "8. Jira workflow, gates and automation", 1)
add_picture_with_caption(doc, JIRA_IMG, "Figure 2. Proposed Jira lifecycle. Status changes are evidence gates, not manual progress labels.", "Jira workflow from backlog to done, with ticket branch, pull request, QA, UAT, release approval, production deployment and post-deployment evidence.")

workflow_rows = [
    ("Backlog", "Request accepted for triage", "Problem/value, requester, affected system", "Product owner"),
    ("Ready for Dev", "Definition of Ready is complete", "Acceptance criteria, design/risk, dependencies, estimate, test and rollback approach", "Product owner + architect"),
    ("In Development", "Ticket branch exists", "Branch link and assignee", "Developer/admin"),
    ("Code Review", "Pull request opened", "PR link, component list, test evidence, risk/rollback", "Reviewer/CODEOWNER"),
    ("Ready for QA", "PR approved and validation passed", "Review approvals and successful required checks", "Dev lead"),
    ("In QA", "Approved commit deployed to QA", "Deployment job ID, commit SHA, scope and test plan", "QA lead"),
    ("Ready for UAT", "QA exit criteria passed", "QA results and defects disposition", "QA lead"),
    ("In UAT", "UAT environment ready", "UAT version/commit, test cases and named testers", "Business owner"),
    ("Ready for Release", "UAT approved", "UAT sign-off, release notes, risk class and rollback", "Business owner + architect"),
    ("Scheduled", "Production approval and change window confirmed", "Release record, approver, runbook and production validation", "Release manager"),
    ("Deployed", "Production job succeeded", "Job ID, commit/tag, scope, tests and operator", "Release manager"),
    ("Done", "Smoke/business checks passed", "Post-deploy evidence, incidents/known issues and final release link", "Product owner"),
]
add_table(doc, ["Status", "Entry condition", "Mandatory evidence", "Transition owner"], workflow_rows, [1450, 2450, 4050, 1410], font_size=7.8)

add_heading(doc, "Transition and approval controls", 2)
for text in [
    "Do not allow unrestricted Any Status transitions into approval-controlled statuses. Atlassian notes that permissive transitions can bypass approval controls. [S9-S10]",
    "If Jira Premium approval steps are available, use distinct approval fields for UAT and Production Release. If not, use validators/conditions plus named approver fields and immutable approval comments.",
    "A failed validation, rejected review, QA defect or declined UAT returns the ticket to In Development with the failure/reason linked.",
    "Automations may advance a ticket only from trusted events (for example, an approved PR or successful deployment webhook) and must never mark Done on deployment success alone.",
    "Subtasks may track build, testing or documentation, but the parent ticket owns the end-to-end release status and evidence.",
]:
    add_list_item(doc, text, BULLET_ID)

add_heading(doc, "Required Jira fields", 2)
fields = [
    ("Affected system/repository", "Controls owner and workflow routing"),
    ("Change class", "Standard / Normal / High risk / Emergency"),
    ("Target release", "Release train or approved production window"),
    ("Component/dependency summary", "Metadata, code, APIs, data/config and downstream dependencies"),
    ("Test strategy", "Automated tests, regression scope, personas and data setup"),
    ("Rollback/forward-fix plan", "Action, decision point and owner"),
    ("UAT approver + approval", "Named business approval evidence"),
    ("Release approver + approval", "Named production approval evidence"),
    ("PR/build/deployment links", "Automated where supported"),
    ("Production commit/tag and job ID", "Immutable release evidence"),
]
add_table(doc, ["Field", "Purpose"], fields, [3200, 6160], font_size=8.8)

page_break(doc)

doc.add_page_break()
add_heading(doc, "9. CI/CD quality gates by stage", 1)
gate_rows = [
    ("Commit/local", "Formatting, lint, focused unit tests, secret pre-check, metadata source sanity", "Author", "Fail blocks PR readiness"),
    ("Pull request", "Jira key, allowed paths, CODEOWNERS, secret scan, static analysis, Apex/LWC tests, metadata validation/check-only, dependency/package review", "CI + reviewer", "All required checks pass"),
    ("Integration", "Deploy approved commit, integration tests, scheduled Apex conflicts, API/contract checks", "Release platform", "Job and test evidence retained"),
    ("QA", "Regression, permission/persona coverage, negative paths, CPQ/OmniStudio/integration scenarios", "QA", "No open release blocker"),
    ("UAT", "Business acceptance criteria against immutable commit; training/release notes checked", "Business owner", "Named approval"),
    ("Production validation", "Same package/commit, target-org validation, prescribed Apex tests, manual-step review", "Release manager", "Validation within approved window"),
    ("Production", "Environment approval, protected credential, approved package only, job logging", "Release manager/operator", "Successful deployment"),
    ("Post-production", "Smoke tests, monitoring, key persona check, evidence write-back", "Release + business", "Done criteria met"),
]
add_table(doc, ["Gate", "Checks", "Owner", "Pass rule"], gate_rows, [1700, 4350, 1450, 1860], font_size=8.0)

add_heading(doc, "Minimum Salesforce pull-request checks", 2)
for text in [
    "Validate the exact changed package against the target integration or validation org using Salesforce CLI or the selected platform; do not treat a source parse as a deployability check. [S11]",
    "Run impacted Apex tests at minimum; require broader regression for shared triggers/frameworks, security changes, critical revenue logic and production validation.",
    "Run LWC Jest tests and lint where relevant; run static analysis with an approved ruleset and documented suppressions.",
    "Check profiles and permission sets for unintended removals or broad access. Treat permissions as whole-file changes where platform behavior requires it.",
    "Detect destructive changes, renamed metadata, missing dependencies, unsupported metadata and environment-specific references; route any destructive change to high-risk approval.",
    "Store validation job IDs, component counts, test counts/results and omissions. A partial or blocked validation is not a pass.",
]:
    add_list_item(doc, text, BULLET_ID)

add_heading(doc, "Evidence schema", 2)
evidence_rows = [
    ("Identity", "Jira key, repository, branch, PR, commit SHA/tag"),
    ("Scope", "Added/changed/deleted components; package/manifest; manual steps; explicit exclusions"),
    ("Target", "Named environment/org alias and organization identifier where policy permits"),
    ("Validation", "Job ID, status, timestamp, test level, tests run/passed/failed and static-analysis outcome"),
    ("Approval", "Reviewer, UAT approver, release approver and approval timestamp"),
    ("Deployment", "Production job ID, operator/service identity, start/end time and result"),
    ("Verification", "Smoke checks, monitoring, incidents, rollback decision and final sign-off"),
]
add_table(doc, ["Evidence category", "Required content"], evidence_rows, [2500, 6860], font_size=8.8)

page_break(doc)

add_heading(doc, "10. Environment, deployment and release controls", 1)
env_rows = [
    ("Developer / scratch", "Isolated build and focused tests", "Developer/admin", "No production credentials; ticket branch only"),
    ("Integration", "Shared dependency and integration validation", "Dev lead", "Automatic after approved merge where safe"),
    ("QA", "Functional and regression testing", "QA lead", "Approved commit; controlled test data"),
    ("UAT", "Business acceptance and release candidate", "Business owner", "Immutable candidate; no untracked fixes"),
    ("Production", "Live service", "Release manager + system owner", "Protected job, explicit approval, least privilege, post-deploy verification"),
]
add_table(doc, ["Environment", "Purpose", "Owner", "Control"], env_rows, [1600, 2850, 1700, 3210], font_size=8.5)

add_heading(doc, "Release readiness checklist", 2)
release_checklist = [
    "All included Jira tickets are Ready for Release and identify the same immutable commit/tag.",
    "Production validation passed for the exact package and target, with current job evidence.",
    "Dependencies, manual steps, data/configuration tasks, integration sequencing and expected downtime are documented.",
    "No unresolved blocker or unapproved scope change exists; exceptions have named acceptance and expiry.",
    "Rollback/forward-fix decision points, owners, scripts and backups are ready and rehearsed for high-risk changes.",
    "Communications, monitoring and business smoke-test owners are scheduled.",
    "The production operator is different from the sole code author and has the correct least-privilege access.",
]
for item in release_checklist:
    add_list_item(doc, item, BULLET_ID)

add_heading(doc, "Production deployment rules", 2)
for text in [
    "Production deployment runs only from the approved release commit/tag; never from a workstation’s uncommitted state or an arbitrary org comparison.",
    "No component may be added after validation without returning to review and revalidation.",
    "Use the platform’s validation/quick-deploy capability only when the validated artifact, tests and validity window meet Salesforce and company policy.",
    "Manual steps require two-person verification for high-risk actions and a timestamped execution record.",
    "After deployment, create the release tag and release notes only when the actual production state is confirmed; if tagging occurs before execution, the release record must distinguish planned from deployed.",
]:
    add_list_item(doc, text, BULLET_ID)

page_break(doc)

add_heading(doc, "11. Security, access and secrets", 1)
security_rows = [
    ("Human access", "Enterprise SSO, MFA, group-based teams, least privilege and periodic recertification."),
    ("Service accounts", "Dedicated non-human identities; no shared personal credentials; named owner and rotation plan."),
    ("Secrets", "Repository/environment secret store or approved enterprise vault; never Git, Jira comments, build logs or documents."),
    ("Production gate", "Protected deployment environment with required reviewer(s), branch restriction and no self-approval where supported. [S1]"),
    ("Actions/workflows", "Allow approved actions only; pin third-party actions to immutable versions/SHAs per enterprise policy. [S3]"),
    ("Audit", "Retain Git, Jira and deployment-platform audit records according to Flywire policy; export when vendor retention is insufficient."),
    ("Data", "Use synthetic or masked non-production data; do not commit Salesforce exports or personal/customer data."),
    ("Break glass", "Small named group, time-bound access, mandatory reason, immediate notification and retrospective review."),
]
add_table(doc, ["Control area", "Required standard"], security_rows, [2300, 7060], font_size=8.7)

add_heading(doc, "Access model", 2)
access_rows = [
    ("Contributor", "Create ticket branches and PRs; no protected-branch push; no production deploy."),
    ("Reviewer / CODEOWNER", "Review owned paths; cannot replace release approval."),
    ("Release operator", "Run approved promotion/deployment jobs; cannot bypass scope/review."),
    ("Release approver", "Approve production environment/job; ideally not the author or operator."),
    ("Repository admin", "Manage policy and teams; no routine bypass; changes are audited."),
    ("Security auditor", "Read audit/configuration evidence; no deployment mutation."),
]
add_table(doc, ["Role", "Access boundary"], access_rows, [2600, 6760], font_size=8.8)

page_break(doc)

add_heading(doc, "12. Rollback, hotfix and drift management", 1)
add_heading(doc, "Rollback strategy", 2)
add_body(doc, "Salesforce rollback is change-specific. The default is a controlled forward fix or a revert commit deployed through the same pipeline; destructive metadata and data changes may require restoration or compensating actions. Every ticket must therefore state the actual rollback boundary rather than using a generic sentence.")
rollback_rows = [
    ("Apex/LWC/Flow/config", "Revert the approved commit or deploy the prior known-good version; validate dependencies and activation state."),
    ("Permissions", "Restore the prior permission source; verify user/persona access and removals before/after deployment."),
    ("Destructive metadata", "Require pre-change backup and explicit restoration/forward-fix runbook; deletion may not be trivially reversible."),
    ("Data/schema migration", "Back up affected data, define transformation reversal or compensating migration, and identify ownership."),
    ("Integration/API", "Feature flag, route/version rollback, credential/config reversal and partner communication as applicable."),
]
add_table(doc, ["Change type", "Rollback expectation"], rollback_rows, [2400, 6960], font_size=8.7)

add_heading(doc, "Emergency hotfix flow", 2)
DECIMAL_ID = new_numbering_instance(doc, DECIMAL_ABSTRACT_ID)
hotfix_steps = [
    "Open or link an incident and emergency Jira ticket; record impact, urgency and approver.",
    "Create hotfix/SALDEV-<key>-<slug> from main and make the smallest safe change.",
    "Run focused automated validation and obtain independent technical review; record any skipped gates and risk acceptance.",
    "Deploy through the protected production path with emergency approval and capture the job evidence.",
    "Verify service recovery, then merge/back-promote the hotfix through UAT/integration branches so future releases retain it.",
    "Complete a retrospective and close any missing test, documentation or automation debt within one business day or the agreed incident SLA.",
]
for item in hotfix_steps:
    add_list_item(doc, item, DECIMAL_ID)

add_heading(doc, "Drift management", 2)
for text in [
    "Run scheduled source-to-org comparison/change monitoring for production and key sandboxes; classify differences as authorized, expected platform behavior or unmanaged drift.",
    "Block routine direct production changes. If an emergency change occurs, retrieve/reconcile it into Git immediately and link the incident/ticket.",
    "Refresh the production baseline only through a reviewed pull request. Never overwrite Git blindly from a broad org retrieve.",
    "Report drift age, unowned differences and repeated bypass patterns as operating metrics.",
]:
    add_list_item(doc, text, BULLET_ID)

page_break(doc)

add_heading(doc, "13. Roles and accountability", 1)
raci_rows = [
    ("Define scope/acceptance", "A", "C", "R", "C", "C", "I"),
    ("Design/risk review", "C", "A", "R", "C", "C", "C"),
    ("Build and local test", "I", "C", "A/R", "C", "I", "I"),
    ("Code/config review", "I", "A", "R", "C", "I", "C"),
    ("QA/regression", "I", "C", "C", "A/R", "I", "I"),
    ("UAT approval", "C", "C", "I", "C", "A/R", "I"),
    ("Release readiness", "C", "C", "C", "C", "C", "A/R"),
    ("Production approval", "C", "C", "I", "I", "A", "R"),
    ("Production deployment", "I", "C", "I", "I", "I", "A/R"),
    ("Post-deploy verification", "A", "C", "C", "R", "R", "R"),
]
add_table(doc, ["Activity", "Product owner", "Architect/dev lead", "Developer/admin", "QA", "Business/UAT", "Release mgr"], raci_rows, [2150, 1200, 1450, 1350, 1000, 1100, 1110], font_size=7.4, center_cols={1,2,3,4,5,6})
add_source_note(doc, "R = Responsible, A = Accountable, C = Consulted, I = Informed. Security and platform administration are mandatory consultees/approvers for changes in their control domains.")

add_heading(doc, "Operating ceremonies", 2)
ceremonies = [
    ("Weekly intake/refinement", "Prioritize requests; establish DoR, dependencies, change class and release target."),
    ("Daily pipeline review", "Review failed validations, blocked promotions, drift alerts and aging PRs."),
    ("Release readiness review", "Confirm evidence, package scope, approvals, runbook and rollback before production window."),
    ("Post-release review", "Confirm business outcome, incidents, metrics, evidence completeness and follow-up actions."),
    ("Monthly control review", "Review access, bypasses, drift, failed change rate, tool health and process exceptions."),
]
add_table(doc, ["Cadence", "Purpose"], ceremonies, [2700, 6660], font_size=8.8)

page_break(doc)

add_heading(doc, "14. Definition of Ready and Definition of Done", 1)
add_heading(doc, "Definition of Ready", 2)
for text in [
    "Business outcome, affected personas and measurable acceptance criteria are clear.",
    "Affected system/repository, component types, dependencies and integration/data impacts are identified.",
    "Change class, architecture/security review needs and release target are set.",
    "Test strategy covers automated tests, QA regression, UAT personas and required test data.",
    "Rollback/forward-fix approach and manual steps are plausible and have owners.",
    "Required environments, access and upstream prerequisites are available or explicitly tracked.",
    "Story is sized and small enough to review, validate and reverse independently where practical.",
]:
    add_list_item(doc, text, BULLET_ID)

add_heading(doc, "Definition of Done", 2)
for text in [
    "Approved source is merged and the ticket, PR, commit and deployment records are linked.",
    "Required automated checks, target-org validation, QA and UAT are complete with exact evidence.",
    "Production deployment succeeded from the approved immutable commit/tag with no unapproved scope changes.",
    "Post-deployment smoke and business checks passed; incidents/known issues are linked and owned.",
    "Documentation, runbooks, release notes and monitoring updates are published where required.",
    "Manual changes are reconciled into Git; temporary access/flags are removed or have an expiry story.",
    "The Jira ticket contains production job ID, component/test evidence, approvers, outcome and final release link.",
]:
    add_list_item(doc, text, BULLET_ID)

add_callout(doc, "Closure rule", "Deployment success alone is not Done. Done requires verified behavior, complete evidence and reconciliation to the source of truth.", "amber")

page_break(doc)

add_heading(doc, "15. Confluence publication structure", 1)
add_body(doc, "Publish this material under the GTM Systems space as a controlled page tree. The operating model should be readable by leadership, while detailed runbooks remain versioned and owned. Add page owners and next-review dates; use the repository as the source for technical artifacts where practical.")
confluence_rows = [
    ("GTM Systems Engineering", "Landing page, ownership, links to standards and current release calendar"),
    ("↳ DevOps & Release Management", "Operating-model home and policy status"),
    ("  ↳ SALDEV-1500 — DevOps Operating Model", "Executive decision, principles, architecture and RACI"),
    ("  ↳ Repository & Branching Standard", "Repo portfolio, folder tree, naming, CODEOWNERS and branch policy"),
    ("  ↳ Jira Delivery Workflow", "Statuses, fields, transition gates, approvals and automation"),
    ("  ↳ CI/CD Quality Gates", "Checks, test strategy, evidence schema and exception process"),
    ("  ↳ Deployment & Rollback Runbook", "Environment promotion, production checklist, rollback and hotfix"),
    ("  ↳ Tool Administration", "GitHub and Gearset/Copado configuration, access, service accounts and audit"),
    ("  ↳ Training & Onboarding", "Contributor, reviewer, QA and release-manager guides"),
    ("  ↳ Metrics & Continuous Improvement", "DORA-style metrics, drift, bypasses, failures and monthly actions"),
]
add_table(doc, ["Page", "Purpose"], confluence_rows, [3300, 6060], font_size=8.6)

add_heading(doc, "Publication controls", 2)
for text in [
    "Page owner: Head of GTM Systems or delegated Platform & DevOps lead.",
    "Review cadence: quarterly and after any tool, environment, workflow or production-control change.",
    "Changes to mandatory controls require a Jira change, architecture/release approval and version-history note.",
    "Runbook commands and configuration examples must be tested in a non-production environment before publication as operational instructions.",
]:
    add_list_item(doc, text, BULLET_ID)

page_break(doc)

add_heading(doc, "16. Implementation roadmap and success measures", 1)
roadmap_rows = [
    ("Wave 0 — Decide", "Weeks 0-2", "Confirm product name, existing contracts, security/procurement, owners and pilot scope", "Signed decision and funded implementation epic"),
    ("Wave 1 — Foundation", "Weeks 2-5", "GitHub organization/teams, repository template, baseline metadata, access/secrets and Jira fields/workflow", "Controlled repository and ready backlog"),
    ("Wave 2 — Automate", "Weeks 5-8", "PR checks, Salesforce validation, release platform pipeline, evidence links and runbooks", "End-to-end non-production pipeline"),
    ("Wave 3 — Pilot", "Weeks 8-10", "One low-risk change through INT, QA, UAT and protected production with rollback rehearsal", "Pilot acceptance and documented lessons"),
    ("Wave 4 — Scale", "Weeks 10-14", "Training, repository migration, change freeze/cutover plan, metrics and control review", "BAU operating model adopted"),
]
add_table(doc, ["Wave", "Indicative timing", "Scope", "Exit"], roadmap_rows, [1900, 1500, 4020, 1940], font_size=8.4)

add_heading(doc, "90-day success measures", 2)
metrics = [
    ("Production traceability", "100% of pilot-scope production changes linked Jira → PR → commit/tag → deployment job."),
    ("Protected-path adoption", "100% of main/production-path merges through approved PRs; zero routine direct pushes."),
    ("Validation evidence", "100% of release candidates have target-org validation status, job ID, test counts and explicit omissions."),
    ("Release quality", "Track failed change rate and rollback/forward-fix events; set the steady-state target after the pilot baseline."),
    ("Lead time", "Measure Ready for Dev to Deployed and identify queue time at review, QA, UAT and release gates."),
    ("Drift", "All detected production drift triaged within two business days; unauthorized drift reconciled through Jira/Git."),
    ("Evidence completeness", "At least 95% of Done tickets pass monthly evidence audit in month one; 100% by month three."),
]
add_table(doc, ["Measure", "Initial target / method"], metrics, [2700, 6660], font_size=8.8)

add_heading(doc, "Pilot exit criteria", 2)
for text in [
    "A low-risk ticket completes the full workflow without direct production changes or manual source substitution.",
    "Required checks block an intentionally failing PR and allow a corrected PR.",
    "QA/UAT/production jobs all reference the same approved source lineage.",
    "The team executes the rollback rehearsal in non-production and validates evidence retrieval.",
    "Users representing developer/admin, reviewer, QA, business approver and release manager confirm the process is usable.",
    "Open gaps have owners, severity and due dates; no critical control gap remains before scale-out.",
]:
    add_list_item(doc, text, BULLET_ID)

add_heading(doc, "17. Jira-ready implementation stories", 1)
add_callout(doc, "Backlog usage", "Create these stories under a new SALDEV implementation epic linked to SALDEV-1500. Estimates are indicative planning placeholders and must be re-estimated by the delivery team after tool and environment discovery.", "blue")
doc.add_page_break()

stories = [
    (
        "SALDEV-1500-A01", "Approve repository and Salesforce DevOps toolchain",
        "As the GTM Systems leadership team, we want an approved source-control and release-tool decision so that implementation proceeds with clear ownership, security approval and funding.",
        [
            "GitHub Enterprise Cloud is confirmed as the repository standard or an exception is formally documented.",
            "The transcript term “GitSense” is confirmed as Gearset or corrected to the intended product.",
            "Existing Copado contracts, licenses, administrators and configured pipelines are inventoried.",
            "Gearset and Copado are scored against the procurement gates in Section 3 using the same pilot scenario.",
            "Security, privacy, data residency, SSO/SCIM, audit retention and vendor risk reviews are complete.",
            "Decision record names the selected platform, accountable owner, license quantities, service-account needs and three-year cost assumptions.",
        ],
        "Head of GTM Systems + Procurement", "None", "5 points", "Wave 0",
    ),
    (
        "SALDEV-1500-A02", "Provision GitHub Enterprise organization, teams and base controls",
        "As Platform & DevOps, we want the GitHub organization and access model configured so that GTM source is private, governed and auditable.",
        [
            "Enterprise organization, billing owner and support contacts are confirmed.",
            "SSO/MFA and team-based access are enabled; contributor, reviewer, release and admin groups are mapped.",
            "Default repository visibility is private and external collaboration policy is configured.",
            "Enterprise/organization rulesets protect main and designated promotion branches from direct push, force push and deletion.",
            "Required PR reviews, CODEOWNER review, status checks, conversation resolution and bypass audit are configured.",
            "Audit-log ownership and retention/export procedure are documented and tested.",
        ],
        "Platform & DevOps", "A01", "8 points", "Wave 1",
    ),
    (
        "SALDEV-1500-A03", "Create GTM repository template and Salesforce repository",
        "As a GTM engineer, I want a standard repository template and Salesforce repository so that every change begins with consistent structure, controls and documentation.",
        [
            "gtm-devops-templates and gtm-salesforce-core repositories are created using the structure in Section 5 and Appendix A.",
            "README, CONTRIBUTING, SECURITY, CODEOWNERS, PR template, ignore files and ownership metadata are present.",
            "Salesforce DX project configuration, force-app path, manifests, scripts and test directories are valid.",
            "Repository contains no credentials, auth URLs, production data, generated artifacts or local CLI state.",
            "Template contains a sample ticket-linked branch and PR that passes policy checks.",
            "Repository creation procedure is documented for future GTM systems.",
        ],
        "Platform & DevOps + Salesforce lead", "A02", "8 points", "Wave 1",
    ),
    (
        "SALDEV-1500-A04", "Create reviewed Salesforce production baseline",
        "As the Salesforce engineering team, we want a reviewed source baseline of the current production-intent metadata so that Git becomes the controlled source of truth without hiding existing drift or omissions.",
        [
            "Approved metadata scope and explicit exclusions are documented before retrieval.",
            "Baseline is retrieved using a dedicated read-only/least-privilege process and committed to a temporary review branch.",
            "Unsupported, missing or partial metadata and source-format conversion issues are listed with owners.",
            "Sensitive/environment-specific values and generated files are excluded or externalized.",
            "Architect and system owner approve the baseline PR after component-count and drift review.",
            "No deployment or production mutation occurs as part of the baseline creation.",
        ],
        "Salesforce lead + Architect", "A03", "13 points", "Wave 1",
    ),
    (
        "SALDEV-1500-A05", "Configure service identities, secrets and environment access",
        "As Security and Platform Operations, we want controlled non-human identities and secret storage so that CI/CD can validate and deploy without shared personal credentials.",
        [
            "One named service identity per required trust boundary is documented with owner, purpose and expiry/rotation process.",
            "Non-production and production credentials are separated; pull-request jobs cannot access production credentials.",
            "Secrets are stored only in the approved vault/repository environment store and are masked from logs.",
            "Production environment requires named approval and prevents self-approval where supported.",
            "Least-privilege Salesforce and repository permissions are verified through an access test.",
            "Break-glass access, alerting and quarterly recertification procedures are documented.",
        ],
        "Security + Platform Operations", "A01-A03", "8 points", "Wave 1",
    ),
    (
        "SALDEV-1500-A06", "Implement pull-request quality and Salesforce validation gates",
        "As a reviewer, I want automated pull-request checks so that unsafe or untraceable changes cannot merge into the promotion path.",
        [
            "Workflow validates Jira key, allowed branch naming, changed paths and required PR fields.",
            "Secret scanning, lint/static analysis, relevant LWC tests and Apex test strategy are configured.",
            "Exact changed Salesforce scope is validated/check-only against the approved target org or platform job.",
            "Required checks report job ID, component counts, tests run/passed/failed, warnings and explicit omissions.",
            "An intentionally failing fixture blocks merge and a corrected fixture passes.",
            "Workflow actions/dependencies are approved, pinned and owned; troubleshooting runbook is published.",
        ],
        "Platform & DevOps + Salesforce lead", "A03-A05", "13 points", "Wave 2",
    ),
    (
        "SALDEV-1500-A07", "Configure Salesforce release pipeline across INT, QA, UAT and Production",
        "As a release manager, I want a controlled Salesforce promotion pipeline so that one approved source lineage moves through every environment with auditable gates.",
        [
            "Selected Gearset or Copado platform is connected to GitHub and the named Salesforce environments using approved service identities.",
            "Branch/environment mapping, metadata filters, permission behavior, test levels and deployment ownership are documented.",
            "Pull-request validation and non-production promotion jobs retain immutable source, target, scope, test and outcome evidence.",
            "UAT and production steps require named approvals; production cannot run from unapproved branches or ad hoc source.",
            "Destructive changes and manual deployment tasks use separate high-risk controls.",
            "A dry-run/non-production test proves promotion, failure handling, evidence export and rollback rehearsal.",
        ],
        "Release manager + Salesforce lead", "A01, A04-A06", "13 points", "Wave 2",
    ),
    (
        "SALDEV-1500-A08", "Implement Jira delivery workflow, fields and automation",
        "As a delivery team, we want Jira statuses and evidence gates aligned with CI/CD so that ticket status accurately reflects delivery state.",
        [
            "Statuses, transitions and rework paths match Section 8 and are reviewed with product, QA, architecture and release roles.",
            "Required fields for change class, system, dependencies, test, rollback, UAT, release approval and deployment evidence are configured.",
            "Approval-controlled statuses cannot be reached through unrestricted transitions.",
            "Trusted branch/PR/build/deployment events update links or status without marking Done prematurely.",
            "Rejected review, failed validation, QA defect and declined UAT return to In Development with evidence.",
            "A sandbox Jira project or safe test issue demonstrates each transition, permission and automation path.",
        ],
        "Jira admin + Delivery lead", "A01 and workflow approval", "13 points", "Wave 1-2",
    ),
    (
        "SALDEV-1500-A09", "Publish deployment, rollback, hotfix and drift runbooks",
        "As an on-call release team, we want tested runbooks so that normal, emergency and recovery actions are repeatable and auditable.",
        [
            "Normal release checklist covers readiness, production validation, approval, execution, smoke test and evidence write-back.",
            "Rollback patterns address code/config, permissions, destructive metadata, data/schema and integrations.",
            "Hotfix runbook requires incident linkage, independent review, protected deployment and back-promotion.",
            "Drift-monitoring procedure defines schedule, triage, ownership, reconciliation and escalation.",
            "Commands and screenshots contain no secrets or production session data.",
            "Runbooks are rehearsed in non-production and published in Confluence with owners and review dates.",
        ],
        "Release manager + Architect", "A07-A08", "8 points", "Wave 2",
    ),
    (
        "SALDEV-1500-A10", "Define GTM test strategy and release evidence standard",
        "As QA and Architecture, we want risk-based test and evidence standards so that release decisions are consistent across code, configuration and integrations.",
        [
            "Test matrix defines minimum Apex/LWC/static/integration/regression/UAT coverage by change class.",
            "Permission/persona, CPQ/OmniStudio, integration and negative-path expectations are documented.",
            "Evidence schema from Section 9 is implemented in Jira templates and deployment reporting.",
            "Partial, blocked, skipped and not-applicable outcomes are distinguishable; only explicit passes satisfy a gate.",
            "Test-data ownership and masking/synthetic-data rules are approved.",
            "Monthly evidence audit procedure and sample size are defined.",
        ],
        "QA lead + Architect", "A06-A08", "8 points", "Wave 2",
    ),
    (
        "SALDEV-1500-A11", "Pilot the operating model with one low-risk Salesforce change",
        "As GTM Systems leadership, we want a controlled end-to-end pilot so that the process is proven before broader adoption.",
        [
            "Pilot ticket meets Definition of Ready and uses the approved repository, branch, PR and Jira workflow.",
            "Required CI checks and Salesforce validation pass; an intentional failure is demonstrated safely before correction.",
            "The same approved source lineage is promoted through INT, QA, UAT and the protected production path.",
            "Rollback is rehearsed in non-production and production smoke evidence is captured.",
            "Pilot meets every exit criterion in Section 16 and records time, failures, manual effort and user feedback.",
            "Critical gaps are resolved before scale-out; non-critical gaps have owners and dates.",
        ],
        "Delivery lead + Release manager", "A02-A10", "8 points", "Wave 3",
    ),
    (
        "SALDEV-1500-A12", "Train GTM roles and transition to business-as-usual governance",
        "As a GTM contributor, reviewer, tester or release manager, I want role-based training and support so that I can use the operating model correctly without relying on tribal knowledge.",
        [
            "Role-based guides and live sessions cover contributor/admin, reviewer, QA, business approver and release manager tasks.",
            "Each participant completes a practical exercise in a non-production repository/project.",
            "Support channel, office hours, escalation path and named tool owners are published.",
            "Legacy direct-deployment paths are disabled or documented as time-bound exceptions after pilot acceptance.",
            "Metrics dashboard and monthly control review begin with named owners.",
            "Operating model and Confluence pages receive final owner approval and next-review dates.",
        ],
        "Delivery lead + Platform & DevOps", "A11", "8 points", "Wave 4",
    ),
]

for story_index, story in enumerate(stories):
    if story_index:
        doc.add_page_break()
    add_story(doc, *story)

doc.add_page_break()
add_heading(doc, "18. Acceptance-criteria traceability and ready-to-paste response", 1)
trace_rows = [
    ("What tool works best?", "GitHub Enterprise Cloud as canonical repository; Gearset preferred greenfield Salesforce release layer; Copado approved when existing investment is decisive.", "Sections 3-4; A01"),
    ("What repository structure?", "Domain-aligned repo portfolio plus standard Salesforce DX tree, ownership and content rules.", "Sections 5-6; Appendix A; A03-A04"),
    ("How will deployment/CI/CD work?", "Ticket-to-production lifecycle, environment gates, validation evidence, approvals, rollback and drift controls.", "Sections 7, 9-12; A05-A07, A09-A11"),
    ("How will Jira tickets move?", "Evidence-gated statuses with transition owners, required fields, rejection paths and automation boundaries.", "Section 8; A08"),
    ("What stories implement the plan?", "Twelve Jira-ready stories covering decision, foundation, automation, pilot and adoption.", "Sections 16-17"),
    ("How is Confluence completed?", "Controlled GTM Systems page tree with owners, reviews and linked runbooks.", "Section 15; A09, A12"),
]
add_table(doc, ["SALDEV-1500 acceptance item", "Resolution", "Evidence"], trace_rows, [2500, 5100, 1760], font_size=8.4)

add_heading(doc, "Ready-to-paste SALDEV-1500 comment", 2)
comment = (
    "Completed the proposed GTM Systems DevOps and Jira operating model for SALDEV-1500. The recommendation is to use GitHub Enterprise Cloud as the canonical source repository, Jira as the authoritative work/approval record, and a Salesforce-aware release platform on top of Git. Gearset is the preferred greenfield option; Copado remains the approved alternative if Flywire’s existing license, skills and operating investment make it lower risk. The document defines the repository portfolio and Salesforce DX folder structure, branch and pull-request controls, end-to-end deployment flow, Jira statuses and evidence gates, CI/CD quality checks, environment/production controls, rollback/hotfix/drift processes, RACI, Confluence page structure, rollout measures and twelve Jira-ready implementation stories. Before procurement, confirm that the transcript reference to “GitSense” means Gearset. No repositories, Jira workflows, Salesforce orgs or production pipelines were changed as part of this planning task."
)
add_callout(doc, "Jira comment", comment, "blue")

add_heading(doc, "Closure recommendation", 2)
add_body(doc, "SALDEV-1500 can be closed after: (1) this operating model is approved and published in the GTM Systems Confluence space; (2) the implementation epic and stories are created and linked; and (3) accountable owners accept A01 through A12. Tool purchase, configuration and pipeline execution remain implementation work and should not be represented as completed by this document alone.")

doc.add_page_break()
add_heading(doc, "Appendix A. Copy-ready repository tree", 1)
tree_lines = [
    "gtm-salesforce-core/",
    "├── .github/",
    "│   ├── CODEOWNERS",
    "│   ├── pull_request_template.md",
    "│   └── workflows/",
    "│       ├── pr-validate.yml",
    "│       ├── integration-promote.yml",
    "│       └── release-evidence.yml",
    "├── force-app/main/default/",
    "│   ├── classes/",
    "│   ├── flows/",
    "│   ├── lwc/",
    "│   ├── objects/",
    "│   ├── permissionsets/",
    "│   └── ...native Salesforce metadata folders",
    "├── packages/<domain>/                  # optional after dependency review",
    "├── manifest/",
    "│   ├── package.xml",
    "│   ├── destructiveChangesPre.xml",
    "│   └── destructiveChangesPost.xml",
    "├── config/",
    "│   └── project-scratch-def.json",
    "├── scripts/",
    "│   ├── ci/",
    "│   └── release/",
    "├── tests/",
    "│   ├── integration/",
    "│   └── ui/",
    "├── data/seed/                          # synthetic or masked definitions only",
    "├── docs/",
    "│   ├── adr/",
    "│   ├── architecture/",
    "│   ├── releases/",
    "│   └── runbooks/",
    "├── .forceignore",
    "├── .gitignore",
    "├── CONTRIBUTING.md",
    "├── README.md",
    "├── SECURITY.md",
    "├── package.json",
    "└── sfdx-project.json",
]
for line in tree_lines:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(line)
    set_run_font(r, name="Consolas", size=8.8, color=INK)

doc.add_page_break()
add_heading(doc, "Standard pull-request body", 2)
pr_fields = [
    ("Jira", "SALDEV-####"),
    ("Business outcome", "What changes for which persona?"),
    ("Scope", "Components/files added, changed and deleted"),
    ("Explicit exclusions", "What is intentionally not included?"),
    ("Risk/change class", "Standard / Normal / High risk / Emergency"),
    ("Validation", "Job ID, target, component counts and test results"),
    ("Manual steps", "Before/after deployment steps and owners"),
    ("Rollback", "Revert/forward-fix/restoration path and decision point"),
    ("Evidence", "Screenshots, logs, QA/UAT links and known limitations"),
]
add_table(doc, ["Field", "Required content"], pr_fields, [2500, 6860], font_size=8.7)

add_heading(doc, "Copy/paste pull-request template", 2)
pr_template_lines = [
    "## Summary",
    "Jira: SALDEV-####",
    "Business outcome: <persona and measurable outcome>",
    "",
    "## Scope and evidence",
    "Added / changed / deleted: <components and files>",
    "Explicit exclusions: <items intentionally omitted>",
    "Validation: <target, job ID, component counts and result>",
    "Tests: <level, run/pass/fail counts and omissions>",
    "QA / UAT evidence: <links and approver>",
    "",
    "## Risk and release",
    "Change class: Standard | Normal | High risk | Emergency",
    "Manual steps: <before/after steps and owner>",
    "Rollback / forward fix: <action, decision point and owner>",
    "Known limitations: <none or linked issue>",
]
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.08)
p.paragraph_format.right_indent = Inches(0.08)
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(0)
p.paragraph_format.line_spacing = Pt(9.6)
for index, line in enumerate(pr_template_lines):
    r = p.add_run(line)
    set_run_font(r, name="Consolas", size=7.8, color=INK)
    if index < len(pr_template_lines) - 1:
        r.add_break()

doc.add_page_break()
add_heading(doc, "Appendix B. Sources and assumptions", 1)
add_body(doc, "Sources were accessed on 28 August 2026. Vendor features, editions, licensing and limits can change; procurement and security teams must confirm current contract terms before purchase. Vendor documentation supports capability statements, while the operating-model recommendation and trade-offs are an architecture judgment for the SALDEV-1500 context.")

sources = [
    ("S1", "GitHub Docs — Deployments and environments", "https://docs.github.com/en/actions/reference/workflows-and-actions/deployments-and-environments"),
    ("S2", "GitHub Enterprise Cloud Docs — Managing protected branches", "https://docs.github.com/en/enterprise-cloud@latest/repositories/configuring-branches-and-merges-in-your-repository/managing-protected-branches"),
    ("S3", "GitHub Enterprise Cloud Docs — Establishing a governance framework", "https://docs.github.com/en/enterprise-cloud@latest/admin/overview/establishing-a-governance-framework-for-your-enterprise"),
    ("S4", "Gearset Help Center — Continuous integration", "https://docs.gearset.com/en/articles/7858348-continuous-integration"),
    ("S5", "Gearset Help Center — Supported software integrations", "https://docs.gearset.com/en/articles/3741573-gearset-supported-software-integrations"),
    ("S6", "Gearset Help Center — Sample source-driven Salesforce workflow", "https://docs.gearset.com/en/articles/2934192-a-sample-source-driven-development-workflow-for-salesforce"),
    ("S7", "Copado Docs — Known limitations in Salesforce Source Format Pipelines", "https://docs.copado.com/articles/?_escaped_fragment_=source-format-pipelines-publication%2Fknown-limitations-in-salesforce-source-format-pipelines"),
    ("S8", "Copado Docs — Structure of a user story", "https://docs.copado.com/articles/?_escaped_fragment_=source-format-pipelines-publication%2Fstructure-of-a-user-story"),
    ("S9", "Atlassian Support — What are Jira workflows?", "https://support.atlassian.com/jira-software-cloud/docs/what-are-jira-workflows/"),
    ("S10", "Atlassian Support — What are approvals?", "https://support.atlassian.com/jira-software-cloud/docs/what-are-approvals/"),
    ("S11", "Salesforce Developers — Salesforce CLI command reference", "https://developer.salesforce.com/docs/platform/salesforce-cli-reference/guide/cli_reference.html"),
    ("S12", "SALDEV-1500 Jira ticket", "https://jira.flywire.tech/browse/SALDEV-1500"),
]
for sid, title, url in sources:
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(f"[{sid}] ")
    set_run_font(r, size=9.5, color=NAVY, bold=True)
    add_hyperlink(p, title, url)

add_heading(doc, "Assumptions requiring confirmation", 2)
assumptions = [
    "“GitSense” means Gearset; if not, A01 must compare the intended product.",
    "GitHub Enterprise Cloud is acceptable under Flywire’s enterprise architecture and vendor-security policies.",
    "Jira edition and project type support the desired approval/automation behavior; if not, validators and explicit approval evidence will be used.",
    "Environment names shown are logical. The implementation stories must map them to Flywire’s actual Salesforce orgs without assuming aliases or access.",
    "CPQ/OmniStudio and other GTM configuration will be versioned only through supported, testable source representations; unsupported artifacts require documented manual steps or a specialized tool capability.",
    "This document does not approve a deployment, license purchase or external-system change.",
]
for item in assumptions:
    add_list_item(doc, item, BULLET_ID)

# Core properties and save
doc.core_properties.title = "SALDEV-1500 — GTM DevOps, Code Repository and Jira Operating Model"
doc.core_properties.subject = "Source control, repository structure, CI/CD, Jira workflow and implementation stories"
doc.core_properties.author = "Flywire GTM Systems"
doc.core_properties.keywords = "SALDEV-1500, DevOps, GitHub, Gearset, Copado, Jira, Salesforce, CI/CD"
doc.core_properties.comments = "Proposed operating standard prepared for stakeholder review."

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
