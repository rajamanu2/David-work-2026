import sys
from docx import Document

sys.stdout.reconfigure(encoding="utf-8")

path = r"C:\Users\LIKKI\Downloads\Pro_Biomedical_Salesforce_Refactoring_Assessment_and_Plan.docx"
doc = Document(path)

print(f"PARAGRAPHS={len(doc.paragraphs)} TABLES={len(doc.tables)} SECTIONS={len(doc.sections)}")
for index, paragraph in enumerate(doc.paragraphs, start=1):
    text = paragraph.text.strip()
    if text:
        print(f"P{index:04d}\t{paragraph.style.name}\t{text}")

for table_index, table in enumerate(doc.tables, start=1):
    print(f"TABLE {table_index} rows={len(table.rows)} cols={len(table.columns)}")
    for row_index, row in enumerate(table.rows, start=1):
        values = [cell.text.strip().replace("\n", " | ") for cell in row.cells]
        print(f"T{table_index:02d}R{row_index:03d}\t" + "\t".join(values))
