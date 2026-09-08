from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


OUT_PATH = Path(
    r"C:\Users\LIKKI\Documents\ChatGPT\David 2 Probo medical\outputs\ge_cpc_parts_shipped_approval\GE_CPC_Parts_Shipped_DevDO_Approval_Note.docx"
)

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
NAVY = RGBColor(11, 37, 69)
GRAY = RGBColor(85, 85, 85)
BLACK = RGBColor(0, 0, 0)
RISK_RED = RGBColor(155, 28, 28)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.10):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def set_style_font(style, size, color=BLACK, bold=False, before=0, after=6, line=1.10):
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold
    ppr = style.paragraph_format
    ppr.space_before = Pt(before)
    ppr.space_after = Pt(after)
    ppr.line_spacing = line
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=BLACK, size=10.5, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, before=0, after=0, line=1.10)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D7DBE2", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_widths(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    new_run.append(r_pr)

    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    set_paragraph_spacing(p)
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        set_run_font(run, bold=True)
        run = p.add_run(text[len(bold_prefix):])
        set_run_font(run)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_callout(doc, title, body, fill="F4F6F9", title_color=NAVY):
    table = doc.add_table(rows=1, cols=1)
    set_table_widths(table, [9360])
    set_table_borders(table, color="D7DBE2", size="6")
    repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, before=0, after=3, line=1.10)
    run = p.add_run(title)
    set_run_font(run, size=11.5, color=title_color, bold=True)
    p = cell.add_paragraph()
    set_paragraph_spacing(p, before=0, after=0, line=1.10)
    run = p.add_run(body)
    set_run_font(run, size=10.5, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def add_key_value_rows(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_widths(table, [2040, 7320])
    set_table_borders(table)
    repeat_table_header(table.rows[0])
    for row_idx, (label, value) in enumerate(rows):
        cells = table.rows[row_idx].cells
        set_cell_shading(cells[0], "F2F4F7")
        set_cell_text(cells[0], label, bold=True, color=DARK_BLUE, size=10.5)
        set_cell_text(cells[1], value, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_matrix(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_widths(table, widths)
    set_table_borders(table)
    repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, "F2F4F7")
        set_cell_text(cell, header, bold=True, color=DARK_BLUE, size=10)
    for row_data in rows:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            align = WD_ALIGN_PARAGRAPH.CENTER if idx == 1 and len(headers) == 3 else None
            set_cell_text(row.cells[idx], value, size=10, align=align)
        set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_numbered_item(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_paragraph_spacing(p, after=4, line=1.167)
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    return p


def build_doc():
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    set_style_font(styles["Normal"], 11, BLACK, before=0, after=6, line=1.10)
    set_style_font(styles["Heading 1"], 16, BLUE, bold=True, before=16, after=8, line=1.10)
    set_style_font(styles["Heading 2"], 13, BLUE, bold=True, before=12, after=6, line=1.10)
    set_style_font(styles["Heading 3"], 12, DARK_BLUE, bold=True, before=8, after=4, line=1.10)

    doc.core_properties.title = "GE CPC Parts Shipped Report - DevDO Completion and Prod Approval Note"
    doc.core_properties.subject = "Salesforce DevDO report completion and Production approval request"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "Salesforce, DevDO, GE CPC, report, approval"

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(footer, before=0, after=0, line=1.0)
    fr = footer.add_run("GE CPC Parts Shipped Report - Approval Note")
    set_run_font(fr, size=9, color=GRAY)

    title = doc.add_paragraph()
    set_paragraph_spacing(title, before=0, after=4, line=1.10)
    tr = title.add_run("GE CPC Parts Shipped Report")
    set_run_font(tr, size=23, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    set_paragraph_spacing(subtitle, before=0, after=14, line=1.10)
    sr = subtitle.add_run("DevDO completion summary and Production approval request")
    set_run_font(sr, size=13, color=GRAY)

    add_key_value_rows(
        doc,
        [
            ("Prepared for", "David"),
            ("Status", "DevDO completed; Production not changed"),
            ("Date", "September 3, 2026"),
            ("DevDO report", "GE CPC Parts Shipped"),
            ("Folder", "Service Operations Reports"),
            ("Report type", "Work Orders with Work Order Line items2"),
        ],
    )

    add_callout(
        doc,
        "Decision Needed",
        "The report build is complete in DevDO. Approval is needed before creating, deploying, saving, or changing anything in Production.",
    )

    add_heading(doc, "What Was Done", 1)
    add_body(
        doc,
        "Created a new Salesforce report in Sandbox (DevDO) named GE CPC Parts Shipped. The report opens successfully in DevDO and is ready for business review.",
    )
    add_key_value_rows(
        doc,
        [
            ("DevDO report ID", "00OiK0000000k9JUAQ"),
            ("Report format", "Tabular"),
            ("Primary purpose", "Track parts shipped against the six GE CPC opportunities by system serial, asset, part, quantity, and cost."),
        ],
    )
    link_para = doc.add_paragraph()
    set_paragraph_spacing(link_para, before=0, after=8, line=1.10)
    set_run_font(link_para.add_run("DevDO link: "), bold=True)
    add_hyperlink(
        link_para,
        "Open GE CPC Parts Shipped report",
        "https://probomedical--devdo.sandbox.my.salesforce.com/lightning/r/Report/00OiK0000000k9JUAQ/view",
    )

    add_heading(doc, "Report Scope", 1)
    add_matrix(
        doc,
        ["Configured Item", "Value"],
        [
            ("Account filter", "GE PRECISION HEALTHCARE LLC"),
            ("Opportunity filter", "CPC - PHILIPS EPIQ"),
            ("Customer PO Number filter", "302808083, 302808091, 302808107, 302808142, 302808182, 302808186"),
            ("Ship Date filter", "Current Fiscal Year"),
            ("Primary objects", "Work Orders with Work Order Line Items"),
        ],
        [2700, 6660],
    )

    add_heading(doc, "Columns Included", 1)
    add_matrix(
        doc,
        ["Business Need", "Report Column"],
        [
            ("Work order reference", "Work Order Number"),
            ("GE CPC opportunity", "Opportunity Name"),
            ("Contract traceability", "Service Contract"),
            ("GE PO tracking", "Customer PO Number"),
            ("Shipment timing", "Ship Date"),
            ("GE system serial", "Asset Serial Number"),
            ("Asset traceability", "Asset Number"),
            ("Part identification", "Part Number"),
            ("Part description", "Product Name and Description"),
            ("Shipped quantity", "Quantity"),
            ("Cost reporting", "Part Used Cost and Total Cost of Assets Used in Repair"),
        ],
        [3000, 6360],
    )

    add_heading(doc, "Production Cross-Check", 1)
    add_body(
        doc,
        "Production was checked in read-only mode only. No Production changes were made. No Production reports, records, metadata, folders, or data values were created or changed.",
    )
    add_matrix(
        doc,
        ["Check", "Result", "Meaning"],
        [
            ("Production identity", "Confirmed", "Org is Probo Medical Production, not sandbox."),
            ("Matching GE CPC opportunities", "6", "The real GE CPC opportunities are in Production."),
            ("Matching Work Orders", "118", "The DevDO report logic should return records after approved Production move."),
            ("Matching Work Order Line Items", "142", "The report should show shipped-part rows in Production."),
            ("Report exists in Production", "No", "GE CPC Parts Shipped was not created in Production."),
        ],
        [2880, 1440, 5040],
    )

    add_heading(doc, "Data Notes", 1)
    add_body(
        doc,
        "These are Production read-only findings to review with the business before final acceptance. They do not block the DevDO report build, but they may affect what users see in the report.",
    )
    add_matrix(
        doc,
        ["Data Point", "Read-Only Finding", "Business Impact"],
        [
            ("Part Used, Part Number, Asset Number, Quantity", "142 of 142 populated", "Core shipped-part fields are available."),
            ("Part Used Cost", "119 of 142 populated", "23 line items may show blank cost."),
            ("Asset Serial Number blank", "2 Work Orders", "Some system serial values may need cleanup."),
            ("Asset Serial Number = N/A", "31 Work Orders", "Business may need to decide whether N/A is acceptable."),
        ],
        [2880, 2160, 4320],
    )

    add_heading(doc, "How To Review In DevDO", 1)
    add_numbered_item(doc, "Open Sandbox (DevDO), then go to Reports.")
    add_numbered_item(doc, "Search for GE CPC Parts Shipped.")
    add_numbered_item(doc, "Open the report from Service Operations Reports.")
    add_numbered_item(doc, "Click Edit and confirm the filters, columns, and Current Fiscal Year ship-date filter.")
    add_numbered_item(doc, "Expected DevDO result is No Results because the real GE CPC data is in Production.")

    add_heading(doc, "Suggested Approval Message", 1)
    message = (
        "Hi David,\n\n"
        "DevDO work is completed for the GE CPC Parts Shipped report.\n\n"
        "Created a new report in DevDO named GE CPC Parts Shipped under the Service Operations Reports folder. "
        "The report uses the Work Orders with Work Order Line items2 report type and includes the requested fields "
        "for Work Order, Opportunity, Service Contract, Customer PO, Ship Date, Asset Serial Number, Asset Number, "
        "Part Number, Description/Product Name, Quantity, and Cost.\n\n"
        "The report opens successfully in DevDO. DevDO shows no results because the real GE CPC records exist in Production. "
        "Production was cross-checked in read-only mode only, and no Production changes were made.\n\n"
        "Production has the six GE CPC opportunities. The same report logic should return 118 Work Orders and 142 Work Order "
        "Line Items once approved and moved to Production.\n\n"
        "Known data cleanup items in Production: 23 line items are missing Part Used Cost, 2 Work Orders have blank Asset Serial "
        "Number, and 31 Work Orders have Asset Serial Number as N/A.\n\n"
        "Please confirm approval to create or move this report in Production. No Production changes will be made until approval is received."
    )
    msg_table = doc.add_table(rows=1, cols=1)
    set_table_widths(msg_table, [9360])
    set_table_borders(msg_table, color="D7DBE2", size="6")
    repeat_table_header(msg_table.rows[0])
    msg_cell = msg_table.cell(0, 0)
    set_cell_shading(msg_cell, "F4F6F9")
    msg_cell.text = ""
    for idx, part in enumerate(message.split("\n\n")):
        p = msg_cell.paragraphs[0] if idx == 0 else msg_cell.add_paragraph()
        set_paragraph_spacing(p, before=0, after=6 if idx < len(message.split("\n\n")) - 1 else 0, line=1.10)
        r = p.add_run(part)
        set_run_font(r, size=10.5)

    doc.save(OUT_PATH)


if __name__ == "__main__":
    build_doc()
