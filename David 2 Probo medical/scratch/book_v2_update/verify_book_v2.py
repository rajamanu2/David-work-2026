import json
import openpyxl

path = r"C:\Users\LIKKI\Documents\ChatGPT\David 2 Probo medical\outputs\book_v2_direct_permission_sets\Book v2 - with Direct Permission Sets.xlsx"
wb = openpyxl.load_workbook(path, data_only=False, read_only=True)

ws = wb["Sheet1"]
ds = wb["Direct Permission Sets"]

bad_refs = []
formula_count = 0
for sheet in wb.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, str):
                if value.startswith("="):
                    formula_count += 1
                if "#REF!" in value:
                    bad_refs.append(f"{sheet.title}!{cell.coordinate}:{value}")
                    if len(bad_refs) >= 10:
                        break
        if len(bad_refs) >= 10:
            break
    if len(bad_refs) >= 10:
        break

result = {
    "sheets": wb.sheetnames,
    "sheet1_rows": ws.max_row,
    "sheet1_cols": ws.max_column,
    "direct_rows": ds.max_row,
    "direct_cols": ds.max_column,
    "sheet1_headers": [ws.cell(1, c).value for c in range(1, min(ws.max_column, 14) + 1)],
    "direct_headers": [ds.cell(1, c).value for c in range(1, ds.max_column + 1)],
    "sample_sheet1": [ws.cell(2, c).value for c in range(1, min(ws.max_column, 14) + 1)],
    "sample_direct": [ds.cell(2, c).value for c in range(1, ds.max_column + 1)],
    "formula_count": formula_count,
    "ref_errors_found": bad_refs,
}

print(json.dumps(result, default=str, indent=2))
