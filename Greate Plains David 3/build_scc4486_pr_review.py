from __future__ import annotations

import shutil
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt

from build_architect_review import (
    AMBER,
    BLUE,
    LIGHT_AMBER,
    LIGHT_BLUE,
    LIGHT_RED,
    LIGHT_TEAL,
    MUTED,
    NAVY,
    RED,
    TEAL,
    add_body,
    add_figure,
    add_heading,
    arrow,
    draw_wrapped,
    load_font,
    rounded,
    set_run_font,
)
from build_scc3387_architect_review import (
    configure_document,
    normalize_table_margins,
    review_callout,
    review_table,
)


ROOT = Path(__file__).resolve().parent
REFERENCE = ROOT / "SCC-4179-architect-review.docx"
ASSET_DIR = ROOT / "doc-assets-scc4486"
OUTPUT = ROOT / "SCC-4486-PR-Review-DevA-Verified.docx"
ASSET_DIR.mkdir(exist_ok=True)


def clear_document_body(document: Document) -> None:
    body = document._body._element
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)


def mark_data_table_headers(document: Document) -> None:
    """Mark only true comparison-table headers; leave metadata/callout tables alone."""
    header_labels = {"#", "Result", "DataPack", "Control"}
    for table in document.tables:
        if not table.rows:
            continue
        first_label = table.rows[0].cells[0].text.strip()
        if first_label not in header_labels or len(table.rows) < 2:
            continue
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        if tr_pr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader") is None:
            tr_pr.append(OxmlElement("w:tblHeader"))


def update_footer_environment(document: Document) -> None:
    """Replace the retained reference-org label in all footer variants."""
    seen = set()
    for section in document.sections:
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            if id(footer._element) in seen:
                continue
            seen.add(id(footer._element))
            paragraphs = list(footer.paragraphs)
            for table in footer.tables:
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs.extend(cell.paragraphs)
            for paragraph in paragraphs:
                for run in paragraph.runs:
                    if "GreatPlainsMerge" in run.text:
                        run.text = run.text.replace("GreatPlainsMerge", "GreatPlainsDevA")


def create_condition_diagram(path: Path) -> None:
    image = Image.new("RGB", (1600, 1000), "#F7F9FC")
    draw = ImageDraw.Draw(image)
    title = load_font(42, bold=True)
    subtitle = load_font(22)
    heading = load_font(24, bold=True)
    body = load_font(18)
    small = load_font(16, bold=True)

    draw.text((70, 50), "SCC-4486 | PR condition and dependency path", font=title, fill="#0B2545")
    draw.text(
        (70, 110),
        "DevA confirms the tenant exclusion is additive and leaves downstream branch behavior unchanged.",
        font=subtitle,
        fill="#5B6777",
    )

    boxes = [
        (
            (55, 225, 350, 610),
            "Tenant signal",
            ["Order.Account.Segment__c", "Restricted picklist", "Active value: MDU Tenant"],
            "#E7F5F1",
            "#16836B",
        ),
        (
            (430, 225, 755, 610),
            "Live condition",
            ["Existing condition retained", "AND Segment != MDU Tenant", "Exact guard appears once"],
            "#E7F5F1",
            "#16836B",
        ),
        (
            (835, 225, 1160, 610),
            "Runtime result",
            ["MDU Tenant: condition false", "Five items designed to skip", "Non-MDU gates remain intact"],
            "#E7F5F1",
            "#16836B",
        ),
        (
            (1240, 225, 1545, 610),
            "Dependency path",
            ["Skip Branch remains false", "Welcome Email / Create WOLI continue", "Capture Tech Details continues"],
            "#EAF2FA",
            "#2E74B5",
        ),
    ]

    for box, box_title, lines, fill, outline in boxes:
        rounded(draw, box, fill, outline=outline, radius=22, width=4)
        draw.text((box[0] + 24, box[1] + 24), box_title, font=heading, fill="#0B2545")
        y = box[1] + 100
        for line in lines:
            draw.ellipse((box[0] + 26, y + 7, box[0] + 38, y + 19), fill=outline)
            wrapped = draw_wrapped(
                draw,
                (box[0] + 52, y),
                line,
                body,
                "#1F2937",
                box[2] - box[0] - 78,
                spacing=6,
            )
            y += max(78, 27 * len(wrapped) + 34)

    for start, end, label, color in (
        ((350, 420), (430, 420), "READ", "#16836B"),
        ((755, 420), (835, 420), "EVALUATE", "#16836B"),
        ((1160, 420), (1240, 420), "CONTINUE", "#2E74B5"),
    ):
        arrow(draw, start, end, color=color, width=6)
        draw.text((start[0] + 3, start[1] - 42), label, font=small, fill=color)

    rounded(draw, (70, 690, 1530, 905), "#E7F5F1", outline="#16836B", radius=22, width=3)
    draw.text((105, 725), "Peer-review decision", font=heading, fill="#0B2545")
    decision = (
        "Approve with Comments. GreatPlainsDevA contains the exact MDU Tenant exclusion on all five Internet Plan "
        "items, preserves the existing fulfillment gates, and keeps Skip Branch false. The repository reviewer should "
        "still confirm the nine-file patch and four ParentKeys files contain no packaging drift."
    )
    draw_wrapped(draw, (105, 785), decision, body, "#1F2937", 1360, spacing=8)
    image.save(path)


def add_metadata_table(document: Document) -> None:
    metadata = [
        ("Review type", "PR / peer review only - no IT or UAT execution"),
        ("Scope", "5 OrchestrationItemDefinition DataPacks | 9 named files"),
        ("Source evidence", "GreatPlainsDevA | 00DEa00000Fc086MAB | Vlocity CMT 900.650.3.1"),
        ("Decision", "Approve with Comments - live configuration logic verified"),
        ("Boundary", "No PR approval/comment, deployment, activation, record write, or org metadata change"),
    ]
    table = document.add_table(rows=len(metadata), cols=2)
    table.style = "Table Grid"
    from build_architect_review import LIGHT_BLUE, set_cell_text, set_table_geometry, shade_cell

    set_table_geometry(table, [1900, 7460])
    normalize_table_margins(table)
    for row, (label, value) in zip(table.rows, metadata):
        shade_cell(row.cells[0], LIGHT_BLUE)
        set_cell_text(row.cells[0], label, size=9.5, color=NAVY, bold=True)
        set_cell_text(row.cells[1], value, size=9.5)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def build_document() -> None:
    diagram_path = ASSET_DIR / "scc4486-pr-condition-path.png"
    create_condition_diagram(diagram_path)

    shutil.copy2(REFERENCE, OUTPUT)
    document = Document(OUTPUT)
    clear_document_body(document)
    configure_document(document)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(2)
    set_run_font(paragraph.add_run("PR REVIEW"), size=10, color=BLUE, bold=True)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(
        paragraph.add_run("SCC-4486 | Skip Book Appointment & Drop Tasks"),
        size=25,
        color=NAVY,
        bold=True,
    )

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(16)
    set_run_font(
        paragraph.add_run("Great Plains | OmniStudio orchestration peer review"),
        size=13,
        color=MUTED,
    )

    add_metadata_table(document)
    review_callout(
        document,
        "PR decision",
        "Approve with Comments. GreatPlainsDevA proves that all five Internet Plan definitions use the authoritative Account Segment value, retain their existing conditions, and keep Skip Branch false. No blocking condition-logic issue was found. The actual nine-file PR patch still needs a repository-level identity and ParentKeys check before merge.",
        fill=LIGHT_TEAL,
        accent=TEAL,
    )

    add_heading(document, "Acceptance and PR outcome", 1)
    acceptance_rows = [
        ("1", "Five named items", "Pass", "Exactly five matching ManualTask definitions exist in Internet Plan; same-name CRC/FSL records were excluded from the review."),
        ("2", "MDU Tenant signal", "Pass", "Order.Account.Segment__c is a restricted Account picklist path; MDU Tenant is an active value and is populated on DevA accounts/orders."),
        ("3", "Automatic Skipped design", "Pass", "Each live ConditionData tree has one exact Segment != MDU Tenant clause under the top-level AND, making the item condition false for that segment."),
        ("4", "Existing behavior", "Pass", "Book Appointment retains appointment-status logic; the four drop tasks retain installation-method, facility-type, and product gates."),
        ("5", "Downstream continuation", "Pass", "IsSkipBranch is false on all five items and the Internet Plan dependency chain remains present."),
        ("6", "PR packaging", "Comment", "The live records verify deployed DataPack state; the actual five DataPack.json and four ParentKeys.json patches were not available for line-level review."),
    ]
    review_table(
        document,
        ["#", "Review scope", "Result", "Evidence"],
        acceptance_rows,
        [500, 2200, 1450, 5210],
        status_column=2,
    )

    add_heading(document, "Condition and dependency design", 1)
    add_body(
        document,
        "The DevA implementation preserves every existing execution condition and adds one explicit non-MDU gate. For an Account whose Segment is MDU Tenant, the additional inequality evaluates false and the item is designed to enter Skipped state while the dependency graph remains intact because Skip Branch is false.",
    )
    add_figure(
        document,
        diagram_path,
        "Figure 1. SCC-4486 PR condition and dependency path",
        "Diagram showing the verified Account Segment discriminator, additive condition, expected skipped state, and downstream dependency continuation.",
        width=6.45,
    )
    review_callout(
        document,
        "Verified logic",
        "Observed on all five items: ExistingCondition AND (Order.Account.Segment__c != 'MDU Tenant'). Each guard uses picklist / field-OrderItem tokens, appears exactly once, and is the final top-level AND clause.",
        fill=LIGHT_TEAL,
        accent=TEAL,
    )

    document.add_page_break()
    add_heading(document, "PR findings and reviewer comments", 1)
    findings_rows = [
        ("Pass", "Authoritative discriminator", "The implementation correctly uses Order.Account.Segment__c with the active restricted-picklist value MDU Tenant; it does not depend on nullable MDU identifier strings."),
        ("Pass", "Condition direction", "The != operator makes the full top-level AND condition false for MDU Tenant, which is the correct direction for an automatic Skipped state."),
        ("Pass", "Existing-condition preservation", "The original appointment or drop gates remain present. No live evidence shows that the SCC-4486 change replaced the existing boolean tree."),
        ("Pass", "Dependency safety", "IsSkipBranch remains false. Book Appointment still gates Triage, Welcome Email, and Create WOLI; the drop chain continues through Capture Tech Details."),
        ("Comment", "Repository diff and ParentKeys", "Before merge, confirm the PR is limited to five DataPack.json and four ParentKeys.json files, keeps stable keys, and contains no environment-specific record IDs."),
        ("Comment", "Blank Segment behavior", "Account.Segment__c is nullable. The PR logic is appropriate, but the normal-order expectation for blank Segment should remain covered by the team's separate test process."),
        ("Comment", "Product relationship paths", "Install Drop uses Product2.ProductCode; the other three drop definitions use vlocity_cmt__Product2Id__r.ProductCode. Preserve this pre-existing distinction unless separately justified."),
        ("Pass", "No code expansion", "No Apex, Flow, trigger, API, field, or permission change is required for this story. Reject unrelated automation or schema changes in the PR."),
    ]
    review_table(
        document,
        ["Result", "Finding", "Reviewer evidence / action"],
        findings_rows,
        [1300, 2700, 5360],
        status_column=0,
    )
    review_callout(
        document,
        "Reviewer disposition",
        "No blocking code or live configuration defect was found in GreatPlainsDevA. Approve with Comments, subject to a final repository check that the nine declared files faithfully represent this verified state.",
        fill=LIGHT_TEAL,
        accent=TEAL,
    )

    document.add_page_break()
    add_heading(document, "Five-DataPack live evidence", 1)
    item_rows = [
        ("Book Appointment", "Appointment status + Segment != MDU Tenant", "Pass | Skip Branch false | 86d713dd-cd8a-7c69-091f-d3d2f4f5f4a9"),
        ("Triage, Prep-Work & Sitewalk", "Install method + facility + product + Segment guard", "Pass | Skip Branch false | 01c083d3-0a02-24ec-f2a4-1e4a1fee347e"),
        ("Notify 811", "Install method + facility + product + Segment guard", "Pass | Skip Branch false | 274e769c-e9a9-ec1f-d782-d4fa8ed280f4"),
        ("Install Drop", "Install method + facility + Product2.ProductCode + Segment guard", "Pass | Skip Branch false | ce652191-615d-84a5-22aa-6ba800141153"),
        ("Splice and Cutover Drop", "Install method + facility + product + Segment guard", "Pass | Skip Branch false | 409ef12f-c6b1-06b7-a79e-655e04fad8da"),
    ]
    review_table(
        document,
        ["DataPack", "Observed DevA condition", "Result / GlobalKey"],
        item_rows,
        [2500, 3100, 3760],
    )

    add_heading(document, "DataPack and ParentKeys controls", 2)
    control_rows = [
        ("Definition identity", "All five records are Internet Plan / ManualTask definitions using the orderItems node. The deprecated Condition field is null."),
        ("ConditionData", "All five use a top-level AND and contain exactly one MDU guard with picklist, field-OrderItem, !=, and MDU Tenant tokens."),
        ("Dependencies", "Internet Plan links remain: Book -> Triage / Welcome Email / Create WOLI; Triage -> Notify 811 -> Install Drop -> Splice -> Capture Tech Details."),
        ("ParentKeys", "Four ParentKeys source files are declared by the ticket but are not separately represented as Salesforce records. Verify them in the actual PR patch."),
        ("Change history", "All five Internet Plan definitions were modified by Naveen Kumar on 25 August 2026; same-name CRC/FSL definitions are separate and unchanged by this review."),
        ("Scope discipline", "No Apex, Flow, trigger, schema, permission, implementation reference, or Skip Branch change is needed for SCC-4486."),
    ]
    review_table(document, ["Control", "Peer-review expectation"], control_rows, [2300, 7060])
    review_callout(
        document,
        "PR-only boundary",
        "This is a PR/configuration review only. Read-only DevA schema and record evidence was inspected, including existing MDU Tenant orders. No IT test, UAT execution, orchestration run, deployment validation, activation, or org write is claimed.",
        fill=LIGHT_AMBER,
        accent=AMBER,
    )

    document.add_page_break()
    add_heading(document, "Reviewer completion checks", 1)
    gate_rows = [
        ("1", "Five live definitions", "Pass", "Correct Internet Plan records, exact MDU guard once per item, original gates retained."),
        ("2", "Tenant discriminator", "Pass", "Order.Account.Segment__c; restricted picklist; active exact value MDU Tenant."),
        ("3", "Branch and dependency safety", "Pass", "IsSkipBranch=false on all five and the expected Internet Plan dependency records remain."),
        ("4", "Nine-file PR scope", "Review comment", "Confirm exactly five DataPack.json plus four ParentKeys.json files and no unrelated regenerated JSON."),
        ("5", "Key and ID integrity", "Review comment", "Match the five GlobalKeys above; ParentKeys must use stable keys and exported files must not introduce DevA record IDs."),
        ("6", "Final disposition", "Approve with Comments", "No blocking live logic issue; complete the repository-only checks before merge."),
    ]
    review_table(
        document,
        ["#", "Review check", "Status", "Evidence / final action"],
        gate_rows,
        [500, 2300, 2200, 4360],
    )

    add_heading(document, "Ready-to-paste PR review response", 2)
    response = (
        "PR review completed for SCC-4486 against GreatPlainsDevA (00DEa00000Fc086MAB) - Approve with Comments. "
        "All five Internet Plan ManualTask definitions contain the exact top-level guard Order.Account.Segment__c != "
        "'MDU Tenant'. Book Appointment retains its appointment-status condition, the four drop tasks retain their "
        "installation, facility, and product conditions, and IsSkipBranch remains false on every item. The expected "
        "Internet Plan dependencies are also present, so no blocking code or live configuration issue was found. Before "
        "merge, confirm the PR contains only the five DataPack.json and four ParentKeys.json files, matches the verified "
        "GlobalKeys, preserves stable ParentKeys, and introduces no DevA record IDs or unrelated JSON churn. No Apex, "
        "Flow, trigger, schema, or permission change is required. This was PR review only: no IT/UAT execution, deployment, "
        "activation, orchestration run, record write, or metadata change was performed."
    )
    review_callout(document, "PR comment", response, fill=LIGHT_BLUE, accent=BLUE)

    document.core_properties.title = "SCC-4486 Skip Book Appointment and Drop Tasks PR Review"
    document.core_properties.subject = "OmniStudio orchestration peer review"
    document.core_properties.author = "Peer Review"
    document.core_properties.keywords = (
        "Salesforce, Great Plains, SCC-4486, OmniStudio, orchestration, DataPack, PR review"
    )
    mark_data_table_headers(document)
    update_footer_environment(document)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
