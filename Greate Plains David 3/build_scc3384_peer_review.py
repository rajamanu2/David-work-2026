from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_architect_review import (
    BLUE,
    GRID,
    INK,
    LIGHT_BLUE,
    LIGHT_GRAY,
    MUTED,
    NAVY,
    add_body,
    add_heading,
    add_page_field,
    add_table,
    rgb,
    set_cell_margins,
    set_cell_text,
    set_run_font,
    set_table_geometry,
    shade_cell,
)


ROOT = Path(__file__).resolve().parent
REFERENCE = ROOT / "SCC-3387-architect-review.docx"
OUTPUT = ROOT / "SCC-3384-peer-review.docx"


def clear_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_light_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9D9D9")


def normalize_table(table) -> None:
    set_light_borders(table)
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=85, start=120, bottom=85, end=120)


def review_table(document, headers, rows, widths, status_column=None):
    table = add_table(document, headers, rows, widths, status_column=status_column)
    normalize_table(table)
    return table


def configure_footer(section) -> None:
    for header in (section.header, section.even_page_header, section.first_page_header):
        paragraph = header.paragraphs[0]
        paragraph.text = ""
        paragraph.paragraph_format.space_after = Pt(0)
    for footer in (section.footer, section.even_page_footer, section.first_page_footer):
        paragraph = footer.paragraphs[0]
        paragraph.text = ""
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph.paragraph_format.space_before = Pt(0)
        run = paragraph.add_run("GreatPlainsDevA  |  Page ")
        set_run_font(run, size=9, color=MUTED)
        add_page_field(paragraph)


def configure_document(document: Document) -> None:
    document.settings.odd_and_even_pages_header_footer = False
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        section.different_first_page_header_footer = False
        configure_footer(section)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = document.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = rgb("000000")
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.keep_with_next = True
    p_pr = title._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)

    for name, size, before, after in (
        ("Heading 1", 16, 12, 6),
        ("Heading 2", 13, 10, 5),
        ("Heading 3", 12, 8, 4),
    ):
        style = document.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb("000000")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_decision(document: Document, label: str, message: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.10
    lead = paragraph.add_run(f"{label}  ")
    set_run_font(lead, size=11, color=NAVY, bold=True)
    body = paragraph.add_run(message)
    set_run_font(body, size=11, color=INK)


def add_metadata(document: Document) -> None:
    rows = [
        ("Environment", "GreatPlainsDevA sandbox | Great Plains Communications | Org 00DEa00000Fc086MAB"),
        ("Review date", "5 September 2026"),
        ("Retrieval", "Succeeded | Jobs 09SEa00000jCLrhMAG, 09SEa00000jCwkfMAC, 09SEa00000jCacFMAS"),
        ("Decision", "Changes Requested - not ready for promotion"),
        ("Boundary", "No records, deployment, activation, or org metadata were changed"),
    ]
    table = document.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [1900, 7460])
    normalize_table(table)
    for row, (label, value) in zip(table.rows, rows):
        shade_cell(row.cells[0], LIGHT_BLUE)
        set_cell_text(row.cells[0], label, size=9.5, color=NAVY, bold=True)
        set_cell_text(row.cells[1], value, size=9.3)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def build_document() -> None:
    document = Document(REFERENCE)
    clear_body(document)
    configure_document(document)

    eyebrow = document.add_paragraph()
    eyebrow.paragraph_format.space_before = Pt(8)
    eyebrow.paragraph_format.space_after = Pt(2)
    set_run_font(eyebrow.add_run("PEER REVIEW"), size=10, color="000000", bold=True)

    title = document.add_paragraph(style="Title")
    title.add_run("SCC 3384 Case Trouble Ticket Peer Review")

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    set_run_font(
        subtitle.add_run("Great Plains DevA Sandbox Read only implementation assessment"),
        size=13,
        color="000000",
    )

    add_metadata(document)
    add_decision(
        document,
        "Decision",
        "Do not promote yet. The Trouble Ticket record type, layout, core fields, and pending-customer timing logic exist, but the runtime picklist controls, status semantics, timestamp coverage, and NOC access model do not satisfy the ticket.",
    )

    add_heading(document, "Acceptance outcome", 1)
    acceptance_rows = [
        (
            "1",
            "Record type and layout",
            "Pass",
            "Trouble_Ticket is active. The administrator profile assigns Case-Trouble Ticket, and the layout contains the requested Case Information, Critical Dates, and Resolution Details fields.",
        ),
        (
            "2",
            "Status process",
            "Fail",
            "The reused Billing business process contains the eight requested values plus Applied, Approved, Pending, and Rejected. Complete, Canceled, and Duplicate are not closed statuses.",
        ),
        (
            "3",
            "Status timestamping",
            "Fail",
            "The active date flow still watches legacy In-Progress and Completed values. Pending-customer entry and exit are timestamped, but the full requested progression is not captured.",
        ),
        (
            "4",
            "Service and subscriber control",
            "Fail",
            "Subscriber Report has 72 values but no Service Type controller. Fixed Wireless exists globally but is not enabled for the Trouble Ticket record type.",
        ),
        (
            "5",
            "Resolution controls",
            "Fail",
            "Primary Resolution has 60 dependency mappings and Secondary Resolution is multi-select, but the record type enables only Cleared Without Repair for each field.",
        ),
        (
            "6",
            "NOC access",
            "Fail",
            "No active NOC-named profile, role, or user was found. The NOC permission set has one System Administrator assignee and omits most SCC-3384 field permissions.",
        ),
    ]
    review_table(document, ["#", "Scope", "Result", "Evidence"], acceptance_rows, [500, 2050, 1350, 5460], status_column=2)

    document.add_page_break()
    add_heading(document, "Field and dependency configuration", 1)
    field_rows = [
        ("Trouble Ticket record type", "Active", "Case-Trouble Ticket", "Pass - administrator assignment is configured"),
        ("Service Type", "Restricted picklist", "8 global; 7 on record type", "Fail - Fixed Wireless is omitted"),
        ("Service Address", "Compound address", "Edit on layout", "Pass - field and UI placement exist"),
        ("Custom Priority", "Restricted picklist", "Low, Med, High", "Pass - requested values are enabled"),
        ("Subscriber Report", "Picklist", "72 values; no controller", "Fail - not dependent on Service Type"),
        ("Primary Resolution", "Dependent picklist", "60 mappings from Subscriber Report", "Fail - record type enables one value"),
        ("Secondary Resolution", "Multi-select picklist", "Shared 64-value set", "Fail - record type enables one value"),
        ("Resolution Detail", "Long Text Area 32768", "Edit on layout", "Pass - field and UI placement exist"),
        ("Critical Dates", "Three custom Date fields", "Created and Closed system dates included", "Pass - section contains required dates"),
        ("Parent Case", "Standard Case lookup", "Edit on layout; two validations", "Pass - circular and child-type controls are active"),
    ]
    review_table(
        document,
        ["Requirement", "Configured type", "Observed setup", "Peer review reading"],
        field_rows,
        [2250, 1850, 2600, 2660],
        status_column=3,
    )
    add_body(
        document,
        "The Primary Resolution metadata maps every Subscriber Report value to at least 16 resolutions, but four values in the shared Case Resolution set are not mapped: CLEAR- TRIP, CUST CAUSED- NO CHG, CUST CAUSED- TRIP CHG, and Fiber Splicing Completed. Several duplicate capitalization and spacing variants should be reconciled before wider enablement.",
        bold_lead="The Primary Resolution metadata",
        after=4,
    )

    document.add_page_break()
    add_heading(document, "Status automation and existing data", 1)
    status_rows = [
        ("Requested status values", "Partial pass", "All eight values exist, but the Trouble Ticket record type reuses the Billing process and exposes four unrelated billing values."),
        ("Terminal behavior", "Fail", "Only Closed has IsClosed true. Complete, Canceled, and Duplicate do not set Closed Date or close the Case."),
        ("In progress and complete dates", "Fail", "Case Date Update based on status changes is active/latest but checks In-Progress and Completed, not In Progress and Complete."),
        ("Pending customer timing", "Partial pass", "The active/latest before-save Case flow stamps entry and exit. Re-entry replaces the prior start/end pair and the logic is not scoped to Trouble Ticket."),
        ("Full status progression", "Fail", "No timestamp fields or transition logic were found for Not Started, Assigned, On Hold GPC, Canceled, or Duplicate."),
    ]
    review_table(document, ["Control", "Result", "Evidence"], status_rows, [2450, 1550, 5360], status_column=1)

    add_heading(document, "Aggregate record evidence", 2)
    data_rows = [
        ("Trouble Ticket cases", "46", "Existing population used only as corroborating evidence"),
        ("Service Type and custom Priority", "21 and 12", "Core classification is only partially populated"),
        ("Subscriber Report and Primary Resolution", "12 and 0", "No reviewed case has a Primary Resolution"),
        ("Problem Occurred, Problem Reported, Repair", "1, 1, and 0", "Critical dates are rarely populated"),
        ("Pending Customer start and end", "9 and 1", "Eight records have no paired end timestamp"),
        ("In Progress and Completed dates", "0 and 1", "The active legacy-value flow does not evidence the requested transitions"),
        ("Complete status", "1 case; 0 Closed Dates", "Complete currently behaves as an open status"),
    ]
    review_table(document, ["Metric", "Observed count", "Interpretation"], data_rows, [3100, 1700, 4560])
    add_body(
        document,
        "Existing records may predate the current configuration, so these counts corroborate the metadata gaps but do not replace named-persona UI testing.",
        italic=True,
        color=MUTED,
        after=2,
    )

    document.add_page_break()
    add_heading(document, "Access and component evidence", 1)
    access_rows = [
        ("System Administrator profile", "Pass", "Trouble Ticket is visible/default, the dedicated layout is assigned, and all reviewed SCC-3384 fields are editable."),
        ("SCC 3380 NOC permission set", "Fail", "Case create/read/edit and Trouble Ticket visibility exist, but only Parent Case, custom Priority, and Related Outage Case field permissions are included."),
        ("Permission set assignment", "Fail", "One active assignee exists, but the user is a System Administrator with no role; this does not prove a least-privilege NOC persona."),
        ("Named NOC persona", "Not proven", "The targeted active-user query returned no NOC match by name, username, profile, or role."),
    ]
    review_table(document, ["Access path", "Result", "Evidence"], access_rows, [2550, 1550, 5260], status_column=1)

    add_heading(document, "Component and evidence boundary", 2)
    component_rows = [
        ("Focused retrieval", "Case object container, Trouble Ticket and default layouts, Case Record Page, Case Resolution value set, three Case flows, NOC permission set, and administrator profile were retrieved successfully."),
        ("Active version check", "Case Before Insert Update Trigger Flow, Case After Trigger Flow, and Case Date Update based on status changes are each active/latest."),
        ("Ticket documentation gap", "The ticket lists New or Updated Objects and Fields, Triggers or Automations, Components, and Permissions as None even though those components implement the story."),
        ("Workbook boundary", "The linked SharePoint workbooks were not available in this session, so exact workbook-to-org matrix parity remains unverified."),
        ("Validation boundary", "No check-only deployment was run because this review inspected current DevA configuration and introduced no change package."),
    ]
    review_table(document, ["Boundary", "Evidence"], component_rows, [2200, 7160])
    add_decision(
        document,
        "Peer review reading",
        "The administrator path demonstrates that the components exist, but the business NOC path and runtime dependency behavior are not ready for approval.",
    )

    document.add_page_break()
    add_heading(document, "Required approval gates", 1)
    gate_rows = [
        ("1", "Configure Subscriber Report dependency", "Admin and Product Owner", "Make Service Type the controller and verify the approved reports for all eight service types, including Fixed Wireless."),
        ("2", "Enable resolution values", "Admin and Product Owner", "Enable the approved Primary and Secondary values for Trouble Ticket; remove or reconcile unmapped duplicate global values."),
        ("3", "Correct status design", "Architect and Product Owner", "Use a Trouble Ticket-specific process or approve the shared process. Define whether Complete, Canceled, and Duplicate must close the Case."),
        ("4", "Complete timestamp automation", "Admin and Developer", "Use exact status values, scope logic to Trouble Ticket, cover every required transition, and define repeated pending-customer intervals."),
        ("5", "Establish NOC least privilege", "Security Owner", "Grant Case CRUD, Trouble Ticket visibility, layout access, and read/edit access to all SCC-3384 fields through a named NOC persona."),
        ("6", "Correct ticket inventory", "Release Manager", "List the record type, business process, layout, fields, value set, flows, validation rules, permission changes, and Lightning page dependency."),
        ("7", "Run named persona UAT", "QA and NOC", "Create a Trouble Ticket, test all dependencies, select multiple secondary resolutions, cycle statuses, repeat pending-customer hold, and verify closure behavior."),
        ("8", "Review existing records", "Data Owner", "Decide whether the 46 existing Trouble Tickets require classification, resolution, date, or status cleanup after the configuration is corrected."),
    ]
    review_table(document, ["#", "Required action", "Owner", "Evidence needed"], gate_rows, [550, 2600, 2100, 4110])

    add_heading(document, "Ready to paste peer review response", 1)
    response = (
        "Peer review completed in GreatPlainsDevA. Changes Requested. The Trouble Ticket record type is active, the dedicated layout is assigned for the administrator profile, and the requested fields and Critical Dates and Resolution Details sections are present. However, Subscriber Report is not dependent on Service Type; Fixed Wireless is not enabled on the record type; Primary and Secondary Resolution each expose only Cleared Without Repair; the active status-date flow still watches legacy In-Progress and Completed values; Complete, Canceled, and Duplicate are not closed statuses; and the NOC permission set does not provide the complete SCC-3384 field-access path to a named NOC persona. Update the component inventory, complete the eight approval gates, and rerun the workflow as a named NOC user before promotion. No records, metadata, activation, deployment, or production changes were made."
    )
    add_body(document, response, after=4)

    document.core_properties.title = "SCC 3384 Case Trouble Ticket Peer Review"
    document.core_properties.subject = "Read only implementation assessment in GreatPlainsDevA"
    document.core_properties.author = "Peer Review"
    document.core_properties.keywords = "Salesforce, Great Plains, SCC-3384, Case, Trouble Ticket, peer review"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
