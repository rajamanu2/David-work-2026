from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from pathlib import Path

out=Path('outputs/business_templates_2026-09-08/Probo_Medical_Business_Requirements_Presentation_Template.docx')
d=Document()
for root in [d.styles.element,d.element]:
 for el in root.xpath('.//w:pBdr'): el.getparent().remove(el)
s=d.sections[0]; s.page_width=Inches(8.5); s.page_height=Inches(11)
s.left_margin=s.right_margin=Inches(.8); s.top_margin=Inches(.7); s.bottom_margin=Inches(.65)
s.header_distance=s.footer_distance=Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=d.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.space_after=Pt(8)
d.styles['Normal'].paragraph_format.line_spacing=1.08
d.styles['Title'].font.size=Pt(36)
d.styles['Title'].paragraph_format.space_after=Pt(18)
d.styles['Heading 1'].font.size=Pt(25)
d.styles['Heading 1'].paragraph_format.space_after=Pt(12)
d.styles['Heading 2'].font.size=Pt(13)
d.styles['Heading 2'].paragraph_format.space_before=Pt(13)
d.styles['Heading 2'].paragraph_format.space_after=Pt(5)
s.different_first_page_header_footer=True
s.header.paragraphs[0].text='PROBO MEDICAL                                      BUSINESS REQUIREMENTS'
s.header.paragraphs[0].style='Caption'
for sectionfooter in [s.footer,s.first_page_footer]:
 f=sectionfooter.paragraphs[0]; f.add_run('Probo Medical  •  Business discussion template').font.size=Pt(9)
 f.add_run('                                             ')
 x=OxmlElement('w:fldSimple');x.set(qn('w:instr'),'PAGE');f._p.append(x)
def p(text):d.add_paragraph(text)
def h(text):d.add_heading(text,2)
def prompt(label,hint,space=10):
 h(label)
 z=d.add_paragraph(); z.paragraph_format.space_after=Pt(space)
 r=z.add_run('['+hint+']');r.italic=True;r.font.color.rgb=RGBColor.from_string('526477')
def page(n,title,intro):
 d.add_page_break(); z=d.add_paragraph('0'+str(n)+'  /  BUSINESS DISCUSSION');z.paragraph_format.space_after=Pt(7)
 z.runs[0].font.size=Pt(9);z.runs[0].bold=True
 d.add_heading(title,1);p(intro)
def table(heads,rows,widths):
 t=d.add_table(rows=1,cols=len(heads));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for c,v in zip(t.rows[0].cells,heads):c.text=v
 for vals in rows:
  for c,v in zip(t.add_row().cells,vals):c.text=v
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr(); x=OxmlElement('w:cantSplit');pr.append(x)
  if ri==0:pr.append(OxmlElement('w:tblHeader'))
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tc=c._tc.get_or_add_tcPr(); shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'203B53' if ri==0 else ('F0F4F7' if ri%2 else 'FFFFFF'));tc.append(shade)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   tc.append(borders);m=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'100');el.set(qn('w:type'),'dxa');m.append(el)
   tc.append(m)
   for z in c.paragraphs:
    z.paragraph_format.space_after=Pt(3);z.paragraph_format.space_before=Pt(3)
    for r in z.runs:
     r.font.size=Pt(10.5)
     if ri==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
 d.add_paragraph().paragraph_format.space_after=Pt(0)

p('PROBO MEDICAL')
d.add_paragraph('Business Requirements\nDiscussion and Approval', 'Title')
p('A reusable template for business stakeholders and leadership')
prompt('Initiative','Enter the business initiative or process name',18)
p('Prepared for  [Business owner and stakeholders]\nPrepared by  [Name and role]\nMeeting date  [Date]     Version  [Number]     Status  [Draft / In review / Approved]')
h('Purpose of the discussion')
p('Agree the business need, the outcome that matters and the requirements the team must meet. Use the completed document to confirm scope, assign ownership and request approval to proceed.')
h('Discussion route')
table(['Section','Business conversation','Page'],[['01','What needs to improve and why','2'],['02','Who is affected and how work happens today','3'],['03','What is included and how work should change','4'],['04','What the business needs the solution to do','5'],['05','How we will confirm it works','6'],['06','What needs a decision or approval','7']],[.65,5.55,.7])
p('Replace the bracketed prompts during preparation or the meeting. Keep open points visible with an owner and due date. All examples are illustrative and require business agreement.')
page(1,'The business need','Start with the business outcome. This page becomes the summary for leadership once the discussion is complete.')
prompt('What problem are we solving','Describe the issue, who experiences it and when it occurs',18)
prompt('Why does this matter now','Explain the impact on customers, teams, turnaround time, quality or cost',18)
prompt('What should be different','Describe the desired outcome in plain business language',18)
h('How success will be measured')
table(['Measure','Current position','Target and date'],[['[Measure 1]','[Baseline and evidence source]','[Agreed target and date]'],['[Measure 2]','[Baseline and evidence source]','[Agreed target and date]']],[1.5,2.6,2.8])
prompt('Decision needed from leadership','State the approval, priority or resource decision and the date needed')
p('Benefit owner  [Name]     Business sponsor  [Name]')
page(2,'People and the current process','Invite the people who perform the work and those who receive its output. Capture how the process operates in practice.')
table(['Stakeholder','Role in the discussion','Name'],[['Business owner','Owns the outcome and priorities','[Name]'],['Process users','Explain tasks and exceptions','[Names]'],['System or data owner','Confirms feasibility and data needs','[Name]'],['Sponsor or approver','Resolves trade-offs and approves scope','[Name]']],[1.6,3.8,1.5])
prompt('Walk through the work today','What starts the work? What are the main steps, systems and handoffs?',20)
prompt('Where does the process become difficult','Identify delays, repeat work, missing information and common exceptions',20)
prompt('What evidence supports the issue','Reference examples, volumes, timings, reports or stakeholder feedback')
h('Prompts for the conversation')
p('Who receives the request? What must be known before work starts? Who makes decisions? What happens when information is missing? How is completion communicated?')
p('For relevant Probo Medical initiatives, consider service intake, equipment identification, evaluation, repair approvals, parts, returns, shipment handoffs or reporting. Confirm which areas apply.')
page(3,'Scope and future process','Agree the boundaries before discussing detailed functionality. Identify the business changes as well as the system changes.')
prompt('Included in this initiative','List the processes, teams, sites, data and channels covered',14)
prompt('Outside this initiative','List exclusions and any work deferred to a later phase',14)
h('The future way of working')
table(['Stage','Responsible role','Required outcome'],[['Receive and validate','[Role]','[Valid inputs and an accepted request]'],['Perform and review','[Role]','[Completed work and required approval]'],['Complete and hand over','[Role]','[Confirmed output and recipient notified]']],[1.6,1.5,3.8])
prompt('Rules and exceptions','What must be checked, who may approve, and when should work stop or escalate?')
prompt('Constraints and dependencies','Record deadlines, budget, resources, other teams and assumptions to confirm')
p('Process owner  [Name]     Supporting process map or work instruction  [Reference]')
page(4,'Business requirement record','Use one record for each distinct requirement. Duplicate this page as needed and give every requirement a stable reference.')
p('Reference  [BR-001]     Owner  [Name]     Priority  [Must / Should / Could / Deferred]')
prompt('Requirement title','Use a short name that business users recognize')
prompt('The business need','As a [role], I need [capability], so that [business outcome]',14)
prompt('Rules data and access','Specify required information, allowed roles, decisions and exception behavior',14)
prompt('Acceptance criteria','Given [starting condition], when [action], then [observable result]',14)
prompt('Source and dependencies','Who requested this? What evidence or other requirements does it depend on?')
h('Illustrative example')
p('A service coordinator needs missing request information identified before handoff so that the receiving team can act on the request. If an agreed required input is absent, submission identifies the missing item and prevents the request from advancing. The business owner defines the required inputs and permitted exceptions.')
p('Review status  [Proposed / Reviewed / Approved / Deferred]\nDecision and date  [Record the agreement or remaining question]')
page(5,'Validation and readiness','Agree the evidence that will demonstrate each requirement is met. Include exceptions and access checks as well as the normal process.')
table(['Requirement','Business test and expected result','Owner'],[['[BR-001]','[Normal scenario and measurable expected result]','[Name]'],['[BR-001]','[Missing input or rejected approval and expected result]','[Name]'],['[BR-002]','[Access restriction or failed handoff and expected result]','[Name]']],[1.05,4.75,1.1])
prompt('Test arrangements','Confirm the environment, representative users, test data, dates and evidence location')
prompt('Acceptance threshold','State which tests must pass, how defects are handled and who accepts the outcome')
h('Business readiness')
table(['Area','What must be ready','Owner and date'],[['Work instructions','[Verified steps, controls and exception routes]','[Name and date]'],['Training and communication','[Audience, materials and delivery approach]','[Name and date]'],['Support and measurement','[Support route and benefit review plan]','[Name and date]']],[1.5,3.8,1.6])
p('Coverage review  [Confirm applicable needs for reporting, data quality, access, integrations, performance, accessibility, audit evidence and retention ownership.]')
page(6,'Decisions and approval','Close the discussion with clear actions and an explicit decision. Record conditions and unresolved items before establishing the requirements baseline.')
h('Open questions actions and risks')
table(['Reference','Question action or risk','Owner and due date'],[['[Q-01]','[Open question and affected requirement]','[Name and date]'],['[A-01]','[Action needed before approval]','[Name and date]'],['[R-01]','[Risk, impact and mitigation]','[Name and date]']],[.9,4.25,1.75])
prompt('Recommendation and decision','Proceed / Rework / Defer; explain the reason and any conditions')
h('Business sign off')
table(['Role','Decision and conditions','Name and date'],[['Business owner','[Approve / Rework / Defer]','[Name and date]'],['Sponsor','[Scope and priority decision]','[Name and date]'],['Delivery or system owner','[Feasibility and dependencies reviewed]','[Name and date]']],[1.6,3.55,1.75])
h('Managing changes after approval')
p('Record each change with its reason, affected requirement IDs, impact on cost and dates, decision maker and approval date. Update the related acceptance criteria, process map and work instructions, then circulate the revised version.')
p('Approved version  [Number]     Baseline date  [Date]\nNext review  [Date]     Document owner  [Name]')
d.core_properties.title='Probo Medical Business Requirements Discussion and Approval'
d.core_properties.author='Probo Medical'
d.core_properties.subject='Editable business stakeholder requirements template'
d.save(out)
print(out.resolve())
