from pathlib import Path
from docx import Document


def extract(path: Path) -> str:
    doc = Document(path)
    lines: list[str] = []
    lines.append(f"FILE: {path}")
    lines.append(f"PARAGRAPHS: {len(doc.paragraphs)}")
    for idx, paragraph in enumerate(doc.paragraphs, start=1):
        text = paragraph.text.strip()
        if text:
            lines.append(f"P{idx} [{paragraph.style.name}]: {text}")
    lines.append(f"TABLES: {len(doc.tables)}")
    for t_idx, table in enumerate(doc.tables, start=1):
        lines.append(f"TABLE {t_idx} rows={len(table.rows)} cols={len(table.columns)}")
        for r_idx, row in enumerate(table.rows, start=1):
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            lines.append(f"T{t_idx}R{r_idx}: " + " | ".join(cells))
    rel_images = [r for r in doc.part.rels.values() if "image" in r.reltype]
    lines.append(f"IMAGES: {len(rel_images)}")
    for idx, rel in enumerate(rel_images, start=1):
        part = rel.target_part
        lines.append(f"IMAGE {idx}: {part.partname} bytes={len(part.blob)} content_type={part.content_type}")
    return "\n".join(lines)


if __name__ == "__main__":
    import sys

    for source in map(Path, sys.argv[1:]):
        print(extract(source))
        print("\n" + "=" * 100 + "\n")
