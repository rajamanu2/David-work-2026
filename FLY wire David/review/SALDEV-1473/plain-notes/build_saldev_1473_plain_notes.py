from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


OUTPUT = Path(r"C:\Users\LIKKI\Documents\ChatGPT\david\review\SALDEV-1473\SALDEV-1473_Plain_Implementation_Notes.docx")


def set_font(run, name="Arial", size=11, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def configure_style(style, size=11, bold=False, before=0, after=6, keep=False):
    style.font.name = "Arial"
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = 1.0
    style.paragraph_format.keep_with_next = keep


def add_plain_paragraph(doc, text="", bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        set_font(p.add_run(bold_prefix), bold=True)
        set_font(p.add_run(text[len(bold_prefix):]), italic=italic)
    else:
        set_font(p.add_run(text), italic=italic)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    set_font(p.add_run(text))
    return p


def new_numbering_instance(doc):
    numbering = doc.part.numbering_part.element
    style_num_id = doc.styles["List Number"]._element.pPr.numPr.numId.val
    source_num = numbering.xpath(f'./w:num[@w:numId="{style_num_id}"]')[0]
    abstract_num_id = source_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    existing_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(existing_ids) + 1

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_num_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return num_id


def add_step(doc, text, num_id):
    p = doc.add_paragraph(style="List Number")
    num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = num_id
    set_font(p.add_run(text))
    return p


def add_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    set_font(p.add_run(text), size=11, bold=True)
    return p


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[index]
            cell.width = Twips(width)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_deployment_table(doc):
    headers = [
        "Component type",
        "API name",
        "What was deployed in FlywirePartial",
        "Release-team instruction",
    ]
    rows = [
        [
            "Flow",
            "QuoteLine_After_Record_Triggered_Flow",
            "Active source version 8. Runs CreateAndUpdate after save. Adds the quote-wide different Ship-To lookup and deterministic Yes/No routing while preserving the existing exception paths. Check-only: 0AfhG000001bMygSAE. Sandbox deployment: 0AfhG000001bPN5SAM.",
            "Deploy this Flow metadata member to the next environment. Confirm the resulting Flow version is active. Run the controlled QLE Quick Save test and verify Multiple Ship to Accounts = Yes on the Quote.",
        ],
        [
            "Not included",
            "None",
            "No Apex class, Custom Field, Permission Set, Production data change, or bulk data backfill was included in SALDEV-1473.",
            "Do not add other metadata to the story deployment package unless the release team identifies an environment-specific dependency.",
        ],
    ]

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        set_font(paragraph.add_run(text), size=8.5, bold=True)

    header_properties = table.rows[0]._tr.get_or_add_trPr()
    header_marker = OxmlElement("w:tblHeader")
    header_marker.set(qn("w:val"), "true")
    header_properties.append(header_marker)

    for row_data in rows:
        row = table.add_row()
        for index, text in enumerate(row_data):
            cell = row.cells[index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            set_font(paragraph.add_run(text), size=8.5)

    set_table_geometry(table, [1200, 2500, 2900, 2760])
    return table


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    configure_style(doc.styles["Normal"], size=11, after=6)
    configure_style(doc.styles["Heading 1"], size=11, bold=True, before=10, after=4, keep=True)
    configure_style(doc.styles["List Bullet"], size=11, after=3)
    configure_style(doc.styles["List Bullet 2"], size=11, after=3)
    configure_style(doc.styles["List Number"], size=11, after=4)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(4)
    set_font(title.add_run("SALDEV-1473 - Implementation and Test Notes"), size=14, bold=True)

    add_plain_paragraph(doc, "Org: FlywirePartial sandbox")
    add_plain_paragraph(doc, "Flow: QuoteLine After Record Triggered Flow")
    add_plain_paragraph(doc, "Test quote: Q-37873")
    add_plain_paragraph(doc, "Date verified: 1 September 2026")

    add_heading(doc, "1. Story requirement")
    add_plain_paragraph(
        doc,
        "The Multiple Ship to Accounts field on the Quote must be set to Yes when quote lines use more than one related Ship-To Account. The behavior must work for stepped-pricing quotes, non-stepped quotes, and quotes containing groups that are not used for stepped pricing. The field is expected to update after Save or Quick Save in the Quote Line Editor."
    )

    add_heading(doc, "2. What we found")
    add_bullet(doc, "The issue was in the existing QuoteLine After Record Triggered Flow.")
    add_bullet(doc, "The previous logic evaluated the quote line being processed instead of calculating the result from the complete set of lines on the quote.")
    add_bullet(doc, "This created a last-writer problem. A line using a different Ship-To Account could set the Quote field to Yes, but a later line matching the Quote Account could set it back to No.")
    add_bullet(doc, "The failure was easier to see on stepped-pricing quotes because Quick Save processes many ramped and grouped lines in one transaction.")
    add_bullet(doc, "The correct decision must be quote-wide and must not depend on quote-line processing order.")

    add_heading(doc, "3. What we changed in the Flow")
    add_plain_paragraph(doc, "A new version of the existing Flow was created and activated. The active version is version 8.")
    add_bullet(doc, "Changed the record-trigger behavior so the Flow runs after create and update.")
    add_bullet(doc, "Added the Get Records element Get Quote Line With Different Ship To.")
    add_bullet(doc, "The lookup searches the full quote, not only the current quote line.")
    add_bullet(doc, "The lookup checks for a quote line with a direct Ship-To Account that is different from the Quote Account.")
    add_bullet(doc, "The lookup also supports the inherited Parent Ship-To Account path used by the existing account-hierarchy logic.")
    add_bullet(doc, "When a different related Ship-To Account is found, the Flow routes to the existing Yes update path.")
    add_bullet(doc, "When no different Ship-To Account is found, the Flow routes to the existing No update path.")
    add_bullet(doc, "Existing hierarchy validation was preserved.")
    add_bullet(doc, "The existing Strategic Partner exception was preserved.")
    add_bullet(doc, "The existing Renewal plus Upsell behavior was preserved.")
    add_bullet(doc, "No Apex class, field, permission set, Production configuration, or bulk data update was included.")

    add_heading(doc, "4. Flow logic after the change")
    flow_num_id = new_numbering_instance(doc)
    add_step(doc, "A Quote Line is created or updated and the after-save Flow starts.", flow_num_id)
    add_step(doc, "The Flow keeps the existing validation and exception checks.", flow_num_id)
    add_step(doc, "The Flow queries all relevant Quote Lines for the same Quote.", flow_num_id)
    add_step(doc, "The query looks for at least one direct or inherited Ship-To Account that differs from the Quote Account.", flow_num_id)
    add_step(doc, "If a matching line is found, the Quote Multiple Ship to Accounts field is updated to Yes.", flow_num_id)
    add_step(doc, "If no matching line is found, the Quote Multiple Ship to Accounts field is updated to No.", flow_num_id)
    add_step(doc, "Because the decision is based on the complete quote, later quote-line processing cannot overwrite the correct result based only on one line.", flow_num_id)

    add_heading(doc, "5. Technical and deployment facts")
    add_bullet(doc, "Target org: FlywirePartial, Flywire Enterprise sandbox.")
    add_bullet(doc, "Organization ID: 00DhG0000000jOXUAY.")
    add_bullet(doc, "Instance: USA1148S.")
    add_bullet(doc, "Active Flow version: 8.")
    add_bullet(doc, "Active version ID: 301hG00000BfrsBQAR.")
    add_bullet(doc, "Flow trigger: CreateAndUpdate, RecordAfterSave.")
    add_bullet(doc, "Previous active version 6 is obsolete. Version 7 is also obsolete.")
    add_bullet(doc, "Check-only job 0AfhG000001bMygSAE completed for one Flow with zero component errors.")
    add_bullet(doc, "Deployment job 0AfhG000001bPN5SAM completed for one Flow with zero component errors.")
    add_bullet(doc, "Read-only sampling of 2,000 recent quote lines covered 155 quotes. Twenty-four quotes showed an incorrect No value, and 15 of those quotes included ramped lines.")

    add_heading(doc, "6. Functional test completed")
    add_bullet(doc, "Selected record: Q-37873.")
    add_bullet(doc, "The Quote was Draft and not ordered.")
    add_bullet(doc, "The Quote had 44 lines in five groups.")
    add_bullet(doc, "Thirty-two lines were ramped and 12 were non-ramped.")
    add_bullet(doc, "The lines used two Ship-To Accounts. Six lines used the alternate account.")
    add_bullet(doc, "Before Quick Save, Multiple Ship to Accounts was No.")
    add_bullet(doc, "The Quote Line Editor was opened, the line/account setup was confirmed, and Quick Save was selected without intentionally changing line values.")
    add_bullet(doc, "After Quick Save, Multiple Ship to Accounts changed to Yes.")
    add_bullet(doc, "An independent Salesforce query confirmed the stored Quote value was Yes.")
    add_bullet(doc, "The primary acceptance criterion passed in FlywirePartial.")

    add_heading(doc, "7. Steps to reproduce and verify in FlywirePartial")
    add_plain_paragraph(doc, "Use an editable test quote or a clone. Do not reuse a shared business-testing record for repeat testing.")
    repro_num_id = new_numbering_instance(doc)
    add_step(doc, "Open the FlywirePartial sandbox.", repro_num_id)
    add_step(doc, "Open a stepped-pricing test Quote that has a primary Account and access to a second related Ship-To Account in the same hierarchy. Q-37873 was the selected record used for this verification.", repro_num_id)
    add_step(doc, "On the Quote detail page, record the Quote number, Account, Multiple Ship to Accounts value, Quote type, and status.", repro_num_id)
    add_step(doc, "Confirm the Quote is editable and open Edit Lines.", repro_num_id)
    add_step(doc, "Confirm the Quote contains ramped or stepped groups and that at least one line uses the alternate related Ship-To Account while other lines use the Quote Account.", repro_num_id)
    add_step(doc, "For Q-37873, confirm 44 total lines, five groups, 32 ramped lines, two Ship-To Accounts, and six alternate-account lines.", repro_num_id)
    add_step(doc, "Select Quick Save. Do not intentionally change pricing or other quote-line values when only verifying the Flow behavior.", repro_num_id)
    add_step(doc, "Return to or refresh the Quote detail page.", repro_num_id)
    add_step(doc, "Confirm Multiple Ship to Accounts is Yes.", repro_num_id)
    add_step(doc, "Run an independent read-only query or record refresh to confirm the stored Quote value is Yes.", repro_num_id)
    add_step(doc, "Record the result and the test Quote used. For any additional negative or reset testing, use a separate clone.", repro_num_id)

    add_heading(doc, "8. Components to deploy")
    add_deployment_table(doc)

    core = doc.core_properties
    core.title = "SALDEV-1473 - Implementation and Test Notes"
    core.subject = "Plain black-and-white implementation, reproduction, and deployment-component notes"
    core.author = "Flywire Salesforce Review"
    core.keywords = "SALDEV-1473, Salesforce, Flow, Multiple Ship to Accounts, deployment"

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
