from __future__ import annotations

from pathlib import Path
from typing import Dict, Iterable, List, Sequence, Tuple

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DIR = ROOT / "outputs" / "next-four-story-documents"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

PRESET = {
    "page_width": 12240,
    "page_height": 15840,
    "margin": 1440,
    "header_footer_distance": 708,
    "content_width": 9360,
    "table_indent": 120,
    "cell_top": 80,
    "cell_bottom": 80,
    "cell_start": 120,
    "cell_end": 120,
    "body_font": "Calibri",
    "body_size": 11,
    "body_after": 6,
    "body_line": 1.10,
    "title_size": 23,
    "subtitle_size": 14,
    "h1_size": 16,
    "h1_before": 16,
    "h1_after": 8,
    "h2_size": 13,
    "h2_before": 12,
    "h2_after": 6,
}

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
GRAY = "555555"
MUTED = "6B7280"
LIGHT_GRAY = "F2F4F7"
LIGHT_BLUE = "E8EEF5"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
GREEN = "1F6D4A"
AMBER = "7A5A00"
RED = "9B1C1C"
BORDER = "C9D2DC"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = tc_mar.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            tc_mar.append(tag)
        tag.set(qn("w:w"), str(value))
        tag.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa: Sequence[int], indent_dxa: int = 120) -> None:
    if sum(widths_dxa) != PRESET["content_width"]:
        raise ValueError(f"Table widths must total {PRESET['content_width']}: {widths_dxa}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(PRESET["content_width"]))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    old_grid = table._tbl.tblGrid
    for child in list(old_grid):
        old_grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        old_grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(
                cell,
                PRESET["cell_top"],
                PRESET["cell_start"],
                PRESET["cell_bottom"],
                PRESET["cell_end"],
            )


def set_table_borders(table, color=BORDER, size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "nil")


def set_run_font(run, size=None, color=None, bold=None, italic=None, name=None) -> None:
    font_name = name or PRESET["body_font"]
    run.font.name = font_name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), font_name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = PRESET["body_font"]
    normal._element.rPr.rFonts.set(qn("w:ascii"), PRESET["body_font"])
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), PRESET["body_font"])
    normal.font.size = Pt(PRESET["body_size"])
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(PRESET["body_after"])
    normal.paragraph_format.line_spacing = PRESET["body_line"]

    title = styles["Title"]
    title.font.name = PRESET["body_font"]
    title._element.rPr.rFonts.set(qn("w:ascii"), PRESET["body_font"])
    title._element.rPr.rFonts.set(qn("w:hAnsi"), PRESET["body_font"])
    title.font.size = Pt(PRESET["title_size"])
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.keep_with_next = True

    subtitle = styles["Subtitle"]
    subtitle.font.name = PRESET["body_font"]
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), PRESET["body_font"])
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), PRESET["body_font"])
    subtitle.font.size = Pt(PRESET["subtitle_size"])
    subtitle.font.color.rgb = RGBColor.from_string(GRAY)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle.paragraph_format.keep_with_next = True

    for name, size, color, before, after in (
        ("Heading 1", PRESET["h1_size"], BLUE, PRESET["h1_before"], PRESET["h1_after"]),
        ("Heading 2", PRESET["h2_size"], BLUE, PRESET["h2_before"], PRESET["h2_after"]),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = PRESET["body_font"]
        style._element.rPr.rFonts.set(qn("w:ascii"), PRESET["body_font"])
        style._element.rPr.rFonts.set(qn("w:hAnsi"), PRESET["body_font"])
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr)
    run._r.append(fld_char_2)


def configure_section(doc: Document, case_number: str) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    left = hp.add_run("PROBO MEDICAL | SALESFORCE")
    set_run_font(left, size=8.5, color=MUTED, bold=True)
    tab_stops = hp.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(6.5))
    right = hp.add_run(f"\tCASE {case_number} | DEVDO VALIDATION")
    set_run_font(right, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.paragraph_format.space_before = Pt(0)
    add_page_field(fp)


def add_metadata_rows(doc: Document, rows: Sequence[Tuple[str, str]]) -> None:
    for label, value in rows:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        label_run = p.add_run(f"{label}: ")
        set_run_font(label_run, size=10.5, color=NAVY, bold=True)
        value_run = p.add_run(value)
        set_run_font(value_run, size=10.5, color="222222")


def add_callout(doc: Document, label: str, text: str, tone: str = "blue") -> None:
    color = {"blue": BLUE, "green": GREEN, "amber": AMBER, "red": RED}[tone]
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.05)
    p.paragraph_format.line_spacing = 1.10
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), CALLOUT)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    p_bdr.append(left)
    p_pr.append(p_bdr)
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, size=10.5, color=color, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color="222222")


def add_labeled_paragraph(doc: Document, label: str, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    label_run = p.add_run(f"{label}: ")
    set_run_font(label_run, bold=True, color=NAVY)
    text_run = p.add_run(text)
    set_run_font(text_run, color="222222")


def add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    widths_dxa: Sequence[int],
    status_column: int | None = None,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    table.rows[0]._tr.get_or_add_trPr()
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=9.5, color=NAVY, bold=True)

    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            color = "222222"
            bold = False
            if status_column is not None and idx == status_column:
                normalized = str(value).upper()
                if any(term in normalized for term in ("PASS", "YES", "COMPLETE", "CLOSED", "ACTIVE")):
                    color = GREEN
                    bold = True
                elif any(term in normalized for term in ("NO", "HOLD", "PENDING", "NOT")):
                    color = AMBER
                    bold = True
            set_run_font(r, size=9.5, color=color, bold=bold)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_page_break(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


SHARED_EVIDENCE = {
    "org": "Probo Medical DevDO sandbox (Org ID 00DiK0000001co1UAA)",
    "deploy": "0AfiK0000000WpWSAU - 16/16 components deployed; 13/13 deployment-gate tests passed",
    "test": "707iK0000005xkHQAQ - 14 tests passed; 0 failures; 100% pass rate",
    "coverage": "P_RMATriggerHandler focused coverage: 83%",
    "change_set": "DEVDO Next 4 High Priority Fixes (0A2iK00000001lh) - 16 components; not uploaded",
}


STORIES: List[Dict] = [
    {
        "case": "00009998",
        "dev_case": "00010703",
        "title": "RMA Updates Must Preserve Asset Qualification Attributes",
        "subject": "System behavior where updating an RMA",
        "priority": "Medium",
        "prod_status": "Claimed",
        "description": "Updating Quality Issue Details, Root Cause, or Corrective Action on an RMA removed key asset attributes.",
        "root_cause": "The RMA receipt/update branch in P_RMATriggerHandler explicitly cleared Cosmetic Rating, Functional Rating, and Review Status while resetting other check-in evaluation fields.",
        "implementation": [
            ("Preservation rule", "Removed the three destructive assignments so Cosmetic_Rating__c, Functional_Rating__c, and Review_Status__c remain unchanged."),
            ("Existing behavior retained", "Repair Evaluation, repairability flags, functional notes, and cosmetic notes continue to reset as before."),
            ("Scope control", "The change is limited to P_RMATriggerHandler; no data backfill and no Production write were performed."),
        ],
        "hardcoding": "No Salesforce IDs or business configuration values were introduced. The change removes destructive assignments and relies on existing object relationships.",
        "repro_preconditions": "Use DevDO test data and an authorized RMA user. Prepare a Product Item with Cosmetic Rating = A, Functional Rating = B, Review Status = In Review, and at least one other check-in evaluation field populated.",
        "repro_steps": [
            "Open or create an RMA linked to the prepared Product Item and confirm Date Item Received is blank before the test.",
            "Populate or update Quality Issue Details, Root Cause, and Corrective Action on the RMA.",
            "Set Date Item Received to the current test date and save the RMA so the receipt/update branch executes.",
            "Refresh the linked Product Item and review Cosmetic Rating, Functional Rating, Review Status, and the other check-in evaluation fields.",
        ],
        "repro_result": "Fixed result: Cosmetic Rating remains A, Functional Rating remains B, and Review Status remains In Review; the other intended check-in evaluation fields reset. Original defect signature: the three preserved fields become blank.",
        "test_method": "P_RMATriggerHandlerTest.testDateItemReceived_TriggersProductItemResetBlock",
        "test_assertion": "Cosmetic Rating remains A, Functional Rating remains B, and Review Status remains In Review after the RMA received-date transition; the other evaluation fields still reset.",
        "components": [
            ("Apex Class", "P_RMATriggerHandler", "Modified"),
            ("Apex Test", "P_RMATriggerHandlerTest", "Regression assertion"),
        ],
        "decision": "DevDO passed. The implementation is ready for Production deployment and business smoke testing. Keep Production case 00009998 open until those two gates complete.",
    },
    {
        "case": "00010158",
        "dev_case": "00010704",
        "title": "RMA Save Must Support a Product Item Without an Opportunity",
        "subject": "Need help with saving an RMA in Salesforce - receiving an error",
        "priority": "Medium",
        "prod_status": "Claimed",
        "description": "RMA A-206654 failed during P_RMATrigger processing when an RMA was saved without an assigned Opportunity.",
        "root_cause": "Opportunity Product synchronization dereferenced RepairOpp.Id and Product Item values before confirming that both records existed.",
        "implementation": [
            ("Null-safe guard", "Opportunity Product synchronization now runs only when both thisPItem and RepairOpp are non-null."),
            ("Valid save path", "Selecting a Product Item remains valid even when the RMA has no Opportunity; the unrelated Opportunity synchronization is skipped."),
            ("Scope control", "No default Opportunity is manufactured, and no relationship is hard-coded or silently inferred."),
        ],
        "hardcoding": "No Salesforce IDs, Opportunity IDs, Product Item IDs, or fallback business values were added. Runtime relationships determine whether synchronization is applicable.",
        "repro_preconditions": "Use DevDO test data and an authorized RMA user. Have a valid Product Item available, and do not assign an Opportunity to the test RMA.",
        "repro_steps": [
            "Create a new RMA, or open an editable test RMA, with Opportunity left blank.",
            "Select the valid Product Item and complete all other required RMA fields.",
            "Enter Quality Issue Notes so the relevant P_RMATrigger processing path is exercised.",
            "Save the RMA and refresh the record.",
            "Confirm the selected Product Item and notes remain saved while Opportunity remains blank.",
        ],
        "repro_result": "Fixed result: the RMA saves successfully without an Opportunity and without a P_RMATrigger exception. Original defect signature: the save fails while the handler attempts Opportunity Product synchronization.",
        "test_method": "P_RMATriggerHandlerTest2.case00010158AllowsProductItemWithoutOpportunity",
        "test_assertion": "An RMA saves with its selected Product Item and Quality Issue Notes while Opportunity__c remains null; the transaction completes without the trigger error.",
        "components": [
            ("Apex Class", "P_RMATriggerHandler", "Modified"),
            ("Apex Test", "P_RMATriggerHandlerTest2", "Regression assertion"),
        ],
        "decision": "DevDO passed. The implementation is ready for Production deployment and a save-path smoke test using an RMA without an Opportunity. Keep Production case 00010158 open until that test succeeds.",
    },
    {
        "case": "00009589",
        "dev_case": "00010705",
        "title": "SRI Completion Must Not Remove a Dispatched Appointment from the Gantt",
        "subject": "SA removed from Gantt when SRI completed",
        "priority": "Medium",
        "prod_status": "New",
        "description": "Completing SRI-00015817 moved SA-17132 from its scheduled state back to Internal, removing it from the Field Service Gantt.",
        "root_cause": "The asynchronous SRI Flow evaluated a relationship snapshot and hard-set the Service Appointment to Internal, allowing an appointment that had since advanced to Dispatched to regress.",
        "implementation": [
            ("Fresh state read", "The Flow now re-queries the current Service Appointment before evaluating whether any status transition is allowed."),
            ("Configurable transition", "An active Service_Appointment_Status_Rule__mdt record maps Pending Customer Authorization to Internal."),
            ("Regression prevention", "If the current appointment is Dispatched, no matching rule exists, so the Flow performs no status update."),
        ],
        "hardcoding": "The trigger and target statuses are stored in Custom Metadata, not embedded in the Flow. The Flow does not contain Salesforce record IDs.",
        "repro_preconditions": "Use DevDO Field Service test data and an authorized scheduling user. Prepare a Service Appointment in Dispatched status that is visible on the Gantt and can be linked to a test SRI.",
        "repro_steps": [
            "Open the Dispatched Service Appointment and confirm it is visible on the Field Service Gantt.",
            "Create or update the linked SRI and complete the SRI action that invokes the asynchronous status Flow.",
            "Allow the asynchronous processing to finish, then refresh the Service Appointment.",
            "Reopen or refresh the Gantt and locate the same Service Appointment.",
            "Confirm the appointment status remains Dispatched and the appointment remains visible on the Gantt.",
        ],
        "repro_result": "Fixed result: a Dispatched appointment is not moved back to Internal and remains on the Gantt. Original defect signature: SRI completion changes the appointment to Internal and removes it from the Gantt.",
        "test_method": "ProboNextFourFlowRegressionTest.case00009589DispatchedAppointmentDoesNotReturnToInternal",
        "test_assertion": "Creating the SRI against a Dispatched Service Appointment leaves the appointment status as Dispatched after asynchronous processing.",
        "components": [
            ("Flow", "On_Create_of_SRI_Check_if_SA_status_needs_to_be_updated", "Active"),
            ("Apex Test", "ProboNextFourFlowRegressionTest", "Regression assertion"),
            ("Custom Metadata Type", "Service_Appointment_Status_Rule__mdt", "Included"),
            ("Custom Fields", "Active__c; Trigger_Status__c; Target_Status__c", "Included"),
            ("Metadata Record", "Pending_Authorization_To_Internal", "Active"),
        ],
        "decision": "DevDO passed. The implementation is ready for Production deployment and a Field Service scheduling smoke test. Keep Production case 00009589 open until the appointment remains visible on the Gantt in Production.",
    },
    {
        "case": "00010514",
        "dev_case": "00010706",
        "title": "Automatically Create One Core Exchange RMA for a Qualifying Work Order",
        "subject": "Core Exchange RMA not automatically opening on Work Order",
        "priority": "Medium",
        "prod_status": "Claimed",
        "description": "GE contract Work Orders required manual Core Exchange RMA creation; the business requested automatic creation.",
        "root_cause": "No active automation created the required Core Exchange RMA when a qualifying GE Work Order became complete with Account, Opportunity, and Service Contract data.",
        "implementation": [
            ("Qualification", "An after-save Work Order Flow runs only when Account, Opportunity, and Service Contract are populated and the record newly meets the criteria."),
            ("Configuration", "An active Core_Exchange_Automation_Rule__mdt record identifies the Account name and the RMA record type developer name."),
            ("Duplicate safety", "The Flow checks for an existing RMA with the same Work Order and Core Exchange record type before creating one."),
            ("Backfill control", "Historical Work Orders are not backfilled; the behavior applies when a Work Order first meets the criteria."),
        ],
        "hardcoding": "The qualifying Account and RMA record type developer name are held in Custom Metadata. RecordType is resolved dynamically; no Salesforce IDs are embedded in the Flow.",
        "repro_preconditions": "Use DevDO test data and an authorized Work Order/RMA user. Confirm the active Core Exchange Custom Metadata rule is available, and prepare a qualifying GE account, Opportunity, Service Contract, and Work Order with no existing Core Exchange RMA.",
        "repro_steps": [
            "Open the test Work Order and populate Account, Opportunity, and Service Contract with the qualifying DevDO records.",
            "Save or update the Work Order so it newly meets the Flow entry criteria.",
            "Open the Work Order related records and confirm exactly one RMA was created with the configured Core Exchange record type.",
            "Confirm the new RMA retains the expected Work Order, Opportunity, and Service Contract relationships.",
            "Edit and save the same qualifying Work Order again, then confirm a second Core Exchange RMA was not created.",
        ],
        "repro_result": "Fixed result: the first qualifying transition creates exactly one configured Core Exchange RMA and later Work Order updates do not create duplicates. Original defect signature: no RMA is created automatically and manual creation is required.",
        "test_method": "ProboNextFourFlowRegressionTest.case00010514CreatesOneConfiguredCoreExchangeRma",
        "test_assertion": "A qualifying Work Order creates exactly one Core Exchange RMA, retains the Opportunity and Service Contract relationships, and does not create a duplicate after update.",
        "components": [
            ("Flow", "Create_Configured_Core_Exchange_RMA_From_Work_Order", "Active"),
            ("Apex Test", "ProboNextFourFlowRegressionTest", "Regression assertion"),
            ("Custom Metadata Type", "Core_Exchange_Automation_Rule__mdt", "Included"),
            ("Custom Fields", "Account_Name__c; Active__c; RMA_Record_Type_Developer_Name__c", "Included"),
            ("Metadata Record", "GE_Precision_Healthcare", "Active"),
        ],
        "decision": "DevDO passed. The implementation is ready for Production deployment and a GE Work Order smoke test. Keep Production case 00010514 open until exactly one Core Exchange RMA is verified in Production.",
    },
]


def build_document(story: Dict) -> Path:
    doc = Document()
    configure_styles(doc)
    configure_section(doc, story["case"])
    doc.core_properties.title = f"Case {story['case']} Architect Closure Review"
    doc.core_properties.subject = story["title"]
    doc.core_properties.author = "Probo Medical Salesforce Team"
    doc.core_properties.keywords = "Salesforce, DevDO, architect review, deployment readiness"

    title = doc.add_paragraph("ARCHITECT CLOSURE REVIEW", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle = doc.add_paragraph(story["title"], style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT

    add_metadata_rows(
        doc,
        [
            ("Production Case", story["case"]),
            ("DevDO Tracking Case", f"{story['dev_case']} - Closed"),
            ("Prepared", "September 1, 2026"),
            ("Environment", "Probo Medical DevDO sandbox"),
            ("Review Status", "DEVDO PASS | READY FOR PRODUCTION DEPLOYMENT/UAT"),
        ],
    )
    add_callout(doc, "Architect determination", story["decision"], "green")

    doc.add_paragraph("1. Production Case Summary", style="Heading 1")
    add_table(
        doc,
        ["Field", "Verified value"],
        [
            ("Subject", story["subject"]),
            ("Production status", story["prod_status"]),
            ("Priority", story["priority"]),
            ("Business issue", story["description"]),
        ],
        [2160, 7200],
    )

    doc.add_paragraph("2. Root Cause and Resolution", style="Heading 1")
    add_labeled_paragraph(doc, "Root cause", story["root_cause"])
    for label, text in story["implementation"]:
        add_labeled_paragraph(doc, label, text)
    add_callout(doc, "No-hardcoding control", story["hardcoding"], "blue")

    doc.add_page_break()
    doc.add_paragraph("3. Steps to Reproduce and Validate", style="Heading 1")
    add_callout(
        doc,
        "Execution boundary",
        "Perform these steps in DevDO with non-production test records. Repeat in Production only after the Change Set is deployed and the business owner authorizes the smoke test.",
        "blue",
    )
    add_labeled_paragraph(doc, "Preconditions", story["repro_preconditions"])
    add_table(
        doc,
        ["Step", "Action"],
        [(str(index), action) for index, action in enumerate(story["repro_steps"], start=1)],
        [900, 8460],
    )
    add_callout(doc, "Expected result and defect signature", story["repro_result"], "green")

    doc.add_paragraph("4. Verification Evidence", style="Heading 1")
    add_labeled_paragraph(doc, "Focused regression", story["test_method"])
    add_labeled_paragraph(doc, "Acceptance assertion", story["test_assertion"])
    add_table(
        doc,
        ["Evidence", "Result"],
        [
            ("Deployment validation", SHARED_EVIDENCE["deploy"]),
            ("Final regression rerun", SHARED_EVIDENCE["test"]),
            ("Focused handler coverage", SHARED_EVIDENCE["coverage"]),
            ("DevDO tracking case", f"{story['dev_case']} - Closed and independently read back"),
        ],
        [2700, 6660],
        status_column=1,
    )

    doc.add_page_break()
    doc.add_paragraph("5. Story-Specific Change Set Scope", style="Heading 1")
    add_table(
        doc,
        ["Component type", "API name", "State"],
        story["components"],
        [2100, 5580, 1680],
        status_column=2,
    )
    add_labeled_paragraph(doc, "Outbound Change Set", SHARED_EVIDENCE["change_set"])
    add_labeled_paragraph(doc, "Production safety", "The Change Set has not been uploaded or deployed. Production metadata and Production case status were not changed.")

    doc.add_paragraph("6. Closure Gate", style="Heading 1")
    add_table(
        doc,
        ["Gate", "Status", "Required evidence"],
        [
            ("DevDO implementation complete", "YES", "Components deployed and active in DevDO"),
            ("Automated regression passed", "YES", "Test job 707iK0000005xkHQAQ"),
            ("Outbound Change Set prepared", "YES", "0A2iK00000001lh; 16 components"),
            ("Production deployment", "NO - PENDING", "Upload, validate, and deploy the Change Set"),
            ("Production business smoke test", "NO - PENDING", "Execute the case-specific acceptance path"),
            ("Close original Production case", "HOLD", "Close only after both Production gates pass"),
        ],
        [3180, 1740, 4440],
        status_column=1,
    )

    add_callout(
        doc,
        "Final recommendation",
        "DevDO work is complete and the sandbox tracking case is correctly closed. Move the original Production case to Ready for Production/UAT; do not mark the Production issue resolved until deployment and smoke-test evidence are attached.",
        "amber",
    )

    output_path = OUTPUT_DIR / f"Case_{story['case']}_Architect_Closure_Review.docx"
    doc.save(output_path)
    return output_path


def main() -> None:
    paths = [build_document(story) for story in STORIES]
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
