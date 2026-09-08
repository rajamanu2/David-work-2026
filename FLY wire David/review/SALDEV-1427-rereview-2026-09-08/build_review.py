from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).parent
doc=Document()
s=doc.sections[0];s.top_margin=s.bottom_margin=Inches(.7);s.left_margin=s.right_margin=Inches(.75)
for name in ['Normal','Title','Heading 1','Heading 2','List Bullet']:
 st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(11);st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.line_spacing=1.05;st.paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(22);doc.styles['Title'].font.bold=True
for name,size in [('Heading 1',15),('Heading 2',12)]:
 st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.space_before=Pt(10);st.paragraph_format.keep_with_next=True
for el in list(doc.styles.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.title='SALDEV-1427 Architect Re-review September 8 2026'

def p(text,bold=None,style=None):
 x=doc.add_paragraph(style=style)
 if bold and text.startswith(bold):x.add_run(bold).bold=True;x.add_run(text[len(bold):])
 else:x.add_run(text)
 x.paragraph_format.widow_control=True
 return x
def h(text,new=False):
 x=doc.add_heading(text,1);x.paragraph_format.page_break_before=new
def sub(text):doc.add_heading(text,2)
def bullet(text):
 x=p(text,style='List Bullet');x.paragraph_format.left_indent=Inches(.16);x.paragraph_format.first_line_indent=Inches(-.16)
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for c,x in zip(t.rows[0].cells,headers):c.text=x
 for row in rows:
  for c,x in zip(t.add_row().cells,row):c.text=str(x)
 for ri,row in enumerate(t.rows):
  rp=row._tr.get_or_add_trPr();rp.append(OxmlElement('w:cantSplit'))
  if ri==0:rp.append(OxmlElement('w:tblHeader'))
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'95');e.set(qn('w:type'),'dxa');mar.append(e)
   cp.append(mar);b=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   cp.append(b)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E5EBF1');cp.append(sh)
   for x in c.paragraphs:
    x.paragraph_format.space_after=Pt(0);x.paragraph_format.line_spacing=1.02
    for r in x.runs:r.font.size=Pt(10);r.bold=ri==0
 return t

doc.add_paragraph('SALDEV-1427 Architect Re-review','Title')
p('Update QCP cloning logic for stepped-up pricing')
p('Review date: 8 September 2026 | Environment: FlywirePartial',bold='Review date:')
p('Decision: Changes Requested. The latest QCP changes are present and the previous Apex coverage blocker is resolved. The story still needs corrections to editability, ARR visibility, field references, and identifier handling before business acceptance.',bold='Decision:')
h('Verified scope and latest changes')
p('The sandbox was independently verified as Flywire, org 00DhG0000000jOXUAY, instance USA1148S. This review compares the current scripts and related Apex with the September 4 review and the retained SALDEV-1427 story requirements.')
table(['Component','Latest modification in IST','Observed change'],[
('QCP_HideFieldInQLE','8 Sep 2026 00:28','ARR constant introduced; Opportunity Line ID now cleared before calculation; duplicate Estimated Volume block removed.'),
('QLE_CloneLineHandler','7 Sep 2026 17:40','ARR dependency list expanded; ARR constant introduced; Opportunity Line ID now cleared after cloning.'),
('QuoteLineTrigger and tests','7 Sep 2026','Current shared trigger and expanded tests retrieved and validated. Lineage handler remains the September 4 version.')],[2.1,1.45,3.45])
h('Fresh validation results')
table(['Validation','Result'],[
('Check-only validation','PASS: 3 components; 9 tests; zero component or test errors.'),
('QuoteLineTrigger coverage','79.49% (31 of 39 lines). Prior 43.59% coverage blocker resolved.'),
('QuoteLineTriggerHandler coverage','96.89% (280 of 289 lines).'),
('JavaScript syntax','Both retrieved scripts passed.'),
('Local behavior probes','36 executions: 24 matched expectations; 12 mismatches representing six cases repeated against source and stored compiled code.')],[2.1,4.9])
p('These JavaScript probes use controlled in-memory records. They reproduce code behavior; they do not constitute a Salesforce QLE business test. No metadata was deployed or activated and no business records were edited.')

h('QCP behavior findings',True)
sub('1 Downstream field locking does not use the available line flag')
p('High priority. The live quote-line schema includes Group_Allow_Product_Ramping__c, a formula reflecting the parent group ramping flag. It is declared in the QCP field list, but isFieldEditable does not read it. Instead, the function relies on parentGroup or relationship objects, plus group source and number values.')
p('Reproduction: call isFieldEditable with a quote-line record containing Group_Allow_Product_Ramping__c = true, a source line, and Group 2. Ship_To_Account__c and Transaction_Currency__c both return true. A supplied wrapper containing parentGroup does return false, demonstrating that the result depends on the input shape.',bold='Reproduction:')
p('Salesforce documents this callback parameter as a quote-line SObject, not a full QuoteLineModel. A true result leaves field settings unchanged; it does not enforce the required lock. Use a reliable record-level indication of both ramping and downstream-clone status, then prove the drawer and main QLE behavior.')
p('Code: QCP_HideFieldInQLE, lines 459-534.',bold='Code:')
sub('2 Usage-rate fields remain locked before the ramping exception')
p('High priority. Usage_Rate__c and Usage_Rate_Amt__c appear in both the global read-only list and the ramping editable list. The global guard returns false first. Both fields remain locked even when a complete downstream parent-group wrapper is supplied.')
p('Controls passed for Quantity, Quoted Price, Adjusted Rate Amount, Adjusted Rate Percent, and Rate Details in that enriched wrapper. Resolve the conflicting rule order for the two usage-rate fields and confirm the permitted price fields with the story owner.')
p('Code: QCP_HideFieldInQLE, lines 427-467 and 523-532.',bold='Code:')
sub('3 ARR visibility is not restricted to the first group')
p('High priority. The visibility function evaluates product, vertical, and Requires_ARR_Fields_Completion__c without a downstream-group exclusion. For an Education Domestic Payments downstream clone with that flag true, Utilization__c returns true. The current plugin therefore does not hide this ARR field as required.')
p('Add the downstream ARR visibility rule before product-specific early returns. Check all ARR fields across supported verticals and verify saved clone values outside QLE; hiding a field does not prove its value was cleared.')
p('Code: QCP_HideFieldInQLE, lines 1-362; Utilization rule at lines 74-100.',bold='Code:')

h('Identifier and dependency findings',True)
sub('4 New identifier clearing is broader than new unsaved clones')
p('High priority. onBeforeCalculate clears NS_ID__c and Opportunity_Line_Id__c whenever SBQQ__Existing__c is false. It does not check whether the record has an Id. A controlled saved record with Existing = false lost both established IDs. The clone hook also clears both IDs for every processed clone.')
p('This makes reliable Apex restoration essential. copyOpportunityLineIdOnRamp restores a source ID where available, otherwise matches by Quote + Product and occurrence order, with a first-ID fallback. That fallback does not distinguish bundle, configuration, Ship-To, or ramping context. The subsequent OLI sync can also replace a different Quote Line ID.')
p('Required correction: limit clearing to the intended clone lifecycle, preserve identifiers where required, and prove deterministic mapping through calculate, save, reload, and OLI synchronization. Do not assume that blanking an ID makes the clone safe.')
p('Code: QCP lines 390-418; clone hook lines 40-60; QuoteLineTriggerHandler lines 278-350; OpportunityLineItemTriggerHandler lines 39-45.',bold='Code:')
sub('5 Field declarations improved but the invalid API remains')
p('The clone script now declares the ARR fields it writes, resolving much of the earlier missing-dependency finding. However, Domestic_Sponsor__c remains in both ARR constants and the clone dependency list, and is absent from the current quote-line describe. Confirm the correct API or an approved removal; do not create a replacement field solely to satisfy the script.')
p('The QCP group dependency list still contains only Allow_Product_Ramping__c although the logic reads group Source and Number. Its explicit line list also omits NS_ID__c and Opportunity_Line_Id__c. Some fields can arrive through other CPQ mechanisms, so an omitted declaration alone is not proof that a runtime payload omits them. Capture the real payload or declare the required valid fields explicitly.')
p('The two ARR constants currently match, but remain separate copies in two script records. Add a consistency check to prevent future drift. Keep unrelated shared-QCP cleanup outside this story unless separately agreed.')
sub('6 Passing Apex tests still do not prove every group outcome')
p('The current ramp test queries two records without ordering, then asserts only evalList[0].Opportunity_Line_Id__c. It does not explicitly assert the saved Group 2 record by Id. Managed CPQ triggers are disabled around its insert. Exception handling in the lineage method logs and suppresses errors.')
p('Add exact saved-record assertions for both source and downstream groups, duplicate products, distinct Ship-To accounts, and bundles. Include the full managed CPQ path and an explicit failure-handling expectation. Coverage is resolved; business correctness is not established by coverage alone.')
p('Code: QuoteLineTriggerHandlerTest lines 349-408; QuoteLineTriggerHandler lines 349-350.',bold='Code:')

h('Required acceptance and next steps',True)
table(['Scenario','Required evidence'],[
('First group and later clones','First group remains normally editable; downstream attributes and Ship-To lock; agreed quantity and price fields remain editable.'),
('ARR clone integrity','ARR hidden downstream; required ARR values null on saved Quote Line detail pages; original line unchanged.'),
('Non-stepped cloning','Clone lines and groups; verify required ARR values outside QLE. Keep linked SALDEV-1371 scope traceable.'),
('Lineage across calculate and save','Corresponding stepped lines retain the same ID after save/reload and OLI sync; duplicate products, bundles and multiple Ship-To accounts do not cross-map.'),
('Mixed groups and regression','Cover multiple verticals, stepped/non-stepped groups, payments and bundles; confirm group edits and added products do not propagate forward.')],[2,5])
p('Scope clarification: the retained story says amendments were removed from the MVP on August 17 and one-time-product rules moved to SALDEV-1441. Do not add those as new SALDEV-1427 acceptance gates. Resolve the story’s remaining NS-ID wording with the owner; coordinate any changes to shared code with the related stories.',bold='Scope clarification:')
h('Suggested Jira comment')
p('Re-review completed against the latest FlywirePartial QCP and clone scripts. The new ARR constants, expanded clone dependencies, and ID-clearing logic are present. Fresh check-only passed for 3 components and 9 tests, with QuoteLineTrigger at 79.49% and QuoteLineTriggerHandler at 96.89%; the previous coverage blocker is resolved. Changes Requested remains for downstream editability, conflicting usage-rate locks, missing downstream ARR visibility, invalid Domestic_Sponsor__c, and overly broad ID clearing with ambiguous fallback mapping. Local probes reproduced the same issues in source and stored compiled code. Please correct these paths, add exact saved Group 2 and bundle/Ship-To assertions, and provide QLE calculate/save/reload evidence before sign-off. No deployment was performed.')
h('Evidence and boundaries')
p('Current script records, field describe, retrieved Apex, JavaScript probes, and successful check-only are retained in the dated review evidence folder. Current Jira comments were not independently retrieved; requirements were taken from the retained story export. Plugin assignment and full browser QLE execution were not independently reverified in this review.')
p('Salesforce reference: Javascript Page Security Plugin - callback record shape and true/false behavior. https://developer.salesforce.com/docs/revenue/cpq-plugins/guide/cpq-javascript-page-security-plugin.html')
p('Salesforce reference: Custom Action Plugin - clone-hook contract. https://developer.salesforce.com/docs/revenue/cpq-plugins/guide/cpq-custom-action-plugin.html')
doc.save(ROOT/'SALDEV-1427_Architect_Re-review_Latest_2026-09-08.docx')
print(ROOT/'SALDEV-1427_Architect_Re-review_Latest_2026-09-08.docx')
