from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"C:\Users\LIKKI\Documents\ChatGPT\David 2 Probo medical\outputs\Probo_Medical_Salesforce_MFA_User_Runbook.docx")

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "20262E"
MUTED = "5B6573"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_BLUE = "F4F7FB"
PALE_GOLD = "FFF6D8"
GOLD = "7A5A00"
GREEN = "1F5A45"
WHITE = "FFFFFF"
BORDER = "C9D3DF"


def set_run_font(run, name="Calibri", size=11, color=INK, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_no_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)


def set_table_geometry(table, widths_dxa, indent_dxa=120, borders=True):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = tr_pr.find(qn("w:cantSplit"))
        if cant_split is None:
            cant_split = OxmlElement("w:cantSplit")
            tr_pr.append(cant_split)
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if borders:
                set_cell_border(cell)
            else:
                set_no_cell_border(cell)


def set_cell_text(cell, text, *, bold=False, color=INK, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    p.clear()
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)


def add_hyperlink(paragraph, text, url, color=BLUE):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_color = OxmlElement("w:color")
    r_color.set(qn("w:val"), color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(r_color)
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, sep, txt, end])
    set_run_font(run, size=9, color=MUTED)


def make_numbering(doc, kind="decimal"):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids or [0]) + 1
    num_id = max(num_ids or [0]) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if kind == "decimal" else "bullet")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1." if kind == "decimal" else "•")
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    # OOXML requires all abstract numbering definitions to appear before
    # concrete w:num instances. Insert before the first w:num so Word honors
    # the decimal/bullet format instead of falling back to an existing list.
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(list(numbering).index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list(doc, items, kind="decimal"):
    num_id = make_numbering(doc, kind)
    paragraphs = []
    for item in items:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        p_pr = p._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id_el = OxmlElement("w:numId")
        num_id_el.set(qn("w:val"), str(num_id))
        num_pr.extend([ilvl, num_id_el])
        p_pr.append(num_pr)
        if isinstance(item, tuple):
            lead, body = item
            r1 = p.add_run(lead)
            set_run_font(r1, bold=True)
            r2 = p.add_run(body)
            set_run_font(r2)
        else:
            r = p.add_run(item)
            set_run_font(r)
        paragraphs.append(p)
    return paragraphs


def add_body(doc, text, *, bold_lead=None, color=INK, after=6):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(after)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True, color=color)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, color=color)
    else:
        r = p.add_run(text)
        set_run_font(r, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_callout(doc, label, text, fill=PALE_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360], indent_dxa=120, borders=True)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, size=10.5, color=accent, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_fact_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_geometry(table, [1700, 7660], indent_dxa=120, borders=True)
    for idx, (label, value) in enumerate(rows):
        shade_cell(table.cell(idx, 0), LIGHT_BLUE)
        set_cell_text(table.cell(idx, 0), label, bold=True, color=DARK_BLUE, size=9.5)
        set_cell_text(table.cell(idx, 1), value, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_user_section(doc, number, name, facts, admin_steps, user_steps, validation, caution=None):
    add_heading(doc, f"{number}. {name}", 1)
    add_fact_table(doc, facts)
    if caution:
        add_callout(doc, "Important", caution, fill=PALE_GOLD, accent=GOLD)
    add_heading(doc, "Administrator steps", 2)
    add_list(doc, admin_steps, "decimal")
    add_heading(doc, "User steps", 2)
    add_list(doc, user_steps, "decimal")
    add_heading(doc, "Completion checks", 2)
    add_list(doc, validation, "bullet")


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for level, size, color, before, after in (
    (1, 16, BLUE, 18, 10),
    (2, 13, BLUE, 14, 7),
    (3, 12, DARK_BLUE, 10, 5),
):
    style = styles[f"Heading {level}"]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

# Quiet running header.
header_table = section.header.add_table(rows=1, cols=2, width=Inches(6.5))
set_table_geometry(header_table, [5900, 3460], indent_dxa=0, borders=False)
set_cell_text(header_table.cell(0, 0), "PROBO MEDICAL | SALESFORCE ACCESS RECOVERY", bold=True, color=MUTED, size=8.5)
set_cell_text(header_table.cell(0, 1), "INTERNAL RUNBOOK", bold=True, color=MUTED, size=8.5, align=WD_ALIGN_PARAGRAPH.RIGHT)

footer_table = section.footer.add_table(rows=1, cols=2, width=Inches(6.5))
set_table_geometry(footer_table, [7000, 2360], indent_dxa=0, borders=False)
set_cell_text(footer_table.cell(0, 0), "Probo Medical production | Prepared 20 August 2026", color=MUTED, size=8.5)
fp = footer_table.cell(0, 1).paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp.clear()
rr = fp.add_run("Page ")
set_run_font(rr, size=8.5, color=MUTED)
add_field(fp, "PAGE")

# Opening masthead.
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("OPERATIONS RUNBOOK")
set_run_font(r, size=9.5, color=BLUE, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(5)
r = p.add_run("Salesforce MFA & Individual Access Recovery")
set_run_font(r, size=26, color=NAVY, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(16)
r = p.add_run("User-by-user steps for Dominick Vena, Karly Sheriff, Isael Sarmiento, and Juan Torres")
set_run_font(r, size=12.5, color=MUTED)

metadata = doc.add_table(rows=4, cols=2)
set_table_geometry(metadata, [1800, 7560], indent_dxa=120, borders=True)
for idx, (label, value) in enumerate((
    ("Target", "Probo Medical production org (Org ID 00DU0000000LaKoMAK)"),
    ("Purpose", "Restore individual Salesforce access and complete MFA/passkey enrollment"),
    ("Change boundary", "No Apex, metadata deployment, profile change, permission-set change, or duplicate user creation"),
    ("Recommended owner", "Salesforce administrator working live with each affected user"),
)):
    shade_cell(metadata.cell(idx, 0), LIGHT_BLUE)
    set_cell_text(metadata.cell(idx, 0), label, bold=True, color=DARK_BLUE, size=9.5)
    set_cell_text(metadata.cell(idx, 1), value, size=9.5)

doc.add_paragraph().paragraph_format.space_after = Pt(0)
add_callout(doc, "Verified decision", "Do not create new users. All four named employees already have active Salesforce accounts. Restore and validate those existing accounts instead.")

add_heading(doc, "Current account snapshot", 1)
snapshot_rows = [
    ("User", "Username", "Profile", "Access finding"),
    ("Dominick Vena", "dvena@probomedical.com", "QC Manager", "Super User; phishing-resistant MFA required"),
    ("Karly Sheriff", "ksheriff@probomedical.com", "QC Manager", "Super User; phishing-resistant MFA required"),
    ("Isael Sarmiento", "isael.sarmiento@probomedical.com", "Shipping", "Active individual account; standard employee MFA"),
    ("Juan Torres", "jtorres@probomedical.com", "QC Manager", "Super User; phishing-resistant MFA required"),
]
table = doc.add_table(rows=len(snapshot_rows), cols=4)
set_table_geometry(table, [1750, 3200, 1650, 2760], indent_dxa=120, borders=True)
for ri, row in enumerate(snapshot_rows):
    for ci, value in enumerate(row):
        if ri == 0:
            shade_cell(table.cell(ri, ci), LIGHT_BLUE)
            set_cell_text(table.cell(ri, ci), value, bold=True, color=DARK_BLUE, size=9)
        else:
            set_cell_text(table.cell(ri, ci), value, size=9)
set_repeat_table_header(table.rows[0])

add_heading(doc, "Before touching any MFA registration", 1)
add_list(doc, [
    ("Coordinate live: ", "Have the named user on a call or at the workstation before disconnecting a verification method."),
    ("Confirm authority: ", "The administrator needs Manage Multi-Factor Authentication in User Interface."),
    ("Use supported software: ", "Use current Microsoft Edge, Google Chrome, or Safari. Do not troubleshoot passkeys in Internet Explorer."),
    ("Protect recovery codes: ", "Share temporary verification codes privately. Never place the code in email, chat history, or the service ticket."),
    ("Preserve access: ", "Do not remove profiles, the Super User permission set, licenses, or business permissions as part of this login recovery."),
    ("Record evidence: ", "Record only the action, administrator, time, method name, and outcome. Never record the passkey, device PIN, biometric data, password, or temporary code."),
], "bullet")

add_heading(doc, "Standard administrator recovery sequence", 1)
add_body(doc, "Use this sequence whenever an existing built-in authenticator is tied to the wrong device or no longer works.")
add_list(doc, [
    ("Open the user: ", "Setup > Users > Users, then select the exact user by name and username."),
    ("Generate short-term access: ", "Find Temporary Verification Code, select Generate, choose a one-hour lifetime, and give the code directly to the user."),
    ("Disconnect only the broken method: ", "On the user detail page, find Built-In Authenticator and select Del beside the device-bound registration that is being replaced."),
    ("Let the user re-register: ", "The user signs in and completes the Create a Passkey prompt. An administrator cannot register the passkey on the user's behalf."),
    ("Add resilience: ", "After the primary method works, have the user register another supported passkey or security key where practical."),
    ("Close recovery access: ", "Expire the temporary verification code immediately after successful enrollment."),
    ("Validate: ", "Sign out and complete fresh logins on every supported browser/device required for the user's job."),
], "decimal")

add_callout(doc, "Temporary-code limitation", "A temporary verification code is valid for MFA recovery only. It might not satisfy a separate device-activation challenge from an unrecognized browser or app. When possible, perform enrollment from a browser that already works for the user.", fill=PALE_GOLD, accent=GOLD)

add_heading(doc, "User passkey registration sequence", 1)
add_list(doc, [
    ("Open Salesforce: ", "Go to https://probomedical.my.salesforce.com and enter the user's own username."),
    ("Complete recovery: ", "Enter the administrator-issued temporary verification code if Salesforce requests MFA verification."),
    ("Create the passkey: ", "At Create a Passkey, select the desired built-in authenticator. To use a phone, choose the option for another device or phone, then scan the QR code and approve with Face ID, Touch ID, or the phone unlock method."),
    ("Name and save it: ", "Use a clear name such as Dominick iPhone or Karly Windows Laptop."),
    ("Add another device if required: ", "From personal Settings, search for Advanced User Details (or Personal Information), find Built-In Authenticators, and select Add."),
    ("Prove the setup: ", "Sign out, close the browser, reopen it, and complete a fresh login."),
], "decimal")

add_user_section(
    doc,
    1,
    "Dominick Vena",
    [
        ("Username", "dvena@probomedical.com"),
        ("Current access", "Active | QC Manager | Super User"),
        ("Observed behavior", "Edge has successful logins; Chrome and other devices continue to require MFA. The original passkey was registered through a Windows PIN."),
        ("Goal", "Replace the laptop-bound registration with a passkey strategy that supports Dominick's approved browsers and devices."),
    ],
    [
        ("Schedule the recovery: ", "Have Dominick present with his work laptop and phone. Confirm he can still open Salesforce in Edge before changing anything."),
        ("Open the correct record: ", "Setup > Users > Users > Dominick Vena. Confirm username dvena@probomedical.com and Active is selected."),
        ("Issue one-hour recovery access: ", "Generate a Temporary Verification Code and share it verbally or through an approved private channel."),
        ("Document the old method: ", "Record the displayed built-in authenticator name for the ticket, but do not record secrets or codes."),
        ("Disconnect the affected registration: ", "Select Del beside the Windows/laptop built-in authenticator that Dominick can no longer use across his required devices."),
        ("Keep authorization unchanged: ", "Do not remove the QC Manager profile or Super User permission set."),
        ("Support enrollment: ", "Stay with Dominick while he completes the user steps below."),
        ("Close recovery access: ", "Expire the temporary code immediately after validation."),
    ],
    [
        ("Start from the working browser: ", "Use current Microsoft Edge on the work laptop and go to the Probo Medical Salesforce login page."),
        ("Sign in as yourself: ", "Use dvena@probomedical.com. Do not allow Isael or Juan to complete this registration."),
        ("Use the recovery code: ", "Enter the temporary verification code if prompted."),
        ("Register the intended passkey: ", "At Create a Passkey, choose another device/phone if that is the preferred method. Scan the QR code with the iPhone and approve with Face ID or the configured phone unlock method."),
        ("Give it a clear name: ", "For example, Dominick iPhone Passkey."),
        ("Add a backup: ", "If permitted by the company, add a second phishing-resistant method from Advanced User Details > Built-In Authenticators > Add, or register a physical security key."),
        ("Test every required surface: ", "Sign out and test Edge Classic, Chrome Lightning, and Dominick's own approved iPad workflow."),
    ],
    [
        "A fresh Edge login succeeds with Dominick's new passkey.",
        "A fresh Chrome login succeeds without reporting that no matching passkey exists.",
        "Dominick's own approved iPad/photo workflow succeeds, or its separate mobile-app limitation is documented.",
        "Login History shows Success after the MFA challenge for Dominick.",
        "The one-hour temporary verification code is expired.",
        "After Isael and Juan validate their own accounts, Dominick's shared password is reset and his old sessions are revoked.",
    ],
    caution="Because Dominick's credentials were shared, do not reset his password and revoke sessions until Isael and Juan have confirmed working individual access. Once they do, complete that containment step immediately.",
)

add_user_section(
    doc,
    2,
    "Karly Sheriff",
    [
        ("Username", "ksheriff@probomedical.com"),
        ("Current access", "Active | QC Manager | Super User"),
        ("Observed behavior", "Recent Edge login successes exist, while Chrome/IE attempts continue to stop at MFA."),
        ("Goal", "Replace or repair the existing passkey and validate supported-browser access."),
    ],
    [
        ("Schedule the recovery: ", "Have Karly present with the Windows computer and the phone or device that will hold the replacement passkey."),
        ("Use the exact account: ", "Setup > Users > Users > Karly Sheriff. Confirm username ksheriff@probomedical.com and Active is selected."),
        ("Generate one-hour recovery access: ", "Create a Temporary Verification Code and share it privately."),
        ("Disconnect only the failed passkey: ", "Under Built-In Authenticator, select Del beside the registration that is not producing a usable prompt."),
        ("Do not treat missing alternatives as a permission bug: ", "Karly is a privileged Super User, so Salesforce can require a phishing-resistant passkey rather than showing standard authenticator choices."),
        ("Keep her access unchanged: ", "Do not remove QC Manager or Super User during troubleshooting."),
        ("Expire recovery access: ", "After successful validation, expire the temporary code."),
    ],
    [
        ("Use Edge or Chrome: ", "Do not use Internet Explorer for passkey enrollment."),
        ("Sign in with your account: ", "Use ksheriff@probomedical.com and the temporary verification code if requested."),
        ("Complete Create a Passkey: ", "Register Windows Hello on the work laptop or choose another device/phone and approve the QR-code request on the phone."),
        ("Name the registration: ", "For example, Karly Work Laptop or Karly iPhone Passkey."),
        ("Register a backup method: ", "Use Advanced User Details > Built-In Authenticators > Add for another supported device, or add an approved physical security key."),
        ("Perform fresh logins: ", "Close all Salesforce tabs and test new sessions in both Edge and Chrome."),
    ],
    [
        "A fresh Edge login succeeds after MFA.",
        "A fresh Chrome login succeeds after MFA.",
        "No testing depends on Internet Explorer.",
        "Login History records Success after Karly's MFA challenge.",
        "A backup phishing-resistant method is registered where company policy permits.",
        "The temporary verification code is expired.",
    ],
    caution="The missing Choose Another Verification Method option can be expected for a privileged user. The objective is a working phishing-resistant passkey, not restoring weaker MFA choices.",
)

add_user_section(
    doc,
    3,
    "Isael Sarmiento",
    [
        ("Username", "isael.sarmiento@probomedical.com"),
        ("Current access", "Active | Shipping profile | individual Salesforce license"),
        ("Observed behavior", "The account exists and was last used in July; there is no need to create another user."),
        ("Goal", "Restore access to Isael's existing account, enroll individual MFA, and stop all use of Dominick's credentials."),
    ],
    [
        ("Do not create a duplicate: ", "Open Setup > Users > Users > Isael Sarmiento and confirm username isael.sarmiento@probomedical.com."),
        ("Confirm email delivery: ", "Verify that Isael can access the mailbox used for the Salesforce username and reset email."),
        ("Reset the password only if needed: ", "Use Reset Password on Isael's user record if he does not know the current password. Salesforce should send the reset link to Isael."),
        ("Recover MFA only if blocked: ", "If an old verification method prevents login, generate a one-hour Temporary Verification Code and disconnect only the obsolete method."),
        ("Preserve business access: ", "Keep the Shipping profile and existing permissions unchanged."),
        ("Support individual enrollment: ", "Remain available while Isael completes the registration and validates the iPad/photo workflow."),
    ],
    [
        ("Use your own username: ", "Sign in as isael.sarmiento@probomedical.com. Never use Dominick's username or password again."),
        ("Set a private password: ", "Complete the Salesforce password-reset link if the administrator issued one."),
        ("Enroll MFA: ", "Complete the Create a Passkey prompt. A passkey is recommended even though Isael does not currently have the privileged Super User assignment found on the other three users."),
        ("Name the method clearly: ", "For example, Isael Work iPad or Isael Work Phone."),
        ("Test the real workflow: ", "Sign out, sign in again, then verify the required Salesforce records and photo-capture workflow using Isael's own account."),
    ],
    [
        "A fresh login succeeds with isael.sarmiento@probomedical.com.",
        "Isael can access the records and actions required by the Shipping role.",
        "The approved iPad/photo workflow succeeds under Isael's identity.",
        "Login History records Isael's individual successful login.",
        "Any temporary verification code is expired.",
        "Isael confirms that Dominick's credentials are deleted from saved passwords and will not be reused.",
    ],
    caution="Isael already has an active account and Salesforce license. Creating another user would duplicate identity, access, and audit history.",
)

add_user_section(
    doc,
    4,
    "Juan Torres",
    [
        ("Username", "jtorres@probomedical.com"),
        ("Current access", "Active | QC Manager | Super User | individual Salesforce license"),
        ("Observed behavior", "The account exists and was last used in July. Super User grants Modify All Data and View All Data."),
        ("Goal", "Restore Juan's own account and complete phishing-resistant MFA enrollment."),
    ],
    [
        ("Do not create a duplicate: ", "Open Setup > Users > Users > Juan Torres and confirm username jtorres@probomedical.com."),
        ("Confirm the mailbox: ", "Verify that Juan can access the email address attached to the Salesforce account."),
        ("Reset the password if needed: ", "Use Reset Password only if Juan cannot authenticate with the current password."),
        ("Prepare MFA recovery: ", "If an inaccessible verification method is registered, generate a one-hour Temporary Verification Code, document the method name, and disconnect only that obsolete method."),
        ("Preserve authorization: ", "Do not remove the QC Manager profile or Super User permission set without separate business and security approval."),
        ("Support phishing-resistant enrollment: ", "Stay with Juan while he registers a built-in authenticator or physical security key."),
    ],
    [
        ("Use your individual identity: ", "Sign in as jtorres@probomedical.com, never as Dominick."),
        ("Complete any password reset: ", "Use the private reset link delivered to Juan's mailbox."),
        ("Register a passkey: ", "Complete Create a Passkey using Windows Hello, Face ID/Touch ID through a supported phone, a compatible password manager, or an approved physical security key."),
        ("Add backup access: ", "Because Juan is a privileged Super User, register a second phishing-resistant method where possible."),
        ("Test required devices: ", "Complete fresh logins on the work browser and approved iPad/photo workflow using Juan's account."),
    ],
    [
        "A fresh login succeeds with jtorres@probomedical.com.",
        "Juan completes the required QC Manager work under his own identity.",
        "The approved iPad/photo workflow succeeds under Juan's identity.",
        "Login History records Juan's individual successful login.",
        "A backup phishing-resistant method is registered where company policy permits.",
        "Any temporary verification code is expired and Dominick's credentials are removed from Juan's devices.",
    ],
    caution="Juan is a privileged user because Super User grants Modify All Data and View All Data. Standard authenticator choices might not satisfy the enforced phishing-resistant MFA requirement.",
)

add_heading(doc, "Final containment and closure", 1)
add_body(doc, "Complete these actions only after Isael and Juan have both demonstrated successful individual access.")
add_list(doc, [
    ("Reset Dominick's password: ", "Because it was shared, issue a password reset and have Dominick set a new private password."),
    ("Revoke old sessions: ", "From Setup > Session Management, remove Dominick's active sessions so previously shared sessions cannot continue."),
    ("Remove saved shared credentials: ", "Confirm Isael and Juan delete Dominick's username/password from browsers, tablets, password managers, and mobile devices."),
    ("Repeat device tests: ", "Each person signs in once more with their own username and completes the job-specific workflow."),
    ("Review Login History: ", "Confirm successful individual logins and investigate any continuing Invalid Password, Password Lockout, or Multi-factor required events."),
    ("Close the ticket with evidence: ", "Record the usernames tested, browsers/devices, result, administrator, date/time, and confirmation that temporary codes were expired."),
], "decimal")

add_heading(doc, "Per-user evidence record", 1)
evidence_rows = [
    ("User", "Admin action", "User validation", "Result / date"),
    ("Dominick Vena", "Passkey reset; temp code expired; shared password reset after migration", "Edge, Chrome, approved iPad", "________________"),
    ("Karly Sheriff", "Passkey reset; temp code expired", "Edge and Chrome", "________________"),
    ("Isael Sarmiento", "Existing account restored; password/MFA recovery if required", "Own account and approved iPad", "________________"),
    ("Juan Torres", "Existing account restored; phishing-resistant passkey registered", "Own account and approved iPad", "________________"),
]
table = doc.add_table(rows=len(evidence_rows), cols=4)
set_table_geometry(table, [1650, 3150, 2500, 2060], indent_dxa=120, borders=True)
for ri, row in enumerate(evidence_rows):
    for ci, value in enumerate(row):
        if ri == 0:
            shade_cell(table.cell(ri, ci), LIGHT_BLUE)
            set_cell_text(table.cell(ri, ci), value, bold=True, color=DARK_BLUE, size=9)
        else:
            set_cell_text(table.cell(ri, ci), value, size=9)
set_repeat_table_header(table.rows[0])

add_heading(doc, "Recovery escalation", 1)
add_list(doc, [
    "If a temporary verification code is rejected, confirm it has not expired and that the prompt is an MFA prompt rather than device activation.",
    "If passkey registration never opens, retry in a current WebAuthn-compatible Edge, Chrome, or Safari session with pop-ups and required operating-system passkey services available.",
    "If an existing method cannot be disconnected, confirm the administrator has Manage Multi-Factor Authentication in User Interface.",
    "If the only administrator is locked out, contact Salesforce Support; do not weaken org-wide MFA controls or remove privileged permissions as an emergency shortcut.",
], "bullet")

add_heading(doc, "Authoritative Salesforce guidance", 1)
sources = [
    ("Resolve MFA Access Issues for Your Users", "https://help.salesforce.com/s/articleView?id=sf.mfa_access_recovery.htm&language=en_US&type=5"),
    ("Generate a Temporary Verification Code", "https://help.salesforce.com/s/articleView?id=xcloud.security_temp_id_verification_code_generate.htm&language=en_US&type=5"),
    ("Disconnect a Built-In Authenticator", "https://help.salesforce.com/s/articleView?id=sf.security_built_in_authenticator_remove.htm&language=en_US&type=5"),
    ("Register a Built-In Authenticator", "https://help.salesforce.com/s/articleView?id=xcloud.register_built_in_authenticator.htm&language=en_US&type=5"),
    ("Phishing-Resistant MFA for Privileged Users", "https://help.salesforce.com/s/articleView?id=005321563&language=en_US&type=1"),
]
num_id = make_numbering(doc, "bullet")
for label, url in sources:
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(4)
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    p_pr.append(num_pr)
    add_hyperlink(p, label, url)

add_callout(doc, "Production safety", "This runbook authorizes only user-access recovery when approved by the responsible administrator and coordinated with the affected user. It does not authorize metadata deployment, permission removal, profile changes, or bulk MFA resets.", fill=LIGHT_GRAY, accent=NAVY)

doc.core_properties.title = "Probo Medical Salesforce MFA & Individual Access Recovery"
doc.core_properties.subject = "User-by-user MFA and account recovery runbook"
doc.core_properties.author = "Probo Medical Salesforce Administration"
doc.core_properties.keywords = "Salesforce, MFA, passkey, access recovery, Probo Medical"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
