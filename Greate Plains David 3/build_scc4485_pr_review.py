from __future__ import annotations

import shutil
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Pt

from build_architect_review import (
    AMBER,
    BLUE,
    INK,
    LIGHT_AMBER,
    LIGHT_BLUE,
    LIGHT_RED,
    LIGHT_TEAL,
    MUTED,
    NAVY,
    RED,
    TEAL,
    add_body,
    add_figure,
    add_heading,
    arrow,
    draw_wrapped,
    load_font,
    rounded,
    set_run_font,
)
from build_scc3387_architect_review import (
    configure_document,
    normalize_table_margins,
    review_callout,
    review_table,
)


ROOT = Path(__file__).resolve().parent
REFERENCE = ROOT / "SCC-4179-architect-review.docx"
ASSET_DIR = ROOT / "doc-assets-scc4485"
OUTPUT = ROOT / "SCC-4485-PR-Review.docx"
ASSET_DIR.mkdir(exist_ok=True)


def clear_document_body(document: Document) -> None:
    body = document._body._element
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)


def create_pr_path_diagram(path: Path) -> None:
    image = Image.new("RGB", (1600, 1000), "#F7F9FC")
    draw = ImageDraw.Draw(image)
    title = load_font(42, bold=True)
    subtitle = load_font(22)
    heading = load_font(24, bold=True)
    body = load_font(18)
    small = load_font(16, bold=True)

    draw.text((70, 50), "SCC-4485 | Required PR correction path", font=title, fill="#0B2545")
    draw.text(
        (70, 110),
        "Classify the address once, pass explicit GL values, and map all three Account fields.",
        font=subtitle,
        fill="#5B6777",
    )

    boxes = [
        (
            (55, 225, 350, 605),
            "Buy-flow input",
            ["Place.buildingType", "Place.Type", "Normalize to isMduTenant"],
            "#FFF6DE",
            "#B7791F",
        ),
        (
            (430, 225, 755, 605),
            "Current defect",
            ["Segment always Residential", "Customer Type not supplied", "No MDU condition"],
            "#FDECEC",
            "#C53030",
        ),
        (
            (835, 225, 1160, 605),
            "Required values",
            ["Line of Business: Residential", "Customer Type: Consumer", "Segment: dynamic"],
            "#E7F5F1",
            "#16836B",
        ),
        (
            (1240, 225, 1545, 605),
            "AccountCreation",
            ["Map all three fields", "Keep Consumer-only scope", "Preserve create-only behavior"],
            "#EAF2FA",
            "#2E74B5",
        ),
    ]

    for box, box_title, lines, fill, outline in boxes:
        rounded(draw, box, fill, outline=outline, radius=22, width=4)
        draw.text((box[0] + 24, box[1] + 24), box_title, font=heading, fill="#0B2545")
        y = box[1] + 98
        for line in lines:
            draw.ellipse((box[0] + 26, y + 7, box[0] + 38, y + 19), fill=outline)
            wrapped = draw_wrapped(
                draw,
                (box[0] + 52, y),
                line,
                body,
                "#1F2937",
                box[2] - box[0] - 78,
                spacing=6,
            )
            y += max(76, 27 * len(wrapped) + 34)

    for start, end, label, color in (
        ((350, 415), (430, 415), "READ", "#B7791F"),
        ((755, 415), (835, 415), "FIX", "#C53030"),
        ((1160, 415), (1240, 415), "MAP", "#16836B"),
    ):
        arrow(draw, start, end, color=color, width=6)
        draw.text((start[0] + 12, start[1] - 42), label, font=small, fill=color)

    rounded(draw, (70, 685, 1530, 905), "#EAF2FA", outline="#2E74B5", radius=22, width=3)
    draw.text((105, 720), "PR reviewer reading", font=heading, fill="#0B2545")
    decision = (
        "The PR must replace the universal Residential segment with one deterministic MDU classification, add an "
        "explicit Consumer customer-type mapping, and keep Service Account and existing-account behavior outside "
        "the change. Approval remains blocked until the actual PR diff demonstrates those edits."
    )
    draw_wrapped(draw, (105, 780), decision, body, "#1F2937", 1360, spacing=8)
    image.save(path)


def add_metadata_table(document: Document) -> None:
    metadata = [
        ("Environment", "GreatPlainsMerge | 00DEa00000GkAsHMAV | Sandbox"),
        ("Review date", "28 August 2026"),
        ("Evidence", "Live OmniStudio configuration | Account schema | Aggregate Account read-back"),
        ("Decision", "Request Changes - do not approve or merge"),
        ("Boundary", "PR review only; no deployment, activation, account creation, or org change"),
    ]
    table = document.add_table(rows=len(metadata), cols=2)
    table.style = "Table Grid"
    from build_architect_review import LIGHT_BLUE, set_cell_text, set_table_geometry, shade_cell

    set_table_geometry(table, [1900, 7460])
    normalize_table_margins(table)
    for row, (label, value) in zip(table.rows, metadata):
        shade_cell(row.cells[0], LIGHT_BLUE)
        set_cell_text(row.cells[0], label, size=9.5, color=NAVY, bold=True)
        set_cell_text(row.cells[1], value, size=9.5)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def build_document() -> None:
    diagram_path = ASSET_DIR / "scc4485-pr-correction-path.png"
    create_pr_path_diagram(diagram_path)

    shutil.copy2(REFERENCE, OUTPUT)
    document = Document(OUTPUT)
    clear_document_body(document)
    configure_document(document)

    # Page 1: merge decision and issue summary.
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(2)
    set_run_font(paragraph.add_run("PR REVIEW"), size=10, color=BLUE, bold=True)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(
        paragraph.add_run("SCC-4485 | Consumer Account GL Classification"),
        size=25,
        color=NAVY,
        bold=True,
    )

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(16)
    set_run_font(
        paragraph.add_run("GreatPlainsMerge | OmniStudio implementation and merge-readiness review"),
        size=13,
        color=MUTED,
    )

    add_metadata_table(document)
    review_callout(
        document,
        "Merge decision",
        "Request Changes. The active pipeline has no MDU classification, hard-codes Segment as Residential, and does not populate Customer Type. The story cannot be approved in its current state.",
        fill=LIGHT_RED,
        accent=RED,
    )

    add_heading(document, "Finding summary", 1)
    summary_rows = [
        ("P0", "SQ/OrderPipeline v5", "Blocker", "CreateConsumerAccount always passes Segment = Residential; no MDU branch exists."),
        ("P0", "AccountCreation v1", "Blocker", "No input/output mapping populates Account.Customer_Type__c with Consumer."),
        ("P1", "Version state", "Changes", "AccountCreation v1 is marked inactive while the active pipeline invokes it by bundle name."),
        ("P1", "Change isolation", "Changes", "The same mapper serves Consumer and Service Account actions; the PR must not broaden the new values to Service Accounts."),
        ("P1", "PR evidence", "Blocked", "No PR URL, branch, commit, or exported DataPack diff was supplied for line-by-line review."),
    ]
    review_table(
        document,
        ["Severity", "Location", "Result", "Reviewer finding"],
        summary_rows,
        [1050, 2050, 1400, 4860],
        status_column=2,
    )

    # Page 2: intended code/configuration path.
    document.add_page_break()
    add_heading(document, "Required implementation path", 1)
    add_body(
        document,
        "The PR should derive one MDU flag from the documented Place inputs, calculate the Account segment once, pass all three GL values into CreateConsumerAccount, and map them explicitly in AccountCreation.",
    )
    add_figure(
        document,
        diagram_path,
        "Figure 1. SCC-4485 required PR correction path",
        "Diagram showing buy-flow address inputs normalized to an MDU flag, the current hard-coded and missing-field defects, required GL values, and explicit AccountCreation mappings.",
        width=6.45,
    )
    review_callout(
        document,
        "Reviewer interpretation",
        "The classification belongs before the Data Mapper call. AccountCreation should receive explicit, final values and map them without relying on Account defaults or implicit record-type behavior.",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )

    # Page 3: detailed review findings.
    document.add_page_break()
    add_heading(document, "Detailed PR findings", 1)
    finding_rows = [
        (
            "PR-01",
            "P0",
            "OrderPipeline v5 > Consumer Account action",
            "additionalInput contains LineofBusiness = Residential and Segment = Residential. It contains no Place.buildingType, Place.Type, Unit, Multi-Family, or MDU condition.",
            "Add deterministic MDU classification and pass a derived segment instead of the hard-coded universal value.",
        ),
        (
            "PR-02",
            "P0",
            "AccountCreation v1 > Account output",
            "LineofBusiness maps to Line_of_Business__c and Segment maps to Segment__c. No mapping targets Customer_Type__c.",
            "Pass CustomerType = Consumer and add CustomerType -> Account.Customer_Type__c.",
        ),
        (
            "PR-03",
            "P1",
            "Customer Type field default",
            "The field default is Business, not Consumer, and recent Consumer Accounts are blank. A field default cannot satisfy this story.",
            "Make the value explicit in the pipeline and mapper; do not depend on schema defaults.",
        ),
        (
            "PR-04",
            "P1",
            "AccountCreation version lifecycle",
            "The only AccountCreation record found is Load v1 and IsActive = false. The ticket does not identify an active replacement.",
            "Include the intended version/state change and prove SQ/OrderPipeline resolves to that version after merge.",
        ),
        (
            "PR-05",
            "P1",
            "Shared mapper and create semantics",
            "CreateConsumerAccount and CreateServiceAccount both invoke AccountCreation. The Account Id output remains disabled, indicating create-only behavior.",
            "Keep Customer Type and MDU Segment Consumer-only; do not enable Id/upsert behavior or update existing Accounts.",
        ),
        (
            "PR-06",
            "P2",
            "Component identity",
            "The ticket calls the component SQ_OrderPipeline; the live key is Type SQ / SubType OrderPipeline, Name SQtoOrderPipeline, active v5.",
            "Use the exact exported component key in the PR and review notes to prevent reviewing or packaging the wrong asset.",
        ),
    ]
    review_table(
        document,
        ["ID", "Severity", "Component / location", "Issue", "Required modification"],
        finding_rows,
        [650, 1100, 2150, 2750, 2710],
    )
    review_callout(
        document,
        "Review conclusion",
        "PR-01 and PR-02 are merge blockers. The remaining findings control version safety, change isolation, and reviewability and should be resolved in the same PR.",
        fill=LIGHT_RED,
        accent=RED,
    )

    # Page 4: change specification and merge checklist.
    document.add_page_break()
    add_heading(document, "Required PR modifications", 1)
    change_rows = [
        ("MDU classifier", "No classifier is referenced", "Create isMduTenant from buildingType in {MDU, MDU Tenant, Multi-Family} OR Place.Type = Unit."),
        ("Derived Segment", "Residential for every Consumer", "Set AccountSegment to MDU Tenant when isMduTenant is true; otherwise Residential."),
        ("Consumer inputs", "LineofBusiness and Segment only", "Pass LineofBusiness = Residential, CustomerType = Consumer, and Segment = AccountSegment."),
        ("Data Mapper", "Maps Line of Business and Segment", "Add CustomerType -> Customer_Type__c; retain existing Line and Segment mappings."),
        ("Scope guard", "Shared mapper", "Apply new classification inputs only under the Consumer record-type branch; preserve Service Account inputs."),
        ("Create-only guard", "Account Id output disabled", "Keep Id/upsert mappings disabled so existing accounts remain outside SCC-4485."),
    ]
    review_table(
        document,
        ["Change area", "Current implementation", "Required PR content"],
        change_rows,
        [1900, 2900, 4560],
    )

    add_heading(document, "Implementation shape", 2)
    review_callout(
        document,
        "Pseudocode - adapt to OmniStudio formula syntax",
        "isMduTenant = allowed buildingType OR Place.Type equals Unit; AccountSegment = IF(isMduTenant, MDU Tenant, Residential); CreateConsumerAccount inputs = Residential / Consumer / AccountSegment.",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )

    add_heading(document, "PR merge checklist", 2)
    gate_rows = [
        ("1", "Pipeline diff removes universal Residential segment", "Required"),
        ("2", "Mapper diff adds Customer_Type__c mapping", "Required"),
        ("3", "Consumer-only scope is explicit", "Required"),
        ("4", "AccountCreation active/version target is unambiguous", "Required"),
        ("5", "PR or exported DataPack diff is attached for re-review", "Blocked"),
    ]
    review_table(
        document,
        ["#", "Merge check", "Status"],
        gate_rows,
        [600, 6900, 1860],
        status_column=2,
    )

    # Page 5: evidence, limitation, and PR comment.
    document.add_page_break()
    add_heading(document, "Review evidence and limitations", 1)
    evidence_rows = [
        ("Target org", "GreatPlainsMerge | Great Plains Communications | 00DEa00000GkAsHMAV | Sandbox"),
        ("Active pipeline", "SQ / OrderPipeline | SQtoOrderPipeline v5 | active | 7 active elements"),
        ("Data Mapper", "AccountCreation | Load v1 | inactive | 46 items"),
        ("Account fields", "Line_of_Business__c, Customer_Type__c, Segment__c; required picklist values exist"),
        ("Current outcome signal", "64 Consumer Accounts since 12 August: Residential / blank / Residential; 8 created by Heroku Applink Service"),
        ("Configuration search", "No MDU, buildingType, Multi-Family, Unit, Place, or CustomerType logic in active pipeline elements"),
    ]
    review_table(document, ["Evidence source", "Confirmed result"], evidence_rows, [2200, 7160])

    review_callout(
        document,
        "Review limitation",
        "No repository PR URL, source branch, commit, or exported DataPack diff was provided. This is a PR-readiness review of the live configuration, not a claim that a specific code diff was inspected.",
        fill=LIGHT_AMBER,
        accent=AMBER,
    )

    add_heading(document, "Ready-to-paste PR review comment", 2)
    response = (
        "Request Changes for SCC-4485. In active SQ/OrderPipeline v5, CreateConsumerAccount still passes "
        "LineofBusiness = Residential and Segment = Residential for every Consumer Account and contains no MDU "
        "classification. AccountCreation maps Line of Business and Segment but has no mapping to Customer_Type__c; "
        "the field default is Business and recent Consumer Accounts are blank. Update the PR to derive one "
        "isMduTenant flag from the documented Place values, pass Segment as MDU Tenant or Residential accordingly, "
        "pass CustomerType = Consumer, and add the Customer_Type__c mapping. Keep the change Consumer-only, preserve "
        "Service Account inputs and create-only behavior, and include the exact SQ/OrderPipeline plus AccountCreation "
        "DataPack diff and intended active version for re-review. No deployment, activation, account creation, or org "
        "write was performed during this review."
    )
    review_callout(document, "PR comment", response, fill=LIGHT_BLUE, accent=BLUE)

    document.core_properties.title = "SCC-4485 Consumer Account GL Classification PR Review"
    document.core_properties.subject = "OmniStudio implementation and merge-readiness review"
    document.core_properties.author = "Solution Architecture Review"
    document.core_properties.keywords = (
        "Salesforce, Great Plains, SCC-4485, PR review, OmniStudio, AccountCreation, OrderPipeline, MDU Tenant"
    )
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
