from pathlib import Path

from docx import Document


FILES = [
    Path(r"C:\Users\LIKKI\Downloads\SCC-3386.docx"),
    Path(r"C:\Users\LIKKI\Documents\ChatGPT\Greate Plains David 3\SCC-3385-review-copy.docx"),
]


for path in FILES:
    document = Document(path)
    print(f"\n==== {path.name} ====")
    print(
        f"paragraphs={len(document.paragraphs)} "
        f"tables={len(document.tables)} sections={len(document.sections)}"
    )
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if text:
            print(f"P{index}: {text}")
    for table_index, table in enumerate(document.tables):
        print(f"--TABLE {table_index} {len(table.rows)}x{len(table.columns)}--")
        for row_index, row in enumerate(table.rows):
            values = [cell.text.strip().replace("\n", " | ") for cell in row.cells]
            print(f"R{row_index}: " + " || ".join(values))
