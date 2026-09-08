from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT.parent.parent / "output" / "documents" / "SALDEV-1500_Gearset_Copado_MuleSoft_Implementation_Plan.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)


# standard_business_brief preset, with customer_pack first-page pattern.
PAGE_W = 8.5
PAGE_H = 11.0
MARGIN = 1.0
CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120

FONT = "Calibri"
INK = "24313F"
NAVY = "163A5F"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "657382"
PALE_BLUE = "E8F0F7"
PALE_GRAY = "F2F4F7"
PALE_GREEN = "EAF5EE"
PALE_AMBER = "FFF4D6"
PALE_RED = "FCEBEC"
GREEN = "26734D"
AMBER = "8A6400"
RED = "A13A3A"
WHITE = "FFFFFF"
BORDER = "C9D2DC"


def set_run_font(run, name=FONT, size=None, color=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    tr_pr.append(flag)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=TABLE_INDENT_DXA):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_keep_with_next(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepNext"))
    if node is None:
        node = OxmlElement("w:keepNext")
        p_pr.append(node)
    node.set(qn("w:val"), "1" if value else "0")


def set_keep_together(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepLines"))
    if node is None:
        node = OxmlElement("w:keepLines")
        p_pr.append(node)
    node.set(qn("w:val"), "1" if value else "0")


def set_row_cant_split(row):
    """Keep a table row together on one page in Word."""
    tr_pr = row._tr.get_or_add_trPr()
    node = tr_pr.find(qn("w:cantSplit"))
    if node is None:
        node = OxmlElement("w:cantSplit")
        tr_pr.append(node)
    node.set(qn("w:val"), "1")


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), MUTED)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    rpr.append(color)
    rpr.append(size)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


def add_numbering(doc, abstract_id, num_id, num_fmt, lvl_text, left=720, hanging=360):
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), num_fmt)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    ppr.append(tabs)
    ppr.append(ind)
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    rpr.append(fonts)
    lvl.extend([start, fmt, text, suff, ppr, rpr])
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True


def add_body(doc, text, *, bold_prefix=None, italic=False, after=6, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, size=11, color=INK, bold=True)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest, size=11, color=INK, italic=italic)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11, color=INK, italic=italic)
    if keep:
        set_keep_together(p)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph()
    if text.startswith("The MuleSoft pipeline packages once"):
        p.paragraph_format.page_break_before = True
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    p.paragraph_format.left_indent = Pt(24)
    p.paragraph_format.first_line_indent = Pt(-14)
    run = p.add_run(f"•\u00a0\u00a0{text}")
    set_run_font(run, size=11, color=INK)
    set_keep_together(p)
    return p


def add_step(doc, text):
    add_step.counter += 1
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    p.paragraph_format.left_indent = Pt(28)
    p.paragraph_format.first_line_indent = Pt(-22)
    marker = p.add_run(f"{add_step.counter}. ")
    set_run_font(marker, size=11, color=BLUE, bold=True)
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)
    set_keep_together(p)
    return p


add_step.counter = 0


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(text)
    set_run_font(run, size={1:16,2:13,3:12}[level], color={1:BLUE,2:BLUE,3:DARK_BLUE}[level], bold=True)
    set_keep_with_next(p)
    set_keep_together(p)
    return p


def add_callout(doc, label, text, tone="blue"):
    colors = {
        "blue": (PALE_BLUE, NAVY),
        "green": (PALE_GREEN, GREEN),
        "amber": (PALE_AMBER, AMBER),
        "red": (PALE_RED, RED),
        "gray": (PALE_GRAY, INK),
    }
    fill, accent = colors[tone]
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(8)
    p.paragraph_format.right_indent = Pt(8)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.10
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "8")
        node.set(qn("w:space"), "5")
        node.set(qn("w:color"), accent)
        borders.append(node)
    p_pr.append(borders)
    r = p.add_run(label)
    set_run_font(r, size=10.5, color=accent, bold=True)
    r.add_break()
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=INK)
    set_keep_together(p)
    return p


def add_table(doc, headers, rows, widths, *, font_size=9, center_cols=None, header_fill=PALE_GRAY):
    center_cols = set(center_cols or [])
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    set_repeat_header(table.rows[0])
    set_row_cant_split(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade_cell(cell, header_fill)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=110, bottom=110)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in center_cols else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(str(header))
        set_run_font(run, size=font_size, color=NAVY, bold=True)
        set_keep_with_next(p)
    for row_data in rows:
        row = table.add_row()
        set_row_cant_split(row)
        for idx, value in enumerate(row_data):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=100, bottom=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in center_cols else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(str(value))
            set_run_font(run, size=font_size, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_metadata_grid(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    set_repeat_header(table.rows[0])
    set_row_cant_split(table.rows[0])
    for idx, text in enumerate(("Plan detail", "Value")):
        cell = table.rows[0].cells[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=80, bottom=80)
        shade_cell(cell, PALE_GRAY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        set_run_font(run, size=9, color=NAVY, bold=True)
        set_keep_with_next(p)
    for label, value in rows:
        row = table.add_row()
        set_row_cant_split(row)
        for idx, text in enumerate((label, value)):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=80, bottom=80)
            if idx == 0:
                shade_cell(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            set_run_font(run, size=9.5, color=NAVY if idx == 0 else INK, bold=(idx == 0))
    set_table_geometry(table, [2200, 7160])
    set_table_borders(table, color=BORDER, size="6")
    return table


def add_link_paragraph(doc, label, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(f"{label}: ")
    set_run_font(r, size=9.5, color=INK, bold=True)
    r2 = p.add_run(url)
    set_run_font(r2, size=9.5, color=BLUE)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(PAGE_W)
section.page_height = Inches(PAGE_H)
section.top_margin = Inches(MARGIN)
section.bottom_margin = Inches(MARGIN)
section.left_margin = Inches(MARGIN)
section.right_margin = Inches(MARGIN)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)
configure_styles(doc)

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hp.paragraph_format.space_after = Pt(0)
hr = hp.add_run("SALDEV-1500 | Implementation Plan")
set_run_font(hr, size=9, color=MUTED, bold=True)
footer = section.footer
fp = footer.paragraphs[0]
add_page_field(fp)


# Cover - customer_pack pattern, built from native Word elements only.
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(12)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("GTM SYSTEMS DELIVERY PLAN")
set_run_font(r, size=10, color=BLUE, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
r = p.add_run("SALDEV-1500 Implementation Plan")
set_run_font(r, size=28, color=NAVY, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(20)
r = p.add_run("Gearset, Copado and MuleSoft CI/CD rollout")
set_run_font(r, size=14, color=MUTED)

add_metadata_grid(doc, [
    ("Prepared for", "Flywire GTM Systems, Salesforce, Integration and DevOps teams"),
    ("Prepared on", "2 September 2026"),
    ("Status", "Draft for stakeholder review and implementation approval"),
    ("Delivery horizon", "12-week controlled rollout, subject to access and licensing"),
    ("Source ticket", "SALDEV-1500"),
])

doc.add_paragraph().paragraph_format.space_after = Pt(8)
add_callout(
    doc,
    "Implementation decision",
    "Use Gearset for Salesforce pull-request validation, metadata analysis and lower-environment delivery; use Copado for governed Salesforce release promotion, approvals and production deployment; use GitHub Actions and Anypoint for MuleSoft build, MUnit testing, packaging and runtime promotion. GitHub remains the shared source of truth and Jira remains the work and evidence record.",
    "blue",
)
add_callout(
    doc,
    "Important control",
    "There will be one production deployment owner per platform. Copado is the only Salesforce production deployer. The protected MuleSoft pipeline is the only Anypoint production deployer. Gearset production jobs remain validation-only. This prevents two tools from deploying different versions of the same change.",
    "amber",
)

add_body(doc, "This is a separate implementation plan. It does not replace the SALDEV-1500 operating model and does not make changes to Jira, GitHub, Gearset, Copado, MuleSoft or Salesforce by itself.", italic=True, after=8)
add_body(doc, "Design note: the document uses only native Word text, tables and styles. No Figma design work, exports or Figma assets are included.", italic=True, after=0)

doc.add_page_break()


add_heading(doc, "1. What this plan is meant to achieve", 1)
add_body(doc, "The goal is not to install three tools and call the program complete. The goal is to give delivery teams one understandable route from a Jira story to production, even when a release includes both Salesforce metadata and a MuleSoft integration.")
add_body(doc, "A developer should know where to commit, which checks will run, who can approve a promotion and where to find the final evidence. A release manager should be able to identify the exact Salesforce commit, MuleSoft artifact version and deployment jobs without reconstructing the release from chat messages.")

add_heading(doc, "Outcomes", 2)
for item in [
    "A controlled Git-based delivery path for Salesforce and MuleSoft work.",
    "Automated Salesforce comparison, dependency analysis, validation and QA deployment through Gearset.",
    "Copado-managed Salesforce release packaging, approval, UAT and production promotion.",
    "MuleSoft CI using Maven and MUnit, with versioned assets published to Exchange and deployed through protected Anypoint environments.",
    "A Jira board whose status reflects real evidence rather than optimistic manual updates.",
    "One cross-platform release manifest linking Jira, Git commits, Gearset jobs, Copado promotions and MuleSoft artifacts.",
    "Documented rollback, back-promotion, support ownership and operating metrics.",
]:
    add_bullet(doc, item)

add_heading(doc, "Success looks like this", 2)
add_callout(doc, "Pilot definition", "One low-risk release moves through Salesforce and MuleSoft development, automated validation, QA, UAT and protected production without direct changes, scope substitution or missing evidence. A deliberate test failure is blocked, corrected and re-run before the pilot is accepted.", "green")


add_heading(doc, "2. Scope, boundaries and assumptions", 1)
add_heading(doc, "In scope", 2)
for item in [
    "Jira workflow, evidence fields, transition ownership and board views for SALDEV delivery.",
    "GitHub repository structure, branch protection, pull-request templates and release manifest.",
    "Gearset team connections, metadata filters, PR validation, QA CI jobs and production validation-only job.",
    "Copado governance/training environment, Salesforce pipeline, user-story/release configuration, approvals, promotions and back-promotion.",
    "MuleSoft repository standards, MUnit quality gates, Exchange publishing and Anypoint deployment jobs.",
    "Cross-platform test strategy, security controls, observability, rollback and handover.",
]:
    add_bullet(doc, item)

add_heading(doc, "Not included in this document", 2)
for item in [
    "Procurement approval or a statement that current Gearset, Copado or MuleSoft licenses are sufficient.",
    "A decision on CloudHub 2.0, Runtime Fabric or on-premises Mule runtime; the current target must be confirmed during discovery.",
    "Bulk migration of every existing repository or pipeline before the pilot proves the approach.",
    "Direct production changes, credential creation, package installation or tool configuration.",
    "Figma prototypes, diagrams, exports or design-system work.",
]:
    add_bullet(doc, item)

add_heading(doc, "Assumptions to validate in week 1", 2)
add_table(doc, ["Assumption", "Why it matters", "Owner"], [
    ("GitHub Enterprise is available and approved", "It is the common source of truth and branch-control layer.", "Platform/DevOps"),
    ("Gearset licenses include required CI and team-shared capabilities", "The design uses shared Git-to-org jobs and PR validation.", "Tool owner"),
    ("Copado edition and pipeline mode are confirmed", "Source Format and Metadata Pipelines have different capabilities and limitations.", "Copado owner"),
    ("A Copado governance or training environment is available", "The team needs a safe place to prove pipeline configuration before live use.", "Salesforce platform"),
    ("Anypoint runtime target and business group are known", "Packaging, credentials and deployment settings depend on the target.", "Integration lead"),
    ("Jira administrators can create fields, statuses and automation", "The evidence workflow cannot be implemented without project-level control.", "Jira owner"),
], [3100, 4560, 1700], font_size=8.7)


add_heading(doc, "3. Tool responsibilities", 1)
add_body(doc, "The tools overlap in places. The plan makes the boundaries explicit so a release cannot be promoted independently by two systems.")
add_table(doc, ["Platform", "Primary responsibility", "What it must not become"], [
    ("Jira", "Scope, acceptance criteria, ownership, risk, approvals, status and final evidence", "A replacement for Git history or deployment logs"),
    ("GitHub", "Canonical source, branch protection, pull requests, immutable commit references and release manifest", "A place for unreviewed environment-specific secrets"),
    ("Gearset", "Salesforce metadata compare, dependency analysis, PR validation, QA delivery and production validation", "A second Salesforce production deployment path"),
    ("Copado", "Salesforce release orchestration, promotion bundle, UAT/production approval, deployment and back-promotion", "A parallel source of truth disconnected from GitHub"),
    ("GitHub Actions", "Repository-native checks and MuleSoft CI/build automation", "An unapproved production bypass"),
    ("MuleSoft Anypoint", "API assets, Mule runtime deployment, environment configuration, policies, logs and monitoring", "A source repository for unreviewed changes"),
    ("Confluence", "Published runbooks, standards, training and operational guidance", "The authoritative status of a live release"),
], [1700, 4650, 3010], font_size=8.2)

add_heading(doc, "Dual-tool Salesforce rule", 2)
add_callout(doc, "Gearset + Copado", "Gearset owns CI and technical validation. Copado owns controlled release promotion. Gearset may deploy automatically to Integration and QA, but its UAT and Production jobs remain disabled or validation-only. Copado deploys the exact approved Git commit to UAT and Production. Every Copado promotion records the corresponding Gearset validation job.", "blue")

add_heading(doc, "MuleSoft release rule", 2)
add_callout(doc, "Copado + Anypoint", "Copado coordinates the release and approval record; the protected MuleSoft pipeline performs the Anypoint deployment. A direct Copado-to-Anypoint trigger is optional and must not be assumed until the selected Copado pipeline supports the integration securely. The default pilot uses paired approvals and shared release evidence rather than an unproven custom callout.", "amber")


add_heading(doc, "4. Target delivery architecture", 1)
add_body(doc, "The delivery chain is deliberately simple enough to explain in one line:")
add_callout(doc, "End-to-end path", "Jira story -> GitHub branch and pull request -> Gearset Salesforce checks + MuleSoft GitHub Actions checks -> approved merge -> Gearset QA + MuleSoft QA deployment -> integrated test -> Copado release approval -> Salesforce and MuleSoft protected production deployment -> Jira evidence and closure", "gray")

add_heading(doc, "Repository model", 2)
add_table(doc, ["Repository", "Contents", "Required controls"], [
    ("gtm-salesforce", "Salesforce DX source, manifests, Apex/LWC tests, permission sets and deployment notes", "CODEOWNERS, PR validation, protected integration/uat/main branches"),
    ("gtm-mulesoft-integrations", "Mule applications, RAML/OAS, DataWeave, MUnit tests, POM files and properties templates", "MUnit and package checks, protected main, versioned release tags"),
    ("gtm-release-config", "Cross-platform release manifests, environment mappings and runbook references", "Release-manager approval and immutable production tags"),
], [2300, 4450, 2610], font_size=8.5)

add_heading(doc, "Release manifest", 2)
add_body(doc, "Each coordinated release gets one manifest. It may be a versioned YAML/JSON file or a controlled Jira release record, but it must capture the same identifiers every time.")
for item in [
    "Jira release and included stories.",
    "Salesforce repository, commit SHA, Gearset validation job and Copado promotion/deployment ID.",
    "MuleSoft repository, commit SHA, packaged artifact version, Exchange asset version and Anypoint deployment job.",
    "Target environments, approvers, test results, manual steps, rollback point and final outcome.",
]:
    add_bullet(doc, item)

add_heading(doc, "Environment mapping", 2)
add_table(doc, ["Stage", "Salesforce path", "MuleSoft path", "Promotion owner"], [
    ("Developer", "Personal sandbox or scratch org", "Local/isolated development", "Developer"),
    ("Integration", "Gearset deploys approved integration branch", "Pipeline deploys snapshot to Mule Dev", "Dev lead"),
    ("QA", "Gearset deploys approved QA candidate", "Pipeline deploys versioned candidate to Mule QA", "QA lead"),
    ("UAT", "Copado deploys approved release candidate", "Protected pipeline promotes same artifact to Mule UAT", "Business owner + release manager"),
    ("Production", "Copado deploys immutable approved commit", "Protected pipeline deploys immutable artifact to Anypoint Prod", "Release manager"),
], [1400, 2860, 3170, 1930], font_size=8.2)


add_heading(doc, "5. Delivery process from story to production", 1)
steps = [
    "Refine the Jira story. Confirm the business outcome, affected Salesforce components, MuleSoft APIs/flows, data and security impact, test personas, dependencies and rollback outline.",
    "Create linked ticket branches. Use the same Jira key in the Salesforce and MuleSoft repositories when both are affected.",
    "Build in isolation. Salesforce work is captured in DX source format; MuleSoft work includes the API contract, Mule configuration, DataWeave, POM changes and MUnit coverage.",
    "Open focused pull requests. The PR lists added, changed and deleted components, manual steps, risk class and cross-repository dependencies.",
    "Run CI. Gearset validates Salesforce metadata and tests against the target org. GitHub Actions runs Mule Maven packaging, MUnit, coverage and API-contract checks.",
    "Review and merge only after both streams pass. The release manifest pins the approved Salesforce SHA and MuleSoft artifact version.",
    "Deploy to Integration and QA. Gearset delivers Salesforce changes; the protected Mule pipeline deploys the packaged Mule artifact. Rebuilding a different artifact during promotion is not allowed.",
    "Run integrated QA. Test the API contract, authentication, retries, error mapping, idempotency, data behavior, permissions and Salesforce user journeys.",
    "Approve UAT and release. Copado groups the Salesforce stories and records the release approval. The paired MuleSoft artifact and Anypoint job are linked in the same release manifest.",
    "Deploy through protected production paths. Copado deploys Salesforce; the MuleSoft production job deploys to Anypoint. Both use named service identities and retain immutable evidence.",
    "Verify and close. Complete smoke checks, monitor logs and alerts, reconcile any manual actions, link known issues and write the evidence back to Jira before Done.",
]
for step in steps:
    add_step(doc, step)


add_heading(doc, "6. Twelve-week implementation roadmap", 1)
add_table(doc, ["Phase", "Weeks", "Main work", "Exit condition"], [
    ("0. Decide and discover", "1-2", "Licenses, current pipelines, runtime target, repositories, environments, security and owner confirmation", "Signed design decisions and pilot scope"),
    ("1. Foundation", "2-4", "GitHub controls, Jira workflow, service identities, baseline metadata and repository templates", "Controlled repositories and ready backlog"),
    ("2. Salesforce CI", "4-6", "Gearset connections, filters, PR validation, Integration/QA jobs and production validation-only job", "Salesforce non-prod path proven"),
    ("3. Governed Salesforce CD", "6-8", "Copado training pipeline, release/user-story model, UAT/Prod promotion, approvals and back-promotion", "Salesforce release path proven"),
    ("4. MuleSoft delivery", "5-8", "Maven/MUnit CI, Exchange publishing, Anypoint Dev/QA/UAT/Prod deployment and monitoring", "Immutable Mule artifact path proven"),
    ("5. Integrated pilot", "9-10", "One low-risk cross-platform change, failure test, QA/UAT, rollback rehearsal and production smoke", "Pilot accepted; critical gaps closed"),
    ("6. Handover and scale", "11-12", "Training, runbooks, support model, metrics, backlog and adoption plan", "BAU ownership accepted"),
], [1700, 900, 4550, 2210], font_size=8.0, center_cols={1})

add_heading(doc, "Implementation sequence", 2)
add_body(doc, "Gearset and MuleSoft CI can be built in parallel after the repository and service-account foundations are ready. Copado production work starts only after the team has confirmed which pipeline mode is licensed and tested the intended features in a training pipeline.")


add_heading(doc, "7. Workstream plan", 1)
add_heading(doc, "Workstream A - Governance and access", 2)
for item in [
    "Name the product owner, Salesforce lead, Gearset owner, Copado owner, MuleSoft lead, QA lead, security approver and release manager.",
    "Create service identities with least privilege and non-personal ownership. Record renewal and offboarding procedures.",
    "Approve repository access, connected applications, Anypoint connected app, IP allowlisting and secret storage.",
    "Document break-glass access, approval rules and monthly access review.",
]:
    add_bullet(doc, item)

add_heading(doc, "Workstream B - GitHub and Jira foundation", 2)
for item in [
    "Create or validate repositories, CODEOWNERS, branch naming and protected branches.",
    "Add PR templates for Salesforce, MuleSoft and coordinated releases.",
    "Configure Jira statuses, evidence fields, blocked/rework paths and release approval fields.",
    "Link Jira keys, branches, pull requests, Gearset jobs, Copado promotions and MuleSoft jobs.",
]:
    add_bullet(doc, item)

add_heading(doc, "Workstream C - Gearset Salesforce CI", 2)
for item in [
    "Create team-shared Git and org connections using approved service accounts.",
    "Create one reviewed metadata filter and explicit permission/destructive-change policy.",
    "Configure PR validation against the correct target org with Apex, LWC and static-analysis gates.",
    "Configure branch-triggered Integration and QA deployment jobs.",
    "Configure Production as validation-only and restrict deployment access.",
    "Store job history, component scope, test results and omissions as release evidence.",
]:
    add_bullet(doc, item)

add_heading(doc, "Workstream D - Copado Salesforce CD", 2)
for item in [
    "Confirm Source Format versus Metadata Pipelines against required capabilities and current licensing.",
    "Set up a governance/training environment before associating live releases.",
    "Configure environments, credentials, branch mapping, user-story/release model and approval owners.",
    "Build UAT and Production promotions from approved Git references only.",
    "Configure deployment steps, manual tasks, back-promotion and evidence retention.",
    "Prove failed promotion, rejected approval, rollback decision and back-promotion behavior.",
]:
    add_bullet(doc, item)

add_heading(doc, "Workstream E - MuleSoft CI/CD", 2)
for item in [
    "Confirm the Anypoint organization, business group, runtime plane, regions, environment names and deployment target.",
    "Standardize POM files, Mule Maven Plugin configuration, MUnit tests, coverage thresholds and properties templates.",
    "Run API specification validation, Maven packaging, MUnit and security/dependency checks on every pull request.",
    "Publish approved versioned assets to Exchange and retain the artifact coordinates in the release manifest.",
    "Deploy the same packaged artifact through Mule Dev, QA, UAT and Production using protected environment approvals.",
    "Configure API Manager policies, Runtime Manager alerts, log access and post-deployment health checks.",
]:
    add_bullet(doc, item)


add_heading(doc, "8. CI/CD quality gates", 1)
add_table(doc, ["Gate", "Salesforce checks", "MuleSoft checks", "Pass evidence"], [
    ("Pull request", "Gearset compare, dependencies, metadata validation, Apex/LWC tests, static analysis", "API spec lint, Maven package, MUnit, coverage, dependency/security scan", "Both checks pass against linked commits"),
    ("Integration", "Gearset deployment and integration checks", "Deploy snapshot to Mule Dev and run contract tests", "Job IDs and results linked"),
    ("QA", "Regression, permissions, negative paths and impacted automation", "Contract, transformation, retry, error and performance smoke tests", "No open release blocker"),
    ("UAT", "Copado candidate with named business acceptance", "Same artifact in Mule UAT with business/API consumer acceptance", "Named sign-off for the paired release"),
    ("Production readiness", "Copado validation of immutable commit and approved scope", "Protected deployment job for immutable artifact and config review", "Release manager approval and rollback readiness"),
    ("Post-production", "Salesforce smoke, monitoring and business validation", "Runtime health, API policy, log/alert and downstream checks", "Evidence complete before Done"),
], [1300, 2920, 2920, 2220], font_size=7.8)

add_heading(doc, "Minimum MuleSoft test expectations", 2)
for item in [
    "MUnit covers core success, error and retry paths; coverage thresholds are agreed by the integration lead and enforced in CI.",
    "API contract tests prove request/response schemas, status codes and backward compatibility for published consumers.",
    "Secrets and environment properties are injected at deployment and are never stored in source or Jira.",
    "The pipeline verifies the deployed application reaches a healthy state; skipping deployment verification requires explicit approval and evidence.",
    "Performance or volume testing is required for high-throughput, batch or latency-sensitive flows.",
]:
    add_bullet(doc, item)


add_heading(doc, "9. Jira workflow and evidence", 1)
add_body(doc, "The board should be readable by people outside DevOps. Status names describe the work state; transition checks provide the technical control.")
add_table(doc, ["Status", "Entry condition", "Required evidence", "Owner"], [
    ("Backlog", "Request accepted for triage", "Problem, value and affected platforms", "Product owner"),
    ("Ready for Development", "Definition of Ready met", "Acceptance, architecture, dependencies, test and rollback", "Product + architect"),
    ("In Progress", "Linked branch exists", "Assignee and repository links", "Developer"),
    ("Code/Peer Review", "Pull request open", "Diff, checks, risk and reviewer", "CODEOWNER"),
    ("Ready for QA", "Both platform CI gates pass", "Gearset and MuleSoft CI evidence", "Dev lead"),
    ("In QA", "Paired candidate deployed", "Salesforce and Mule job IDs", "QA lead"),
    ("Ready for UAT", "QA exit met", "Results and defect disposition", "QA lead"),
    ("In UAT", "Immutable candidate available", "Named testers and version manifest", "Business owner"),
    ("Ready for Release", "UAT approved", "Sign-off, risk, runbook and rollback", "Business + architect"),
    ("Scheduled", "Window and approvals confirmed", "Copado release and Mule production job reference", "Release manager"),
    ("Deployed", "Both production jobs succeed", "Exact job IDs, commits and artifact version", "Release manager"),
    ("Done", "Smoke and business checks pass", "Final evidence, known issues and sign-off", "Product owner"),
], [1550, 2460, 3760, 1590], font_size=7.6)

add_callout(doc, "Closure rule", "Deployment success is not Done. Done requires verified behavior, complete cross-platform evidence and reconciliation to Git and the release manifest.", "amber")


add_heading(doc, "10. Security and access model", 1)
add_table(doc, ["Control", "Implementation requirement", "Evidence"], [
    ("Service identities", "Separate Gearset, Copado and Anypoint/GitHub identities; no personal production credentials", "Owner, purpose and expiry review"),
    ("Least privilege", "Validation identities cannot deploy; production jobs have only required target access", "Permission review and test"),
    ("Secrets", "GitHub environments or approved vault; masked output; no secrets in POM, repo, Jira or deployment notes", "Secret scan and configuration record"),
    ("Approvals", "Protected branches, Copado release approval and GitHub protected MuleSoft environment", "Named approver and timestamp"),
    ("Network", "Review connected apps, callback URLs, IP allowlists and Anypoint control-plane access", "Security sign-off"),
    ("Audit", "Retain PR, Gearset, Copado, GitHub Actions and Anypoint logs according to policy", "Retrievable pilot evidence"),
], [1800, 5150, 2410], font_size=8.2)


add_heading(doc, "11. Pilot, cutover and rollback", 1)
add_heading(doc, "Pilot selection", 2)
add_body(doc, "Choose a change that is small enough to reverse but real enough to exercise the integration. It should include at least one Salesforce metadata change and one MuleSoft contract or transformation change, without billing, authentication or destructive data impact.")

add_heading(doc, "Pilot execution", 2)
for item in [
    "Run the full story, branch, PR, CI, QA, UAT, release and production path.",
    "Introduce one safe CI failure and confirm the pipeline blocks promotion.",
    "Rehearse Salesforce rollback or forward-fix and MuleSoft artifact rollback in non-production.",
    "Confirm both production deployments reference the approved manifest and no artifact was rebuilt.",
    "Measure lead time, waiting time, manual touches, failure points and evidence completeness.",
]:
    add_bullet(doc, item)

add_heading(doc, "Rollback model", 2)
add_table(doc, ["Platform", "Primary rollback", "Decision point"], [
    ("Salesforce", "Copado rollback/back-promotion or approved forward-fix from the production tag; reconcile Git immediately", "Smoke failure, data risk or release-manager decision"),
    ("MuleSoft", "Redeploy the last known-good immutable artifact and restore approved properties/policies", "Health check, API error rate or downstream failure"),
    ("Coordinated release", "Rollback both platforms when the interface contract is incompatible; otherwise isolate only the failed side after architect review", "Release manager + architect + business owner"),
], [1700, 4970, 2690], font_size=8.3)

add_heading(doc, "Production cutover checklist", 2)
for item in [
    "Release manifest approved and frozen.",
    "Gearset production validation and Copado Salesforce validation are successful and current.",
    "MuleSoft artifact is published, tested and promotable without rebuild.",
    "UAT sign-off, change window, operator, monitoring and rollback owner are confirmed.",
    "Manual steps are ordered, timed and assigned.",
    "Post-deployment smoke tests and business contacts are ready.",
]:
    add_bullet(doc, item)


add_heading(doc, "12. Roles and operating cadence", 1)
add_table(doc, ["Activity", "Product", "Architect", "Salesforce", "MuleSoft", "QA", "Release"], [
    ("Refinement and priority", "A/R", "C", "C", "C", "C", "I"),
    ("Technical design", "C", "A", "R", "R", "C", "I"),
    ("Build and unit test", "I", "C", "R", "R", "C", "I"),
    ("CI and review", "I", "A", "R", "R", "C", "I"),
    ("Integrated QA", "I", "C", "C", "C", "A/R", "I"),
    ("UAT approval", "A/R", "C", "I", "I", "C", "I"),
    ("Release readiness", "C", "C", "C", "C", "C", "A/R"),
    ("Production deployment", "I", "C", "R", "R", "I", "A"),
    ("Post-deploy verification", "A", "C", "R", "R", "R", "R"),
], [2350, 1000, 1100, 1280, 1280, 1000, 1350], font_size=7.4, center_cols={1,2,3,4,5,6})
add_body(doc, "R = Responsible, A = Accountable, C = Consulted, I = Informed.", italic=True, after=8)

add_heading(doc, "Cadence", 2)
add_table(doc, ["Meeting", "Purpose", "Participants"], [
    ("Weekly implementation review", "Remove setup blockers, confirm decisions and track workstream exits", "Workstream owners"),
    ("Daily pipeline review during pilot", "Review failed jobs, drift, approvals and cross-platform dependencies", "Dev, QA, release"),
    ("Release readiness review", "Confirm scope, evidence, rollback and operators before UAT/Prod", "Product, architect, QA, release"),
    ("Post-release review", "Confirm outcome, incidents, metrics and backlog actions", "All accountable owners"),
    ("Monthly control review", "Access, bypasses, evidence completeness, drift and tool health", "Platform, security, audit"),
], [2300, 4350, 2710], font_size=8.3)


add_heading(doc, "13. Risks, decisions and mitigations", 1)
add_table(doc, ["Risk", "Impact", "Mitigation", "Owner"], [
    ("Gearset and Copado both deploy to Production", "Conflicting source and incomplete audit trail", "Copado-only Salesforce production deployment; Gearset validation-only", "Release manager"),
    ("Copado pipeline mode does not support a required feature", "Custom work or delayed rollout", "Capability matrix and training pipeline before design approval", "Copado owner"),
    ("MuleSoft production integration is assumed rather than proven", "Unsafe or unsupported trigger path", "Use protected paired approval first; automate only after proof", "Integration lead"),
    ("Salesforce and MuleSoft releases drift", "API contract mismatch", "Cross-platform release manifest and integrated QA gate", "Architect"),
    ("Secrets appear in repositories or logs", "Security incident", "Connected apps, secret store, masking and scans", "Security"),
    ("Permissions are deployed unintentionally", "Access regression or overexposure", "Explicit Gearset filter and Copado permission policy; persona tests", "Salesforce lead"),
    ("Team cannot retrieve evidence later", "Audit and support failure", "Evidence schema, retention test and monthly sampling", "Delivery lead"),
], [2300, 2250, 3490, 1320], font_size=7.8)

add_heading(doc, "Decisions required before build", 2)
for item in [
    "Confirm Gearset edition, team ownership and allowed environments.",
    "Confirm Copado pipeline type, package versions, governance environment and licensing.",
    "Confirm MuleSoft runtime target, region, business group, API Manager model and deployment identity.",
    "Approve which tool owns Salesforce UAT; this plan assigns UAT and Production to Copado.",
    "Approve the pilot story, release window and named accountable owners.",
]:
    add_bullet(doc, item)


add_heading(doc, "14. Implementation backlog", 1)
backlog = [
    ("IMP-01", "Confirm tool editions, pipeline modes and runtime targets", "Architecture + tool owners", "3"),
    ("IMP-02", "Create GitHub repository controls and release manifest", "Platform/DevOps", "8"),
    ("IMP-03", "Configure Jira workflow, evidence fields and automation", "Jira admin", "13"),
    ("IMP-04", "Provision service identities, secrets and access", "Security + platform", "8"),
    ("IMP-05", "Configure Gearset shared connections and metadata policy", "Salesforce lead", "8"),
    ("IMP-06", "Implement Gearset PR validation and non-prod CI jobs", "Salesforce + DevOps", "13"),
    ("IMP-07", "Build Copado training pipeline and user-story/release model", "Copado owner", "13"),
    ("IMP-08", "Implement Copado UAT, production and back-promotion path", "Release + Salesforce", "13"),
    ("IMP-09", "Standardize MuleSoft repository, Maven and MUnit controls", "MuleSoft lead", "8"),
    ("IMP-10", "Implement Anypoint Dev-to-Prod deployment pipeline", "MuleSoft + DevOps", "13"),
    ("IMP-11", "Implement cross-platform release manifest and evidence links", "Architecture + Jira", "8"),
    ("IMP-12", "Define integrated QA, security and rollback tests", "QA + security", "8"),
    ("IMP-13", "Run and accept the cross-platform pilot", "Delivery team", "13"),
    ("IMP-14", "Publish runbooks, train users and transition to BAU", "Delivery lead", "8"),
]
add_table(doc, ["Draft ID", "Story", "Owner", "Points"], backlog, [1200, 5200, 1960, 1000], font_size=8.2, center_cols={0,3})
add_body(doc, "The draft IDs are planning references, not existing Jira keys. Create the stories under an implementation epic linked to SALDEV-1500 and re-estimate them with the delivery team.", italic=True)


add_heading(doc, "15. Acceptance and handover", 1)
add_heading(doc, "Program acceptance criteria", 2)
for item in [
    "The same Jira key links the Salesforce and MuleSoft work when both platforms are changed.",
    "Protected branches reject a failing Salesforce or MuleSoft pull request.",
    "Gearset validates and deploys Salesforce to lower environments with complete job evidence.",
    "Copado promotes the approved Salesforce commit through UAT and Production with named approvals.",
    "The MuleSoft pipeline packages once and promotes the same immutable artifact through Anypoint environments.",
    "Integrated QA proves the Salesforce-to-MuleSoft contract, security, error behavior and negative paths.",
    "A rollback rehearsal succeeds in non-production and its decision path is documented.",
    "Production evidence can be retrieved from Jira without relying on individual inboxes or chat history.",
    "Runbooks, access owners, support contacts and monthly control measures are accepted by BAU owners.",
]:
    add_bullet(doc, item)

add_heading(doc, "Handover package", 2)
add_table(doc, ["Deliverable", "Minimum content"], [
    ("Architecture decision record", "Tool boundaries, pipeline mode, environment map and exception decisions"),
    ("Configuration register", "Connections, jobs, filters, branches, credentials, approvals and owners"),
    ("Runbooks", "Normal release, hotfix, rollback, back-promotion, failed job and access recovery"),
    ("Test pack", "CI negative tests, integrated QA, UAT, smoke and rollback evidence"),
    ("Operating dashboard", "Lead time, failed change rate, blocked work, evidence completeness and drift"),
    ("Training materials", "Role-based steps for developers, reviewers, QA, approvers and release managers"),
], [2500, 6860], font_size=8.7)

add_callout(doc, "Recommended approval", "Approve the dual-tool Salesforce model and MuleSoft workstream for a controlled pilot, subject to week-1 confirmation of licenses, Copado pipeline mode, Anypoint runtime target and named owners. Do not authorize production configuration from this document alone.", "green")


add_heading(doc, "Appendix A - Definition of Ready and Done", 1)
add_heading(doc, "Definition of Ready", 2)
for item in [
    "Business outcome, impacted personas and measurable acceptance criteria are clear.",
    "Salesforce components, MuleSoft APIs/flows, dependencies, data and security impacts are identified.",
    "Architecture review, change class, release target, test strategy and rollback owner are set.",
    "Required environments, service identities and upstream prerequisites are available or tracked.",
    "The story is small enough to review, test and reverse independently where practical.",
]:
    add_bullet(doc, item)

add_heading(doc, "Definition of Done", 2)
for item in [
    "Approved source is merged and Jira, PRs, commits, Gearset jobs, Copado promotion and MuleSoft deployment are linked.",
    "Automated tests, QA, UAT and production validation passed with exact results and explicit omissions.",
    "Production used the approved Salesforce commit and MuleSoft artifact without unapproved scope changes.",
    "Post-deployment Salesforce, API and integration smoke checks passed; known issues are linked and owned.",
    "Documentation, runbooks, monitoring and release notes are updated.",
    "Manual changes are reconciled to Git and temporary access or flags are removed or tracked to expiry.",
]:
    add_bullet(doc, item)


add_heading(doc, "Appendix B - Source notes", 1)
add_body(doc, "The implementation details below were checked against official product documentation accessed on 2 September 2026. Product editions, supported integrations and limitations can change; confirm current contracts and supported features before configuration.")
sources = [
    ("Gearset - Setting up your first CI job", "https://docs.gearset.com/en/articles/8333138-setting-up-your-first-ci-job-in-gearset"),
    ("Gearset - Creating CI jobs in Pipelines", "https://docs.gearset.com/en/articles/12460112-creating-ci-jobs-in-gearset-pipelines"),
    ("Copado - Structure of a User Story", "https://docs.copado.com/articles/?_escaped_fragment_=source-format-pipelines-publication/structure-of-a-user-story"),
    ("Copado - Known limitations in Salesforce Source Format Pipelines", "https://docs.copado.com/articles/?_escaped_fragment_=source-format-pipelines-publication/known-limitations-in-salesforce-source-format-pipelines"),
    ("Copado - Upgrade to Source Format Pipelines: Planning and Recommendations", "https://docs.copado.com/articles/?_escaped_fragment_=copado-pipelines-publication/upgrade-to-source-format-pipelines-planning-and-recommendations"),
    ("MuleSoft - Mule Maven Plugin", "https://docs.mulesoft.com/mule-runtime/4.6/mmp-concept"),
    ("MuleSoft - MUnit Maven coverage", "https://docs.mulesoft.com/munit/latest/coverage-maven-concept"),
    ("MuleSoft - Publishing assets using Maven", "https://docs.mulesoft.com/exchange/to-publish-assets-maven"),
]
for label, url in sources:
    add_link_paragraph(doc, label, url)

add_heading(doc, "Appendix C - Ready-to-paste stakeholder update", 1)
add_callout(
    doc,
    "Stakeholder update",
    "A separate implementation plan has been prepared for SALDEV-1500. The proposed delivery model uses Gearset for Salesforce CI and lower-environment validation, Copado for governed Salesforce UAT/Production promotion, and GitHub Actions plus Anypoint for MuleSoft build, MUnit testing, artifact publication and runtime deployment. GitHub remains the source of truth and Jira remains the work, approval and evidence record. The plan includes a 12-week roadmap, role boundaries, security controls, cross-platform quality gates, pilot and rollback steps, risks, acceptance criteria and a Jira-ready implementation backlog. No system configuration or production change was performed as part of the document preparation.",
    "blue",
)


doc.core_properties.title = "SALDEV-1500 - Gearset, Copado and MuleSoft Implementation Plan"
doc.core_properties.subject = "Controlled Salesforce and MuleSoft CI/CD implementation plan"
doc.core_properties.author = "Flywire GTM Systems"
doc.core_properties.keywords = "SALDEV-1500, Gearset, Copado, MuleSoft, GitHub, Jira, Salesforce, CI/CD"
doc.core_properties.comments = "Separate implementation plan; no system changes performed."

doc.save(OUT)
print(OUT)
