import copy
import json
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

base_dir = Path(r"C:\Users\LIKKI\Documents\ChatGPT\David 2 Probo medical")
source_path = base_dir / "scratch" / "book_v2_update" / "Book v2 source.xlsx"
query_path = base_dir / "scratch" / "book_v2_update" / "direct_permission_sets_all_active.json"
output_dir = base_dir / "outputs" / "book_v2_direct_permission_sets"
output_path = output_dir / "Book v2 - with Direct Permission Sets.xlsx"

output_dir.mkdir(parents=True, exist_ok=True)

with query_path.open("r", encoding="utf-8-sig") as f:
    payload = json.load(f)

records = payload["result"]["records"]

detail_rows = []
by_username = defaultdict(list)

for record in records:
    assignee = record.get("Assignee") or {}
    profile = assignee.get("Profile") or {}
    permission_set = record.get("PermissionSet") or {}
    row = {
        "User Name": assignee.get("Name"),
        "Username": assignee.get("Username"),
        "Email": assignee.get("Email"),
        "User Type": assignee.get("UserType"),
        "Profile": profile.get("Name"),
        "Permission Set Label": permission_set.get("Label"),
        "Permission Set API Name": permission_set.get("Name"),
        "Permission Set Type": permission_set.get("Type"),
        "Modify All Data": bool(permission_set.get("PermissionsModifyAllData")),
        "View All Data": bool(permission_set.get("PermissionsViewAllData")),
        "Assignment Source": "Direct Permission Set",
        "User Id": record.get("AssigneeId"),
        "Permission Set Id": record.get("PermissionSetId"),
    }
    detail_rows.append(row)
    if row["Username"]:
        by_username[row["Username"]].append(row["Permission Set Label"] or row["Permission Set API Name"] or "")

headers = [
    "User Name",
    "Username",
    "Email",
    "User Type",
    "Profile",
    "Permission Set Label",
    "Permission Set API Name",
    "Permission Set Type",
    "Modify All Data",
    "View All Data",
    "Assignment Source",
    "User Id",
    "Permission Set Id",
]

wb = openpyxl.load_workbook(source_path)

if "Direct Permission Sets" in wb.sheetnames:
    del wb["Direct Permission Sets"]

summary_ws = wb["Sheet1"]
detail_ws = wb.create_sheet("Direct Permission Sets")

detail_ws.append(headers)
for row in detail_rows:
    detail_ws.append([row.get(header) for header in headers])

header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(bold=True, color="FFFFFF")
thin_border = Border(bottom=Side(style="thin", color="D9E2F3"))

for cell in detail_ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border

for row in detail_ws.iter_rows(min_row=2, max_row=detail_ws.max_row):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=False)
        cell.border = thin_border

detail_ws.freeze_panes = "A2"
detail_ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{detail_ws.max_row}"

table_ref = f"A1:{get_column_letter(len(headers))}{detail_ws.max_row}"
table = Table(displayName="DirectPermissionSetsTable", ref=table_ref)
table.tableStyleInfo = TableStyleInfo(
    name="TableStyleMedium2",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False,
)
detail_ws.add_table(table)

widths = {
    "A": 24,
    "B": 34,
    "C": 34,
    "D": 18,
    "E": 28,
    "F": 45,
    "G": 40,
    "H": 18,
    "I": 16,
    "J": 16,
    "K": 24,
    "L": 20,
    "M": 22,
}
for col, width in widths.items():
    detail_ws.column_dimensions[col].width = width

detail_ws.sheet_view.showGridLines = False
summary_ws.sheet_view.showGridLines = False

summary_header_by_col = {summary_ws.cell(1, col).value: col for col in range(1, summary_ws.max_column + 1)}
username_col = summary_header_by_col.get("Username", 1)
direct_list_col = summary_header_by_col.get("Direct Permission Sets", 3)
direct_count_col = summary_header_by_col.get("Direct Permission Set Count", 10)
group_count_col = summary_header_by_col.get("Permission Set Group Count", 12)

detail_last_row = detail_ws.max_row
for row_idx in range(2, summary_ws.max_row + 1):
    username = summary_ws.cell(row_idx, username_col).value
    labels = [label for label in by_username.get(username, []) if label]
    summary_ws.cell(row_idx, direct_list_col).value = "; ".join(labels) if labels else None
    summary_ws.cell(row_idx, direct_count_col).value = (
        f'=COUNTIF(\'Direct Permission Sets\'!$B$2:$B${detail_last_row},A{row_idx})'
    )
    if group_count_col:
        summary_ws.cell(row_idx, group_count_col).value = None

for col in [direct_list_col, direct_count_col, group_count_col]:
    if not col:
        continue
    header_cell = summary_ws.cell(1, col)
    header_cell.fill = copy.copy(summary_ws.cell(1, 1).fill)
    header_cell.font = copy.copy(summary_ws.cell(1, 1).font)
    header_cell.alignment = copy.copy(summary_ws.cell(1, 1).alignment)

summary_ws.column_dimensions[get_column_letter(direct_list_col)].width = 70
summary_ws.column_dimensions[get_column_letter(direct_count_col)].width = 20
summary_ws.column_dimensions[get_column_letter(group_count_col)].width = 20
summary_ws.auto_filter.ref = f"A1:{get_column_letter(summary_ws.max_column)}{summary_ws.max_row}"
summary_ws.freeze_panes = "A2"

summary_ws["N1"] = "Direct Permission Set Source"
summary_ws["N2"] = f"Added from ProboMedical SOQL export. Rows: {len(detail_rows)}. Assignment Source = Direct Permission Set, excluding profile-owned permission sets and permission set group assignments."
summary_ws["N1"].fill = header_fill
summary_ws["N1"].font = header_font
summary_ws["N1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
summary_ws["N2"].alignment = Alignment(vertical="top", wrap_text=True)
summary_ws.column_dimensions["N"].width = 55

wb.save(output_path)

print(
    json.dumps(
        {
            "output": str(output_path),
            "direct_assignment_rows": len(detail_rows),
            "users_with_direct_permission_sets": len(by_username),
            "summary_rows": summary_ws.max_row - 1,
            "detail_sheet_rows": detail_ws.max_row,
        },
        indent=2,
    )
)
