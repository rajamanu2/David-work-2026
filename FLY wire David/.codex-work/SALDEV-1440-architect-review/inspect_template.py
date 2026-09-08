from pathlib import Path
import json
import sys

from docx import Document
from docx.oxml.ns import qn


def color_value(run):
    color = run.font.color.rgb
    return str(color) if color is not None else None


def run_info(run):
    rpr = run._element.rPr
    fonts = rpr.rFonts if rpr is not None else None
    return {
        "text": run.text[:120],
        "font": run.font.name,
        "ascii_font": fonts.get(qn("w:ascii")) if fonts is not None else None,
        "size_pt": run.font.size.pt if run.font.size is not None else None,
        "bold": run.bold,
        "italic": run.italic,
        "color": color_value(run),
    }


def paragraph_info(paragraph):
    fmt = paragraph.paragraph_format
    return {
        "style": paragraph.style.name,
        "text": paragraph.text[:180],
        "alignment": str(paragraph.alignment),
        "space_before_pt": fmt.space_before.pt if fmt.space_before else None,
        "space_after_pt": fmt.space_after.pt if fmt.space_after else None,
        "line_spacing": fmt.line_spacing,
        "left_indent_in": fmt.left_indent.inches if fmt.left_indent else None,
        "right_indent_in": fmt.right_indent.inches if fmt.right_indent else None,
        "runs": [run_info(run) for run in paragraph.runs if run.text],
    }


def main():
    doc = Document(Path(sys.argv[1]))
    data = {
        "paragraphs": [paragraph_info(p) for p in doc.paragraphs],
        "tables": [],
    }
    for table in doc.tables:
        data["tables"].append({
            "rows": len(table.rows),
            "columns": len(table.columns),
            "widths_in": [cell.width.inches if cell.width else None for cell in table.rows[0].cells],
            "sample_cells": [paragraph_info(cell.paragraphs[0]) for cell in table.rows[0].cells],
        })
    print(json.dumps(data, indent=2))


if __name__ == "__main__":
    main()
