from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_architect_review import (
    AMBER,
    BLUE,
    GRID,
    INK,
    LIGHT_AMBER,
    LIGHT_BLUE,
    LIGHT_GRAY,
    LIGHT_RED,
    LIGHT_TEAL,
    MUTED,
    NAVY,
    RED,
    TEAL,
    add_body,
    add_callout,
    add_figure,
    add_heading,
    add_table,
    arrow,
    configure_header_footer,
    draw_wrapped,
    load_font,
    rgb,
    rounded,
    set_cell_margins,
    set_cell_text,
    set_run_font,
    set_table_geometry,
    shade_cell,
)


ROOT = Path(__file__).resolve().parent
ASSET_DIR = ROOT / "doc-assets-scc3387"
OUTPUT = ROOT / "SCC-3387-architect-review.docx"
ASSET_DIR.mkdir(exist_ok=True)


def normalize_table_margins(table) -> None:
    """Apply the decision_memo table token: 80/120/80/120 DXA."""
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=80, start=120, bottom=80, end=120)


def review_table(document, headers, rows, widths, status_column=None):
    table = add_table(document, headers, rows, widths, status_column=status_column)
    normalize_table_margins(table)
    return table


def review_callout(document, label, message, *, fill=LIGHT_AMBER, accent=AMBER):
    add_callout(document, label, message, fill=fill, accent=accent)
    table = document.tables[-1]
    normalize_table_margins(table)
    return table


def create_outage_diagram(path: Path) -> None:
    image = Image.new("RGB", (1600, 1000), "#F7F9FC")
    draw = ImageDraw.Draw(image)
    title = load_font(42, bold=True)
    subtitle = load_font(22)
    heading = load_font(25, bold=True)
    body = load_font(19)
    small = load_font(17, bold=True)

    draw.text((70, 50), "SCC-3387 | Outage creation and ownership path", font=title, fill="#0B2545")
    draw.text(
        (70, 110),
        "Core Incident automation works, but required-field, severity, access, and queue-membership gates remain.",
        font=subtitle,
        fill="#5B6777",
    )

    boxes = [
        ((55, 230, 350, 590), "NOC rep", [
            "No named NOC profile or role",
            "Only admin profile can create",
            "NOC permission set has 0 assignments",
        ], "#FDECEC", "#C53030"),
        ((430, 230, 755, 590), "New Outage", [
            "Incident record type active",
            "Outage is default for admin",
            "Five named fields exist",
        ], "#E7F5F1", "#16836B"),
        ((835, 230, 1160, 590), "Data quality", [
            "Story fields remain optional",
            "Impact is High / Medium / Low",
            "Requested severity values absent",
        ], "#FDECEC", "#C53030"),
        ((1240, 230, 1545, 590), "NOC queue", [
            "Before-save flow active/latest",
            "Incident queue enabled",
            "0 direct queue members",
        ], "#FFF6DE", "#B7791F"),
    ]

    for box, box_title, lines, fill, outline in boxes:
        rounded(draw, box, fill, outline=outline, radius=22, width=4)
        draw.text((box[0] + 24, box[1] + 24), box_title, font=heading, fill="#0B2545")
        y = box[1] + 92
        for line in lines:
            draw.ellipse((box[0] + 26, y + 7, box[0] + 38, y + 19), fill=outline)
            draw_wrapped(draw, (box[0] + 52, y), line, body, "#1F2937", box[2] - box[0] - 78, spacing=6)
            y += 76

    for start, end, label, color in (
        ((350, 410), (430, 410), "CREATE", "#C53030"),
        ((755, 410), (835, 410), "SAVE", "#8A97A8"),
        ((1160, 410), (1240, 410), "ASSIGN", "#B7791F"),
    ):
        arrow(draw, start, end, color=color, width=6)
        draw.text((start[0] + 10, start[1] - 42), label, font=small, fill=color)

    rounded(draw, (70, 680, 1530, 900), "#EAF2FA", outline="#2E74B5", radius=22, width=3)
    draw.text((105, 715), "Architect reading", font=heading, fill="#0B2545")
    decision = (
        "The Outage record type, queue, and assignment flow are structurally present. Promotion should wait until "
        "the six story fields are enforced, the severity design is corrected, a real NOC persona is authorized, "
        "and the queue has operational members."
    )
    draw_wrapped(draw, (105, 775), decision, body, "#1F2937", 1360, spacing=8)
    image.save(path)


def configure_document(document: Document) -> None:
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        configure_header_footer(section)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 12, 6),
        ("Heading 2", 13, BLUE, 10, 5),
        ("Heading 3", 12, NAVY, 8, 4),
    ):
        style = document.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def build_document() -> None:
    diagram_path = ASSET_DIR / "scc3387-outage-path.png"
    create_outage_diagram(diagram_path)

    document = Document()
    configure_document(document)

    # Page 1: decision memo.
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(2)
    set_run_font(paragraph.add_run("ARCHITECT REVIEW"), size=10, color=BLUE, bold=True)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(paragraph.add_run("SCC-3387 | Create Outage Ticket"), size=25, color=NAVY, bold=True)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(16)
    set_run_font(
        paragraph.add_run("Great Plains Merge Sandbox | Read-only dry-run and readiness assessment"),
        size=13,
        color=MUTED,
    )

    metadata = [
        ("Environment", "GreatPlainsMerge (Sandbox) | Org 00DEa00000GkAsHMAV"),
        ("Review date", "26 August 2026"),
        ("Retrieval", "Succeeded | Job 09SEa00000iccfmMAA"),
        ("Decision", "Changes Requested - not ready for promotion"),
        ("Boundary", "No records, deployment, activation, or org metadata were changed"),
    ]
    table = document.add_table(rows=len(metadata), cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [1900, 7460])
    normalize_table_margins(table)
    for row, (label, value) in zip(table.rows, metadata):
        shade_cell(row.cells[0], LIGHT_BLUE)
        set_cell_text(row.cells[0], label, size=9.5, color=NAVY, bold=True)
        set_cell_text(row.cells[1], value, size=9.5)
    document.add_paragraph().paragraph_format.space_after = Pt(2)

    review_callout(
        document,
        "Architect decision",
        "Do not promote yet. The Incident-based Outage path exists, but required data enforcement, severity values, NOC persona access, and queue membership are incomplete.",
        fill=LIGHT_RED,
        accent=RED,
    )

    add_heading(document, "Acceptance outcome", 1)
    acceptance_rows = [
        ("1", "Incident and Outage type", "Partial pass", "Incident is available and Outage is active. It is visible/default only for System Administrator; the business NOC persona is not identified."),
        ("2", "Required outage fields", "Fail", "Five custom fields plus Detected Date Time exist, but all six are optional in schema, layout, and Dynamic Forms."),
        ("3", "Severity and narrative", "Fail", "Description and Created Date exist; Impact is required but still High/Medium/Low, not Minor/Major/Critical."),
        ("4", "NOC assignment", "Partial pass", "Active/latest before-save flow assigns new Outages to NOC. One existing Outage is NOC-owned, but the queue has zero direct members."),
        ("5", "NOC access", "Fail", "No NOC profile/role user is identifiable. Only System Administrator has usable create access and Outage record-type visibility."),
    ]
    review_table(document, ["#", "Scope", "Result", "Evidence"], acceptance_rows, [500, 2050, 1450, 5360], status_column=2)

    # Page 2: visual system map.
    document.add_page_break()
    add_heading(document, "Solution and control path", 1)
    add_body(document, "The platform objects and automation are present, but the intended NOC user cannot be proven and the record can be saved without the story's required operational data.")
    add_figure(
        document,
        diagram_path,
        "Figure 1. SCC-3387 outage creation and ownership path",
        "Diagram showing a NOC rep creating an Incident Outage, passing through data-quality controls, then a before-save assignment flow routing the record to the NOC queue. NOC access, required fields, severity values, and queue membership are marked as blockers or incomplete.",
        width=6.45,
    )
    review_callout(document, "Design reading", "Green is configured. Amber is structurally present but operationally incomplete. Red blocks acceptance.", fill=LIGHT_BLUE, accent=BLUE)

    # Page 3: field evidence.
    document.add_page_break()
    add_heading(document, "Field and page configuration", 1)
    field_rows = [
        ("Affected Area/Region", "Text(255)", "Optional", "Fail - required by story"),
        ("Equipment/Facility ID", "Text(255)", "Optional", "Fail - required by story"),
        ("Estimated Customers Impacted", "Number(18,0)", "Optional", "Type passes; requiredness fails"),
        ("Detected Date Time", "Date/Time", "Optional", "Fail - required by story"),
        ("NOC Ticket Number Reference", "Text(255)", "Optional", "Fail - required by story"),
        ("Estimated Restoration Hours", "Number(18,0)", "Optional", "Type passes; requiredness fails"),
        ("Impact", "Restricted picklist", "Required", "Fail - High/Medium/Low, not Minor/Major/Critical"),
        ("Description", "Text Area(32,000)", "Optional", "Present; can hold description/cause if product confirms"),
        ("Created Date", "System Date/Time", "Read-only", "Pass - platform timestamps creation"),
    ]
    review_table(document, ["Field", "Configured type", "UI behavior", "Acceptance reading"], field_rows, [2550, 1900, 1300, 3610], status_column=3)

    add_heading(document, "Requiredness evidence", 2)
    required_rows = [
        ("Schema", "All five SCC-3387 custom fields have required=false; DetectedDateTime is nillable."),
        ("Page layout", "Only Subject, Status, and Impact are marked Required. The six story fields use Edit."),
        ("Dynamic Forms", "Only Subject, Status, and Impact use uiBehavior=required. The six story fields use uiBehavior=none."),
        ("Severity", "Outage record type enables Impact values High, Medium, and Low. Low is the default."),
    ]
    review_table(document, ["Layer", "Observed configuration"], required_rows, [1800, 7560])
    review_callout(
        document,
        "Data quality blocker",
        "A user can save an Outage without the six operational fields named as required in SCC-3387. Enforce the requirement at the Dynamic Forms/page layer or with a record-type-scoped validation rule.",
        fill=LIGHT_RED,
        accent=RED,
    )

    # Page 4: persona, queue, automation, and package evidence.
    document.add_page_break()
    add_heading(document, "Persona, queue, and automation evidence", 1)
    access_rows = [
        ("System Administrator", "41 active", "Create/read/edit; Outage visible/default; fields editable", "Technically works, but it is not a defined NOC least-privilege persona."),
        ("Standard User", "0 active", "No Incident object permission; Outage invisible; fields read-only", "Cannot execute the story."),
        ("System Administrator - API Only", "1 active", "Incident create/edit; Outage invisible; fields editable", "Integration persona, not a NOC business user."),
        ("SCC-3380 NOC permission set", "0 assignments", "No Incident object permission and none of the five SCC-3387 fields", "Does not close the SCC-3387 access gap."),
    ]
    review_table(document, ["Access path", "Count", "Observed permission", "Architect interpretation"], access_rows, [2500, 1350, 2900, 2610])

    automation_rows = [
        ("Assignment flow", "Pass", "Assign_Outage_Incident_to_NOC is active/latest (version id 301Ea00001jiGElIAM)."),
        ("Trigger logic", "Pass", "Before-save on Incident create; checks Incident.Outage, looks up NOC queue, assigns OwnerId."),
        ("Existing evidence", "Pass", "One Outage exists and the aggregate ownership check shows it is NOC-owned."),
        ("Queue operations", "Fail", "NOC supports Incident and emails members, but it has zero direct members."),
        ("Runtime UAT", "Not run", "No test record was created or changed during this read-only review."),
    ]
    add_heading(document, "Automation and operational readiness", 2)
    review_table(document, ["Area", "Result", "Evidence"], automation_rows, [2200, 1500, 5660], status_column=1)

    package_rows = [
        ("Focused retrieval", "16 metadata files including the Incident object container; job 09SEa00000iccfmMAA succeeded."),
        ("Ticket omission", "Queue:NOC and FlexiPage:Incident_Record_Page1 are required to reproduce the reviewed solution but are absent from the ticket component list."),
        ("Ticket extras", "Case.SendEmail and Case_Before_Insert_Update_Trigger_Flow do not implement the stated outage-creation acceptance criteria."),
        ("Validation boundary", "No check-only deployment was run because this review retrieved current sandbox configuration and introduced no change package."),
    ]
    add_heading(document, "Component and evidence boundary", 2)
    review_table(document, ["Boundary", "Evidence"], package_rows, [1900, 7460])

    # Page 5: approval gates and stakeholder response.
    document.add_page_break()
    add_heading(document, "Required approval gates", 1)
    gate_rows = [
        ("1", "Enforce required data", "Admin / Product owner", "Make the six story fields required for Incident.Outage and prove blank-save prevention."),
        ("2", "Resolve severity design", "Architect / Product owner", "Use Minor/Major/Critical on Impact or approve a documented alternative; regression-check Incident logic."),
        ("3", "Create NOC access model", "Security owner", "Grant named NOC reps least-privilege Incident create/read/edit, field edit, tab, and Outage record-type access."),
        ("4", "Populate NOC queue", "NOC manager / Admin", "Add approved users or groups and verify notification, list-view visibility, and ownership work routing."),
        ("5", "Correct component inventory", "Release manager", "Include Queue:NOC and Incident_Record_Page1; remove or justify unrelated Case components."),
        ("6", "Run named-persona UAT", "QA / NOC", "Create an Outage from Incidents, complete all fields, save, verify NOC owner, and capture record ID/screenshots."),
    ]
    review_table(document, ["#", "Required action", "Owner", "Evidence needed"], gate_rows, [600, 2600, 2100, 4060])

    add_heading(document, "Ready-to-paste architect response", 1)
    response = (
        "Architect review completed in the GreatPlainsMerge sandbox. Changes Requested: the Incident Outage record type, "
        "Incident-enabled NOC queue, and Assign Outage Incident to NOC flow are present, and the flow is active/latest. "
        "However, all six operational fields named as required can still be left blank; Impact remains High/Medium/Low "
        "instead of Minor/Major/Critical; only the System Administrator profile has a usable create path to Outage; the "
        "NOC queue has zero direct members; and the existing NOC permission set has zero assignments and does not grant "
        "the SCC-3387 object/field access. The ticket package also omits Queue:NOC and Incident_Record_Page1 while listing "
        "unrelated Case components. Complete the six gates above and rerun the story with a named NOC user before promotion. "
        "No records, metadata, activation, deployment, or production changes were made."
    )
    review_callout(document, "Story comment", response, fill=LIGHT_BLUE, accent=BLUE)

    document.core_properties.title = "SCC-3387 Create Outage Ticket Architect Review"
    document.core_properties.subject = "Read-only dry-run and readiness assessment"
    document.core_properties.author = "Solution Architecture Review"
    document.core_properties.keywords = "Salesforce, Great Plains, SCC-3387, Incident, Outage, NOC, architect review, dry run"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
