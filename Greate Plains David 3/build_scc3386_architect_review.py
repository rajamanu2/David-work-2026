from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

from build_architect_review import (
    AMBER,
    BLUE,
    GRID,
    INK,
    LIGHT_AMBER,
    LIGHT_BLUE,
    LIGHT_GRAY,
    LIGHT_RED,
    LIGHT_TEAL,
    MUTED,
    NAVY,
    RED,
    TEAL,
    add_body,
    add_callout,
    add_figure,
    add_heading,
    add_table,
    arrow,
    configure_header_footer,
    draw_wrapped,
    load_font,
    rgb,
    rounded,
    set_cell_text,
    set_run_font,
    set_table_geometry,
    shade_cell,
)


ROOT = Path(__file__).resolve().parent
ASSET_DIR = ROOT / "doc-assets-scc3386"
OUTPUT = ROOT / "SCC-3386-architect-review.docx"
ASSET_DIR.mkdir(exist_ok=True)


def create_access_diagram(path: Path) -> None:
    image = Image.new("RGB", (1600, 1000), "#F7F9FC")
    draw = ImageDraw.Draw(image)
    title = load_font(42, bold=True)
    subtitle = load_font(22)
    heading = load_font(26, bold=True)
    body = load_font(20)
    small = load_font(18)

    draw.text((70, 52), "SCC-3386 | Technician context delivery path", font=title, fill="#0B2545")
    draw.text(
        (70, 112),
        "Configured source fields exist, but two downstream gates prevent the intended Field Service Mobile experience.",
        font=subtitle,
        fill="#5B6777",
    )

    boxes = [
        ((70, 220, 440, 560), "Trouble Ticket Case", [
            "5 new context fields",
            "Restricted 8-value contact picklist",
            "Related Outage lookup to Incident",
            "Field Technician Context section",
        ], "#E7F5F1", "#16836B"),
        ((610, 220, 990, 560), "Work Order record page", [
            "Quick action defines 9 read-only fields",
            "Active SCC-3385 Create Work Order v1",
            "Missing Related Record component",
            "Quick action is not referenced",
        ], "#FDECEC", "#C53030"),
        ((1160, 220, 1530, 560), "Field technician", [
            "2 active Field Service Mobile users",
            "Profile: GPC-Residential CRC Team",
            "No read access to the 5 new fields",
            "No compensating permission set",
        ], "#FDECEC", "#C53030"),
    ]

    for box, box_title, lines, fill, outline in boxes:
        rounded(draw, box, fill, outline=outline, radius=24, width=4)
        draw.text((box[0] + 28, box[1] + 26), box_title, font=heading, fill="#0B2545")
        y = box[1] + 92
        for line in lines:
            draw.ellipse((box[0] + 30, y + 7, box[0] + 42, y + 19), fill=outline)
            draw_wrapped(draw, (box[0] + 58, y), line, body, "#1F2937", box[2] - box[0] - 92, spacing=6)
            y += 58

    arrow(draw, (440, 390), (610, 390), color="#8A97A8", width=6)
    arrow(draw, (990, 390), (1160, 390), color="#C53030", width=7)
    draw.text((480, 345), "Case lookup", font=small, fill="#5B6777")
    draw.text((1028, 336), "BLOCKED", font=heading, fill="#C53030")

    rounded(draw, (70, 650, 760, 900), "#FFF6DE", outline="#B7791F", radius=22, width=3)
    draw.text((100, 682), "Related outage evidence", font=heading, fill="#0B2545")
    related_lines = [
        "Lookup relationship is valid and labeled Trouble Tickets.",
        "The only Incident layout does not contain that related list.",
    ]
    y = 742
    for line in related_lines:
        draw.ellipse((102, y + 7, 114, y + 19), fill="#B7791F")
        draw_wrapped(draw, (130, y), line, body, "#1F2937", 585, spacing=6)
        y += 72

    rounded(draw, (840, 650, 1530, 900), "#EAF2FA", outline="#2E74B5", radius=22, width=3)
    draw.text((870, 682), "Architect decision", font=heading, fill="#0B2545")
    decision = (
        "Changes Requested. Add the page wiring and technician permissions, then rerun the four story tests "
        "with a named Field Service Mobile user."
    )
    draw_wrapped(draw, (870, 752), decision, body, "#1F2937", 605, spacing=8)

    image.save(path)


def configure_document(document: Document) -> None:
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        configure_header_footer(section)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 12, 6),
        ("Heading 2", 13, BLUE, 10, 5),
        ("Heading 3", 12, NAVY, 8, 4),
    ):
        style = document.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def build_document() -> None:
    diagram_path = ASSET_DIR / "scc3386-technician-context-flow.png"
    create_access_diagram(diagram_path)

    document = Document()
    configure_document(document)

    # Page 1: decision memo.
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(2)
    set_run_font(paragraph.add_run("ARCHITECT REVIEW"), size=10, color=BLUE, bold=True)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(paragraph.add_run("SCC-3386 | Field Technician Context"), size=25, color=NAVY, bold=True)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(16)
    set_run_font(
        paragraph.add_run("Great Plains Merge Sandbox | Read-only dry-run and readiness assessment"),
        size=13,
        color=MUTED,
    )

    metadata = [
        ("Environment", "GreatPlainsMerge (Sandbox) | Org 00DEa00000GkAsHMAV"),
        ("Review date", "26 August 2026"),
        ("Retrieval", "Succeeded | Job 09SEa00000ickQLMAY"),
        ("Decision", "Changes Requested - not ready for promotion"),
        ("Boundary", "No records, deployment, activation, or org metadata were changed"),
    ]
    table = document.add_table(rows=len(metadata), cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [1900, 7460])
    for row, (label, value) in zip(table.rows, metadata):
        shade_cell(row.cells[0], LIGHT_BLUE)
        set_cell_text(row.cells[0], label, size=9.5, color=NAVY, bold=True)
        set_cell_text(row.cells[1], value, size=9.5)
    document.add_paragraph().paragraph_format.space_after = Pt(2)

    add_callout(
        document,
        "Architect decision",
        "Do not promote yet. The Case-side fields are present, but the Work Order page wiring, technician field access, and Incident related-list configuration are incomplete.",
        fill=LIGHT_RED,
        accent=RED,
    )

    add_heading(document, "Acceptance outcome", 1)
    acceptance_rows = [
        ("1", "Case fields and layout", "Partial pass", "Five fields exist; Contact Preference is restricted to 8 values; section order matches the test. Technician access fails."),
        ("2", "Related Outage", "Partial fail", "Lookup to Incident exists, but the only Incident layout lacks the Trouble Tickets related list."),
        ("3", "Create Work Order regression", "Partial pass", "SCC-3385 flow v1 is active/latest. Runtime mapping and duplicate behavior were not executed in this read-only review."),
        ("4", "Read-only detail on Work Order", "Fail", "Nine-field read-only quick action exists, but the only Work Order record page does not reference it or include a Related Record component."),
        ("5", "Field Service Mobile persona", "Fail", "Two active mobile users use GPC-Residential CRC Team; that profile and its assigned permission sets grant none of the required field access."),
    ]
    add_table(document, ["#", "Scope", "Result", "Evidence"], acceptance_rows, [500, 2100, 1500, 5260], status_column=2)

    # Page 2: visual system map.
    document.add_page_break()
    add_heading(document, "Solution and access path", 1)
    add_body(document, "The configured Case fields are usable by administrators, but the intended technician path breaks at both the Work Order page and field-security layers.")
    add_figure(
        document,
        diagram_path,
        "Figure 1. SCC-3386 technician context delivery path",
        "Diagram showing Case context fields flowing to a Work Order record page and then a field technician. The Work Order Related Record component and technician field access are both marked as blocked. A related-outage note shows that the Incident layout lacks the Trouble Tickets related list.",
        width=6.45,
    )
    add_callout(document, "Design reading", "Green is configured. Amber is incomplete. Red blocks an acceptance criterion.", fill=LIGHT_BLUE, accent=BLUE)

    # Page 3: component and access evidence.
    document.add_page_break()
    add_heading(document, "Detailed configuration findings", 1)
    add_heading(document, "1. Case field implementation", 2)
    field_rows = [
        ("Contact Preference", "Restricted picklist", "8 approved values; placed first in Field Technician Context."),
        ("Special Instructions", "Long Text Area (4,000)", "Placed second; site-access guidance."),
        ("Tier 2 Notes", "Long Text Area (32,768)", "Placed third; full troubleshooting history capacity."),
        ("Equipment Information", "Long Text Area (4,000)", "Placed fourth; equipment identifiers and context."),
        ("Related Outage", "Lookup to Incident", "Placed fifth; relationship label is Trouble Tickets."),
    ]
    add_table(document, ["Field", "Type", "Architect evidence"], field_rows, [2400, 2200, 4760])

    add_heading(document, "2. Field-level access", 2)
    access_rows = [
        ("System Administrator (Admin)", "Readable / editable", "Configured; 16 active Field Service Mobile assignments also use this profile."),
        ("System Administrator - API Only", "Readable / editable", "Configured; not the business technician persona."),
        ("Standard User (Standard)", "Readable / read-only", "Configured, but there are zero active users on this profile."),
        ("GPC-Residential CRC Team", "Not readable", "Two active Field Service Mobile users; no assigned permission set grants any of the 5 new fields, Service Type, or Subscriber Report."),
    ]
    add_table(document, ["Profile", "SCC-3386 field access", "Interpretation"], access_rows, [2500, 2200, 4660], status_column=1)
    add_callout(
        document,
        "Security blocker",
        "The only non-admin profile with active Field Service Mobile users cannot read the story fields. A read-only quick action cannot override field-level security.",
        fill=LIGHT_RED,
        accent=RED,
    )

    # Page 4: UI wiring and evidence boundary.
    document.add_page_break()
    add_heading(document, "UI wiring and regression evidence", 1)
    ui_rows = [
        ("Quick action", "Present", "Case.SCC_3386_Trouble_Ticket_Details contains the 9 expected fields, all read-only; Case Status is omitted."),
        ("Work Order record page", "Missing wiring", "Work_Order_Record_Page contains highlights, chatter, detail, and related lists only; no force:relatedRecord component or SCC-3386 action reference."),
        ("Work Order page inventory", "One page", "Tooling API found exactly one WorkOrder RecordPage; last modified 10 March 2026, before SCC-3386."),
        ("Incident layout", "Missing related list", "Outage Incident Layout has Case Related Issues and Files, but not the Related_Outage__c Trouble Tickets relationship."),
        ("Create Work Order flow", "Active/latest v1", "SCC_3385_Create_Trouble_Ticket_Work_Order is active; runtime regression remains unproven."),
    ]
    add_table(document, ["Area", "Result", "Evidence"], ui_rows, [2400, 1800, 5160], status_column=1)

    add_heading(document, "Read-only evidence boundary", 2)
    boundary_rows = [
        ("Metadata retrieved", "13 named components plus the Case object container (14 source files)."),
        ("Named UAT data", "No Case records exist for Account 'SCC-3386 UAT Customer'."),
        ("Tests not executed", "No Account, Case, Incident, Work Order, Task, or mobile-session records were created or changed."),
        ("Deployment validation", "Not performed; no check-only deployment or production action was requested."),
    ]
    add_table(document, ["Boundary", "Evidence"], boundary_rows, [2500, 6860])

    # Page 5: gates and stakeholder response.
    document.add_page_break()
    add_heading(document, "Required approval gates", 1)
    gate_rows = [
        ("1", "Wire the Work Order page", "Salesforce admin", "Add a Related Record component using the Case lookup and SCC-3386 read-only action; verify desktop and mobile assignment."),
        ("2", "Grant technician read access", "Security owner / Admin", "Give the actual technician profile or permission set read access to all 9 displayed Case fields, without Case Status edit access."),
        ("3", "Add Incident related list", "Salesforce admin", "Place Trouble Tickets on Outage Incident Layout and confirm navigation from Case to Incident and back."),
        ("4", "Run named-persona tests", "QA / Field Service", "Execute all four story tests using a named Field Service Mobile user and capture record IDs/screenshots."),
        ("5", "Confirm deployment inventory", "Release manager", "Add Work_Order_Record_Page and Incident-Outage Incident Layout to the story component list/package."),
    ]
    add_table(document, ["Gate", "Required action", "Owner", "Evidence needed"], gate_rows, [600, 2700, 2100, 3960])

    add_heading(document, "Ready-to-paste architect response", 1)
    response = (
        "Architect review completed in the GreatPlainsMerge sandbox. Changes Requested: the five SCC-3386 Case fields and the nine-field read-only quick action are present, but the only Work Order record page does not reference the quick action or contain a Related Record component. The two active Field Service Mobile users on GPC-Residential CRC Team cannot read the new Case fields, and no assigned permission set compensates. The Related Outage lookup exists, but Outage Incident Layout does not include the Trouble Tickets related list. The SCC-3385 Create Work Order flow is active/latest v1, but runtime regression was not executed because this was a no-write review and the named SCC-3386 UAT data does not exist. Please complete the five gates above and rerun the story tests with a named technician before promotion. No org data, metadata, activation, or deployment changes were made."
    )
    add_callout(document, "Story comment", response, fill=LIGHT_BLUE, accent=BLUE)

    document.core_properties.title = "SCC-3386 Field Technician Context Architect Review"
    document.core_properties.subject = "Read-only dry-run and readiness assessment"
    document.core_properties.author = "Solution Architecture Review"
    document.core_properties.keywords = "Salesforce, Great Plains, SCC-3386, Field Service, architect review, dry run"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
