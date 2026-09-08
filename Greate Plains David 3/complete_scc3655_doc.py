from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


SOURCE = Path(r"C:\Users\LIKKI\Downloads\SCC-3655.docx")
OUTPUT = Path(r"C:\Users\LIKKI\Documents\ChatGPT\Greate Plains David 3\SCC-3655-completed.docx")


def set_text_preserve_format(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        first = paragraph.runs[0]
        first.text = text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def insert_after(paragraph: Paragraph, text: str, num_id: int | None = None) -> Paragraph:
    new_xml = OxmlElement("w:p")
    paragraph._p.addnext(new_xml)
    new_paragraph = Paragraph(new_xml, paragraph._parent)
    new_paragraph.add_run(text)
    if num_id is not None:
        apply_numbering(new_paragraph, num_id)
    return new_paragraph


def create_numbering(document: Document, kind: str) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    level.append(level_text)
    level_justification = OxmlElement("w:lvlJc")
    level_justification.set(qn("w:val"), "left")
    level.append(level_justification)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    p_pr.append(indent)
    level.append(p_pr)
    abstract.append(level)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        first_num.addprevious(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph: Paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)


def find_paragraph(document: Document, exact_text: str) -> Paragraph:
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == exact_text:
            return paragraph
    raise ValueError(f"Paragraph not found: {exact_text}")


document = Document(SOURCE)
bullet_num_id = create_numbering(document, "bullet")
number_num_id = create_numbering(document, "number")

# Fill the original placeholder description.
description = find_paragraph(document, "Add a description...")
set_text_preserve_format(
    description,
    "Use standard Salesforce Case Comments so reps can post and read case updates. "
    "Leave the Public checkbox unchecked for internal-only comments; select it only "
    "when a comment is intended to be externally available.",
)

# Make the acceptance criteria explicit and testable.
acceptance = next(
    paragraph
    for paragraph in document.paragraphs
    if paragraph.text.strip().startswith("Comments are identified as Internal")
)
set_text_preserve_format(acceptance, "Acceptance criteria:")
anchor = acceptance
for criterion in (
    "A comment is internal when the standard Public checkbox is unchecked; no custom Internal field is used.",
    "Salesforce automatically records the comment Created Date/Time.",
    "Salesforce automatically records the user who posted the comment.",
    "The Case Comments related list displays multiple comments in chronological order throughout the Case lifecycle.",
):
    anchor = insert_after(anchor, criterion, bullet_num_id)

# Replace the architecture prose with the verified standard-field description.
architecture = next(
    paragraph
    for paragraph in document.paragraphs
    if paragraph.text.strip().startswith("This should be standard out of the box functionality")
)
set_text_preserve_format(
    architecture,
    "Use standard Salesforce Case Comments. The standard IsPublished field is the "
    "Public/Published checkbox in the user interface: unchecked means internal-only; "
    "checked means externally available. No custom field, trigger, Flow, or component is required.",
)

# Complete the previously blank Deployment section with live read-only evidence.
deployment = next(
    paragraph for paragraph in document.paragraphs if paragraph.text.strip() == "Deployment"
)
set_text_preserve_format(deployment, "Deployment")
anchor = deployment
for item in (
    "Disposition: no SCC-3655-specific deployment is required.",
    "GreatPlainsMerge (sandbox 00DEa00000GkAsHMAV): RelatedCommentsList is already present on Case-Case Layout and Case-Trouble Ticket.",
    "GreatPlainsUAT (sandbox 00DEa00000FZlLBMA1): RelatedCommentsList is already present on Case-Case Layout.",
    "Dependency: GreatPlainsUAT does not yet contain the active Trouble_Ticket Case record type or Case-Trouble Ticket layout. That broader Trouble Ticket prerequisite must be promoted separately before SCC-3655 can be tested specifically on Trouble Ticket records.",
    "Read-only review completed 26 August 2026. No Salesforce records, metadata, activation, or deployment changes were made.",
):
    anchor = insert_after(anchor, item, bullet_num_id)

# Replace the run-on recommendation with a reproducible UAT procedure.
testing = next(
    paragraph
    for paragraph in document.paragraphs
    if paragraph.text.strip().startswith("1.From the app launcher")
)
steps = (
    "Sign in to GreatPlainsUAT as a named rep (recommended persona: GPC-Residential CRC Team) and open an approved UAT Case.",
    "Open the Related tab, locate Case Comments, and select New.",
    "Enter a uniquely labeled internal test comment, leave Public unchecked, and save.",
    "Add a second internal comment at least one minute later.",
    "Confirm both comments display the posting user and Created Date/Time and appear in chronological order.",
    "If external visibility is in scope, add a separate approved test comment with Public checked and verify it through the approved external test channel.",
    "Record the Case number, rep, timestamps, expected visibility, actual visibility, and screenshots in the SCC-3655 test evidence.",
)
set_text_preserve_format(testing, steps[0])
apply_numbering(testing, number_num_id)
anchor = testing
for step in steps[1:]:
    anchor = insert_after(anchor, step, number_num_id)
pending = insert_after(
    anchor,
    "Runtime status: pending. Both GreatPlainsMerge and GreatPlainsUAT currently contain zero CaseComment records, so posting, author display, ordering, and internal/external visibility were not proven by this read-only review.",
)
pending.runs[0].bold = True

# Replace the generic None entries where standard platform artifacts are used.
sections = {paragraph.text.strip(): i for i, paragraph in enumerate(document.paragraphs)}

def replace_section_answer(section_name: str, answer: str) -> None:
    paragraphs = document.paragraphs
    start = next(i for i, paragraph in enumerate(paragraphs) if paragraph.text.strip() == section_name)
    for paragraph in paragraphs[start + 1 :]:
        if paragraph.text.strip() == "None":
            set_text_preserve_format(paragraph, answer)
            return
    raise ValueError(f"Answer paragraph not found after section: {section_name}")


replace_section_answer(
    "New or Updated Objects/Fields",
    "None. Standard CaseComment fields used: ParentId, CommentBody, IsPublished, CreatedDate, and CreatedById.",
)
replace_section_answer(
    "Components",
    "Standard Case Comments related list (metadata name RelatedCommentsList); no custom component.",
)
replace_section_answer(
    "Settings Modifications",
    "None for SCC-3655. The required standard related list is already present on the verified Case layouts.",
)

document.core_properties.title = "SCC-3655 Case Comments - Internal"
document.core_properties.subject = "Completed deployment disposition, UAT steps, and technical documentation"
document.core_properties.keywords = "Salesforce, SCC-3655, Case Comments, GreatPlainsMerge, GreatPlainsUAT"
document.save(OUTPUT)
print(OUTPUT)
