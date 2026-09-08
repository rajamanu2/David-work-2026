from pathlib import Path
from docx import Document
from docx.document import Document as _Document
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from zipfile import ZipFile
from xml.etree import ElementTree as ET


def iter_blocks(parent):
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("Unsupported parent")

    for child in parent_elm.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield Table(child, parent)


path = Path(r"C:\Users\LIKKI\Downloads\Doc1.docx")
media_dir = Path(r"C:\Users\LIKKI\Documents\ChatGPT\david\.codex-work\doc1-media")
media_dir.mkdir(parents=True, exist_ok=True)
doc = Document(path)

print(f"FILE: {path}")
print(f"PARAGRAPHS: {len(doc.paragraphs)}")
print(f"TABLES: {len(doc.tables)}")
print(f"SECTIONS: {len(doc.sections)}")
print("\n=== DOCUMENT ORDER ===")

table_no = 0
for block in iter_blocks(doc):
    if isinstance(block, Paragraph):
        text = block.text.strip()
        if text:
            print(f"P[{block.style.name}]: {text}")
    else:
        table_no += 1
        print(f"TABLE {table_no} ({len(block.rows)}x{len(block.columns)})")
        for row_no, row in enumerate(block.rows, 1):
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            print(f"  R{row_no}: " + " || ".join(cells))

print("\n=== HEADERS AND FOOTERS ===")
for section_no, section in enumerate(doc.sections, 1):
    header = " | ".join(p.text.strip() for p in section.header.paragraphs if p.text.strip())
    footer = " | ".join(p.text.strip() for p in section.footer.paragraphs if p.text.strip())
    print(f"SECTION {section_no} HEADER: {header}")
    print(f"SECTION {section_no} FOOTER: {footer}")

print("\n=== OOXML TEXT AND EMBEDDED PARTS ===")
with ZipFile(path) as package:
    names = package.namelist()
    media = [name for name in names if name.startswith("word/media/")]
    print(f"MEDIA PARTS: {len(media)}")
    for name in media:
        info = package.getinfo(name)
        print(f"  {name} ({info.file_size} bytes)")
        (media_dir / Path(name).name).write_bytes(package.read(name))

    for name in names:
        if not name.startswith("word/") or not name.endswith(".xml"):
            continue
        root = ET.fromstring(package.read(name))
        texts = []
        for elem in root.iter():
            local = elem.tag.rsplit("}", 1)[-1]
            if local in {"t", "instrText", "delText"} and elem.text:
                value = " ".join(elem.text.split())
                if value:
                    texts.append(value)
            for key, value in elem.attrib.items():
                attr = key.rsplit("}", 1)[-1]
                if attr in {"descr", "title", "name"} and value.strip():
                    texts.append(f"[{attr}: {value.strip()}]")
        if texts:
            print(f"PART {name}")
            for text in texts:
                print(f"  {text}")
