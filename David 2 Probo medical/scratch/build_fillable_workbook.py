from pathlib import Path
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from copy import deepcopy
import re, zipfile

# Reuse the typography and table design from the presentation document.
source=Path('scratch/build_presentation_brd.py').read_text(encoding='utf-8')
exec(source.split("p('PROBO MEDICAL')")[0])
out=Path('outputs/business_templates_2026-09-08/Probo_Medical_Full_Fillable_Requirements_Workbook.docx')
d.styles['Heading 1'].font.size=Pt(24)
d.styles['Caption'].font.color.rgb=RGBColor(0,0,0)
s.header.paragraphs[0].text='PROBO MEDICAL                                      REQUIREMENTS WORKBOOK'
for ft in [s.footer,s.first_page_footer]:
 ft.paragraphs[0].runs[0].text='Probo Medical  |  Fillable business workbook'
def field(title,hint,space=22):
 h(title)
 z=d.add_paragraph(hint);z.paragraph_format.space_after=Pt(5)
 z.runs[0].font.size=Pt(10);z.runs[0].italic=True
 z=d.add_paragraph('[Click here to enter your response]');z.paragraph_format.space_after=Pt(space)
def newpage(n,title,intro):
 d.add_page_break();z=d.add_paragraph(f'{n:02d}  /  COMPLETE WITH THE BUSINESS TEAM');z.paragraph_format.space_after=Pt(6)
 z.runs[0].font.size=Pt(9);z.runs[0].bold=True
 d.add_heading(title,1);p(intro)

p('PROBO MEDICAL')
d.add_paragraph('Business Requirements\nCompletion Workbook','Title')
p('A complete fillable template for business teams and leadership')
field('Initiative name','Enter the project, improvement or process being discussed.',10)
p('Business owner  [Enter name and role]\nPrepared by  [Enter name and role]\nVersion  [Enter version]     Date  [Enter date]\nStatus  [Enter Draft, In review or Approved]')
h('How to complete the workbook')
p('Open this document in Microsoft Word. Click a bracketed field and type your answer. Text fields expand as you type. Table rows and pages may grow to fit your responses. Use Save As to create a separate copy for each initiative.')
p('Write Not applicable with a reason when a section does not apply. For an unknown answer, record an open question, an owner and a due date. Duplicate the requirement and work instruction pages for additional items.')
h('Completion route')
table(['Part','What the team completes'],[['01 to 04','Business outcome, stakeholders, current process and future scope'],['05 to 08','Requirements, data, reporting and service expectations'],['09 to 11','Work instructions, acceptance testing, decisions and approvals']],[1.1,5.8])
p('The prompts guide the discussion; they do not state approved Probo Medical processes. Complete and review the workbook with the relevant business owners before sign off.')

newpage(1,'Business outcome','Describe the business need before selecting a solution. Leadership should be able to understand the case for change from this page.')
field('Problem statement','What is happening today, who is affected and how often does it happen?')
field('Business impact and urgency','Describe delays, cost, repeat work, customer impact or missed opportunities. Reference supporting evidence.')
field('Desired outcome','What should improve, and what would success look like to the business?')
h('Success measures')
table(['Measure','Current value and source','Target and date'],[['[Enter measure]','[Enter baseline and evidence]','[Enter target and date]'],['[Enter measure]','[Enter baseline and evidence]','[Enter target and date]']],[1.65,2.75,2.5])
p('Benefit owner  [Enter name]\nLeadership decision required  [Enter decision and required date]')

newpage(2,'Stakeholders and ownership','Identify who supplies information, who owns the process and who can approve requirements or resolve conflicting priorities.')
table(['Role','Name and team','Responsibility'],[['Sponsor','[Enter name and team]','[Enter decision authority]'],['Business owner','[Enter name and team]','[Enter process ownership]'],['Business users','[Enter names and teams]','[Enter tasks represented]'],['System or data owner','[Enter name and team]','[Enter review responsibility]'],['UAT and training lead','[Enter name and team]','[Enter validation responsibility]']],[1.55,2.35,3])
field('Affected teams sites and users','List the locations, user groups and approximate number of people affected.',18)
field('Workshop record','Record the meeting date, participants, source documents and any absent stakeholders whose input is still required.',18)
field('Communication and decision route','How will updates be shared? Who resolves disagreements and confirms decisions?',12)

newpage(3,'Current process','Capture how work happens today, including manual activity, system use and handoffs. Use real examples with sensitive details removed.')
p('Process name  [Enter name]\nTrigger  [Enter what starts the work]\nCompletion point  [Enter what ends the process]')
table(['Step','Owner','Current action and system','Output or handoff'],[['1','[Enter role]','[Enter action and system]','[Enter output]'],['2','[Enter role]','[Enter action and system]','[Enter output]'],['3','[Enter role]','[Enter action and system]','[Enter output]'],['4','[Enter role]','[Enter action and system]','[Enter output]']],[.5,1.3,3,2.1])
field('Pain points and exceptions','Where does work wait, repeat or fail? What happens when inputs are incomplete?',18)
field('Volumes and supporting evidence','Record typical and peak volumes, elapsed times, frequency and evidence references.',18)
p('Existing procedure or process map  [Enter reference or Not available]')

newpage(4,'Scope and future process','Agree what will change and what will remain outside this initiative. Confirm ownership at each future handoff.')
field('In scope','List included activities, teams, sites, channels and business records.',16)
field('Out of scope','List explicit exclusions and items deferred to a later phase.',16)
table(['Future step','Owner','Required result and handoff'],[['1','[Enter role]','[Enter activity, result and receiver]'],['2','[Enter role]','[Enter activity, result and receiver]'],['3','[Enter role]','[Enter activity, result and receiver]']],[.9,1.4,4.6])
field('Business rules and exceptions','State approval conditions, required checks, rejection paths and escalation rules.',16)
field('Constraints assumptions and dependencies','Record budget or date constraints, assumptions to validate and work needed from others.',10)

newpage(5,'Detailed requirement','Complete one page for each requirement. Use the same requirement ID in the register, test scenarios and work instructions.')
p('Requirement ID  [Enter BR number]\nTitle  [Enter short business title]\nOwner  [Enter name]     Priority  [Enter Must, Should, Could or Deferred]')
field('Requirement statement','As a role, I need a capability so that a business outcome is achieved.',16)
field('Rules inputs outputs and exceptions','What information is required, what should happen and what should happen when the normal path cannot continue?',16)
field('Acceptance criteria','Given a starting condition, when an action occurs, then what observable result must occur?',16)
field('Source rationale and dependencies','Who requested this, why is it needed and which other requirements does it depend on?',16)
p('Status  [Enter Proposed, Reviewed, Approved or Deferred]\nDecision or unresolved point  [Enter decision or question]\nRelated test and work instruction IDs  [Enter references]')

newpage(6,'Requirement register','Use this register to track the full set of requirements. Add rows as needed and retain the IDs when scope or priorities change.')
table(['ID','Requirement summary','Priority','Owner','Status'],[['[BR ID]','[Enter summary]','[Enter priority]','[Enter name]','[Enter status]'] for _ in range(5)],[.7,2.7,1.1,1.2,1.2])
field('Priority decisions','Which requirements are essential for the first release? What can be deferred and why?',24)
field('Conflicts and dependencies','Identify requirements that conflict, overlap or depend on another decision or delivery.',24)
h('Priority definitions')
p('Must: required for the agreed business outcome. Should: important but a temporary alternative is acceptable. Could: useful if capacity allows. Deferred: explicitly outside the current delivery scope.')
p('Register owner  [Enter name]     Last reviewed  [Enter date]')

newpage(7,'Data access and integrations','Capture what information is needed, who may use it and how it moves between systems. Confirm details with the relevant owners.')
table(['Data item','Source and owner','Required checks'],[['[Enter data item]','[Enter source and owner]','[Enter validation and quality rules]'],['[Enter data item]','[Enter source and owner]','[Enter validation and quality rules]']],[1.6,2.65,2.65])
field('Access and approval rights','Which roles may view, create, edit, approve or complete work? What restrictions apply?',18)
field('Integration and failure handling','Name the sending and receiving systems, event or timing, information transferred and the owner of failed transfers.',18)
field('Data migration audit and retention','Is existing data needed? Who validates it? What evidence and approved retention requirements apply?',18)
p('Data owner review  [Enter reviewer, date and remaining actions]\nLinked requirement IDs  [Enter references]')

newpage(8,'Reporting and service expectations','Describe the information needed to manage the process and the measurable operating expectations the solution must meet.')
table(['Report or measure','Audience and purpose','Definition and frequency'],[['[Enter name]','[Enter users and decision supported]','[Enter calculation, filters and frequency]'],['[Enter name]','[Enter users and decision supported]','[Enter calculation, filters and frequency]']],[1.5,2.7,2.7])
field('Report visibility and delivery','Who needs access, at which level of detail, and how will report access be checked?',18)
field('Performance availability and recovery','Specify expected response times, operating hours, peak demand and acceptable recovery times with owners.',18)
field('Usability accessibility and support','Describe user needs, device or location constraints, accessibility requirements and the support route.',18)
p('Related requirement IDs  [Enter references]\nAgreed service targets reviewed by  [Enter name and date]')

newpage(9,'Process and work instructions','Complete one page per task. Validate the steps with a person who performs the work before treating the instruction as approved.')
p('Instruction ID and title  [Enter WI number and title]\nPerformer  [Enter role]     Related requirement  [Enter BR ID]\nTrigger  [Enter starting event]\nPrerequisites  [Enter access, inputs and prior approvals]')
table(['Step','Action to perform','Expected result or evidence'],[['1','[Enter one specific action]','[Enter how success is checked]'],['2','[Enter one specific action]','[Enter how success is checked]'],['3','[Enter one specific action]','[Enter how success is checked]'],['4','[Enter one specific action]','[Enter how success is checked]']],[.5,3.4,3])
field('Exception and escalation','When should the user stop? Who resolves the problem and how is the resolution recorded?',14)
field('Completion and handoff','What must be verified, where is evidence stored and who receives the completed work?',14)
p('Validated by  [Enter name and date]\nProcedure owner and review date  [Enter name and date]')

newpage(10,'Acceptance testing and readiness','Define the business tests and the people who will confirm the outcome. Include the normal path, exceptions and access checks.')
table(['Test and BR ID','Scenario and expected result','Tester and outcome'],[['[Enter IDs]','[Enter normal path and expected result]','[Enter name, result and evidence]'],['[Enter IDs]','[Enter exception and expected result]','[Enter name, result and evidence]'],['[Enter IDs]','[Enter access test and expected result]','[Enter name, result and evidence]']],[1.1,3.7,2.1])
field('Test arrangements and acceptance threshold','Confirm environment, dates, representative data, required test coverage, defect thresholds and acceptance authority.',18)
field('Training communication and support readiness','Who needs training? What materials and support arrangements must be ready, and who owns them?',18)
field('Release and benefit review','Record required readiness approvals, the planned handover and when benefits will be measured.',18)
p('UAT owner  [Enter name]     Business acceptance date  [Enter date]')

newpage(11,'Actions decisions and approval','Resolve or explicitly accept outstanding issues. Approval establishes the business requirements baseline for the agreed scope.')
table(['Type and ID','Item and impact','Owner and due date'],[['[Question ID]','[Enter question and affected requirement]','[Enter name and date]'],['[Risk ID]','[Enter risk, impact and mitigation]','[Enter name and date]'],['[Action ID]','[Enter action and expected outcome]','[Enter name and date]']],[1.1,3.8,2])
field('Decision and approval conditions','Record Approve, Rework or Defer, the reason and any conditions that must be satisfied.',14)
table(['Approver role','Name and decision','Date and approval reference'],[['Business owner','[Enter name and decision]','[Enter date and reference]'],['Sponsor','[Enter name and decision]','[Enter date and reference]'],['System or delivery owner','[Enter name and review outcome]','[Enter date and reference]']],[1.65,2.65,2.6])
h('Version and change record')
p('Approved version  [Enter version]     Next review  [Enter date]\nChange summary  [Enter affected requirements and reason]\nImpact and approval  [Enter cost or date impact, approver and date]')
p('After approval, record scope changes before implementation. Update affected tests and work instructions and circulate the revised version to the named stakeholders.')

# Convert every response placeholder into a native Word plain-text content control.
# Separate multiline runs so the metadata fields on each line are also fillable.
for run in list(d.element.xpath('.//w:r')):
 if len(run.findall(qn('w:t'))) <= 1: continue
 parent=run.getparent(); position=parent.index(run)
 for child in run:
  if child.tag==qn('w:rPr'): continue
  replacement=OxmlElement('w:r')
  if run.find(qn('w:rPr')) is not None: replacement.append(deepcopy(run.find(qn('w:rPr'))))
  replacement.append(deepcopy(child));parent.insert(position,replacement);position+=1
 parent.remove(run)
counter=0
for para in list(d.element.xpath('.//w:p')):
 for run in list(para.findall(qn('w:r'))):
  ts=run.findall(qn('w:t'))
  if len(ts)!=1 or '[' not in (ts[0].text or ''):continue
  text=ts[0].text; parent=run.getparent(); index=parent.index(run)
  for piece in re.split(r'(\[[^\]]+\])',text):
   if not piece:continue
   copy=deepcopy(run);copy.find(qn('w:t')).text=piece
   copy.find(qn('w:t')).set(qn('xml:space'),'preserve')
   if piece.startswith('[') and piece.endswith(']'):
    counter+=1;sdt=OxmlElement('w:sdt');pr=OxmlElement('w:sdtPr')
    for name,val in [('alias',piece[1:-1]),('tag',f'RESPONSE_{counter:03d}'),('id',str(1000+counter))]:
     el=OxmlElement('w:'+name);el.set(qn('w:val'),val);pr.append(el)
    el=OxmlElement('w:text');el.set(qn('w:multiLine'),'1');pr.append(el);sdt.append(pr)
    content=OxmlElement('w:sdtContent');content.append(copy);sdt.append(content);parent.insert(index,sdt)
   else:parent.insert(index,copy)
   index+=1
  parent.remove(run)
d.core_properties.title='Probo Medical Full Fillable Business Requirements Workbook'
d.core_properties.subject='Native Word content controls for business requirements and work instructions'
d.core_properties.author='Probo Medical'
d.save(out)
with zipfile.ZipFile(out) as z:
 assert z.testzip() is None
 xml=z.read('word/document.xml').decode()
 assert xml.count('<w:sdt>')==counter
 assert 'frontdoor' not in xml and 'sid=' not in xml
print(f'{out.resolve()}\nNative fillable fields: {counter}')
