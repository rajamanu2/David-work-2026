from __future__ import annotations

import difflib
import hashlib
import subprocess
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(r"C:\Users\LIKKI\Documents\ChatGPT\David 2 Probo medical")
OLD = ROOT / "scratch" / "case-00010716-uat" / "force-app" / "main" / "default" / "classes"
NEW = ROOT / "scratch" / "case-00010716-uat-after" / "force-app" / "main" / "default" / "classes"
DOCX = ROOT / "outputs" / "case-00010716" / "Case_00010716_Apex_Code_Changes_Old_vs_New.docx"
REFERENCE = Path(
    r"C:\Users\LIKKI\.codex\plugins\cache\openai-curated-remote\openai-templates"
    r"\0.1.1\skills\artifact-template-strategy-memorandum\assets\reference.docx"
)
EXPECTED = {
    "cAuthURIForEval": (11, 1),
    "cGoogleAppAuthenticationWithSalesforce": (153, 134),
    "GoogleAuthTestClass": (92, 7),
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


doc = Document(DOCX)
document_text = "\n".join(
    [paragraph.text for paragraph in doc.paragraphs]
    + [paragraph.text for table in doc.tables for row in table.rows for cell in row.cells for paragraph in cell.paragraphs]
)

failures: list[str] = []
for class_name, expected in EXPECTED.items():
    old_lines = (OLD / f"{class_name}.cls").read_text(encoding="utf-8").splitlines()
    new_lines = (NEW / f"{class_name}.cls").read_text(encoding="utf-8").splitlines()
    matcher = difflib.SequenceMatcher(None, old_lines, new_lines)
    changed_nonblank: list[str] = []
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag in {"replace", "delete"}:
            changed_nonblank.extend(line.strip() for line in old_lines[i1:i2] if line.strip())
        if tag in {"replace", "insert"}:
            changed_nonblank.extend(line.strip() for line in new_lines[j1:j2] if line.strip())
    git_diff = subprocess.run(
        ["git", "diff", "--no-index", "--numstat", "--", str(OLD / f"{class_name}.cls"), str(NEW / f"{class_name}.cls")],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )
    fields = git_diff.stdout.strip().split("\t")
    added, removed = int(fields[0]), int(fields[1])
    if (added, removed) != expected:
        failures.append(f"{class_name}: expected {expected}, found {(added, removed)}")
    missing = [line for line in changed_nonblank if line not in document_text]
    if missing:
        failures.append(f"{class_name}: {len(missing)} changed nonblank lines missing from DOCX")
    old_meta = (OLD / f"{class_name}.cls-meta.xml").read_bytes()
    new_meta = (NEW / f"{class_name}.cls-meta.xml").read_bytes()
    if old_meta != new_meta:
        failures.append(f"{class_name}: metadata XML changed unexpectedly")

for required in (
    "OLD UAT CODE",
    "NEW UAT CODE",
    "Simple explanation:",
    "Code-only comparison. No Google Drive folder setup is included.",
    "0AfjH0000000T6rSAE",
    "4/4 tests passed",
):
    if required not in document_text:
        failures.append(f"required text missing: {required}")

with zipfile.ZipFile(DOCX) as package:
    names = set(package.namelist())
    xml = package.read("word/document.xml")
    if any(name.startswith("word/comments") for name in names):
        failures.append("comments remain in DOCX package")
    if b"<w:ins" in xml or b"<w:del" in xml:
        failures.append("tracked revisions remain in document.xml")

expected_reference_hash = "13BD3AE7AEF4B3AE76C5D65200ACD654C922782974DAC18672E973FAD93BD453"
if sha256(REFERENCE) != expected_reference_hash:
    failures.append("retained reference template hash changed")

if failures:
    print("FAIL")
    for failure in failures:
        print(f"- {failure}")
    raise SystemExit(1)

print("PASS")
for class_name, (added, removed) in EXPECTED.items():
    print(f"{class_name}: +{added}/-{removed}; all changed nonblank lines present; meta XML unchanged")
print("No comments or tracked revisions")
print("Reference template unchanged")
print(f"DOCX SHA256: {sha256(DOCX)}")
