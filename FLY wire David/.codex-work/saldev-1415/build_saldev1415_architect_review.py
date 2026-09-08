from pathlib import Path
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REFERENCE = Path(r"C:\Users\LIKKI\Documents\ChatGPT\Greate Plains David 3\SCC-4179-architect-review.docx")
OUTPUT = Path(r"C:\Users\LIKKI\Documents\ChatGPT\david\flywire-docgen\deliverables\SALDEV-1415_Architect_Review.docx")

NAVY = "0B2545"
BLUE = "2E74B5"
GRAY = "5B6777"
TEXT = "202B3C"
HEADER_FILL = "F1F3F6"
META_FILL = "E7F1FB"
RED = "D32F2F"
RED_FILL = "FCE8E8"
GREEN = "11856F"
GREEN_FILL = "E4F3EF"
AMBER = "B8750C"
AMBER_FILL = "FFF4D6"
WHITE = "FFFFFF"


def set_font(run, size=10, bold=False, color=TEXT, italic=False):
    run.font.name = "Arial"
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def format_paragraph(p, before=0, after=4, line=1.06, keep_next=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=95, start=120, bottom=95, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[idx] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = tr_pr.find(qn("w:tblHeader"))
    if node is None:
        node = OxmlElement("w:tblHeader")
        tr_pr.append(node)
    node.set(qn("w:val"), "true")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def clear_cell(cell):
    p = cell.paragraphs[0]
    for run in list(p.runs):
        p._p.remove(run._r)
    return p


def add_cell_text(cell, text, size=9.2, bold=False, color=TEXT, align=None):
    p = clear_cell(cell)
    if align is not None:
        p.alignment = align
    format_paragraph(p, after=0, line=1.02)
    set_font(p.add_run(text), size=size, bold=bold, color=color)
    return p


def make_table(headers, rows, widths, font_size=9.1, status_col=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    repeat_header(table.rows[0])
    for idx, value in enumerate(headers):
        shade(table.rows[0].cells[idx], HEADER_FILL)
        add_cell_text(table.rows[0].cells[idx], value, size=9.4, bold=True, color=NAVY)
    for row_data in rows:
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for idx, value in enumerate(row_data):
            color = TEXT
            bold = False
            if status_col is not None and idx == status_col:
                status = str(value).lower()
                bold = True
                if status.startswith("pass"):
                    shade(cells[idx], GREEN_FILL)
                    color = GREEN
                elif status.startswith("fail") or status.startswith("gap"):
                    shade(cells[idx], RED_FILL)
                    color = RED
                elif status.startswith("not") or status.startswith("gate"):
                    shade(cells[idx], AMBER_FILL)
                    color = AMBER
            add_cell_text(cells[idx], str(value), size=font_size, bold=bold, color=color,
                          align=WD_ALIGN_PARAGRAPH.CENTER if status_col == idx else WD_ALIGN_PARAGRAPH.LEFT)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_callout(label, body, fill, accent, size=11.0):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = clear_cell(cell)
    format_paragraph(p, after=3, line=1.0)
    set_font(p.add_run(label.upper()), size=9.5, bold=True, color=accent)
    p2 = cell.add_paragraph()
    format_paragraph(p2, after=0, line=1.08)
    set_font(p2.add_run(body), size=size, bold=True, color=NAVY)
    prevent_row_split(table.rows[0])
    set_table_geometry(table, [9360])
    spacer = doc.add_paragraph()
    format_paragraph(spacer, after=0)
    return table


def add_heading(text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    format_paragraph(p, before=8 if level == 1 else 6, after=5, line=1.0, keep_next=True)
    return p


def add_body(text, size=10.4, after=7, bold=False, color=TEXT):
    p = doc.add_paragraph()
    format_paragraph(p, after=after, line=1.12)
    set_font(p.add_run(text), size=size, bold=bold, color=color)
    return p


def add_page_break():
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_real_bullet(cell, text, color=TEXT):
    p = cell.add_paragraph(style="List Bullet")
    format_paragraph(p, after=2, line=1.0)
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.first_line_indent = Inches(-0.11)
    set_font(p.add_run(text), size=8.2, color=color)


def add_hyperlink(paragraph, text, url):
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.extend([color, underline])
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


OUTPUT.parent.mkdir(parents=True, exist_ok=True)
shutil.copy2(REFERENCE, OUTPUT)
doc = Document(OUTPUT)

# Retain the template package and recurring page furniture, but replace its body.
body = doc._element.body
for child in list(body):
    if child.tag != qn("w:sectPr"):
        body.remove(child)

section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)

# Replace the environment text in every footer while preserving PAGE fields.
for footer in (section.footer, section.first_page_footer, section.even_page_footer):
    for p in footer.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for node in p._p.iter(qn("w:t")):
            if node.text:
                node.text = node.text.replace("GreatPlainsMerge", "FlywirePartial")

# Keep the source style system but make its tokens explicit.
normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10)
normal.font.color.rgb = RGBColor.from_string(TEXT)
for name, size in (("Heading 1", 16), ("Heading 2", 13)):
    style = doc.styles[name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(BLUE)
    style.paragraph_format.keep_with_next = True

# PAGE 1 - decision and acceptance.
p = doc.add_paragraph()
format_paragraph(p, before=12, after=2, line=1.0, keep_next=True)
set_font(p.add_run("ARCHITECT REVIEW"), size=10, bold=True, color=BLUE)
p = doc.add_paragraph()
format_paragraph(p, after=4, line=1.0, keep_next=True)
set_font(p.add_run("SALDEV-1415 | Stepped-Up Pricing Date Automation"), size=25, bold=True, color=NAVY)
p = doc.add_paragraph()
format_paragraph(p, after=16, line=1.0, keep_next=True)
set_font(p.add_run("FlywirePartial sandbox | Read-only implementation and release-readiness assessment"), size=13, color=GRAY)

metadata = [
    ("Environment", "FlywirePartial | Org 00DhG0000000jOXUAY"),
    ("Review date", "26 August 2026"),
    ("Evidence", "Scoped metadata retrieval | SOQL and schema read-back | focused Apex tests"),
    ("Decision", "Changes Requested - not ready as the full original SALDEV-1415 scope"),
    ("Boundary", "No record, metadata, assignment, deployment, activation, or Production changes"),
]
table = doc.add_table(rows=len(metadata), cols=2)
table.style = "Table Grid"
set_table_geometry(table, [1900, 7460])
for ridx, (label, value) in enumerate(metadata):
    shade(table.rows[ridx].cells[0], META_FILL)
    add_cell_text(table.rows[ridx].cells[0], label, size=9.5, bold=True, color=NAVY)
    add_cell_text(table.rows[ridx].cells[1], value, size=9.3)
    prevent_row_split(table.rows[ridx])
set_table_geometry(table, [1900, 7460])
doc.add_paragraph().paragraph_format.space_after = Pt(0)

add_callout(
    "Architect decision",
    "Do not promote as complete SALDEV-1415 yet. The active Partial implementation validates New Business date automation, but Renewal is intentionally bypassed and several original acceptance paths remain unproven.",
    RED_FILL,
    RED,
    11.2,
)

add_heading("Acceptance outcome", 1)
acceptance_rows = [
    ("1", "New Business group dates", "Pass", "Ramped groups start from the Quote start date and calculate whole-month end dates."),
    ("2", "Continuous cascade", "Pass", "Each subsequent ramped group starts one day after the prior end date; tests cover reallocation."),
    ("3", "Grouped quote-line sync", "Pass", "QuoteLineDateSyncTrigger stamps group start, end, and term onto grouped lines."),
    ("4", "Read-only group drawer", "Pass", "Formula dates replace editable managed dates in the CPQ Line Editor field set."),
    ("5", "Term-total validation", "Pass", "Under- and over-allocation tests pass; current messages differ from the ticket's requested text."),
    ("6", "Renewal behavior", "Fail", "Current trigger and test deliberately bypass Renewal; original story text included Renewal."),
    ("7", "Remaining runtime proof", "Not run", "Ungrouped-line full-term and Opportunity Contract Start Date cascade were not independently exercised."),
]
make_table(("#", "Scope", "Result", "Evidence"), acceptance_rows, [500, 2150, 1450, 5260], font_size=8.55, status_col=2)

# PAGE 2 - control path.
add_page_break()
add_heading("Solution and control path", 1)
add_body(
    "The implemented control path is Apex-backed: quote dates and terms drive ramped group periods, an asynchronous helper shifts sibling groups, and grouped quote lines inherit the resulting period. Read-only formula fields expose dates in QLE without inviting edits.",
    size=10.6,
    after=10,
)

path_table = doc.add_table(rows=1, cols=7)
path_table.style = "Table Grid"
path_widths = [2100, 240, 2100, 240, 2100, 240, 2340]
set_table_geometry(path_table, path_widths)
boxes = [
    (0, "Quote inputs", ("Start Date", "End Date / header term", "New Business gate"), META_FILL, BLUE),
    (2, "Group automation", ("Sequential periods", "Sibling cascade", "Term-chain validation"), GREEN_FILL, GREEN),
    (4, "Quote-line sync", ("Start and end dates", "Subscription term", "Grouped lines only"), GREEN_FILL, GREEN),
    (6, "Release proof", ("Focused check-only", "Runtime UAT", "Post-deploy read-back"), AMBER_FILL, AMBER),
]
for idx, title, bullets, fill, accent in boxes:
    cell = path_table.rows[0].cells[idx]
    shade(cell, fill)
    p = clear_cell(cell)
    format_paragraph(p, after=5, line=1.0)
    set_font(p.add_run(title), size=10.2, bold=True, color=NAVY)
    for bullet in bullets:
        add_real_bullet(cell, bullet, color=TEXT)
for idx in (1, 3, 5):
    p = clear_cell(path_table.rows[0].cells[idx])
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p, after=0)
    set_font(p.add_run(">"), size=16, bold=True, color=BLUE)
prevent_row_split(path_table.rows[0])
set_table_geometry(path_table, path_widths)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
format_paragraph(p, before=6, after=8)
set_font(p.add_run("Figure 1. SALDEV-1415 stepped-up pricing date control path"), size=9, italic=True, color=GRAY)

add_callout(
    "Design reading",
    "Green paths are implemented and test-backed for New Business. Amber is the release evidence gate. Renewal remains a red scope decision until the story is formally narrowed or the trigger is extended.",
    META_FILL,
    BLUE,
    10.8,
)

add_heading("Architecture observations", 2)
observation_rows = [
    ("Recursion control", "QuoteLineGroupAsyncHelper.isExecuting prevents the helper's sibling updates from re-entering the trigger."),
    ("Unused handler", "QuoteLineGroupTriggerHandler contains only an isExecuting flag and is not referenced by the current retrieved trigger."),
    ("Validation key", "Term totals are grouped by product-code chains with name/proximity fallbacks; this is more complex than a simple quote-wide sum."),
    ("UI control", "Formula fields create a read-only experience; the code does not emit the ticket's manual-date-edit error message."),
]
make_table(("Observation", "Architect reading"), observation_rows, [2300, 7060], font_size=9.0)

# PAGE 3 - exact evidence.
add_page_break()
add_heading("Current implementation evidence", 1)
inventory_rows = [
    ("Apex Trigger", "QuoteLineGroupTrigger", "Active | API 67", "01qhG0000004ON7QAM | 17 Aug 2026"),
    ("Apex Trigger", "QuoteLineDateSyncTrigger", "Active | API 67", "01qhG0000004OOjQAM | 12 Aug 2026"),
    ("Apex Class", "QuoteLineGroupAsyncHelper", "Active | API 67", "01phG000000aImPQAU | 100% target coverage"),
    ("Apex Class", "QuoteLineGroupTriggerHandler", "Active | API 67", "01phG000000aDErQAM | one-line flag holder"),
    ("Apex Tests", "QuoteLineGroupTriggerTest; QuoteLineGroupAsyncHelperTest", "7 / 7 passed", "6 methods + 1 method | synchronous runs"),
    ("Custom Fields", "Effective_Start_Date__c; Effective_EndDate__c", "Formula (Date)", "Read-only mirrors of managed group dates"),
    ("Field Set", "SBQQ__QuoteLineGroup__c.SBQQ__LineEditor", "Verified live", "Name; Ramping; Effective Start; Term; Effective End"),
    ("Permission Set", "CPQ_Sales_Permissions", "Read = true; Edit = false", "0PSPb0000007QqfOAE | both formula fields"),
]
make_table(("Type", "Component", "State", "Evidence"), inventory_rows, [1800, 3100, 1600, 2860], font_size=8.35)

add_heading("Validation read-back", 2)
validation_rows = [
    ("Focused retrieval", "Pass", "Retrieve ID 09ShG000006G5JdUAK; exact Apex, fields, and permission set retrieved locally."),
    ("Trigger tests", "Pass", "QuoteLineGroupTriggerTest: 6/6; test-run coverage 83%; target trigger coverage 96%."),
    ("Helper test", "Pass", "QuoteLineGroupAsyncHelperTest: 1/1; test-run coverage 95%; helper 100%."),
    ("Live New Business", "Pass", "Q-37942: three 12-month ramped groups form a continuous 36-month chain; five lines match their groups."),
    ("Renewal / Amendment", "Gap", "Q-37894 Renewal and Q-37884 Amendment groups show null dates, consistent with the current bypass."),
]
make_table(("Evidence", "Result", "Read-back"), validation_rows, [2100, 1400, 5860], font_size=8.75, status_col=1)

add_callout(
    "Scope blocker",
    "The Partial implementation is internally consistent for New Business, but it is not equivalent to the original New + Renewal requirement. A release owner must approve the narrowed scope or sponsor the Renewal extension before promotion.",
    RED_FILL,
    RED,
    10.6,
)

# PAGE 4 - package and release boundary.
add_page_break()
add_heading("Implementation and release boundary", 1)
scope_rows = [
    ("Apex triggers", "QuoteLineGroupTrigger; QuoteLineDateSyncTrigger", "Include together"),
    ("Runtime classes", "QuoteLineGroupAsyncHelper; QuoteLineGroupTriggerHandler", "Review handler use; include only approved scope"),
    ("Test classes", "QuoteLineGroupTriggerTest; QuoteLineGroupAsyncHelperTest", "Run specified tests in validation"),
    ("Formula fields", "Effective_Start_Date__c; Effective_EndDate__c", "Include both"),
    ("Managed field-set override", "SBQQ__QuoteLineGroup__c.SBQQ__LineEditor", "Preserve order; remove editable dates"),
    ("Permission set", "CPQ_Sales_Permissions", "Retain read=true, edit=false for both formula fields"),
    ("Reference-only data", "Q-37942; Q-37892; Q-37894; Q-37884", "Do not deploy records"),
    ("Excluded", "Local scripts, retrieval ZIPs, QA renders, Jira export", "Engineering evidence only"),
]
make_table(("Component type", "Focused scope", "Required action"), scope_rows, [2100, 3900, 3360], font_size=8.65)

add_heading("Release sequence", 2)
release_rows = [
    ("Pre-deployment", "Confirm approved scope; retrieve target-org baselines; build a focused 10-member manifest; inspect permission-set delta; run check-only validation with both tests."),
    ("Deployment", "Promote only the approved metadata bundle in dependency order: fields and field set, permission set, runtime classes, triggers, then tests/verification."),
    ("Post-deployment", "Read back component status, formula definitions, field-set order, FLS, test outcome, coverage, and named New Business quote results."),
    ("Rollback", "Restore the approved metadata backup if QLE becomes unusable or date automation regresses; verify the prior field-set and Apex state."),
]
make_table(("Phase", "Required evidence"), release_rows, [1900, 7460], font_size=8.9)

add_callout(
    "No-deploy gate",
    "This review performed scoped retrieval, schema/SOQL read-back, and focused tests only. No deployment, activation, record edit, assignment change, or Production action was performed. Explicit approval is required before any promotion.",
    AMBER_FILL,
    AMBER,
    10.7,
)

# PAGE 5 - business UAT and approvals.
add_page_break()
add_heading("UAT and approval gates", 1)
uat_rows = [
    ("Baseline QLE", "Open Q-37942 in Partial, click Edit Lines, and confirm Effective Start Date and Effective EndDate are visible but not editable."),
    ("Date and line continuity", "Confirm SUP Year 1/2/3 are continuous from 24 Sep 2026 to 23 Sep 2029, and representative lines match each parent group's start, end, and term."),
    ("Cascade update", "On a cloned test quote, change an earlier group's term, Save/Quick Save, and confirm later groups shift with no gap or overlap."),
    ("Term validation", "On a cloned test quote, test both under- and over-allocation; capture the actual error text and confirm no invalid save."),
    ("Unproven regressions", "Prove the ungrouped line inherits the full Quote period and that an approved Contract Start Date update shifts Quote, group, and line periods before ordering."),
    ("Renewal decision", "Either document Renewal as excluded for this release or execute equivalent Renewal tests after the trigger is extended."),
]
make_table(("Test", "Pass evidence"), uat_rows, [2100, 7260], font_size=8.55)

p = doc.add_paragraph()
format_paragraph(p, before=1, after=7)
set_font(p.add_run("Baseline record: "), size=9, bold=True, color=NAVY)
add_hyperlink(p, "Q-37942 in FlywirePartial", "https://flywire--partial.sandbox.my.salesforce.com/a2NhG000003o60XUAQ")

add_heading("Required approval gates", 2)
gate_rows = [
    ("1", "Confirm scope", "Product Owner / Architect", "Written decision on New Business-only versus Renewal support."),
    ("2", "Validate focused package", "Release Manager", "Check-only job ID, exact 10-member scope, both test results, coverage."),
    ("3", "Run business UAT", "QA / Business Owner", "Record IDs, before/after dates, actual error text, screenshots."),
    ("4", "Approve promotion", "Architect / Release Owner", "Explicit approval after all acceptance gaps are closed or waived."),
    ("5", "Promote and read back", "DevOps / Salesforce Admin", "Deployment receipt, component/FLS read-back, audit trail, sign-offs."),
]
make_table(("#", "Required action", "Owner", "Evidence needed"), gate_rows, [500, 2400, 2200, 4260], font_size=8.25)

add_heading("Ready-to-paste architect response", 2)
story_comment = (
    "Architect review completed for SALDEV-1415 in FlywirePartial. Changes Requested. The implementation is test-backed for New Business: 7/7 focused methods passed; QuoteLineGroupTrigger is 96% covered; QuoteLineDateSyncTrigger and QuoteLineGroupAsyncHelper are 100%; and Q-37942 shows three continuous groups with five lines matching their parent periods. "
    "Effective dates are read-only in QLE. Renewal is bypassed despite the original scope, while the requested manual-date error text, ungrouped-line behavior, and Contract Start Date cascade still require decision or proof. Confirm scope, run focused check-only validation and documented UAT, then capture post-deployment read-back. No deployment, activation, record, assignment, or Production changes were made."
)
add_callout("Story comment", story_comment, META_FILL, BLUE, 8.7)

doc.core_properties.title = "SALDEV-1415 Architect Review"
doc.core_properties.subject = "Stepped-up pricing date automation implementation and release readiness"
doc.core_properties.keywords = "SALDEV-1415; Salesforce CPQ; stepped-up pricing; architect review; FlywirePartial"
doc.save(OUTPUT)
print(OUTPUT)
