from __future__ import annotations

import importlib.util
from pathlib import Path
from typing import Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[2]
BASE_BUILDER = ROOT / "scratch" / "next-four-documents" / "build_story_docs.py"
REFERENCE = ROOT / "scratch" / "case-00010716-uat-team-guide" / "autofast-reference.docx"
OUTPUT = ROOT / "outputs" / "case-00010716" / "Case_00010716_UAT_Photo_Import_Reproduction_and_Change_Guide_AutoFast.docx"


UAT_BASE = "https://probomedical--uat.sandbox.lightning.force.com"
PRODUCT_ITEM_URL = f"{UAT_BASE}/lightning/r/ProductItem__c/a06jH0000004E2rQAE/view"
RMA_URL = f"{UAT_BASE}/lightning/r/RMA__c/a0mjH0000000BWXQA2/view"
PRIMARY_EVAL_URL = f"{UAT_BASE}/lightning/r/Evaluation_Natalie__c/a1tjH0000003n98QAA/view"
CONTROL_EVAL_URL = f"{UAT_BASE}/lightning/r/Evaluation_Natalie__c/a1tjH0000003n99QAA/view"
AUTH_URL = "https://probomedical--uat.sandbox.my.salesforce.com/apex/AuthenticationGoogleDrive?id=a1tjH0000003n98QAA"


def load_helpers():
    spec = importlib.util.spec_from_file_location("autofast_helpers", BASE_BUILDER)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def clear_document_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        paragraph._p.remove(child)


def prepare_header_footer(doc: Document) -> None:
    for section in doc.sections:
        for part in (section.header, section.footer):
            while len(part.paragraphs) > 1:
                p = part.paragraphs[-1]._element
                p.getparent().remove(p)
            clear_paragraph(part.paragraphs[0])


def add_hyperlink(paragraph, text: str, url: str, h, *, bold: bool = False, color: str | None = None) -> None:
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    run_fonts = OxmlElement("w:rFonts")
    run_fonts.set(qn("w:ascii"), h.PRESET["body_font"])
    run_fonts.set(qn("w:hAnsi"), h.PRESET["body_font"])
    run_properties.append(run_fonts)
    run_color = OxmlElement("w:color")
    run_color.set(qn("w:val"), color or h.BLUE)
    run_properties.append(run_color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "19")
    run_properties.append(size)
    if bold:
        run_properties.append(OxmlElement("w:b"))
    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_rich_table(doc: Document, h, headers: Sequence[str], rows: Sequence[Sequence[object]], widths_dxa: Sequence[int], status_column: int | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    h.set_table_geometry(table, widths_dxa)
    h.set_table_borders(table)
    h.set_repeat_table_header(table.rows[0])

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        h.set_cell_shading(cell, h.LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        h.set_run_font(r, size=9.5, color=h.NAVY, bold=True)

    for row_index, row_values in enumerate(rows):
        row = table.add_row()
        if row_index % 2 == 1:
            for cell in row.cells:
                h.set_cell_shading(cell, "F8FAFC")
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if isinstance(value, tuple):
                label, url = value
                add_hyperlink(p, str(label), str(url), h, bold=(idx == 0))
            else:
                value_text = str(value)
                color = "222222"
                bold = False
                if status_column is not None and idx == status_column:
                    normalized = value_text.upper()
                    if any(term in normalized for term in ("PASS", "READY", "YES", "ACTIVE", "COMPLETE")):
                        color = h.GREEN
                        bold = True
                    elif any(term in normalized for term in ("PENDING", "REQUIRED", "HOLD", "NO")):
                        color = h.AMBER
                        bold = True
                r = p.add_run(value_text)
                h.set_run_font(r, size=9.5, color=color, bold=bold)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def add_link_line(doc: Document, h, label: str, link_text: str, url: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    left = p.add_run(f"{label}: ")
    h.set_run_font(left, bold=True, color=h.NAVY)
    add_hyperlink(p, link_text, url, h)


def build() -> Path:
    h = load_helpers()
    doc = Document(REFERENCE)
    clear_document_body(doc)
    prepare_header_footer(doc)
    h.configure_styles(doc)
    h.configure_section(doc, "00010716")

    # Make the retained header UAT-specific without changing its established layout.
    header = doc.sections[0].header.paragraphs[0]
    for run in header.runs:
        if "DEVDO VALIDATION" in run.text:
            run.text = run.text.replace("DEVDO VALIDATION", "UAT TEST GUIDE")

    doc.core_properties.title = "Case 00010716 UAT Photo Import Reproduction and Change Guide"
    doc.core_properties.subject = "AutoFast-style Probo Medical UAT reproduction, verification, and technical change guide"
    doc.core_properties.author = "Probo Medical Salesforce Team"
    doc.core_properties.keywords = "Salesforce, UAT, AutoFast, Google Drive, OAuth, Evaluation, photo import, regression"
    doc.core_properties.comments = "Contains no deployment IDs, validation IDs, test-run IDs, images, or session credentials."

    title = doc.add_paragraph("AUTOFAST UAT REPRODUCTION & CHANGE GUIDE", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle = doc.add_paragraph(
        "Evaluation Google Drive Photo Import Must Return to the Intended Evaluation",
        style="Subtitle",
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT

    h.add_metadata_rows(
        doc,
        [
            ("Case", "00010716"),
            ("Prepared", "September 4, 2026"),
            ("Environment", "Probo Medical UAT sandbox"),
            ("Audience", "Salesforce, Repair Operations, QA, and UAT testers"),
            ("Status", "READY FOR MANUAL UAT | DEDICATED GOOGLE DRIVE FOLDER REQUIRED"),
        ],
    )
    h.add_callout(
        doc,
        "Current position",
        "The focused code correction and dedicated UAT records are ready. Automated regression testing passed. The remaining activity is a manual Google Drive photo-import test using a dedicated UAT folder link; no deployment action is required for this test.",
        "green",
    )

    doc.add_paragraph("1. Purpose and Environment Boundary", style="Heading 1")
    h.add_labeled_paragraph(
        doc,
        "Purpose",
        "Give the team one repeatable procedure to understand the original regression, confirm what changed, run the live UAT photo-import test, and record evidence that only the intended Evaluation was updated.",
    )
    add_rich_table(
        doc,
        h,
        ["Item", "Confirmed value"],
        [
            ("Organization", "Probo Medical"),
            ("Environment", "UAT sandbox"),
            ("Organization ID", "00DjH0000000rYzUAI"),
            ("Instance", "USA1310S"),
            ("Sandbox domain", "probomedical--uat.sandbox"),
            ("Manual-test readiness", "READY - folder link required"),
        ],
        [2700, 6660],
        status_column=1,
    )
    h.add_callout(
        doc,
        "Security note",
        "Use the standard UAT login and approved Google account. Do not place frontdoor URLs, access tokens, refresh tokens, or authentication codes in test notes, email, or this document.",
        "blue",
    )
    h.add_labeled_paragraph(
        doc,
        "Environment impact",
        "UAT contains the focused correction and dedicated test data. DevDO was not changed during this UAT activity. Production was not changed during this UAT activity. A separate Production rollback remains pending because its authentication session timed out.",
    )

    h.add_page_break(doc)
    doc.add_paragraph("2. Original Defect and Reproduction Scenario", style="Heading 1")
    h.add_labeled_paragraph(
        doc,
        "Original symptom",
        "After Google authentication, Salesforce could lose the Evaluation where the user started. The callback could then raise a null-reference error or select the most recently modified Evaluation instead of the intended record.",
    )
    h.add_labeled_paragraph(
        doc,
        "Root cause",
        "The OAuth request did not preserve the originating Evaluation ID through the redirect. The callback attempted to recover by querying the latest Evaluation modified by the current user, which was unsafe when another Evaluation had been modified more recently or when the shared execution user had no usable record.",
    )
    h.add_callout(
        doc,
        "Important testing distinction",
        "The steps below describe how the old regression was triggered. Because the correction is already active in UAT, the current expected result is the protected behavior in the final column. Reproducing the old wrong-record outcome would require an approved pre-fix environment or code rollback and is not part of this UAT.",
        "amber",
    )
    add_rich_table(
        doc,
        h,
        ["Step", "Historical trigger", "Old result", "Current UAT result"],
        [
            ("1", "Open primary Evaluation A with a linked RMA, Product Item, and Drive folder.", "Authentication starts from A.", "Authentication starts from A."),
            ("2", "Modify a separate Evaluation B after A so B becomes the latest modified record.", "B can become the fallback candidate.", "B is irrelevant to callback selection."),
            ("3", "Start Google authentication from A and complete sign-in.", "The redirect does not retain A.", "The redirect carries A in encoded OAuth state."),
            ("4", "Return to Salesforce and save imported image references.", "A null error may occur, or B may be selected.", "A is restored and validated before processing."),
            ("5", "Refresh both Evaluations.", "The intended record may be unchanged while another record is at risk.", "A receives the intended references; B remains unchanged."),
        ],
        [650, 2970, 2730, 3010],
    )

    doc.add_paragraph("3. What Changed", style="Heading 1")
    add_rich_table(
        doc,
        h,
        ["Apex component", "Change", "Result"],
        [
            ("cAuthURIForEval", "Adds an Evaluation-aware constructor. The original Evaluation ID is appended to OAuth state and the complete state value is URL-encoded.", "Google can return Salesforce to the exact Evaluation where authentication started."),
            ("cGoogleAppAuthenticationWithSalesforce", "Reads the page ID first, then safely restores the Evaluation ID from OAuth state. It validates the ID type and removes the latest-modified-Evaluation fallback.", "The controller stops rather than guessing when an Evaluation cannot be identified."),
            ("cGoogleAppAuthenticationWithSalesforce", "Uses safe relationship and record checks for Evaluation, RMA, Product Item, Product Test, and Drive folder prerequisites.", "Missing data produces a readable page error instead of a null-pointer failure."),
            ("cGoogleAppAuthenticationWithSalesforce", "Normalizes either a raw folder ID or a full Drive folder URL, removes query-string values such as resourcekey, and ignores unsafe blank or malformed file mappings.", "Folder and filename processing is more tolerant and predictable."),
            ("GoogleAuthTestClass", "Adds a competing Evaluation, callback-state restoration, full-folder-URL normalization, dated-folder handling, and negative tests for missing or invalid state.", "Automated coverage now proves the intended Evaluation is updated and the competing Evaluation stays unchanged."),
        ],
        [2600, 4040, 2720],
    )
    h.add_callout(
        doc,
        "Scope",
        "Only cAuthURIForEval, cGoogleAppAuthenticationWithSalesforce, and GoogleAuthTestClass were changed for this correction. No Visualforce page, Flow, object, field, permission, profile, or Production configuration change is included in this guide.",
        "blue",
    )

    h.add_page_break(doc)
    doc.add_paragraph("4. Dedicated UAT Test Records", style="Heading 1")
    h.add_labeled_paragraph(
        doc,
        "Why dedicated records are required",
        "Salesforce record IDs and auto-numbers differ by environment. The DevDO records cannot be reused in UAT, so the records below were created specifically for this test.",
    )
    add_rich_table(
        doc,
        h,
        ["Record", "UAT record", "Test purpose"],
        [
            ("Product Item", ("552082", PRODUCT_ITEM_URL), "Dedicated source Product Item. Serial number UAT-00010716-20260904. Image/folder values are blank until the tester supplies the UAT folder URL."),
            ("RMA", ("A-218876", RMA_URL), "Linked to Product Item 552082; tagged UAT-00010716-20260904; Customer Repair Evaluation record type; Depot Repair enabled."),
            ("Primary Evaluation", ("E-2026-036025", PRIMARY_EVAL_URL), "This is Evaluation A and must receive the imported image references."),
            ("Control Evaluation", ("E-2026-036026", CONTROL_EVAL_URL), "This is Evaluation B and must remain unchanged after the import."),
        ],
        [1800, 2100, 5460],
    )
    add_rich_table(
        doc,
        h,
        ["Evaluation", "Lens ID", "Whole ID", "SN ID", "Folder ID"],
        [
            ("E-2026-036025 - Primary", "Blank", "Blank", "Blank", "Blank"),
            ("E-2026-036026 - Control", "Blank", "Blank", "Blank", "Blank"),
        ],
        [2700, 1665, 1665, 1665, 1665],
    )
    add_link_line(doc, h, "Authentication page", "Open the UAT authentication page for E-2026-036025", AUTH_URL)
    h.add_callout(
        doc,
        "Do not substitute records",
        "Use the named primary and control Evaluations unless the Salesforce owner approves new dedicated UAT records and records their IDs in the execution log.",
        "amber",
    )

    doc.add_paragraph("5. Manual UAT Preconditions", style="Heading 1")
    add_rich_table(
        doc,
        h,
        ["Requirement", "Expected setup"],
        [
            ("Google Drive folder", "A dedicated UAT folder URL that is accessible to the approved Google account."),
            ("Recognizable files", "C-Lens.jpg, C-Cap.jpg, C-Whole.jpg, and C-SN.jpg. Additional test images may be included if their names follow the existing process."),
            ("Salesforce access", "An authorized UAT user who can edit Product Item 552082 and view/edit the primary Evaluation."),
            ("Browser", "Use a normal UAT browser session that can complete the Google sign-in and return to Salesforce."),
            ("Baseline", "Confirm both Evaluations have blank test image-reference fields before starting."),
        ],
        [2700, 6660],
    )

    h.add_page_break(doc)
    doc.add_paragraph("6. Manual UAT Procedure", style="Heading 1")
    h.add_callout(
        doc,
        "Execution boundary",
        "Run this procedure only in Probo Medical UAT using the dedicated records. This is a data-level functional test; it does not require another deployment.",
        "blue",
    )
    add_rich_table(
        doc,
        h,
        ["Step", "Tester action", "Expected result / evidence"],
        [
            ("1", ("Open Product Item 552082", PRODUCT_ITEM_URL), "The dedicated Product Item opens in UAT."),
            ("2", "Click Edit and enter the dedicated UAT Google Drive folder URL in the Image field used by this process. Save.", "The Product Item saves successfully and displays the test folder value."),
            ("3", ("Open primary Evaluation E-2026-036025", PRIMARY_EVAL_URL), "The primary Evaluation opens and is linked to RMA A-218876 / Product Item 552082."),
            ("4", "Record the current Lens ID, Whole ID, SN ID, and Folder ID values in the execution record. They should be blank at baseline.", "A before-test baseline is captured."),
            ("5", ("Open control Evaluation E-2026-036026", CONTROL_EVAL_URL), "The control image-reference fields are blank. Do not edit this record."),
            ("6", ("Open the authentication page for the primary Evaluation", AUTH_URL), "The page is scoped to primary Evaluation ID a1tjH0000003n98QAA."),
            ("7", "Click Authenticate and complete the Google sign-in with the approved account.", "Google returns the browser to the Salesforce UAT authentication page without a null-reference error."),
            ("8", "After returning to Salesforce, review the listed files and click Save.", "The save completes for the primary Evaluation."),
            ("9", ("Refresh primary Evaluation E-2026-036025", PRIMARY_EVAL_URL), "The expected image/reference fields are populated from the recognizable UAT files."),
            ("10", ("Refresh control Evaluation E-2026-036026", CONTROL_EVAL_URL), "All control image-reference fields remain unchanged and blank."),
            ("11", "Record the after-test values, result, tester, date/time, and any message shown. Do not copy OAuth codes or tokens.", "The test evidence is complete and safe to share with the team."),
        ],
        [650, 5040, 3670],
    )
    h.add_callout(
        doc,
        "Pass condition",
        "The primary Evaluation receives the intended photo references, the control Evaluation remains unchanged, the full folder URL is accepted without retaining resourcekey as part of the folder ID, and no null-reference error occurs.",
        "green",
    )

    h.add_page_break(doc)
    doc.add_paragraph("7. Acceptance Criteria and Automated Coverage", style="Heading 1")
    add_rich_table(
        doc,
        h,
        ["Acceptance check", "Pass requirement", "Current status"],
        [
            ("Exact Evaluation context", "Callback returns to E-2026-036025 and does not select another record.", "PENDING MANUAL UAT"),
            ("Primary image mapping", "Recognizable files populate the expected primary Evaluation image/reference fields.", "PENDING MANUAL UAT"),
            ("Control protection", "E-2026-036026 remains unchanged.", "PENDING MANUAL UAT"),
            ("Folder normalization", "A full Drive folder URL is accepted and resourcekey is excluded from the folder ID.", "PENDING MANUAL UAT"),
            ("Safe error behavior", "Missing or invalid prerequisites show a readable error and do not update another Evaluation.", "AUTOMATED PASS"),
            ("Focused regression suite", "All four dedicated test methods complete without failure.", "AUTOMATED PASS"),
        ],
        [3000, 4440, 1920],
        status_column=2,
    )
    h.add_labeled_paragraph(
        doc,
        "Successful automated test methods",
        "mytest; mytest2; mytest3; case00010716RejectsMissingOrInvalidCallbackState.",
    )
    h.add_labeled_paragraph(
        doc,
        "Automated protection proved",
        "The tests restore the intended Evaluation from callback state, normalize complete folder URLs, handle dated subfolders, reject missing or invalid state, update the intended Evaluation, and leave a second Evaluation unchanged.",
    )

    doc.add_paragraph("8. Troubleshooting", style="Heading 1")
    add_rich_table(
        doc,
        h,
        ["Observed message or condition", "Likely cause", "Tester action"],
        [
            ("Evaluation could not be identified", "The authentication page was opened without the primary Evaluation ID or callback state was invalid.", "Return to E-2026-036025 and use the supplied authentication link."),
            ("Associated RMA or Product Item unavailable", "The primary Evaluation relationship is missing or no longer accessible.", "Stop the test and ask the Salesforce owner to verify the dedicated records."),
            ("No valid Google Drive folder", "The Product Item folder value is blank, malformed, or inaccessible.", "Confirm the complete folder URL, access permission, and saved Product Item value."),
            ("No valid image folder or images listed", "File names do not match the process or the expected dated folder cannot be identified.", "Use recognizable names and confirm the folder structure with the process owner."),
            ("Control Evaluation changed", "The regression protection did not hold.", "Stop, capture only Salesforce field values and timestamps, and escalate as a failed UAT result. Do not proceed to Production."),
        ],
        [3000, 3060, 3300],
    )

    h.add_page_break(doc)
    doc.add_paragraph("9. Test Execution Record", style="Heading 1")
    h.add_labeled_paragraph(
        doc,
        "Instructions",
        "Complete this table after the manual test. Record field values and user-visible messages only. Do not paste session tokens, authorization codes, refresh tokens, or cookies.",
    )
    add_rich_table(
        doc,
        h,
        ["Evidence item", "Tester entry"],
        [
            ("Tester name", ""),
            ("Execution date and time", ""),
            ("Google Drive folder label or approved reference", ""),
            ("Primary baseline - Lens / Whole / SN / Folder", ""),
            ("Primary after test - Lens / Whole / SN / Folder", ""),
            ("Control baseline - Lens / Whole / SN / Folder", ""),
            ("Control after test - Lens / Whole / SN / Folder", ""),
            ("User-visible messages", ""),
            ("Overall result", "PASS / FAIL"),
            ("Notes or follow-up owner", ""),
        ],
        [3600, 5760],
    )
    h.add_callout(
        doc,
        "Team handoff",
        "The UAT fix is ready for functional confirmation. Supply the dedicated folder URL, execute Section 6, confirm the primary Evaluation is updated and the control Evaluation is untouched, and return the completed execution record to the Salesforce team.",
        "green",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
