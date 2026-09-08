from __future__ import annotations

import importlib.util
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


ROOT = Path(__file__).resolve().parents[2]
BASE_BUILDER = ROOT / "scratch" / "next-four-documents" / "build_story_docs.py"
OUTPUT = ROOT / "outputs" / "case-00010716" / "Case_00010716_Evaluation_Google_Drive_Authentication_AutoFast_Review.docx"


def load_helpers():
    spec = importlib.util.spec_from_file_location("autofast_helpers", BASE_BUILDER)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def build() -> Path:
    h = load_helpers()
    doc = Document()
    h.configure_styles(doc)
    h.configure_section(doc, "00010716")

    doc.core_properties.title = "Case 00010716 Evaluation Google Drive Authentication AutoFast Review"
    doc.core_properties.subject = "DevDO implementation, validation, and UAT reproduction guide"
    doc.core_properties.author = "Probo Medical Salesforce Team"
    doc.core_properties.keywords = "Salesforce, DevDO, AutoFast, Google Drive, OAuth, Evaluation, UAT"

    title = doc.add_paragraph("AUTOFAST DEVDO VALIDATION & UAT GUIDE", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle = doc.add_paragraph(
        "Evaluation Google Drive Authentication Must Return to the Correct Evaluation",
        style="Subtitle",
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT

    h.add_metadata_rows(
        doc,
        [
            ("Primary Production Case", "00010716"),
            ("Related Production Case", "00010710"),
            ("Prepared", "September 3, 2026"),
            ("Environment", "Probo Medical DevDO sandbox"),
            ("Review Status", "DEVDO PASS | READY FOR BUSINESS UAT"),
        ],
    )
    h.add_callout(
        doc,
        "Architect determination",
        "The focused fix is deployed and active in DevDO. Automated regression passed. Manual Google authentication UAT remains required before a Production change set or Production case closure.",
        "green",
    )

    doc.add_paragraph("1. Case Summary", style="Heading 1")
    h.add_table(
        doc,
        ["Case", "Asset", "Evaluation", "RMA", "Reported result"],
        [
            ("00010710", "554076", "E-2026-036222", "A-219278", "Authentication returned a null-reference error."),
            ("00010716", "554470", "E-2026-036262", "A-219744", "Authentication returned a null-reference error."),
        ],
        [1200, 1200, 1800, 1500, 3660],
    )
    h.add_labeled_paragraph(
        doc,
        "Business impact",
        "A repair user could not reliably complete Evaluation Google Drive authentication and continue image processing. Under concurrent activity, the callback could also select a different recently modified Evaluation.",
    )

    doc.add_paragraph("2. What Happened and What We Changed", style="Heading 1")
    h.add_labeled_paragraph(
        doc,
        "Why it happened",
        "When the user returned from Google, Salesforce had lost the Id of the Evaluation where authentication started. The page tried to recover by choosing the latest Evaluation modified by the current user. In a shared site/guest context, that can be a different Evaluation or no usable record at all.",
    )
    h.add_labeled_paragraph(
        doc,
        "How we fixed the callback",
        "The authentication request now carries the originating Evaluation Id securely through OAuth state. When Google returns the user to Salesforce, the page restores that Id and confirms it belongs to an Evaluation before doing any work.",
    )
    h.add_labeled_paragraph(
        doc,
        "Wrong-record protection",
        "The latest-record-by-user fallback was removed. If Salesforce cannot identify the original Evaluation, processing stops and shows a clear message instead of guessing.",
    )
    h.add_labeled_paragraph(
        doc,
        "Clearer error handling",
        "If the Evaluation, RMA, Product Item, Product Test, or Google Drive folder is missing, the user now receives a readable page message instead of a null-reference error.",
    )
    h.add_labeled_paragraph(
        doc,
        "Cleaner folder handling",
        "Salesforce now extracts the folder Id correctly when the saved Google Drive value is a full folder URL or includes a resource-key query string.",
    )
    h.add_callout(
        doc,
        "No-hardcoding control",
        "This fix introduces no Salesforce record Ids or business configuration constants. Existing legacy Google OAuth credentials predate this change and remain technical debt; a separate Named Credential migration is recommended and is not part of this case fix.",
        "blue",
    )

    h.add_page_break(doc)
    doc.add_paragraph("3. Steps to Reproduce and Validate", style="Heading 1")
    h.add_callout(
        doc,
        "Execution boundary",
        "Run these steps in DevDO with sandbox test records. Do not use the Production record Ids above as DevDO test data, and do not repeat in Production until an approved deployment is completed.",
        "blue",
    )
    h.add_labeled_paragraph(
        doc,
        "Preconditions",
        "Use an authorized DevDO repair user. Prepare Evaluation A with an associated RMA and Product Item whose Google Drive folder value is populated. Prepare Evaluation B as a separate competing record. The tester must be able to complete the external Google sign-in.",
    )
    h.add_table(
        doc,
        ["Step", "Action"],
        [
            ("1", "Open Evaluation A in DevDO and confirm its RMA, Product Item, and Google Drive folder prerequisites."),
            ("2", "In a second browser tab, open Evaluation B and save a harmless test-data update so it is modified after Evaluation A."),
            ("3", "Return to Evaluation A and start the Google Drive authentication action."),
            ("4", "Complete the Google sign-in and consent flow, then allow the browser to return to Salesforce."),
            ("5", "Confirm the callback returns to Evaluation A without an 'Attempt to de-reference a null object' message."),
            ("6", "Refresh both records. Confirm Evaluation A received the expected Drive image references and Evaluation B was not changed."),
        ],
        [900, 8460],
    )
    h.add_callout(
        doc,
        "What the tester should see",
        "Salesforce returns to Evaluation A, processes its configured Drive folder, and leaves Evaluation B untouched. If a required relationship or folder is missing, the page explains what is missing and does not update another Evaluation.",
        "green",
    )
    h.add_labeled_paragraph(
        doc,
        "What happened before the fix",
        "The callback showed a null-reference error or selected the most recently modified Evaluation instead of the Evaluation where authentication started.",
    )

    doc.add_paragraph("4. Verification Evidence", style="Heading 1")
    h.add_table(
        doc,
        ["Evidence", "Verified result"],
        [
            ("Check-only validation", "0AfiK0000000iNdSAI - Succeeded; 3/3 components; 4/4 tests."),
            ("DevDO deployment", "0AfiK0000000iPFSAY - Succeeded; 3/3 components; 4/4 tests."),
            ("Post-deployment regression", "707iK000000793S - 4 tests passed; 0 failures; 100% pass rate."),
            ("Focused coverage", "Controller 301/397 lines (75.8%); OAuth URI helper 15/17 lines (88.2%)."),
            ("Metadata read-back", "All three Apex classes Active; modified 2026-09-02 21:43:16 UTC."),
        ],
        [2700, 6660],
        status_column=1,
    )

    h.add_page_break(doc)
    doc.add_paragraph("5. Deployment Scope", style="Heading 1")
    h.add_table(
        doc,
        ["Component type", "API name", "State"],
        [
            ("Apex Class", "cGoogleAppAuthenticationWithSalesforce", "Active - Modified"),
            ("Apex Class", "cAuthURIForEval", "Active - Modified"),
            ("Apex Test", "GoogleAuthTestClass", "Active - Regression"),
        ],
        [2100, 5580, 1680],
        status_column=2,
    )
    h.add_labeled_paragraph(
        doc,
        "Production safety",
        "Only DevDO was changed. No Production metadata, Production data, or Production case status was modified.",
    )
    h.add_labeled_paragraph(
        doc,
        "Change set status",
        "A story-specific outbound change set has not yet been created or uploaded for this fix.",
    )

    doc.add_paragraph("6. UAT Acceptance and Closure Gate", style="Heading 1")
    h.add_table(
        doc,
        ["Gate", "Status", "Required evidence"],
        [
            ("DevDO implementation deployed", "YES", "Deployment 0AfiK0000000iPFSAY"),
            ("Automated regression passed", "YES", "Test run 707iK000000793S"),
            ("Manual DevDO Google OAuth test", "PENDING", "Tester confirms correct callback and image mapping"),
            ("Competing-Evaluation UAT", "PENDING", "Evaluation B remains unchanged"),
            ("Outbound change set", "NO - PENDING", "Create only after UAT acceptance"),
            ("Production deployment", "NO - PENDING", "Validate and deploy after approval"),
            ("Close Production cases", "HOLD", "Close after Production smoke-test evidence"),
        ],
        [3180, 1740, 4440],
        status_column=1,
    )
    h.add_callout(
        doc,
        "Message to tester",
        "The Evaluation Google Drive authentication fix for cases 00010710 and 00010716 is now in DevDO and the automated tests passed. Please follow Section 3 and let us know whether Salesforce returns to the Evaluation where you started, leaves the second Evaluation unchanged, and completes without the null-reference error.",
        "amber",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
