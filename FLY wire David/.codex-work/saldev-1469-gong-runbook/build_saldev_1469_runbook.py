from __future__ import annotations

import importlib.util
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt


ROOT = Path(r"C:\Users\LIKKI\Documents\ChatGPT\david")
TASK = ROOT / ".codex-work" / "saldev-1469-gong-runbook"
REFERENCE = Path(r"C:\Users\LIKKI\Downloads\SALDEV-1413-architect-review.docx")
OUTPUT = ROOT / "flywire-docgen" / "deliverables" / "SALDEV-1469_Gong_ECA_Verification_Runbook.docx"
DIAGRAM = TASK / "saldev-1469-verification-path.png"
BASE_BUILDER = ROOT / ".codex-work" / "saldev-1499-review" / "build_saldev_1499_review.py"
REFERENCE_SHA256 = "9C1A9F936367232B5DD8F9C08C98A6C801259257B220C84E76CE3CFA3937C13F"

INSTALLED_PACKAGES_URL = "https://flywire--partial.sandbox.lightning.force.com/lightning/setup/ImportedPackage/home"
GONG_SIGNIN_URL = "https://us-4592.app.gong.io/welcome/sign-in?targeturl=%2Fhome"
ZOOM_MESSAGE = (
    "Hi David, please open the Salesforce Installed Packages link and confirm that Gong ECA v0.1, "
    "namespace gongeca, is installed in Sandbox (Partial). Then open the Gong sign-in link and select "
    "Sign in with Salesforce. If prompted, choose Use Custom Domain and enter flywire--partial.sandbox. "
    "Please share your screen and stop at any Allow/Authorization page so we can verify the scopes before proceeding."
)


def load_base():
    spec = importlib.util.spec_from_file_location("saldev1499_base", BASE_BUILDER)
    module = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    spec.loader.exec_module(module)
    return module


b = load_base()


def add_hyperlink(paragraph, text: str, url: str):
    rel_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), b.BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "21")
    rpr.extend([fonts, color, underline, size])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([rpr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_link_block(doc, label: str, url: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label + ": ")
    b.set_run_font(r, size=10.7, bold=True, color=b.NAVY)
    add_hyperlink(p, url, url)
    return p


def add_bullets(doc, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        r = p.add_run(item)
        b.set_run_font(r, size=10.5, color=b.TEXT)


def make_diagram(path: Path):
    width, height = 1500, 770
    image = Image.new("RGB", (width, height), "#F6F8FB")
    draw = ImageDraw.Draw(image)
    arial = r"C:\Windows\Fonts\arial.ttf"
    arial_bold = r"C:\Windows\Fonts\arialbd.ttf"
    title_font = ImageFont.truetype(arial_bold, 34)
    subtitle_font = ImageFont.truetype(arial, 19)
    box_title = ImageFont.truetype(arial_bold, 22)
    body_font = ImageFont.truetype(arial, 17)
    label_font = ImageFont.truetype(arial_bold, 15)
    note_font = ImageFont.truetype(arial, 18)

    draw.text((65, 42), "SALDEV-1469 | controlled verification path", font=title_font, fill="#0B2545")
    draw.text((65, 88), "Confirm the package, inspect authorization, prove sign-in, and preserve the existing CRM connection.", font=subtitle_font, fill="#5E6B7E")
    boxes = [
        (55, 165, 330, 485, "Package evidence", ["Gong ECA v0.1", "Namespace gongeca", "Exact package ID", "Partial sandbox"], "#168A73", "#E5F3EF"),
        (405, 165, 680, 485, "Admin verification", ["David opens Setup", "Shares screen", "Confirms identity", "Captures evidence"], "#2E75B6", "#EAF2FB"),
        (755, 165, 1030, 485, "Authorization review", ["Use custom domain", "Stop before Allow", "Review app + scopes", "Cancel if unexpected"], "#B97816", "#FFF4D9"),
        (1105, 165, 1380, 485, "Completion proof", ["Gong home loads", "Existing import stays on", "No writeback enabled", "Attach evidence"], "#168A73", "#E5F3EF"),
    ]
    for x1, y1, x2, y2, title, lines, edge, fill in boxes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=20, fill=fill, outline=edge, width=5)
        draw.text((x1 + 20, y1 + 28), title, font=box_title, fill="#0B2545")
        y = y1 + 100
        for line in lines:
            draw.ellipse((x1 + 22, y + 6, x1 + 34, y + 18), fill=edge)
            draw.text((x1 + 48, y), line, font=body_font, fill="#263445")
            y += 46
    for start, end, label in [(330, 405, "VERIFY"), (680, 755, "REVIEW"), (1030, 1105, "PROVE")]:
        y = 330
        draw.line((start + 8, y, end - 14, y), fill="#2E75B6", width=6)
        draw.polygon([(end - 14, y - 11), (end + 3, y), (end - 14, y + 11)], fill="#2E75B6")
        text_width = draw.textlength(label, font=label_font)
        draw.text(((start + end - text_width) / 2, y - 34), label, font=label_font, fill="#0B2545")
    draw.rounded_rectangle((70, 560, 1430, 710), radius=20, fill="#FCEAEA", outline="#D13232", width=4)
    note = ("STOP CONDITIONS: wrong org or Production, unexpected app name or scopes, package identity mismatch, "
            "request to disconnect/uninstall the existing integration, or any session/token URL exposed in chat.")
    draw.text((105, 595), "Security gate", font=box_title, fill="#D13232")
    words, lines, current = note.split(), [], ""
    for word in words:
        candidate = (current + " " + word).strip()
        if draw.textlength(candidate, font=note_font) > 1230:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    y = 638
    for line in lines[:2]:
        draw.text((105, y), line, font=note_font, fill="#263445")
        y += 29
    image.save(path)


def build():
    if b.sha256(REFERENCE) != REFERENCE_SHA256:
        raise RuntimeError("Reference DOCX hash changed; fresh template distillation required.")
    TASK.mkdir(parents=True, exist_ok=True)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    b.clear_body(doc)
    b.configure_styles(doc)
    doc.core_properties.title = "SALDEV-1469 Gong ECA Verification Runbook"
    doc.core_properties.subject = "Gong External Client App installation verification and Salesforce sign-in test"
    doc.core_properties.author = "Flywire Salesforce Architecture Review"
    doc.core_properties.keywords = "SALDEV-1469, Gong ECA, Salesforce, OAuth, External Client App, FlywirePartial"

    # PAGE 1
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("IMPLEMENTATION VERIFICATION & TEAM RUNBOOK")
    b.set_run_font(r, size=11.2, bold=True, color=b.BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 0.95
    r = p.add_run("SALDEV-1469 | Gong ECA\nSalesforce Security Update")
    b.set_run_font(r, size=25.5, bold=True, color=b.NAVY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("Flywire Partial sandbox | Package verification, authorization review, and team handoff")
    b.set_run_font(r, size=13.2, color=b.GRAY)
    b.add_table(doc, ["Environment", "FlywirePartial | Org 00DhG0000000jOXUAY | Sandbox"], [
        ["Review date", "28 August 2026"],
        ["Deadline", "Complete required action before 31 August 2026"],
        ["Current decision", "Package installed; admin sign-in/authorization proof still pending"],
        ["Boundary", "No Production change; no uninstall, disconnect, reconnect, or session-specific OAuth link sharing"],
    ], [1.34, 5.16])
    b.add_callout(doc, "Team decision",
                  "Use the Zoom session with David to verify the installed Gong ECA package, start Sign in with Salesforce against the Partial sandbox, and stop at the Allow/Authorization page so the team can inspect the app identity and scopes before proceeding.",
                  "risk")
    b.add_heading(doc, "Current outcome", 1)
    b.add_table(doc, ["Scope", "Status", "Evidence / next action"], [
        ["Flywire Partial org connection", "Confirmed", "Sandbox org identity was established for this work."],
        ["Gong ECA package installation", "Confirmed", "Gong ECA v0.1 / gongeca / 0337y000005iObdAAE observed in Installed Packages."],
        ["Existing Gong CRM connection", "Confirmed", "Gong showed connected as ada+gong@flywire.com with current imports."],
        ["Salesforce sign-in test", "Pending", "David must run the controlled test while sharing his screen."],
        ["Authorization scope review", "Pending", "Stop before Allow and verify the displayed app name and scopes."],
        ["Production readiness", "Not evaluated", "This runbook covers Sandbox (Partial) only."],
    ], [2.05, 1.10, 3.35], status_col=1)

    # PAGE 2
    b.add_page_break(doc)
    b.add_heading(doc, "Story scope and work completed", 1)
    b.add_body(doc, "SALDEV-1469 was raised because Gong is replacing its legacy Salesforce connected app with an External Client App to support Salesforce security requirements, including PKCE, refresh-token rotation, refresh-token idle timeout, and refresh-token IP controls. Because Flywire uses Sign in with Salesforce for Gong, an administrator must install and validate the replacement before the deadline.")
    b.add_heading(doc, "Completed for this story", 2)
    b.add_table(doc, ["Step", "What was completed", "Evidence state"], [
        ["1", "Connected to and identified the Flywire Partial sandbox used for the dry run.", "Confirmed"],
        ["2", "Reviewed Gong Admin Center > Salesforce > Connection.", "Confirmed"],
        ["3", "Confirmed the existing integration user ada+gong@flywire.com and recent Account/Contact imports.", "Confirmed"],
        ["4", "Installed/verified Gong ECA in Salesforce Setup > Installed Packages.", "Confirmed"],
        ["5", "Prepared safe package-verification and Salesforce-login instructions for David.", "Confirmed"],
        ["6", "Defined the authorization stop point, evidence checklist, troubleshooting, and closure criteria.", "Ready"],
    ], [0.42, 4.60, 1.48], status_col=2)
    make_diagram(DIAGRAM)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    shape = p.add_run().add_picture(str(DIAGRAM), width=Inches(6.42))
    shape._inline.docPr.set("descr", "Four-stage SALDEV-1469 verification path from package evidence through authorization review to completion proof")
    shape._inline.docPr.set("title", "SALDEV-1469 controlled verification path")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Figure 1. Controlled package and sign-in verification path")
    b.set_run_font(r, size=10.1, italic=True, color=b.GRAY)
    b.add_callout(doc, "Important distinction",
                  "The Gong ECA supports Salesforce authentication/security. The existing Gong-to-Salesforce CRM import connection uses ada+gong@flywire.com. Do not disconnect, reconnect, uninstall, or change that integration while testing user sign-in.",
                  "info")

    # PAGE 3
    b.add_page_break(doc)
    b.add_heading(doc, "Installed package evidence and David verification", 1)
    add_link_block(doc, "Open FlywirePartial Installed Packages", INSTALLED_PACKAGES_URL)
    b.add_heading(doc, "Observed package identity", 2)
    b.add_table(doc, ["Field", "Expected / observed value", "Verification rule"], [
        ["Package name", "Gong ECA", "Must match exactly."],
        ["Version", "0.1", "Must match the installed-package detail page."],
        ["Namespace prefix", "gongeca", "Must match exactly; do not infer from the display name."],
        ["Package ID", "0337y000005iObdAAE", "Use the full 18-character ID for the record."],
        ["Package type", "Managed", "Observed on the installed-package page."],
        ["Packaging model", "2GP", "Observed on the installed-package page."],
        ["Installed by / time", "David Okolo | 27 Aug 2026, 3:08 PM", "Screenshot evidence; confirm if the live page differs."],
        ["Target org", "Sandbox (Partial)", "The sandbox banner/domain must be visible."],
    ], [1.45, 2.15, 2.90])
    b.add_heading(doc, "What David should do on Zoom", 2)
    add_bullets(doc, [
        "Open the Installed Packages link while signed in to the Flywire Partial sandbox.",
        "Open Gong ECA and share the package detail page on screen.",
        "Confirm the package name, version, namespace, package ID, and Sandbox (Partial) context.",
        "Capture a screenshot that includes the org banner and package values but excludes session tokens, browser storage, or sensitive URLs.",
        "If any value differs, stop and record the mismatch; do not uninstall or reinstall without a separate approved action plan.",
    ])
    b.add_callout(doc, "Package checkpoint",
                  "Pass only when all four identity fields match: Gong ECA, version 0.1, namespace gongeca, and package ID 0337y000005iObdAAE in Sandbox (Partial).",
                  "pass")

    # PAGE 4
    b.add_page_break(doc)
    b.add_heading(doc, "Controlled Gong Salesforce sign-in test", 1)
    add_link_block(doc, "Open Gong sign-in", GONG_SIGNIN_URL)
    b.add_table(doc, ["#", "Action", "Expected evidence"], [
        ["1", "Open the Gong sign-in link in a new tab while screen sharing.", "The URL is the stable Gong sign-in page, not an OAuth callback URL."],
        ["2", "Select Sign in with Salesforce.", "Salesforce authentication begins."],
        ["3", "If prompted, choose Use Custom Domain.", "A Salesforce custom-domain field is displayed."],
        ["4", "Enter flywire--partial.sandbox.", "The browser routes to the Flywire Partial sandbox domain."],
        ["5", "Confirm the org is Sandbox (Partial), not Production.", "Sandbox branding/domain is visible before continuing."],
        ["6", "Stop at the Allow/Authorization page.", "Do not click Allow yet; the team can inspect the app name and scopes."],
        ["7", "Review the displayed client/app identity and requested scopes.", "No unexpected scope or Production target is present."],
        ["8", "Proceed only after David and the reviewing admin/architect agree.", "Approval is explicit and captured in meeting notes."],
        ["9", "Confirm return to Gong home.", "Gong loads successfully and the existing CRM import connection remains unchanged."],
    ], [0.34, 3.16, 3.00])
    b.add_heading(doc, "Mandatory stop conditions", 2)
    b.add_table(doc, ["Stop condition", "Required response"], [
        ["Production or an unexpected Salesforce org is shown", "Cancel authentication; do not continue."],
        ["Unexpected app/client name or scopes", "Stop before Allow; capture a sanitized screenshot and escalate."],
        ["Package values do not match", "Do not reinstall or remove anything; document the mismatch."],
        ["Prompt requests CRM disconnect/reconnect", "Stop; the existing ada+gong@flywire.com integration is outside this sign-in action."],
        ["URL contains a long source= value or token-like material", "Do not paste or forward it; return to the stable Gong sign-in link."],
        ["Any error page appears", "Capture the error text, timestamp, browser, and step—without tokens or session URLs."],
    ], [2.64, 3.86])

    # PAGE 5
    b.add_page_break(doc)
    b.add_heading(doc, "Acceptance criteria and evidence record", 1)
    b.add_table(doc, ["#", "Acceptance criterion", "Evidence to attach", "Status"], [
        ["1", "Gong ECA is present in Sandbox (Partial).", "Installed Packages screenshot with org context.", "Confirmed"],
        ["2", "Version, namespace, and package ID match exactly.", "0.1 / gongeca / 0337y000005iObdAAE.", "Confirmed"],
        ["3", "Sign in with Salesforce reaches the Partial sandbox.", "Sanitized screenshot or meeting note.", "Pending"],
        ["4", "App identity and scopes are reviewed before Allow.", "Reviewer/approver name and authorization screenshot.", "Pending"],
        ["5", "Gong home loads after successful authentication.", "Post-login screenshot and timestamp.", "Pending"],
        ["6", "Existing CRM connection remains connected/importing.", "Gong Admin Center connection/import status after test.", "Pending"],
        ["7", "No Gong writeback is enabled by this story.", "Configuration confirmation / meeting note.", "Pending"],
        ["8", "No Production or unrelated Salesforce change occurs.", "Change record explicitly states sandbox-only scope.", "Required"],
    ], [0.32, 2.35, 2.75, 1.08], status_col=3)
    b.add_heading(doc, "Evidence log template", 2)
    b.add_table(doc, ["Evidence item", "Record this value"], [
        ["Meeting", "Date/time, Zoom participants, facilitator, Salesforce admin, reviewer"],
        ["Package", "Name, version, namespace, 18-character package ID, org banner/domain"],
        ["Authorization", "Displayed client/app name, requested scopes, reviewer decision, Allow/Cancel outcome"],
        ["Sign-in", "Success/failure, timestamp, resulting Gong page, sanitized error if any"],
        ["Regression check", "Existing integration user, connection status, last import indicators, writeback unchanged"],
        ["Attachments", "Sanitized screenshots and links to Jira/meeting notes; no tokens or session-specific URLs"],
    ], [1.70, 4.80])
    b.add_callout(doc, "Closure rule",
                  "Do not close SALDEV-1469 based on package installation alone. Close only after the controlled sign-in, scope review, successful Gong return, and post-test CRM-connection check are evidenced—or document an approved exception/blocker.",
                  "risk")
    b.add_page_break(doc)
    b.add_heading(doc, "Ownership", 2)
    b.add_table(doc, ["Owner", "Responsibility", "Evidence"], [
        ["David / Salesforce Admin", "Verify package, run sign-in, stop at authorization, and approve only after review.", "Screen-share confirmation and screenshots."],
        ["Facilitator", "Send stable links, guide the sequence, record results, and protect tokens/session URLs.", "Meeting notes and evidence checklist."],
        ["Architect / Security reviewer", "Confirm app identity, scopes, sandbox target, and whether Allow is acceptable.", "Named approval or documented concern."],
        ["Gong / Integration owner", "Confirm final Gong access and unchanged CRM import connection; coordinate support.", "Post-test connection/import status."],
    ], [1.35, 3.40, 1.75])
    b.add_heading(doc, "Approval record", 2)
    b.add_table(doc, ["Checkpoint", "Decision owner", "Record in Jira / meeting notes"], [
        ["Package identity", "David / Salesforce Admin", "Pass/fail, timestamp, screenshot, and any mismatch."],
        ["Authorization app and scopes", "Architect / Security reviewer", "Allow, cancel, or escalate; include the reason."],
        ["Successful Gong return", "David / Gong owner", "Post-login result and sanitized evidence."],
        ["Existing connection unchanged", "Gong / Integration owner", "Connection user and import-status confirmation."],
        ["Story closure", "Story owner / Architect", "All evidence attached or approved exception documented."],
    ], [1.70, 1.75, 3.05])
    b.add_callout(doc, "Decision authority",
                  "David can perform the sandbox admin steps, but clicking Allow should follow the team’s review of the displayed app identity and scopes. Production authorization is not part of this runbook.",
                  "info")

    # PAGE 6
    b.add_page_break(doc)
    b.add_heading(doc, "Troubleshooting and escalation", 1)
    b.add_table(doc, ["Symptom", "Safe diagnostic", "Next action"], [
        ["Gong ECA is not listed", "Confirm Sandbox (Partial), admin access, and package filters.", "Stop and notify the Salesforce admin; do not reinstall without approval."],
        ["Package version or ID differs", "Read all identity fields from the detail page.", "Capture evidence and compare with the vendor-provided installation instruction."],
        ["Wrong Salesforce login domain", "Cancel and choose Use Custom Domain.", "Enter flywire--partial.sandbox and reconfirm the sandbox banner."],
        ["No Allow page appears", "Check whether the user/app was already authorized and which org session is active.", "Verify current connected-app authorization before treating the test as passed."],
        ["OAuth or redirect error", "Record the sanitized message, timestamp, browser, and exact step.", "Retry only from the stable Gong sign-in URL; escalate persistent errors to Gong."],
        ["Gong redirects to setup", "Reconfirm package identity and Salesforce authentication target.", "Attach evidence and open a Gong support case if the package is correct."],
        ["CRM import status changes", "Review Gong Admin Center > Salesforce > Connection and import indicators.", "Stop; do not reconnect. Escalate to the Gong/Salesforce integration owner."],
    ], [1.60, 2.55, 2.35])
    b.add_heading(doc, "Security and change controls", 2)
    add_bullets(doc, [
        "Never paste Salesforce frontdoor URLs, sid values, access tokens, refresh tokens, OAuth callback URLs, or long source= links into Zoom/Jira/chat.",
        "Use only the stable Installed Packages and Gong sign-in links printed in this runbook.",
        "Do not uninstall Gong ECA, disconnect/reconnect ada+gong@flywire.com, change integration scopes, or enable writeback as part of this verification.",
        "Do not repeat the action in Production until the sandbox test is complete and a separately approved Production plan exists.",
    ])
    b.add_callout(doc, "Escalation package",
                  "Send Gong Support the sanitized error, timestamp/time zone, Gong workspace URL, Salesforce sandbox domain, package name/version/namespace/ID, and the failed step. Exclude all tokens and session-specific URLs.",
                  "info")
    # PAGE 7
    b.add_page_break(doc)
    b.add_heading(doc, "Team handoff and copy-paste communications", 1)
    b.add_heading(doc, "Links to send David in Zoom", 2)
    add_link_block(doc, "1. Verify the installed package", INSTALLED_PACKAGES_URL)
    add_link_block(doc, "2. Test Gong Salesforce login", GONG_SIGNIN_URL)
    b.add_heading(doc, "Exact Zoom chat message", 2)
    b.add_callout(doc, "Copy and paste", ZOOM_MESSAGE, "info", body_size=9.7)
    b.add_heading(doc, "Meeting talk track", 2)
    add_bullets(doc, [
        "We are validating SALDEV-1469 only in Sandbox (Partial).",
        "First, confirm the exact Gong ECA package identity in Salesforce Setup.",
        "Next, start Sign in with Salesforce and use flywire--partial.sandbox if a custom domain is requested.",
        "Please stop before clicking Allow so we can verify the External Client App and requested scopes together.",
        "After login, we will confirm Gong opens and that the existing ada+gong@flywire.com CRM import connection remains unchanged.",
    ])
    b.add_heading(doc, "Ready-to-paste Jira update", 2)
    jira = (
        "SALDEV-1469 update: Gong ECA is installed in Flywire Sandbox (Partial) and the package identity was observed as "
        "version 0.1, namespace gongeca, package ID 0337y000005iObdAAE. The existing Gong Salesforce CRM connection "
        "remains associated with ada+gong@flywire.com and showed active imports. Next, David will verify the package on screen "
        "and run Sign in with Salesforce using custom domain flywire--partial.sandbox. We will stop at the Allow/Authorization "
        "page to review the External Client App identity and scopes before proceeding, then confirm Gong home access and recheck "
        "the existing CRM connection. No Production change, package uninstall/reinstall, CRM disconnect/reconnect, writeback change, "
        "or session-specific OAuth URL sharing is included in this story verification."
    )
    b.add_callout(doc, "Story comment", jira, "info", body_size=9.1)

    doc.save(OUTPUT)
    if b.sha256(REFERENCE) != REFERENCE_SHA256:
        raise RuntimeError("Reference DOCX was modified unexpectedly.")
    print(OUTPUT)


if __name__ == "__main__":
    build()
