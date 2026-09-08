from docx import Document
from pathlib import Path
import sys

path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(r"C:\Users\LIKKI\Documents\ChatGPT\David 2 Probo medical\scratch\case-00010716-uat-team-guide\reference.docx")
doc = Document(path)

print("PARAGRAPHS")
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip().replace("\n", " | ")
    if text:
        print(f"{i:03d}\t{p.style.name}\t{text}")

print("TABLES")
for ti, table in enumerate(doc.tables):
    print(f"TABLE {ti}: {len(table.rows)}x{len(table.columns)} style={table.style.name if table.style else ''}")
    for ri, row in enumerate(table.rows):
        vals = [cell.text.strip().replace("\n", " | ") for cell in row.cells]
        print(f"  {ri:02d}\t" + " || ".join(vals))

print("STYLES")
for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
    s = doc.styles[name]
    f = s.font
    pf = s.paragraph_format
    print(name, f.name, f.size.pt if f.size else None, f.bold, f.italic, f.color.rgb if f.color and f.color.rgb else None,
          pf.space_before.pt if pf.space_before else None, pf.space_after.pt if pf.space_after else None,
          pf.line_spacing)

print("HEADERS_FOOTERS")
for si, section in enumerate(doc.sections):
    print("SECTION", si)
    for label, part in [("header", section.header), ("footer", section.footer),
                        ("first_header", section.first_page_header), ("first_footer", section.first_page_footer)]:
        vals = [p.text.strip() for p in part.paragraphs if p.text.strip()]
        print(label, vals)
