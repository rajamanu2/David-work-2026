from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "outputs" / "next_four_architect_reviews_2026-09-01"
OUT_DIR.mkdir(parents=True, exist_ok=True)

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
MUTED = "667085"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
GREEN = "1F5E46"
AMBER = "7A5A00"
RED = "9B1C1C"
WHITE = "FFFFFF"
BLACK = "000000"


STORIES = [
    {
        "case": "00009998",
        "title": "RMA edits remove key asset attributes",
        "owner": "Blake Nielsen",
        "case_status": "Claimed",
        "priority": "Medium",
        "classification": "Bug - RMA",
        "decision": "APPROVE FOR DEVDO UAT",
        "decision_detail": (
            "The data-loss defect is addressed in DevDO. Production release remains conditional on business UAT, "
            "target-org validation, and confirmation that the retained reset behavior is still correct for the other check-in fields."
        ),
        "problem": (
            "Editing or receiving an RMA could remove Product Item values that are operationally important for asset qualification. "
            "The reported impact covers Quality Issue Details, Root Cause, and Corrective Action updates that caused key asset attributes to disappear."
        ),
        "repro_steps": [
            "Open an RMA linked to a Product Item that already has Cosmetic Rating, Functional Rating, and Review Status populated.",
            "Record the three existing Product Item values so the before-and-after result can be compared.",
            "Edit the RMA and update Quality Issue Details, Root Cause, or Corrective Action.",
            "Save the RMA and allow the RMA trigger processing to complete.",
            "Open or refresh the related Product Item and inspect Cosmetic Rating, Functional Rating, and Review Status."
        ],
        "expected_result": "The RMA saves and the three pre-existing Product Item qualification attributes remain unchanged.",
        "actual_result": "Before the fix, the trigger cleared Cosmetic Rating, Functional Rating, and Review Status to blank values.",
        "root_cause": (
            "The legacy P_RMATriggerHandler explicitly assigned blank strings to Cosmetic_Rating__c, "
            "Functional_Rating__c, and Review_Status__c while processing the Product Item reset block. "
            "Those destructive assignments were unconditional within that branch, so a valid RMA edit could erase pre-existing asset qualification data."
        ),
        "solution": [
            "Removed the hard-coded blank assignments for Cosmetic_Rating__c, Functional_Rating__c, and Review_Status__c.",
            "Preserved the existing reset behavior for repair-evaluation, repairability, and note fields so the change remains narrowly scoped.",
            "Retained the existing location update behavior, including the Equipment Storage exception.",
            "Added regression assertions proving that the three protected attributes survive the RMA receipt/edit path."
        ],
        "hardcoding": (
            "Resolved for the reported data-loss path. The hard-coded empty-string resets for the three protected asset attributes were removed. "
            "The remaining field resets are still explicit legacy logic and should be confirmed during UAT as intentional business behavior."
        ),
        "components": [
            ("Apex Class", "P_RMATriggerHandler", "Removes destructive attribute resets"),
            ("Apex Test", "P_RMATriggerHandlerTest", "Adds preservation assertions"),
            ("Apex Test", "P_RMATriggerHandlerTest2", "Included with shared RMA regression scope"),
        ],
        "test_method": "P_RMATriggerHandlerTest.testDateItemReceived_TriggersProductItemResetBlock",
        "test_result": "Passed in deployment 0AfiK0000000X8rSAE; shared run completed 13/13 tests with zero failures.",
        "uat": [
            "Open an RMA whose Product Item already has Cosmetic Rating, Functional Rating, and Review Status populated.",
            "Edit Quality Issue Details, Root Cause, and Corrective Action, then save.",
            "Verify the RMA saves and all three Product Item attributes retain their original values.",
            "Verify the intended repair-evaluation and note fields still reset only when the existing business conditions are met.",
            "Repeat for Equipment Storage and a non-Equipment Storage record to confirm location behavior is unchanged."
        ],
        "risks": [
            "P_RMATriggerHandler is a large shared legacy handler; regression testing must cover other RMA types and receipt paths.",
            "The story preserves three fields but leaves other explicit resets in place; business ownership should confirm those fields are intentionally cleared.",
            "Production data should be sampled before release to quantify affected records and define any separate restoration activity."
        ],
        "change_set": (
            "Draft change set 'Story 00009998 RMA Asset Attributes' exists in DevDO with the handler and both RMA test classes. "
            "It is Open and has not been uploaded."
        ),
    },
    {
        "case": "00010158",
        "title": "P_RMATTrigger error prevents saving an RMA",
        "owner": "Lisa Litviak",
        "case_status": "Claimed",
        "priority": "Medium",
        "classification": "Bug - Salesforce error / RMA",
        "decision": "APPROVE FOR DEVDO UAT",
        "decision_detail": (
            "The transaction-blocking null path is guarded in DevDO. Production release remains conditional on UAT across RMA types "
            "and confirmation that Opportunity Product synchronization still occurs when an Opportunity is present."
        ),
        "problem": (
            "Users could not save an RMA, including reported RMA A-206654, because P_RMATTrigger raised an error when a Product Item "
            "was selected before an Opportunity was assigned."
        ),
        "repro_steps": [
            "Open RMA A-206654, or a comparable RMA with a valid Product Item.",
            "Confirm that the RMA has a Product Item selected and that the Opportunity field is blank.",
            "Update an editable field such as Quality Issue Notes.",
            "Select Save.",
            "Observe the save result and review the displayed P_RMATTrigger error."
        ],
        "expected_result": "The RMA saves successfully even when an Opportunity has not yet been assigned.",
        "actual_result": "Before the fix, P_RMATTrigger entered Opportunity Product synchronization with no Opportunity and blocked the save with an error.",
        "root_cause": (
            "The trigger handler dereferenced RepairOpp.Id and attempted Opportunity Product synchronization without first confirming "
            "that RepairOpp existed. A valid RMA record with a Product Item but no Opportunity therefore entered a null-reference path and blocked the transaction."
        ),
        "solution": [
            "Added explicit null guards requiring both the Product Item and Repair Opportunity before Opportunity Product synchronization runs.",
            "Allowed the RMA and Product Item relationship to save when no Opportunity has yet been assigned.",
            "Kept the existing synchronization path unchanged when an Opportunity is present.",
            "Added a dedicated regression test that saves an RMA without an Opportunity and confirms the update persists."
        ],
        "hardcoding": (
            "No new story-specific hard-coded business value was introduced. The fix is relationship-driven and uses null checks. "
            "A legacy default of USD remains elsewhere inside the Opportunity Product synchronization branch; it is pre-existing technical debt and outside this defect's scope."
        ),
        "components": [
            ("Apex Class", "P_RMATriggerHandler", "Adds null-safe synchronization gates"),
            ("Apex Test", "P_RMATriggerHandlerTest2", "Adds the no-Opportunity save regression"),
            ("Apex Test", "P_RMATriggerHandlerTest", "Included with shared RMA regression scope"),
        ],
        "test_method": "P_RMATriggerHandlerTest2.case00010158AllowsProductItemWithoutOpportunity",
        "test_result": "Passed in deployment 0AfiK0000000X8rSAE; shared run completed 13/13 tests with zero failures.",
        "uat": [
            "Create or open an RMA with a valid Product Item and no Opportunity.",
            "Update Quality Issue Notes and save; verify no P_RMATTrigger error occurs.",
            "Reopen the RMA and verify the Product Item and edited notes were retained.",
            "Assign an Opportunity and repeat the save; verify Opportunity Product synchronization still works.",
            "Exercise at least one Customer Repair Eval and one additional RMA type used by operations."
        ],
        "risks": [
            "The shared handler contains multiple RMA branches, so null safety must not mask a genuinely required Opportunity in another process.",
            "The existing USD default in the synchronization branch should be reviewed separately for multi-currency behavior.",
            "UAT should use representative Product Items and Pricebook Entries to prove the guarded path and normal synchronization path."
        ],
        "change_set": (
            "Draft change set 'Story 00010158 RMA Save Trigger Fix' exists in DevDO with the handler and both RMA test classes. "
            "It is Open and has not been uploaded."
        ),
    },
    {
        "case": "00009589",
        "title": "Service Appointment removed from Gantt when SRI is completed",
        "owner": "Trudy Gregory",
        "case_status": "New",
        "priority": "Medium",
        "classification": "Bug - Service Appointment / Work Order",
        "decision": "APPROVE FOR DEVDO UAT; PACKAGING ACTION REQUIRED",
        "decision_detail": (
            "The status-regression defect is addressed and tested in DevDO. Before the outbound change set is ready, the Service Appointment Status Rule "
            "custom metadata type and Pending Authorization To Internal record must be added to the story-specific draft."
        ),
        "problem": (
            "Completing an SRI could move an already scheduled or dispatched Service Appointment back to Internal, removing it from the Gantt and creating a risk of a missed visit."
        ),
        "repro_steps": [
            "Create or identify a Service Appointment that is Scheduled or Dispatched and visible on the Dispatcher Console Gantt.",
            "Assign the Service Appointment to an engineer and confirm it remains visible on the intended date.",
            "Create or use a related SRI for a part or service item. The reported example used SA-17132 and SRI-00015817.",
            "Complete the related SRI.",
            "Refresh the Service Appointment and the Dispatcher Console Gantt, then inspect the appointment status and visibility."
        ],
        "expected_result": "Completing the SRI does not change an active Scheduled or Dispatched appointment, and the appointment remains on the Gantt.",
        "actual_result": "Before the fix, the appointment could be moved back to Internal and disappear from the Gantt.",
        "root_cause": (
            "The asynchronous SRI Flow relied on a related-record status reference and hard-coded both the trigger status "
            "'Pending Customer Authorization' and target status 'Internal'. The decision could operate on stale or inappropriate status context and regress an active appointment."
        ),
        "solution": [
            "Retrieves the current Service Appointment record before evaluating any status change.",
            "Looks up an active Service Appointment Status Rule custom metadata record whose Trigger Status matches the current appointment status.",
            "Uses the metadata Target Status rather than a hard-coded Flow assignment.",
            "Performs no status update when no active rule matches, preserving scheduled or dispatched appointments on the Gantt.",
            "Provides a configured Pending Customer Authorization to Internal rule for the intended legacy transition."
        ],
        "hardcoding": (
            "Resolved in Flow logic. The trigger and target status strings were moved from Flow elements into a configurable custom metadata rule. "
            "Future allowed transitions can be managed as metadata records without editing the Flow."
        ),
        "components": [
            ("Flow", "On_Create_of_SRI_Check_if_SA_status_needs_to_be_updated", "Current-state lookup and metadata-driven transition"),
            ("Apex Test", "ProboNextFourFlowRegressionTest", "Protects dispatched appointments from regression"),
            ("Custom Metadata Type", "Service_Appointment_Status_Rule__mdt", "Configurable status-transition schema"),
            ("Custom Fields", "Active__c; Trigger_Status__c; Target_Status__c", "Rule enablement and transition values"),
            ("Custom Metadata Record", "Pending_Authorization_To_Internal", "Pending Customer Authorization to Internal"),
        ],
        "test_method": "ProboNextFourFlowRegressionTest.case00009589DispatchedAppointmentDoesNotReturnToInternal",
        "test_result": "Passed in deployment 0AfiK0000000X8rSAE; shared run completed 13/13 tests with zero failures.",
        "uat": [
            "Use a Service Appointment that is visible on the Gantt and has status Dispatched; complete the related SRI.",
            "Verify the Service Appointment remains Dispatched and stays on the Gantt.",
            "Repeat for Scheduled or another operational status that has no active transition rule; verify no regression occurs.",
            "Use a Service Appointment in Pending Customer Authorization; complete the related SRI and verify it transitions to Internal.",
            "Deactivate the metadata rule in a test context and verify the Flow performs no transition."
        ],
        "risks": [
            "The Flow depends on the custom metadata type, its three fields, and the active rule record; omitting any of them from the change set will break the intended configuration.",
            "Exact status spelling must match Salesforce status values; governance is required for new rules.",
            "The browser session timed out while packaging the metadata, so story-level outbound change-set completeness is not yet verified."
        ],
        "change_set": (
            "Draft change set 'Story 00009589 Service Appointment Gantt' is Open and not uploaded. The Flow and regression test were confirmed. "
            "The Service Appointment Status Rule type and Pending_Authorization_To_Internal record still require addition and final UI verification."
        ),
    },
    {
        "case": "00010514",
        "title": "Core Exchange RMA not automatically opening from Work Order",
        "owner": "Ashley Madson",
        "case_status": "Claimed",
        "priority": "Medium",
        "classification": "Salesforce Improvement - Work Order / RMA",
        "decision": "APPROVE FOR DEVDO UAT",
        "decision_detail": (
            "The GE Core Exchange automation is implemented and tested in DevDO. Production release remains conditional on contract-owner UAT, "
            "confirmation of the account-matching key, and standard deployment validation."
        ),
        "problem": (
            "Work Orders for the GE contract required users to create Core Exchange RMAs manually, adding operational effort and increasing the risk that a required return process would be missed."
        ),
        "repro_steps": [
            "Create or open a Work Order for the GE Precision Healthcare account under the applicable GE contract process.",
            "Populate the Work Order Account, Opportunity, and Service Contract relationships.",
            "Save the Work Order so it first meets all qualifying conditions.",
            "Open the Work Order's related records and inspect the RMA relationship or related list.",
            "Verify whether a Core Exchange RMA was created automatically."
        ],
        "expected_result": "Exactly one Core Exchange RMA is created automatically and linked to the qualifying Work Order, Opportunity, and Service Contract.",
        "actual_result": "Before the fix, no Core Exchange RMA was created automatically and operations had to create it manually.",
        "root_cause": (
            "No active, scoped automation existed to create a Core Exchange RMA when a qualifying Work Order gained the required Account, Opportunity, and Service Contract relationships."
        ),
        "solution": [
            "Added an active after-save Work Order Flow that runs only when Account, Opportunity, and Service Contract are all present and the record first meets the criteria.",
            "Reads an active Core Exchange Automation Rule custom metadata record instead of embedding GE-specific values directly in Flow logic.",
            "Resolves the configured active RMA record type by DeveloperName.",
            "Checks for an existing Core Exchange RMA by Work Order and Record Type before creating a record, preventing duplicates.",
            "Does not backfill historical Work Orders; the automation is forward-looking."
        ],
        "hardcoding": (
            "GE-specific account and RMA record-type values were externalized into custom metadata. The Flow itself is configuration-driven. "
            "Current matching uses Account.Name, which is editable and therefore less stable than an Account ID or external key; confirm this design before Production."
        ),
        "components": [
            ("Flow", "Create_Configured_Core_Exchange_RMA_From_Work_Order", "Creates one configured RMA after criteria are met"),
            ("Apex Test", "ProboNextFourFlowRegressionTest", "Proves one RMA and duplicate prevention"),
            ("Custom Metadata Type", "Core_Exchange_Automation_Rule__mdt", "Configurable account-to-record-type schema"),
            ("Custom Fields", "Account_Name__c; Active__c; RMA_Record_Type_Developer_Name__c", "Configuration values"),
            ("Custom Metadata Record", "GE_Precision_Healthcare", "GE PRECISION HEALTHCARE LLC / Core_Exchange"),
        ],
        "test_method": "ProboNextFourFlowRegressionTest.case00010514CreatesOneConfiguredCoreExchangeRma",
        "test_result": "Passed in deployment 0AfiK0000000X8rSAE; shared run completed 13/13 tests with zero failures.",
        "uat": [
            "Create a GE qualifying Work Order with Account, Opportunity, and Service Contract populated.",
            "Verify exactly one RMA is created, linked to the Work Order, Opportunity, and Service Contract, with the Core Exchange record type.",
            "Edit the Work Order again and verify a second RMA is not created.",
            "Use a nonconfigured Account and verify no Core Exchange RMA is created.",
            "Remove or deactivate the metadata rule in a test context and verify no RMA is created.",
            "Confirm historical Work Orders are not backfilled unless a separate approved migration is executed."
        ],
        "risks": [
            "Account.Name matching is sensitive to renames, spacing, and duplicate account names; a stable identifier is preferable for long-term governance.",
            "The automation intentionally requires all three Work Order relationships; users must understand when the record becomes eligible.",
            "The flow does not backfill existing Work Orders, so operations must decide whether any separate data remediation is needed."
        ],
        "change_set": (
            "Draft change set 'Story 00010514 Core Exchange RMA' exists in DevDO with the Flow, regression test, Core Exchange Automation Rule type, "
            "and GE_Precision_Healthcare record. It is Open and has not been uploaded."
        ),
    },
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_font(run, name="Calibri", size=None, color=BLACK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, name, size, color=BLACK, bold=None):
    style.font.name = name
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold


def set_spacing(style, before, after, line):
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_numbering_definition(doc, abstract_id=42, num_id=42, fmt="bullet", text="•"):
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_id = OxmlElement("w:abstractNumId")
    abs_id.set(qn("w:val"), str(abstract_id))
    num.append(abs_id)
    numbering.append(num)


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)


def add_page_field(paragraph):
    paragraph.add_run("Page ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


def configure_document(doc, story):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, "Calibri", 11, BLACK)
    set_spacing(normal, 0, 6, 1.10)
    normal.paragraph_format.widow_control = True

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        set_style_font(style, "Calibri", size, color, True)
        set_spacing(style, before, after, 1.0)
        style.paragraph_format.keep_with_next = True

    if "Table Citation" not in [s.name for s in styles]:
        citation = styles.add_style("Table Citation", WD_STYLE_TYPE.PARAGRAPH)
    else:
        citation = styles["Table Citation"]
    set_style_font(citation, "Calibri", 9, MUTED, False)
    set_spacing(citation, 4, 4, 1.0)

    if "Callout" not in [s.name for s in styles]:
        callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
    else:
        callout = styles["Callout"]
    set_style_font(callout, "Calibri", 10.5, NAVY, False)
    set_spacing(callout, 4, 8, 1.10)

    add_numbering_definition(doc, 42, 42, "bullet", "•")
    add_numbering_definition(doc, 43, 43, "decimal", "%1.")
    add_numbering_definition(doc, 44, 44, "decimal", "%1.")

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run(f"PROBO MEDICAL  |  SALESFORCE ARCHITECT REVIEW  |  CASE {story['case']}")
    set_font(hr, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    set_font(fp.add_run("Internal - DevDO review  |  "), size=8.5, color=MUTED)
    add_page_field(fp)
    for run in fp.runs:
        set_font(run, size=8.5, color=MUTED)


def add_title_block(doc, story):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("ARCHITECT REVIEW")
    set_font(r, size=10, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"Case {story['case']}")
    set_font(r, size=24, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(story["title"])
    set_font(r, size=14, color=MUTED, bold=False)

    table = doc.add_table(rows=6, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Review field"
    table.rows[0].cells[1].text = "Value"
    for cell in table.rows[0].cells:
        set_cell_shading(cell, LIGHT_GRAY)
        for run in cell.paragraphs[0].runs:
            set_font(run, size=9, color=NAVY, bold=True)
    set_repeat_table_header(table.rows[0])
    metadata = [
        ("Environment", "Probo Medical DevDO sandbox"),
        ("Case", f"{story['case']} | {story['case_status']} | {story['priority']}"),
        ("Owner", story["owner"]),
        ("Classification", story["classification"]),
        ("Evidence date", "September 1, 2026"),
    ]
    for row, (label, value) in zip(table.rows[1:], metadata):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], LIGHT_GRAY)
        for run in row.cells[0].paragraphs[0].runs:
            set_font(run, size=9.5, color=NAVY, bold=True)
        for run in row.cells[1].paragraphs[0].runs:
            set_font(run, size=9.5, color=BLACK)
    set_table_geometry(table, [2700, 6660])


def add_callout(doc, label, text, color=GREEN):
    p = doc.add_paragraph(style="Callout")
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), CALLOUT)
    p_pr.append(shd)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "160")
    ind.set(qn("w:right"), "160")
    p_pr.append(ind)
    r = p.add_run(f"{label}: ")
    set_font(r, size=10.5, color=color, bold=True)
    r = p.add_run(text)
    set_font(r, size=10.5, color=NAVY)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph()
        apply_num(p, 42)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.167
        p.add_run(item)


def add_numbered(doc, items, num_id=43):
    for item in items:
        p = doc.add_paragraph()
        apply_num(p, num_id)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.167
        p.add_run(item)


def add_component_table(doc, components):
    p = doc.add_paragraph("Story-specific deployment package", style="Table Citation")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["Type", "Component", "Purpose"]
    for idx, value in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = value
        set_cell_shading(cell, LIGHT_GRAY)
        for run in cell.paragraphs[0].runs:
            set_font(run, size=9, color=NAVY, bold=True)
    set_repeat_table_header(table.rows[0])
    for comp_type, name, purpose in components:
        cells = table.add_row().cells
        values = [comp_type, name, purpose]
        for idx, value in enumerate(values):
            cells[idx].text = value
            for run in cells[idx].paragraphs[0].runs:
                set_font(run, size=8.8, color=BLACK)
    set_table_geometry(table, [1900, 3500, 3960])


def add_test_table(doc, story):
    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Evidence"
    table.rows[0].cells[1].text = "Verified result"
    for cell in table.rows[0].cells:
        set_cell_shading(cell, LIGHT_GRAY)
        for run in cell.paragraphs[0].runs:
            set_font(run, size=9, color=NAVY, bold=True)
    set_repeat_table_header(table.rows[0])
    rows = [
        ("Deployment", "0AfiK0000000X8rSAE - Succeeded; 16/16 components; zero component errors"),
        ("Tests", "13/13 completed; zero failures; zero test errors"),
        ("Story regression", story["test_method"] + " - Passed"),
    ]
    for row, (label, value) in zip(table.rows[1:], rows):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], LIGHT_GRAY)
        for run in row.cells[0].paragraphs[0].runs:
            set_font(run, size=9, color=NAVY, bold=True)
        for run in row.cells[1].paragraphs[0].runs:
            set_font(run, size=9, color=BLACK)
    set_table_geometry(table, [2100, 7260])


def build_story(story):
    doc = Document()
    doc.core_properties.title = f"Case {story['case']} Architect Review"
    doc.core_properties.subject = story["title"]
    doc.core_properties.author = "Probo Medical Salesforce Architecture"
    doc.core_properties.keywords = "Salesforce, DevDO, architect review, outbound change set"
    configure_document(doc, story)
    add_title_block(doc, story)

    doc.add_heading("Executive decision", level=1)
    add_callout(doc, story["decision"], story["decision_detail"], GREEN if "PACKAGING" not in story["decision"] else AMBER)

    doc.add_heading("Story and operational impact", level=1)
    doc.add_paragraph(story["problem"])

    doc.add_heading("Steps to reproduce", level=1)
    add_numbered(doc, story["repro_steps"], num_id=44)
    add_callout(doc, "Expected result", story["expected_result"], GREEN)
    add_callout(doc, "Pre-fix actual result", story["actual_result"], RED)

    doc.add_heading("Root cause", level=1)
    doc.add_paragraph(story["root_cause"])

    doc.add_heading("Implemented design", level=1)
    add_bullets(doc, story["solution"])

    doc.add_heading("Hard-coded value assessment", level=1)
    add_callout(doc, "Assessment", story["hardcoding"], DARK_BLUE)

    doc.add_heading("Metadata scope", level=1)
    add_component_table(doc, story["components"])

    doc.add_heading("Verified DevDO evidence", level=1)
    add_test_table(doc, story)
    doc.add_paragraph(story["test_result"])

    doc.add_heading("Outbound change-set status", level=1)
    doc.add_paragraph(story["change_set"])

    doc.add_heading("UAT acceptance checklist", level=1)
    add_numbered(doc, story["uat"])

    doc.add_heading("Risks and production gates", level=1)
    add_bullets(doc, story["risks"])

    doc.add_heading("Release recommendation", level=1)
    if story["case"] == "00009589":
        recommendation = (
            "Proceed with DevDO UAT after the story-specific outbound change set is completed with the metadata type and record. "
            "Do not upload or deploy to Production until UAT evidence, component completeness, and target-org validation are recorded."
        )
    else:
        recommendation = (
            "Proceed with DevDO UAT. Do not upload or deploy to Production until business acceptance, target-org validation, "
            "story-specific component review, and an approved rollback plan are recorded."
        )
    add_callout(doc, "Recommendation", recommendation, GREEN)

    doc.add_heading("Evidence basis", level=2)
    evidence = [
        "Verified DevDO deployment report: 0AfiK0000000X8rSAE (Succeeded, 16/16 components, 13/13 tests).",
        "Production and DevDO scoped source snapshots under scratch/next-four-analysis.",
        "Production case snapshot under outputs/probo_case_story_dry_run_2026-08-31.",
        "Outbound Change Set browser review performed before document preparation; no change set was uploaded."
    ]
    add_bullets(doc, evidence)

    path = OUT_DIR / f"Case_{story['case']}_Architect_Review.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    outputs = [build_story(story) for story in STORIES]
    for output in outputs:
        print(output)
