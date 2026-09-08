from pathlib import Path
from shutil import copy2

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_architect_review import (
    AMBER,
    BLUE,
    INK,
    LIGHT_AMBER,
    LIGHT_BLUE,
    LIGHT_GRAY,
    LIGHT_TEAL,
    MUTED,
    NAVY,
    TEAL,
    add_body,
    add_callout,
    add_figure,
    add_heading,
    add_page_field,
    add_table,
    arrow,
    draw_wrapped,
    load_font,
    rgb,
    rounded,
    set_cell_text,
    set_run_font,
    set_table_geometry,
    shade_cell,
)


ROOT = Path(__file__).resolve().parent
REFERENCE = ROOT / "SCC-3386-architect-review.docx"
ASSET_DIR = ROOT / "doc-assets-scc3655"
OUTPUT = ROOT / "SCC-3655-architect-review.docx"
ASSET_DIR.mkdir(exist_ok=True)


def clear_document_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_document(document: Document) -> None:
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        for header in (section.header, section.even_page_header, section.first_page_header):
            paragraph = header.paragraphs[0]
            paragraph.text = ""
            paragraph.paragraph_format.space_after = Pt(0)

        for footer in (section.footer, section.even_page_footer, section.first_page_footer):
            paragraph = footer.paragraphs[0]
            paragraph.text = ""
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            paragraph.paragraph_format.space_before = Pt(0)
            run = paragraph.add_run("GreatPlainsMerge + UAT  |  Page ")
            set_run_font(run, size=9, color=MUTED)
            add_page_field(paragraph)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 12, 6),
        ("Heading 2", 13, BLUE, 10, 5),
        ("Heading 3", 12, NAVY, 8, 4),
    ):
        style = document.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def create_case_comment_diagram(path: Path) -> None:
    image = Image.new("RGB", (1600, 1000), "#F7F9FC")
    draw = ImageDraw.Draw(image)
    title = load_font(42, bold=True)
    subtitle = load_font(22)
    heading = load_font(26, bold=True)
    body = load_font(20)
    small = load_font(18)

    draw.text((70, 52), "SCC-3655 | Internal Case Comment path", font=title, fill="#0B2545")
    draw.text(
        (70, 112),
        "The standard platform path is configured; a named-rep UAT write test is the remaining evidence gate.",
        font=subtitle,
        fill="#5B6777",
    )

    boxes = [
        (
            (70, 220, 440, 560),
            "Rep on a Case",
            [
                "Open the Related tab",
                "Use the Case Comments list",
                "Select New and enter Body",
                "Named UAT rep is required",
            ],
            "#E7F5F1",
            "#16836B",
        ),
        (
            (610, 220, 990, 560),
            "Standard CaseComment",
            [
                "ParentId links the Case",
                "CommentBody stores the update",
                "IsPublished controls visibility",
                "CreatedBy and CreatedDate are system-set",
            ],
            "#E7F5F1",
            "#16836B",
        ),
        (
            (1160, 220, 1530, 560),
            "Visibility and history",
            [
                "Public unchecked is internal",
                "Public checked is external",
                "Related list holds lifecycle history",
                "Runtime ordering is not yet proven",
            ],
            "#FFF6DE",
            "#B7791F",
        ),
    ]

    for box, box_title, lines, fill, outline in boxes:
        rounded(draw, box, fill, outline=outline, radius=24, width=4)
        draw.text((box[0] + 28, box[1] + 26), box_title, font=heading, fill="#0B2545")
        y = box[1] + 92
        for line in lines:
            draw.ellipse((box[0] + 30, y + 7, box[0] + 42, y + 19), fill=outline)
            draw_wrapped(draw, (box[0] + 58, y), line, body, "#1F2937", box[2] - box[0] - 92, spacing=6)
            y += 58

    arrow(draw, (440, 390), (610, 390), color="#16836B", width=6)
    arrow(draw, (990, 390), (1160, 390), color="#B7791F", width=7)
    draw.text((472, 345), "Create comment", font=small, fill="#5B6777")
    draw.text((1010, 336), "UAT EVIDENCE", font=heading, fill="#B7791F")

    rounded(draw, (70, 650, 760, 900), "#EAF2FA", outline="#2E74B5", radius=22, width=3)
    draw.text((100, 682), "Configuration evidence", font=heading, fill="#0B2545")
    config_lines = [
        "RelatedCommentsList is present on each verified Case layout.",
        "The five standard fields support body, visibility, author, date, and Case linkage.",
    ]
    y = 742
    for line in config_lines:
        draw.ellipse((102, y + 7, 114, y + 19), fill="#2E74B5")
        draw_wrapped(draw, (130, y), line, body, "#1F2937", 585, spacing=6)
        y += 72

    rounded(draw, (840, 650, 1530, 900), "#FFF6DE", outline="#B7791F", radius=22, width=3)
    draw.text((870, 682), "UAT sign-off gate", font=heading, fill="#0B2545")
    gate = (
        "A named CRC Team rep must post two internal comments on an approved UAT Case, then capture "
        "author, date/time, order, and visibility evidence."
    )
    draw_wrapped(draw, (870, 752), gate, body, "#1F2937", 605, spacing=8)
    image.save(path)


def build_document() -> None:
    copy2(REFERENCE, OUTPUT)
    document = Document(OUTPUT)
    clear_document_body(document)
    configure_document(document)

    diagram_path = ASSET_DIR / "scc3655-case-comments-flow.png"
    create_case_comment_diagram(diagram_path)

    # Page 1: executive decision and acceptance coverage.
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(2)
    set_run_font(paragraph.add_run("ARCHITECT REVIEW"), size=10, color=BLUE, bold=True)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(paragraph.add_run("SCC-3655 | Internal Case Comments"), size=25, color=NAVY, bold=True)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(16)
    set_run_font(
        paragraph.add_run("Great Plains Sandboxes | Read-only configuration and UAT readiness assessment"),
        size=13,
        color=MUTED,
    )

    metadata = [
        ("Environments", "GreatPlainsMerge 00DEa00000GkAsHMAV | GreatPlainsUAT 00DEa00000FZlLBMA1"),
        ("Review date", "26 August 2026"),
        ("Retrieval", "Succeeded | Merge 09SEa00000id0hxMAA | UAT 09SEa00000icYhBMAU"),
        ("Decision", "Conditional approval - configuration complete; runtime evidence pending"),
        ("Boundary", "No records, deployment, activation, or org metadata were changed"),
    ]
    table = document.add_table(rows=len(metadata), cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [1900, 7460])
    for row, (label, value) in zip(table.rows, metadata):
        shade_cell(row.cells[0], LIGHT_BLUE)
        set_cell_text(row.cells[0], label, size=9.5, color=NAVY, bold=True)
        set_cell_text(row.cells[1], value, size=9.5)
    document.add_paragraph().paragraph_format.space_after = Pt(2)

    add_callout(
        document,
        "Architect decision",
        "Do not deploy redundant SCC-3655 metadata. The standard Case Comments configuration is present. "
        "Approve runtime behavior only after a named UAT rep completes the two-comment evidence run; "
        "Trouble Ticket-specific testing also depends on its separate UAT record type and layout promotion.",
        fill=LIGHT_AMBER,
        accent=AMBER,
    )

    add_heading(document, "Acceptance outcome", 1)
    acceptance_rows = [
        ("1", "Internal designation", "Config pass", "IsPublished is createable and updateable; leave Public/Published unchecked for internal-only comments."),
        ("2", "Date and time", "Config pass", "CreatedDate is a non-editable system datetime field on CaseComment."),
        ("3", "Posting user", "Config pass", "CreatedById is a non-editable system reference that records the posting user."),
        ("4", "Lifecycle chronology", "Runtime pending", "The related list is configured, but both sandboxes currently contain zero CaseComment records."),
        ("5", "Rep can post and read", "Partial pass", "UAT has the related list and 46 active CRC Team users; a named-user UI test was not executed."),
    ]
    add_table(document, ["#", "Scope", "Result", "Evidence"], acceptance_rows, [500, 2100, 1500, 5260], status_column=2)

    # Page 2: visual path.
    document.add_page_break()
    add_heading(document, "Solution and visibility path", 1)
    add_body(
        document,
        "SCC-3655 uses only the standard Case Comments related list and CaseComment fields. The configuration path is complete; the remaining gate is persona-level runtime evidence in UAT.",
    )
    add_figure(
        document,
        diagram_path,
        "Figure 1. SCC-3655 internal Case Comment path",
        "Diagram showing a rep opening a Case related list, creating a standard CaseComment, and using the Public or Published checkbox to determine internal or external visibility. Configuration is complete and the named-rep UAT evidence step is pending.",
        width=6.45,
    )
    add_callout(
        document,
        "Design reading",
        "Green is verified platform configuration. Amber is the remaining UAT evidence gate.",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )

    # Page 3: detailed layout, field, and persona evidence.
    document.add_page_break()
    add_heading(document, "Detailed configuration findings", 1)
    add_heading(document, "1. Case layout presence", 2)
    layout_rows = [
        ("GreatPlainsMerge", "Case-Case Layout", "Present", "RelatedCommentsList is present."),
        ("GreatPlainsMerge", "Case-Trouble Ticket", "Present", "RelatedCommentsList is present."),
        ("GreatPlainsUAT", "Case-Case Layout", "Present", "RelatedCommentsList is present."),
        ("GreatPlainsUAT", "Case-Trouble Ticket", "Dependency missing", "Trouble_Ticket record type and dedicated layout are not yet present."),
    ]
    add_table(document, ["Org", "Layout", "Result", "Evidence"], layout_rows, [1900, 2300, 1600, 3560], status_column=2)

    add_heading(document, "2. Standard CaseComment behavior", 2)
    field_rows = [
        ("ParentId", "Reference", "Createable; required", "Links each comment to its Case."),
        ("CommentBody", "Text area", "Createable and updateable", "Stores the rep's comment body."),
        ("IsPublished", "Boolean", "Createable and updateable", "Unchecked is internal; checked is externally available."),
        ("CreatedById", "Reference", "System-set", "Records who posted the comment."),
        ("CreatedDate", "Date/Time", "System-set", "Records when the comment was created."),
    ]
    add_table(document, ["Field", "Type", "Write behavior", "Architect evidence"], field_rows, [1600, 1500, 2250, 4010])

    add_callout(
        document,
        "Persona evidence gate",
        "GreatPlainsUAT has 46 active GPC-Residential CRC Team users. The profile query alone does not prove "
        "that the intended rep can post and read Case Comments, so one named-user UI run is required.",
        fill=LIGHT_AMBER,
        accent=AMBER,
    )

    # Page 4: evidence boundary and reproducible UAT steps.
    document.add_page_break()
    add_heading(document, "Runtime evidence and UAT execution", 1)
    add_heading(document, "Read-only evidence boundary", 2)
    boundary_rows = [
        ("Metadata retrieved", "Three named Case layouts across the two sandboxes; both successful retrievals recorded."),
        ("Schema checked", "ParentId, CommentBody, IsPublished, CreatedById, and CreatedDate were verified live."),
        ("Available UAT Cases", "20 Case records exist in GreatPlainsUAT; no Case record was selected or changed."),
        ("Comment baseline", "GreatPlainsMerge and GreatPlainsUAT each contain zero CaseComment records."),
        ("Org write boundary", "No records, metadata, activation, deployment, or check-only deployment were changed or run."),
    ]
    add_table(document, ["Boundary", "Evidence"], boundary_rows, [2400, 6960])

    add_heading(document, "Named-rep UAT runbook", 2)
    runbook_rows = [
        ("1", "Assign test context", "Name one CRC Team rep and approve one UAT Case; record both identifiers."),
        ("2", "Post internal comment A", "Open Related > Case Comments > New; enter a unique body; leave Public unchecked."),
        ("3", "Post internal comment B", "Wait at least one minute, then add a second internal comment."),
        ("4", "Verify audit trail", "Confirm both entries display posting user and Created Date/Time."),
        ("5", "Verify lifecycle order", "Confirm the two comments appear in a consistent chronological sequence."),
        ("6", "Capture sign-off", "Save Case number, rep, timestamps, visibility result, ordering result, and screenshots."),
    ]
    add_table(document, ["Step", "Action", "Expected result / evidence"], runbook_rows, [700, 2700, 5960])

    # Page 5: decision gates and story-ready response.
    document.add_page_break()
    add_heading(document, "Required approval gates", 1)
    gate_rows = [
        ("1", "Name the UAT rep and Case", "QA / CRC lead", "Provide the approved Case number and the named GPC-Residential CRC Team tester."),
        ("2", "Run internal comment test", "Named rep", "Create two comments with Public unchecked and capture both saved records."),
        ("3", "Verify acceptance evidence", "QA", "Confirm author, Created Date/Time, internal visibility, and chronological order."),
        ("4", "Resolve external scope", "Product owner", "If external comments are in scope, approve a separate Public-checked test and external channel."),
        ("5", "Resolve Trouble Ticket dependency", "Release manager", "Promote the Trouble_Ticket record type and layout separately if SCC-3655 must be tested on that record type."),
    ]
    add_table(document, ["Gate", "Required action", "Owner", "Evidence needed"], gate_rows, [800, 2700, 1900, 3960])

    add_heading(document, "Ready-to-paste architect response", 1)
    response = (
        "Architect review completed across GreatPlainsMerge and GreatPlainsUAT. Conditional Approval: "
        "SCC-3655 uses standard Salesforce Case Comments, and no SCC-3655-specific deployment is required. "
        "RelatedCommentsList is present on the GreatPlainsMerge Case-Case Layout and Case-Trouble Ticket layout "
        "and on the GreatPlainsUAT Case-Case Layout. Live schema confirms CommentBody and IsPublished are writable, "
        "while CreatedById and CreatedDate are system-set. Both sandboxes currently contain zero CaseComment records, "
        "so rep posting, author display, timestamp display, chronological order, and internal/external visibility remain "
        "runtime-pending. GreatPlainsUAT has 20 Cases and 46 active GPC-Residential CRC Team users, but no named tester "
        "or approved Case was supplied. Complete the five gates above before final story approval. If testing must use a "
        "Trouble Ticket, promote that separate record type and layout prerequisite first. No org data, metadata, "
        "activation, or deployment changes were made."
    )
    add_callout(document, "Story comment", response, fill=LIGHT_BLUE, accent=BLUE)

    document.core_properties.title = "SCC-3655 Internal Case Comments Architect Review"
    document.core_properties.subject = "Read-only configuration and UAT readiness assessment"
    document.core_properties.author = "Solution Architecture Review"
    document.core_properties.keywords = "Salesforce, Great Plains, SCC-3655, Case Comments, architect review, UAT"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
