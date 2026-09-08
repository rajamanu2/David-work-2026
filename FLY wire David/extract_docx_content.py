from pathlib import Path
from docx import Document
import sys


def extract(path: Path) -> str:
    doc = Document(path)
    out = [f"FILE: {path}", f"PARAGRAPHS: {len(doc.paragraphs)}", f"TABLES: {len(doc.tables)}"]
    for i, p in enumerate(doc.paragraphs, 1):
        text = p.text.strip()
        if text:
            out.append(f"P{i} [{p.style.name}]: {text}")
    for ti, table in enumerate(doc.tables, 1):
        out.append(f"TABLE {ti} ({len(table.rows)}x{len(table.columns)})")
        for ri, row in enumerate(table.rows, 1):
            cells = [" ".join(c.text.split()) for c in row.cells]
            out.append(f"R{ri}: " + " | ".join(cells))
    for si, section in enumerate(doc.sections, 1):
        header = " ".join(p.text.strip() for p in section.header.paragraphs if p.text.strip())
        footer = " ".join(p.text.strip() for p in section.footer.paragraphs if p.text.strip())
        if header:
            out.append(f"SECTION {si} HEADER: {header}")
        if footer:
            out.append(f"SECTION {si} FOOTER: {footer}")
    return "\n".join(out)


if __name__ == "__main__":
    for arg in sys.argv[1:]:
        p = Path(arg)
        target = Path.cwd() / (p.stem + ".txt")
        target.write_text(extract(p), encoding="utf-8")
        print(target)
