from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'SALDEV-1441_Architect_Review.docx'
d = Document()
s = d.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.65)
s.left_margin = s.right_margin = Inches(.7)
normal = d.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(6)
for name in ['Heading 1', 'Heading 2']:
    d.styles[name].font.name = 'Calibri'
    d.styles[name].font.color.rgb = RGBColor.from_string('17365D')
d.styles['Heading 1'].font.size = Pt(17)
d.styles['Heading 2'].font.size = Pt(12)
header = s.header.paragraphs[0]
header.text = 'FLYWIRE  /  SALESFORCE CPQ                                      ARCHITECT REVIEW'
header.style = d.styles['Caption']
footer = s.footer.paragraphs[0]
footer.text = 'SALDEV-1441  |  Document-based review  |  08 September 2026                  '
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
footer.runs[0].font.size = Pt(8)

def p(text, style=None): return d.add_paragraph(text, style)
def h(text): d.add_heading(text, 2)
def table(headers, rows, widths):
    t = d.add_table(rows=1, cols=len(headers)); t.style = 'Light Shading Accent 1'
    t.autofit = False
    for i, text in enumerate(headers): t.rows[0].cells[i].text = text
    trpr = t.rows[0]._tr.get_or_add_trPr(); repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
    for row in rows:
        for cell, text in zip(t.add_row().cells, row): cell.text = text
    for row in t.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i])
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(4)
                for run in para.runs: run.font.size = Pt(9)
        pr = row._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
    return t
def page(title): d.add_page_break(); d.add_heading(title, 1)

p('ARCHITECT REVIEW', 'Subtitle')
d.add_heading('SALDEV-1441', 0)
p('Update business logic to work as needed for Step up product groupings vs Non SUP', 'Subtitle')
table(['Review item', 'Assessment'], [
('Review date', '08 September 2026'),
('Evidence reviewed', 'SALDEV-1441 (1).docx: user story, requirements, design, implementation notes, acceptance criteria and comments.'),
('Review disposition', 'Changes Requested — implementation clarification and test evidence required before architect sign-off.'),
('Review boundary', 'Document-based assessment only. No source-code inspection, org verification, test execution, check-only validation or deployment was performed.')], [1.4, 5.7])
h('Executive assessment')
p('The described approach is directionally sound: Step-Up lines should be validated within their effective Quote Line Group, while non-Step-Up lines retain existing quote-wide validation behavior. Reusing the existing group-context helpers can keep these rules consistent across nested bundles.')
p('The supplied notes do not establish that all acceptance criteria are complete. In particular, they do not describe the one-time-product cloning change or provide executed test results. This is an evidence gap, not proof that the implementation is defective.')
h('What the business is asking for')
table(['Scenario', 'Required outcome'], [
('Step-Up pricing', 'Allow products to repeat across ramping groups. Evaluate required pairs and applicable duplicate restrictions within each Step-Up group.'),
('Standard pricing', 'Preserve existing errors, warnings and duplicate restrictions on both grouped and ungrouped quotes.'),
('One-time products', 'Do not copy one-time products, such as implementation fees, into subsequent Quote Line Groups.'),
('Automatic additions', 'Ensure applicable automatically added products remain correct after grouping and cloning.')], [1.4, 5.7])
h('Example')
p('If Group 1 contains an Amex DOM line and Group 2 contains an Amex XB line, those lines must not satisfy each other’s pairing requirement. Each Step-Up group needs its own valid pair. A valid pair in one group must not hide an error in another.')

page('Design and implementation assessment')
p('The following items are reported by the ticket author. Their deployed state and runtime behavior have not been independently verified.')
table(['Component / rule', 'Reported change and architect assessment'], [
('validateAmexPairing', 'Reported to isolate Amex DOM/XB pairing by Step-Up group and support nested bundle lines. Confirm that the separate attribute-mismatch hard stop is also covered; pairing alone does not demonstrate attribute equality.'),
('checkDuplicateQuoteLineUSLoanProduct', 'Reported to check FAD pairing and duplicates by group and Ship To Account. Verify both the positive matching path and duplicate rejection within that same context.'),
('checkPaymentAndCountryMatrixRules', 'Reported to make CC Country and MC/Visa Credit/Debit pairing group-aware while retaining existing matching dimensions.'),
('checkIdenticalSurchargePairRule', 'Reported to require Surcharging Yes/No pairs in the same Step-Up group. Verify that a No line in a different group cannot satisfy the rule.'),
('Shared helpers', 'isStepUpLine, getGroupIdentity and getValidationContext are reportedly reused without changes. Confirm correct group classification and parent-group resolution, including unsaved or newly cloned lines.'),
('Quote-level Boolean fields', 'Existing Product Rule fields are reportedly retained. Verify that any invalid group keeps the applicable quote error flag set; a later valid group must not clear it.'),
('UK/Canada auto-add rule', 'Ticket reports removal of XB Payments (FWXB1000) applicability from Product Rule a1thG0000004EsUQAU. Verify that intended SFS behavior is preserved.')], [2.05, 5.05])
h('Explicit scope distinctions')
p('The design notes say HC Professional Services and applicable Payments XB automatic additions need no additional group adjustment because cloning supplies them. This remains a behavior to test. The two Payex automatic-add rules are described as deactivated; the story does not request their reactivation.')
p('The implementation notes state that onAfterCalculate, shared helper behavior and non-Step-Up quote-wide validation were not changed. These are author statements rather than verified regression results.')
h('Related work')
p('SALDEV-1417 is listed as the prerequisite for distinguishing group types. The one-time-product requirement was explicitly moved from SALDEV-1427 into SALDEV-1441 on 17 August. Confirm ownership and inclusion of that change in this story’s review evidence.')

page('Findings and required follow-up')
table(['ID / priority', 'Finding', 'Evidence required to close'], [
('AR-01 / High', 'One-time-product exclusion is required by acceptance criterion 1.a, but no cloning implementation is described.', 'Identify the changed cloning logic and show a one-time fee present in the first group and absent from all subsequent cloned groups.'),
('AR-02 / High', 'No executed acceptance or regression results are included in the supplied document.', 'Provide results for Step-Up, standard grouped and ungrouped quotes, covering each applicable rule with both valid and invalid examples.'),
('AR-03 / Medium', 'Amex attribute-mismatch coverage is not explicit in the implementation notes.', 'Map the attribute-mismatch hard stop to its code/field logic and provide same-group mismatch and corrected-match results.'),
('AR-04 / Medium', 'Group-context and quote-flag behavior are described but not demonstrated.', 'Provide a focused code diff and tests for nested bundles, newly cloned groups, any-invalid-group flag aggregation and error clearing after correction.'),
('AR-05 / Medium', 'Automatic-add behavior relies partly on cloning, and the UK/Canada configuration has changed.', 'Show expected first/subsequent group contents and confirm that removing XB Payments applicability does not remove required SFS behavior.')], [1.05, 2.8, 3.25])
h('Architect recommendations for verification')
p('Use the established group identity helper consistently across all four validation methods. The requirement’s suggestion to use a group start date is not sufficient evidence of unique identity; verify that distinct groups sharing a date remain separate.')
p('For quotes containing both Step-Up and standard groups, confirm that standard lines retain their intended quote-wide context and that Step-Up lines cannot incorrectly satisfy standard validation or another Step-Up group. Treat this as an additional architecture edge-case check.')
p('When users correct an invalid pair, remove a line or recalculate, verify that the existing quote-level Boolean fields reflect the current result and do not retain stale errors.')
h('Recommended responsibility')
p('Developer: provide the scoped QCP/configuration diff and identify the one-time cloning change. QA: execute and record the matrix on the next page. Architect: reassess the findings against that evidence. Product owner: resolve any ambiguity about one-time product identification or mixed-group behavior.')

page('Acceptance and regression test matrix')
p('All tests below are required or recommended verification; none were executed in this review. Record the environment, quote ID, setup, expected/actual result and screenshots or equivalent evidence for each.')
table(['Test', 'Scenario and expected result'], [
('T01 — Amex', 'In Step-Up groups, split DOM/XB across groups: error remains. Add the valid pair within each group: pairing error clears. Mismatched Amex attributes still trigger the hard stop.'),
('T02 — FAD', 'Require the matching DOM/XB line in the same group and Ship To Account. Reject duplicate FAD in that context. Allow valid repetitions across separate Step-Up groups.'),
('T03 — CC Country', 'Locally Issued and Non-Locally Issued pairing must satisfy the existing criteria within each Step-Up group. A match from another group must not count.'),
('T04 — MC/Visa', 'Credit/Debit pairing must satisfy all existing matching dimensions within each Step-Up group. Test both valid and missing pairs.'),
('T05 — Surcharging', 'A Surcharging Yes line requires the matching No line in the same Step-Up group. Another group’s No line must not satisfy the rule.'),
('T06 — One-time fees', 'With Allow Product Ramping checked, clone multiple groups. One-time products remain in the first group and are absent from every subsequent group; recurring products copy as intended.'),
('T07 — Standard groups', 'With Allow Product Ramping unchecked, repeat all applicable rule cases. Existing quote-wide errors, warnings and duplicate restrictions remain unchanged.'),
('T08 — No groups', 'Repeat applicable valid/invalid cases on an ungrouped quote. Confirm existing behavior remains unchanged.'),
('T09 — Auto-add', 'Verify applicable HC and Payments XB additions and cloned contents. For UK/Canada, confirm removed FWXB1000 applicability and retained intended SFS behavior.'),
('T10 — Edge cases', 'Verify nested bundles, newly cloned groups, equal start dates in distinct groups and mixed group types. An invalid group must keep the error active even when another is valid.'),
('T11 — Correction', 'Correct or delete invalid lines and recalculate. Errors clear only when all applicable contexts are valid; no stale Boolean flags remain.')], [1.2, 5.9])
h('Ready-to-paste architect response')
p('Document-based architect review completed for SALDEV-1441. The proposed group-aware QCP approach aligns with the story intent, but sign-off remains pending. Please identify the one-time-product cloning change, confirm Amex attribute-mismatch coverage, and attach the scoped implementation diff plus acceptance/regression results for Step-Up, standard grouped and ungrouped quotes. Include nested-group resolution, quote-level error aggregation and auto-add behavior. No deployment or live-org validation was performed as part of this review.')
p('Source: SALDEV-1441 (1).docx supplied by the requester. Review statements about implementation are based solely on that attachment.', 'Caption')
d.core_properties.title = 'SALDEV-1441 Architect Review'
d.core_properties.subject = 'Document-based architecture assessment; no deployment'
d.core_properties.author = 'Architecture Review'
ROOT.mkdir(parents=True, exist_ok=True)
d.save(OUT)
print(OUT)
