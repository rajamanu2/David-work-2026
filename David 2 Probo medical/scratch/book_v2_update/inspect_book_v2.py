import json
import openpyxl

path = r"C:\Users\LIKKI\Documents\ChatGPT\David 2 Probo medical\scratch\book_v2_update\Book v2 source.xlsx"
wb = openpyxl.load_workbook(path, read_only=True, data_only=False)

summary = []
for ws in wb.worksheets:
    rows = []
    for row in ws.iter_rows(
        min_row=1,
        max_row=min(ws.max_row, 8),
        max_col=min(ws.max_column, 12),
        values_only=True,
    ):
        rows.append(list(row))
    summary.append(
        {
            "title": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "sample": rows,
        }
    )

print(json.dumps(summary, default=str, indent=2))
