from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import build_architect_review as base
from build_architect_review import (
    AMBER,
    BLUE,
    GRID,
    INK,
    LIGHT_AMBER,
    LIGHT_BLUE,
    LIGHT_GRAY,
    LIGHT_RED,
    LIGHT_TEAL,
    MUTED,
    NAVY,
    RED,
    TEAL,
    add_body,
    add_callout,
    add_figure,
    add_heading,
    add_table,
    arrow,
    configure_header_footer,
    draw_wrapped,
    load_font,
    rgb,
    rounded,
    set_cell_text,
    set_run_font,
    set_table_geometry,
    shade_cell,
)


ROOT = Path(__file__).resolve().parent
ASSET_DIR = ROOT / "doc-assets-scc4166"
OUTPUT = ROOT / "SCC-4166-architect-review-final.docx"
ASSET_DIR.mkdir(exist_ok=True)


def preset_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    """decision_memo table token override: 80/120/80/120 DXA."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


# The shared helper is intentionally patched to the selected decision_memo preset.
base.set_cell_margins = preset_cell_margins


def create_logic_diagram(path: Path) -> None:
    image = Image.new("RGB", (1600, 1000), "#F7F9FC")
    draw = ImageDraw.Draw(image)
    title = load_font(42, bold=True)
    subtitle = load_font(22)
    heading = load_font(25, bold=True)
    body = load_font(19)
    small = load_font(17)

    draw.text((65, 50), "SCC-4166 | MTTR calculation and control path", font=title, fill="#0B2545")
    draw.text(
        (65, 110),
        "The formula design is sound in principle; two implementation defects prevent the required behavior.",
        font=subtitle,
        fill="#5B6777",
    )

    nodes = [
        ((60, 205, 405, 485), "Case opens", ["Endpoint: CreatedDate", "No Work Order: use ClosedDate", "Work Order: use completion stamp"], "#EAF2FA", "#2E74B5"),
        ((505, 205, 865, 485), "Pending customer", ["Enter: save Start", "Exit: save End", "Only most recent interval retained"], "#FFF6DE", "#B7791F"),
        ((965, 205, 1540, 485), "MTTR result", ["Elapsed minutes", "minus most recent hold interval", "render as Hours and Minutes"], "#E7F5F1", "#16836B"),
    ]
    for box, label, lines, fill, outline in nodes:
        rounded(draw, box, fill, outline=outline, radius=22, width=4)
        draw.text((box[0] + 25, box[1] + 25), label, font=heading, fill="#0B2545")
        y = box[1] + 90
        for line in lines:
            draw.ellipse((box[0] + 27, y + 7, box[0] + 39, y + 19), fill=outline)
            draw_wrapped(draw, (box[0] + 55, y), line, body, "#1F2937", box[2] - box[0] - 85, spacing=5)
            y += 55
    arrow(draw, (405, 345), (505, 345), color="#8A97A8", width=6)
    arrow(draw, (865, 345), (965, 345), color="#8A97A8", width=6)

    rounded(draw, (60, 580, 760, 900), "#FDECEC", outline="#C53030", radius=22, width=4)
    draw.text((90, 612), "Defect 1 | Exit branch never matches", font=heading, fill="#C53030")
    defect_one = [
        "Entered and Exited branches use the same condition.",
        "Pending_Customer_End__c remains blank after returning to In Progress.",
        "The formula deducts hold time through closure instead of through hold exit.",
    ]
    y = 680
    for line in defect_one:
        draw.ellipse((92, y + 7, 104, y + 19), fill="#C53030")
        draw_wrapped(draw, (120, y), line, body, "#1F2937", 590, spacing=6)
        y += 66

    rounded(draw, (840, 580, 1540, 900), "#FDECEC", outline="#C53030", radius=22, width=4)
    draw.text((870, 612), "Defect 2 | Blank displays as zero", font=heading, fill="#C53030")
    defect_two = [
        "MTTR__c uses BlankAsZero while testing ISBLANK(MTTR_Minutes__c).",
        "20 open source-sandbox cases display 0 Hours 0 Minutes with blank numeric MTTR.",
        "Local repair changes the text formula to BlankAsBlank.",
    ]
    y = 680
    for line in defect_two:
        draw.ellipse((872, y + 7, 884, y + 19), fill="#C53030")
        draw_wrapped(draw, (900, y), line, body, "#1F2937", 590, spacing=6)
        y += 66

    draw.text((65, 948), "Green = intended path  |  Amber = tracked hold interval  |  Red = release blocker", font=small, fill="#5B6777")
    image.save(path)


def configure_document(document: Document) -> None:
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        configure_header_footer(section)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 12, 6),
        ("Heading 2", 13, BLUE, 10, 5),
        ("Heading 3", 12, NAVY, 8, 4),
    ):
        style = document.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def build_document() -> None:
    diagram_path = ASSET_DIR / "scc4166-mttr-control-path.png"
    create_logic_diagram(diagram_path)

    document = Document()
    configure_document(document)

    # Page 1: decision memo.
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(2)
    set_run_font(p.add_run("ARCHITECT REVIEW"), size=10, color=BLUE, bold=True)

    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run("SCC-4166 | Mean Time to Repair"), size=25, color=NAVY, bold=True)

    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    set_run_font(p.add_run("Great Plains Merge Sandbox | Local repair and UAT check-only assessment"), size=13, color=MUTED)

    metadata = [
        ("Source", "GreatPlainsMerge | Org 00DEa00000GkAsHMAV | Sandbox"),
        ("Target", "GreatPlainsUAT | Org 00DEa00000FZlLBMA1 | Sandbox"),
        ("Review date", "26 August 2026"),
        ("Source retrieval", "Succeeded | Job 09SEa00000icqSVMAY"),
        ("Decision", "Changes Requested - do not promote the full scope"),
        ("Boundary", "No deployment, activation, or Salesforce record changes"),
    ]
    table = document.add_table(rows=len(metadata), cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [1900, 7460])
    for row, (label, value) in zip(table.rows, metadata):
        shade_cell(row.cells[0], LIGHT_BLUE)
        set_cell_text(row.cells[0], label, size=9.5, color=NAVY, bold=True)
        set_cell_text(row.cells[1], value, size=9.5)
    document.add_paragraph().paragraph_format.space_after = Pt(2)

    add_callout(
        document,
        "Architect decision",
        "Do not promote the full SCC-4166 package. The corrected 8-component calculation core passes check-only validation, but the source implementation has two functional defects and the full 13-component scope still fails on five UAT prerequisite/profile errors.",
        fill=LIGHT_RED,
        accent=RED,
    )

    add_heading(document, "Acceptance outcome", 1)
    acceptance_rows = [
        ("1", "No-work-order MTTR", "Partial fail", "Numeric formula waits for ClosedDate, but open cases display 0 Hours 0 Minutes instead of blank."),
        ("2", "Work-order MTTR", "Partial pass", "Work Order flow stamps the first Completed/Closed transition and marks the Case as having a Work Order."),
        ("3", "Exclude pending-customer time", "Fail", "Exit branch duplicates the entry condition, so Pending Customer End is never set on exit."),
        ("4", "Hours and minutes", "Pass", "Text formula formats the numeric result as Hours and Minutes."),
        ("5", "UAT promotion readiness", "Fail", "Full scope validates 8/13; two layouts and three profiles fail."),
    ]
    add_table(document, ["#", "Requirement", "Result", "Evidence"], acceptance_rows, [500, 2300, 1500, 5060], status_column=2)

    # Page 2: visual logic map.
    document.add_page_break()
    add_heading(document, "MTTR calculation and control path", 1)
    add_body(document, "SCC-4166 uses two endpoint strategies and subtracts the most recent pending-customer interval. The two red defects below are independent and both must be corrected.")
    add_figure(
        document,
        diagram_path,
        "Figure 1. SCC-4166 MTTR calculation and control path",
        "Diagram showing Case open time flowing through a pending-customer interval to an MTTR result. Two red defect boxes identify the unreachable exit branch and the blank-as-zero text formula behavior.",
        width=6.45,
    )
    add_callout(document, "Interpretation", "The architecture needs no new code. A surgical Flow condition repair and formula blank-handling correction restore the intended core behavior.", fill=LIGHT_BLUE, accent=BLUE)

    # Page 3: detailed defects and local repair.
    document.add_page_break()
    add_heading(document, "Detailed findings and local repair", 1)
    defect_rows = [
        ("Case hold exit", "Blocking", "Exited Pending Customer repeats current=hold and prior!=hold.", "Change to current!=hold and prior=hold."),
        ("Open-case display", "Blocking", "MTTR__c uses BlankAsZero; 20 existing open Cases show 0 Hours 0 Minutes while MTTR Minutes is blank.", "Use BlankAsBlank so the ISBLANK guard returns an empty string."),
        ("Work Order endpoint", "Pass", "Active/latest Work_Order_Update_Case_MTTR_Details stamps the first Completed or Closed transition.", "Retain; regression-test both statuses."),
        ("Most recent hold only", "Pass after fix", "Entering hold overwrites Start and clears End, matching the approved most-recent-interval design.", "Retain assignment behavior."),
    ]
    add_table(document, ["Area", "Result", "Source evidence", "Local action"], defect_rows, [1700, 1250, 3440, 2970], status_column=1)

    add_heading(document, "Field and automation inventory", 2)
    inventory_rows = [
        ("Case.Has_Work_Order__c", "Checkbox", "Set true by Work Order after-save flow."),
        ("Case.Work_Order_Completion_Date__c", "Date/Time", "First Completed/Closed Work Order timestamp."),
        ("Case.Pending_Customer_Start__c", "Date/Time", "Most recent hold-entry timestamp."),
        ("Case.Pending_Customer_End__c", "Date/Time", "Most recent hold-exit timestamp; broken in source."),
        ("Case.MTTR_Minutes__c", "Formula Number", "Elapsed minutes minus most recent hold interval."),
        ("Case.MTTR__c", "Formula Text", "Hours/minutes display; blank handling repaired locally."),
    ]
    add_table(document, ["Component", "Type", "Role"], inventory_rows, [3800, 1600, 3960])

    add_callout(
        document,
        "Component-list correction",
        "The story document duplicates MTTR_Minutes__c and omits both active Flows plus Case-Case Layout. The architecture-complete review scope is 13 unique components: 6 fields, 2 Flows, 2 layouts, and 3 profiles.",
        fill=LIGHT_AMBER,
        accent=AMBER,
    )

    # Page 4: validation evidence.
    document.add_page_break()
    add_heading(document, "UAT check-only validation evidence", 1)
    validation_rows = [
        ("Corrected core", "0AfEa00000bbRYbKAM", "Succeeded", "8/8 components; 6 fields + 2 Flows; 0 tests; checkOnly=true."),
        ("Corrected full scope", "0AfEa00000bbRbpKAE", "Failed", "8/13 succeeded; 5 errors; 0 tests; rollbackOnError=true."),
    ]
    add_table(document, ["Scope", "Job", "Result", "Evidence"], validation_rows, [1900, 2450, 1250, 3760], status_column=2)

    add_heading(document, "Five full-scope blockers", 2)
    blocker_rows = [
        ("Case-Case Layout", "Missing Case.Service_Address__c in GreatPlainsUAT."),
        ("Case-Trouble Ticket", "Missing Case.Service_Address__c in GreatPlainsUAT."),
        ("Admin profile", "Unsupported target user permission ArchiveArticles."),
        ("Standard profile", "Missing Case.Trouble_Ticket record type in GreatPlainsUAT."),
        ("System Administrator - API Only", "Unsupported target user permission ArchiveArticles."),
    ]
    add_table(document, ["Failed component", "Validation error"], blocker_rows, [3300, 6060])

    add_heading(document, "Evidence boundary", 2)
    boundary_rows = [
        ("Source state", "Both SCC-4166 Flows are active and latest in GreatPlainsMerge."),
        ("Live read-back", "20 open source-sandbox Cases show text MTTR of 0 Hours 0 Minutes with blank numeric MTTR; zero Cases have a pending-start timestamp."),
        ("Target read-back", "After validation, GreatPlainsUAT still has none of the six fields or two Flows."),
        ("No runtime mutation", "No Cases, Work Orders, metadata, Flow activations, or org settings were changed."),
    ]
    add_table(document, ["Boundary", "Evidence"], boundary_rows, [2500, 6860])

    # Page 5: gates, test plan, stakeholder response.
    document.add_page_break()
    add_heading(document, "Required approval gates", 1)
    gate_rows = [
        ("1", "Accept the two local repairs", "Solution architect", "Review Flow exit logic and MTTR blank behavior."),
        ("2", "Promote SCC-3384 prerequisites first", "Release manager", "Case.Service_Address__c and Case.Trouble_Ticket must exist in UAT."),
        ("3", "Replace broad profile payloads", "Security owner", "Use a focused permission set or retrieve target-aligned profiles; remove unrelated ArchiveArticles drift."),
        ("4", "Rerun full validation", "Release manager", "Require 13/13 success before any deployment approval."),
        ("5", "Run business tests", "QA", "Capture Case/Work Order IDs and field read-back for both endpoint paths."),
    ]
    add_table(document, ["Gate", "Required action", "Owner", "Evidence needed"], gate_rows, [800, 2800, 1900, 3860])

    add_heading(document, "Business-user regression sequence", 2)
    test_rows = [
        ("A", "No Work Order", "Create Trouble Ticket; set In Progress; enter hold; exit hold; close Case.", "MTTR is blank until Closed and excludes only the hold interval."),
        ("B", "With Work Order", "Create Work Order; move to Completed; leave Case open.", "Has Work Order=true; completion timestamp set once; MTTR populated."),
        ("C", "Repeated hold", "Enter/exit hold twice before closure.", "Only the second interval is excluded, per approved design."),
        ("D", "Negative regression", "Open Case with no completion endpoint.", "MTTR text and minutes are both blank, never 0 Hours 0 Minutes."),
    ]
    add_table(document, ["Test", "Path", "Steps", "Expected"], test_rows, [800, 1600, 3760, 3200])

    add_heading(document, "Ready-to-paste architect response", 2)
    response = (
        "Architect review completed in GreatPlainsMerge with check-only validation against GreatPlainsUAT. "
        "Changes Requested: the active Case before-save Flow has a duplicated pending-customer entry condition, so the exit timestamp is never set, and MTTR__c treats blank numeric MTTR as zero; 20 open source-sandbox Cases currently display 0 Hours 0 Minutes. "
        "I prepared local surgical repairs for both defects. The corrected 8-component core passes check-only validation (job 0AfEa00000bbRYbKAM, 8/8). "
        "The corrected full 13-component scope remains blocked (job 0AfEa00000bbRbpKAE, 8/13) because UAT lacks Case.Service_Address__c and Case.Trouble_Ticket, and the broad profile payloads contain ArchiveArticles drift. "
        "Deploy SCC-3384 prerequisites first, replace or trim the profile changes, rerun 13/13 validation, then execute the four business tests. No metadata, activation, settings, or records were changed."
    )
    add_callout(document, "Story comment", response, fill=LIGHT_BLUE, accent=BLUE)

    document.core_properties.title = "SCC-4166 Mean Time to Repair Architect Review"
    document.core_properties.subject = "Local repair and UAT check-only readiness assessment"
    document.core_properties.author = "Solution Architecture Review"
    document.core_properties.keywords = "Salesforce, Great Plains, SCC-4166, MTTR, Flow, check-only, architect review"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
