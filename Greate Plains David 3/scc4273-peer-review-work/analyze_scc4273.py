from __future__ import annotations

import json
import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document


NS = {"m": "http://soap.sforce.com/2006/04/metadata"}


def text(node: ET.Element | None, path: str, default: str = "") -> str:
    if node is None:
        return default
    value = node.findtext(path, default=default, namespaces=NS)
    return value or default


def normalize(value: str) -> str:
    return re.sub(r"[^A-Z0-9]+", "", value.upper())


def source_lists(source_docx: Path) -> dict[str, list[str]]:
    doc = Document(source_docx)
    return {
        "primary_resolution": [v.strip() for v in doc.paragraphs[25].text.splitlines() if v.strip()],
        "subscriber_report": [v.strip() for v in doc.paragraphs[28].text.splitlines() if v.strip()],
    }


def field_metadata(path: Path) -> dict:
    root = ET.parse(path).getroot()
    values = [
        text(value, "m:fullName")
        for value in root.findall(".//m:valueSetDefinition/m:value", NS)
    ]
    settings = []
    for setting in root.findall(".//m:valueSettings", NS):
        settings.append(
            {
                "value": text(setting, "m:valueName"),
                "controllers": [
                    node.text or ""
                    for node in setting.findall("m:controllingFieldValue", NS)
                ],
            }
        )
    return {
        "full_name": text(root, "m:fullName"),
        "label": text(root, "m:label"),
        "type": text(root, "m:type"),
        "controller": text(root, ".//m:valueSet/m:controllingField"),
        "value_set_name": text(root, ".//m:valueSet/m:valueSetName"),
        "values": values,
        "settings": settings,
    }


def compare(desired: list[str], actual: list[str]) -> dict:
    actual_norm = {normalize(v): v for v in actual}
    desired_norm = {normalize(v): v for v in desired}
    return {
        "desired_count": len(desired),
        "actual_count": len(actual),
        "missing_normalized": [v for v in desired if normalize(v) not in actual_norm],
        "extra_normalized": [v for v in actual if normalize(v) not in desired_norm],
        "label_variants": [
            {"source": v, "merge": actual_norm[normalize(v)]}
            for v in desired
            if normalize(v) in actual_norm and v != actual_norm[normalize(v)]
        ],
    }


def record_type_metadata(path: Path) -> dict:
    root = ET.parse(path).getroot()
    picklists: dict[str, list[str]] = {}
    for item in root.findall("m:picklistValues", NS):
        name = text(item, "m:picklist")
        picklists[name] = [
            text(value, "m:fullName")
            for value in item.findall("m:values", NS)
        ]
    return {
        "active": text(root, "m:active"),
        "label": text(root, "m:label"),
        "business_process": text(root, "m:businessProcess"),
        "picklists": picklists,
    }


def layout_metadata(path: Path) -> list[dict]:
    root = ET.parse(path).getroot()
    sections = []
    for section in root.findall("m:layoutSections", NS):
        columns = []
        for column in section.findall("m:layoutColumns", NS):
            columns.append(
                [
                    {
                        "behavior": text(item, "m:behavior"),
                        "field": text(item, "m:field"),
                    }
                    for item in column.findall("m:layoutItems", NS)
                ]
            )
        sections.append(
            {
                "label": text(section, "m:label"),
                "style": text(section, "m:style"),
                "columns": columns,
            }
        )
    return sections


def profile_metadata(path: Path, fields: set[str]) -> list[dict]:
    root = ET.parse(path).getroot()
    out = []
    for item in root.findall("m:fieldPermissions", NS):
        field = text(item, "m:field")
        if field in fields:
            out.append(
                {
                    "field": field,
                    "readable": text(item, "m:readable"),
                    "editable": text(item, "m:editable"),
                }
            )
    return out


def main() -> None:
    source_docx = Path(sys.argv[1])
    root = Path(sys.argv[2])
    source = source_lists(source_docx)
    primary = field_metadata(root / "objects/Case/fields/Primary_Resolution__c.field-meta.xml")
    subscriber = field_metadata(root / "objects/Case/fields/Subscriber_Report__c.field-meta.xml")
    record_type = record_type_metadata(root / "objects/Case/recordTypes/Trouble_Ticket.recordType-meta.xml")
    layout = layout_metadata(root / "layouts/Case-Trouble Ticket.layout-meta.xml")
    review_fields = {
        "Case.Problem_Occurred_Date__c",
        "Case.Problem_Reported_Date__c",
        "Case.Repair_Date__c",
        "Case.Resolution_Detail__c",
        "Case.Service_Type__c",
        "Case.Service_Address__c",
        "Case.Subscriber_Report__c",
        "Case.Primary_Resolution__c",
        "Case.Secondary_Resolution__c",
    }
    profile_path = root / "profiles/System Administrator - API Only.profile-meta.xml"
    result = {
        "source": source,
        "primary_resolution": primary,
        "subscriber_report": subscriber,
        "primary_comparison": compare(source["primary_resolution"], primary["values"]),
        "subscriber_comparison": compare(source["subscriber_report"], subscriber["values"]),
        "record_type": record_type,
        "layout": layout,
        "api_only_profile_permissions": profile_metadata(profile_path, review_fields),
    }
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
