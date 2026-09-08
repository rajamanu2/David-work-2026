from __future__ import annotations

import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document


DOCX = Path(r"C:\Users\LIKKI\Documents\ChatGPT\David 2 Probo medical\outputs\case-00010716\Case_00010716_UAT_Photo_Import_Reproduction_and_Change_Guide_AutoFast.docx")

doc = Document(DOCX)
assert len(doc.sections) == 1
assert len(doc.tables) == 10
assert len([p for p in doc.paragraphs if p.style.name == "Heading 1"]) == 9

with ZipFile(DOCX) as zf:
    xml_names = [name for name in zf.namelist() if name.endswith(".xml") or name.endswith(".rels")]
    package_text = "\n".join(zf.read(name).decode("utf-8", errors="replace") for name in xml_names)

plain_xml_text = " ".join(re.findall(r">([^<>]+)<", package_text))

for forbidden in (
    "Lorem ipsum",
    "[Name]",
    "[Audience]",
    "Strategic decision or question",
    "DEVDO PASS",
    "0AfiK0000000iNdSAI",
    "0AfiK0000000iPFSAY",
    "707iK000000793S",
    "0AfjH0000000T5FSAU",
    "0AfjH0000000T6rSAE",
    "frontdoor.jsp?sid=",
    "�",
):
    assert forbidden not in package_text, forbidden

assert not re.search(r"\b0Af[A-Za-z0-9]{12,15}\b", plain_xml_text), "Deployment/validation ID found"
assert not re.search(r"\b707[A-Za-z0-9]{12,15}\b", plain_xml_text), "Apex test-run ID found"

for required in (
    "AUTOFAST UAT REPRODUCTION &amp; CHANGE GUIDE",
    "00DjH0000000rYzUAI",
    "cAuthURIForEval",
    "cGoogleAppAuthenticationWithSalesforce",
    "GoogleAuthTestClass",
    "E-2026-036025",
    "E-2026-036026",
    "552082",
    "A-218876",
    "C-Lens.jpg",
    "C-Cap.jpg",
    "C-Whole.jpg",
    "C-SN.jpg",
    "case00010716RejectsMissingOrInvalidCallbackState",
):
    assert required in package_text, required

assert "w:drawing" not in package_text
assert "w:pict" not in package_text
assert "a:blip" not in package_text

for url in (
    "a06jH0000004E2rQAE",
    "a0mjH0000000BWXQA2",
    "a1tjH0000003n98QAA",
    "a1tjH0000003n99QAA",
    "AuthenticationGoogleDrive?id=a1tjH0000003n98QAA",
):
    assert url in package_text, url

print("CONTENT AUDIT: PASS")
print(f"size={DOCX.stat().st_size} bytes tables={len(doc.tables)} headings=9")
