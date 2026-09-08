from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

base=Path(__file__).parent
source=Path(r'C:/Users/LIKKI/Downloads/ref do not used SALDEV-1420-20260907T161927Z-1-001/ref do not used SALDEV-1420/Story Test Execution Document for SALDEV-1420.docx')
assets=base.parent/'SALDEV-1420-all-docs'
doc=Document(source)
body=doc._element.body
for child in list(body):
    if child.tag != qn('w:sectPr'): body.remove(child)
for sec in doc.sections:
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    for part in (sec.header,sec.footer):
        for p in part.paragraphs:p.text=''
normal=doc.styles['Normal']
normal.font.name='Arial';normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.line_spacing=1.06
for name,size in [('Title',22),('Heading 1',16),('Heading 2',12)]:
    s=doc.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_after=Pt(8);s.paragraph_format.space_before=Pt(4)
    s.paragraph_format.keep_with_next=True
    if name!='Title':s.font.bold=True
    for borders in list(s.element.xpath('.//w:pBdr')):borders.getparent().remove(borders)
doc.core_properties.title='SALDEV 1420 Story Test Execution'
doc.core_properties.subject='Email approval testing in Partial'
doc.core_properties.author=''
doc.core_properties.last_modified_by=''
doc.core_properties.comments=''
doc.core_properties.keywords=''

def p(t,bold=False):
    x=doc.add_paragraph();r=x.add_run(t);r.bold=bold;return x
def h(t):doc.add_paragraph(t,'Heading 2')
def page(title):doc.add_page_break();doc.add_paragraph(title,'Heading 1')
def steps(items):
    for n,t in enumerate(items,1):p(f'{n}. {t}')
def bullets(items):
    for t in items:p('• '+t)
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for i,tx in enumerate(headers):t.rows[0].cells[i].text=tx
    for row in rows:
        cells=t.add_row().cells
        for c,tx in zip(cells,row):c.text=str(tx)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();x=OxmlElement('w:cantSplit');trpr.append(x)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for ci,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[ci])
            pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
            pr.append(b)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EEF2' if ri==0 else 'FFFFFF');pr.append(shade)
            for pa in c.paragraphs:
                pa.paragraph_format.space_after=Pt(4);pa.paragraph_format.space_before=Pt(4)
                for run in pa.runs:run.font.size=Pt(9);run.bold=(ri==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t
def pic(rel,caption,maxh=3.65):
    from PIL import Image
    fp=assets/rel;w,hh=Image.open(fp).size
    crops={
        'doc6/image13.png':(.15,.53,.38,.01),
        'doc6/image23.png':(.12,.335,.575,.005),
        'doc6/image25.png':(.15,.38,.48,.01),
        'doc6/image41.png':(.155,.40,.49,.04),
        'doc6/image12.png':(.12,.335,.48,.005),
    }
    crop=crops.get(rel,(0,0,0,0))
    l,t,r,b=crop
    ratio=w*(1-l-r)/(hh*(1-t-b))
    width=min(6.5,maxh*ratio)
    x=doc.add_paragraph();x.paragraph_format.space_after=Pt(3)
    shape=x.add_run().add_picture(str(fp),width=Inches(width),height=Inches(width/ratio))
    if rel in crops:
        fill=shape._inline.graphic.graphicData.pic.blipFill
        rect=OxmlElement('a:srcRect')
        for key,value in zip(('l','t','r','b'),crop):rect.set(key,str(round(value*100000)))
        fill.insert(1,rect)
    shape._inline.docPr.set('descr',caption)
    cap=p(caption);cap.paragraph_format.space_after=Pt(7)
    for r in cap.runs:r.italic=True;r.font.size=Pt(9)
def result():p('Result: Not run / Pass / Fail / Blocked     Tester: __________________     Date: __________',False)

doc.add_paragraph('SALDEV 1420\nStory Test Execution','Title')
p('Approval email testing in Partial',True)
p('Use this document to reproduce the email-table scenarios and record the business test results. Confirm group order, product order, descriptions and approval-template behavior against the saved quote.')
table(['Test information','Details'],[
('Environment','Flywire Partial sandbox'),('Test user','Record the CPQ Sales user and approver used for each case'),('Review date','7 September 2026'),('Execution status','Business retest to be completed'),('Evidence','Earlier screenshots are references. Attach a fresh output screenshot for each executed case.')],[1.4,5.1])
h('Before testing')
steps(['Confirm the browser address contains flywire--partial.sandbox.',
'Use a dedicated test quote or an agreed test copy. Open Edit Lines, save and calculate, then check the persisted dates, terms and line order.',
'Use the appropriate existing approval record for a template preview. If an email must be sent, use the agreed sandbox test recipient and normal approval process.',
'Record the quote link, approval record, tester, expected result and actual result. Do not mark a case passed from an earlier screenshot.'])
h('Expected ordering')
p('Groups appear by start date, earliest first. When start dates match, the longer subscription term appears first. Equal-date and equal-term groups retain QLE group sequence. Products retain QLE sequence, with bundle options beneath their parent.')
p('Earlier screenshots may show Group names, date banners or an older row order. The expected results in each test case define the behavior to verify now.')

template_cases=[
('01 Quote approval email','Request - Quote Approval Request - AQS','Quote Approval','doc6/image13.png',[
'Use a quote with grouped lines, a bundle and standalone products. Put at least two products in nonalphabetical QLE order.',
'Save and calculate. Use its Quote Approval record and select the Request - Quote Approval Request - AQS template.',
'Open the available preview. If using the normal submission workflow, inspect the email received by the agreed test approver.',
'Compare the group headers, product sequence, prices and quantities with the saved quote.'],
'Confirm Billing Frequency appears before Subscription Type. Quote Line Type is shown. Ship-To appears only when Multiple Ship-to Accounts is Yes. Group and bundle order must follow the expected ordering rule.'),
('02 Legal approval email','Legal - Approval Request','Legal','doc6/image23.png',[
'Open the Opportunity linked to the test quote and confirm the intended quote is its Primary Quote.',
'Include one product with both Description and Rate Details populated, and another with Rate Details blank.',
'Use the Opportunity Legal Approval record and the Legal - Approval Request template to preview the output, or follow the agreed sandbox Legal approval process.',
'Compare every displayed product with the quote and inspect the Description text.'],
'Confirm Product Name, Price, Quantity, Billing Frequency and Description. Rate Details is part of Description, not a separate column. Blank Rate Details must not leave NULL text or a dangling separator.'),
('03 Credit Risk approval email','Credit Risk - Approval Request','Legal','doc6/image25.png',[
'Open the test Opportunity and verify its Primary Quote and saved quote lines.',
'Use the applicable Credit Risk Approval record and Credit Risk - Approval Request template.',
'Preview the email or inspect the approved sandbox test email. Use the Credit Risk workflow so the correct approval record is supplied.',
'Compare the output with the saved QLE and repeat the group-order and Description checks.'],
'Credit Risk uses the same table format as Legal. Confirm group order, product order, merged Description and Rate Details, and the conditional Ship-To column.'),
('04 Revenue approval email','Revenue - Approval Request','Revenue Approval','doc6/image41.png',[
'Use an agreed sandbox quote that meets the normal Revenue approval criteria. Record any discount or other input used to trigger that approval.',
'Save and calculate, then use the applicable Revenue Approval record and Revenue - Approval Request template.',
'Preview the output or inspect the approved sandbox test email.',
'Compare applied discounts, prices, quantities, groups and product order against the saved quote.'],
'Confirm Product Name, Applied Discount, Price, Quantity, Billing Frequency and Description. Description and Rate Details should be combined. Group and product ordering must match the expected rule.')]
for title,name,context,img,ss,expected in template_cases:
    page(title);p('Template: '+name,True);h('Steps to reproduce');steps(ss)
    pic(img,'Earlier test screenshot — '+name+'. Capture a fresh result after repeating these steps.',3.5)
    h('Expected result');p(expected)
    if context=='Quote Approval':p('Open item: the earlier test record includes Rate Details in Quote Approval Description. Confirm the agreed requirement before accepting Description-only output.')
    result()

page('05 Mixed term group order')
p('Run this case across Quote Approval, Legal, Credit Risk and Revenue.',True)
h('Test data')
p('Use a 72-month test quote beginning 1 October 2026. Give each group a distinct description marker so it is easy to identify in the output.')
table(['QLE group','Ramping','Start date','Term','Output position'],[
('G1','Yes','1 Oct 2026','24 months','2'),('G2','Yes','1 Oct 2028','12 months','3'),('G3','Yes','1 Oct 2029','36 months','4'),('G4','No','1 Oct 2026','72 months','1')],[.85,.8,1.45,1.25,2.15])
h('Steps to reproduce')
steps(['Create or use the agreed test quote and enter the four groups above in Edit Lines.',
'Add products to each group, including a bundle and standalone products. Save and calculate.',
'Reopen the quote and verify saved group dates, subscription terms and QLE sequence.',
'Preview each of the four approval templates using its appropriate approval record.',
'Match group markers and product rows to the source quote. Save the QLE and output screenshots together.'])
h('Expected result')
p('G4 → G1 → G2 → G3',True)
p('Term 1: 72 months\nTerm 2: 24 months\nTerm 3: 12 months\nTerm 4: 36 months')
p('The later 36-month group must remain after the earlier 12-month group. Sorting all groups by term length is not acceptable.')
result()

page('06 Grouping regression cases')
p('For each row, save the quote, verify its inputs and repeat all four template previews. Record one result per template.')
table(['Case','Steps to reproduce','Expected result'],[
('All ramped','Create three successive ramped groups with terms 12, 6 and 18 months.','Groups follow their start dates: 12, 6, 18.'),
('All non-ramped','Create three equal-date, equal-term groups. Put a Z-named group before an A-named group in QLE.','QLE group order is retained; names do not reorder the groups.'),
('One non-ramped and two ramped','Use a full-term group and two successive ramped groups. Give the full-term group the first start date.','Full-term group first, then ramped groups by start date.'),
('Two non-ramped and two ramped','Use two full-term groups starting with the quote and two successive ramped groups.','Full-term groups in QLE tie order, then ramped groups chronologically.'),
('No groups','Use an ungrouped quote with nonalphabetical product order.','One flat table with no Term or Group header; QLE row order retained.'),
('Later non-ramped','Where CPQ permits it, start a non-ramped group after an earlier ramped group.','Earlier start first, regardless of the ramping flag.')],[1.25,2.65,2.6])
pic('doc4/image2.png','Earlier ungrouped Legal email reference. No Term or Group header is shown.',1.4)
p('If a CPQ rule prevents a test setup, record Blocked and the rule message. Do not disable unrelated rules to force the scenario.')
result()

page('07 Product and bundle sequence')
h('Steps to reproduce')
steps(['Open the retained SALDEV-1420 Test Opp quote, or use an agreed copy. Verify its current contents before testing.',
'In Edit Lines, place Domestic Payments before Cross Border Payments. For eStore, place Usage before Implementation Fee.',
'Save and calculate. Preview each approval template.',
'Repeat with a supported nested bundle and a standalone product. Confirm sibling order and parent-child placement.'])
table(['Group','Expected sequence for the retained example'],[
('Term 1: 72 months','eStore → eStore (Usage) → eStore - Implementation Fee'),('Term 2: 72 months','Domestic Payments → Cross Border Payments')],[1.5,5])
h('What to check')
bullets(['Each expected line appears once in the correct group.',
'Parent products appear before their options. Nested options remain beneath their direct parent.',
'Sibling and standalone products follow QLE sequence rather than alphabetical order.',
'Prices, quantities and descriptions match the saved quote.'])
p('Find the quote from the SALDEV-1420 Test Opp Opportunity. Use it as a reference only if its contents still match the example.')
result()

page('08 Description and Rate Details')
p('Use Legal and Credit Risk for the Legal description checks. Repeat the combined-text check for Revenue.',True)
h('Steps to reproduce')
steps(['Create four test lines: both fields populated; Description only; Rate Details only; and both fields blank.',
'Save and calculate. Preview the appropriate approval email.',
'Read the complete Description cell and compare it with the source values.',
'Have the Legal reviewer confirm the text is suitable to copy into the contract document.'])
table(['Input','Expected result'],[
('Description and Rate Details','Both values appear once in one Description cell with readable separation.'),
('Description only','Description appears without a trailing separator or missing-value marker.'),
('Rate Details only','Rate Details appears without a leading separator.'),
('Both blank','Cell is blank; no NULL or placeholder text.')],[2,4.5])
h('Quote Approval confirmation')
p('The previous test record shows Description combined with Rate Details. Confirm with the story owner whether this is required for Quote Approval. Record the decision and compare the output with that agreed expectation. Leave this item open until the difference is resolved.')
result()

page('09 Amendment email')
h('Steps to reproduce')
steps(['Use the normal CPQ amendment process on an approved sandbox test contract.',
'Include unchanged lines and supported changed, new or cancelled lines. Save and calculate.',
'Inspect the calculated Quote Line Type values; do not manually enter them.',
'Preview all four approval templates. Reconcile the included lines with the quote and intended Opportunity Products.'])
pic('doc6/image12.png','Earlier amendment Quote Approval screenshot. Use calculated line types and current quote data for the new result.',3.6)
h('Expected result')
p('Lines with Quote Line Type = Amendment are excluded. Eligible changed lines remain, with correct groups, bundle placement, prices and quantities. No expected line is lost or duplicated.')
result()

page('10 Renewal Ship To and user access')
h('Renewal steps')
steps(['Create a renewal through the normal CPQ process using a suitable sandbox test contract.',
'Save and calculate. Record the expected renewal products, groups, prices and quantities.',
'Preview all four approval outputs and compare them with the saved quote.'])
p('Expected: eligible renewal lines appear and follow the required group and product order. If suitable contract data is unavailable, record the test as Blocked.')
h('Ship To steps')
steps(['Use a valid quote setup with Multiple Ship-to Accounts = Yes and populated Ship-To information. Preview the outputs.',
'Repeat with the flag set to No through the normal supported workflow.',
'Repeat with an ungrouped quote if that setup is supported.'])
p('Expected: Ship-To visibility follows the Yes/No flag. No groups alone does not determine whether the column is hidden. Automatic population of the flag is handled separately.')
h('Business user access steps')
steps(['Run a preview as a representative CPQ Sales user with the intended permission set.',
'Repeat with the relevant approver where the workflow uses another user.',
'Confirm the table is populated and no access error appears. Record the users tested.'])
result()

page('11 Test results and acceptance')
p('Complete this record after executing the cases. Add fresh screenshots beside each case or in an appendix, labelled with the case number, template and quote.')
table(['Case','Quote Approval','Legal','Credit Risk','Revenue'],[
('01–04 Template checks','Not run','Not run','Not run','Not run'),
('05 Mixed terms','Not run','Not run','Not run','Not run'),
('06 Grouping cases','Not run','Not run','Not run','Not run'),
('07 Bundle and line order','Not run','Not run','Not run','Not run'),
('08 Description and rates','Open item','Not run','Not run','Not run'),
('09 Amendment','Not run','Not run','Not run','Not run'),
('10 Renewal','Not run','Not run','Not run','Not run'),
('10 Ship-To and access','Not run','Not run','Not run','Not run')],[1.7,1.2,1.2,1.2,1.2])
h('Record for each executed case')
p('Tester: __________________________     Date: __________________\nQuote link: __________________________________________________\nApproval record or email subject: ________________________________\nExpected result: ______________________________________________\nActual result: _________________________________________________\nScreenshot reference: __________________________________________\nResult and defect reference: ____________________________________')
h('Acceptance')
p('Confirm all required cases have an explicit Pass, Fail or Blocked result. Resolve the Quote Approval Description decision and any failures before closure. Obtain Legal acceptance of the current email text and formatting.')
p('Business reviewer: __________________     Date: ________________\nLegal reviewer: _____________________     Date: ________________')

# Strip unused source relationships/media so the rewritten copy contains only its selected evidence.
used={el.get(qn('r:embed')) for el in doc._element.xpath('.//a:blip')}
for rid,rel in list(doc.part.rels.items()):
    if rel.reltype.endswith('/image') and rid not in used:doc.part.drop_rel(rid)
out=base/'SALDEV-1420_Story_Test_Execution_Revised.docx';doc.save(out)
print(out)
