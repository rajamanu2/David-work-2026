from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT=Path('outputs/business_templates_2026-09-08'); OUT.mkdir(parents=True,exist_ok=True)
def make(title, purpose):
 d=Document()
 for root in [d.styles.element,d.element]:
  for element in root.xpath('.//w:pBdr'): element.getparent().remove(element)
 s=d.sections[0]; s.page_width=Inches(8.5); s.page_height=Inches(11)
 s.top_margin=s.bottom_margin=Inches(.7); s.left_margin=s.right_margin=Inches(.75)
 for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
  st=d.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
 d.styles['Normal'].font.size=Pt(11); d.styles['Normal'].paragraph_format.space_after=Pt(7)
 d.styles['Title'].font.size=Pt(26); d.styles['Heading 1'].font.size=Pt(18); d.styles['Heading 2'].font.size=Pt(12)
 d.core_properties.author='Probo Medical'; d.core_properties.title=title
 d.add_paragraph('Probo Medical', 'Subtitle'); d.add_paragraph(title,'Title'); p(d,purpose)
 f=s.footer.paragraphs[0]; f.add_run('Probo Medical  |  Reusable template  |  ')
 fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)
 return d
def p(d,t): d.add_paragraph(t)
def h(d,t): d.add_heading(t,2)
def page(d,t): d.add_page_break(); d.add_heading(t,1)
def fields(d,items):
 for label,prompt in items: 
  z=d.add_paragraph(); z.add_run(label+'  ').bold=True; z.add_run('['+prompt+']')
def table(d,heads,rows,widths=None):
 t=d.add_table(rows=1,cols=len(heads)); t.style='Table Grid'
 for c,v in zip(t.rows[0].cells,heads): c.text=v
 tr=t.rows[0]._tr.get_or_add_trPr(); x=OxmlElement('w:tblHeader'); tr.append(x)
 for c in t.rows[0].cells:
  sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E7E6E6'); c._tc.get_or_add_tcPr().append(sh)
  for r in c.paragraphs[0].runs:r.bold=True
 for row in rows:
  cells=t.add_row().cells
  for c,v in zip(cells,row):c.text=v
 for row in t.rows:
  x=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(x)
  for i,c in enumerate(row.cells):
   if widths:c.width=Inches(widths[i])
   for z in c.paragraphs:
    z.paragraph_format.space_after=Pt(5); z.paragraph_format.space_before=Pt(5)
 return t

d=make('Business Requirements Gathering Template','Use this template to agree the business problem, desired outcomes, scope and acceptance criteria before solution design begins. Business teams supply the operational detail; leadership confirms priorities, ownership and funding decisions.')
fields(d,[('Initiative','Name and reference'),('Sponsor and business owner','Names and roles'),('Author and document owner','Name and role'),('Version and date','Version, date and next review'),('Status','Draft / In review / Approved'),('Business areas and locations','Teams, sites and countries affected')])
h(d,'How to use this template')
p(d,'Create a copy for each initiative. Replace bracketed prompts, duplicate requirement records as needed, and remove unused sections only after recording why they do not apply. Mark unknown items as open questions with an owner and due date. Illustrative examples are discussion aids and require business validation.')
table(d,['Stage','Required outcome'],[['Discovery','Agree the problem, stakeholders and current process.'],['Definition','Record measurable requirements and acceptance criteria.'],['Review','Resolve gaps and confirm scope, priorities and dependencies.'],['Approval','Baseline the requirements and authorize the next stage.']],[1.25,5.75])
page(d,'Leadership brief and scope')
fields(d,[('Decision requested','What leadership must approve and by when'),('Business problem','Who is affected, how often and the operational impact'),('Evidence','Reference interviews, records or measures supporting the problem'),('Desired outcome','Describe the business result without prescribing technology'),('Options considered','Keep current process, improve process or change system'),('Recommended option','Rationale, indicative cost, effort and dependencies')])
table(d,['Measure','Baseline','Target and deadline','Owner'],[['[e.g. Time to assign a request]','[Value and source]','[Value and date]','[Name]'],['[Quality or rework measure]','[Value and source]','[Value and date]','[Name]']])
h(d,'Scope boundaries')
fields(d,[('In scope','Processes, user groups, locations, data and channels'),('Out of scope','Explicit exclusions and reason'),('Constraints','Budget, dates, resources, technology or operating restrictions'),('Assumptions','What must be true and how it will be confirmed'),('Dependencies','Other teams, vendors, decisions or projects required')])
page(d,'Stakeholders and discovery workshop')
table(d,['Role','Named person','Contribution'],[['Sponsor','[Name]','Resolve priorities and approve business investment.'],['Process owner','[Name]','Own the end-to-end process and business acceptance.'],['Business users','[Names]','Explain actual tasks, exceptions and pain points.'],['System and data owners','[Names]','Assess feasibility, data definitions and access.'],['UAT and training leads','[Names]','Plan validation and adoption.']])
h(d,'Workshop agenda')
p(d,'Suggested 90 minutes: outcomes and scope 10; current process 25; future needs and exceptions 25; priorities and acceptance 20; actions and decisions 10. Capture participant names, date, evidence references and unresolved disagreements.')
h(d,'Questions for the business team')
p(d,'What starts the work? What inputs are required? Who does each step? Which systems or spreadsheets are used? Where does work wait or get repeated? What happens when information is missing? Who approves exceptions? How do you know the work is complete?')
h(d,'Questions for leadership')
p(d,'Which outcome matters most? What is the cost of doing nothing? Which date is a real constraint? What trade-offs are acceptable? Who owns the benefits? What evidence is needed to approve scope and confirm success?')
h(d,'Probo Medical discussion prompts')
p(d,'Where relevant, explore equipment identification, service intake, evaluation, repair approvals, parts availability, returns, shipment handoffs and operational reporting. Confirm which activities apply to the initiative and the names used by the business.')
page(d,'Current process and future process')
fields(d,[('Process boundaries','Trigger, starting point, ending point and customer'),('Volume and timing','Typical and peak volumes, frequency and time constraints')])
table(d,['Step and owner','Current action and system','Problem or evidence'],[['1  [Role]','[Input, action and output]','[Delay, rework or source]'],['2  [Role]','[Handoff and decision]','[Problem and impact]'],['3  [Role]','[Completion and communication]','[Problem and impact]']])
h(d,'Future process')
table(d,['Step and owner','Required future behavior','Control and exception'],[['1  [Role]','[Trigger and valid inputs]','[Missing or invalid inputs]'],['2  [Role]','[Action, routing and approval]','[Rejection or escalation]'],['3  [Role]','[Completion and notification]','[Verification and reopening]']])
fields(d,[('Business rules','Rule ID, condition, result, owner and exceptions'),('Handoffs','Sender, receiver, required information and acknowledgement'),('Process changes','Roles, steps or policies that must change')])
p(d,'Attach a process map when useful. Every decision must show its possible outcomes, and every exception must have an owner and an agreed resolution route.')
page(d,'Requirement record and coverage')
p(d,'Duplicate this record for each requirement. Use one testable business need per record and retain its ID throughout design, testing and change control.')
fields(d,[('Requirement ID and title','BR-001 and concise title'),('Business statement','As a role, I need a capability so that an outcome is achieved'),('Source and rationale','Stakeholder, evidence reference and business value'),('Priority and owner','Must / Should / Could / Deferred; accountable person'),('Rules and dependencies','Conditions, exceptions and related requirement IDs'),('Data and access','Inputs, outputs, required fields and permitted roles'),('Acceptance criteria','Given a condition, when an action occurs, then a measurable result'),('Status and decision','Proposed / Reviewed / Approved / Deferred; date and reason')])
h(d,'Illustrative requirement')
p(d,'EXAMPLE ONLY — BR-001: A service coordinator needs incomplete requests identified before handoff so the receiving team can act on them. Given a required input is absent, when the request is submitted, then the user receives a clear message naming the missing input and the request does not advance. The business owner must define required inputs and permitted exceptions.')
h(d,'Coverage checklist')
p(d,'Confirm functional behavior; data definitions and quality; role-based access; audit evidence; integrations and failure handling; reports and report visibility; measurable performance and availability; accessibility; retention ownership; support; and training. For each applicable area, create requirement IDs and acceptance criteria. Record agreed targets rather than assuming service levels.')
page(d,'Validation and business approval')
table(d,['Requirement','UAT scenario and expected result','Owner and evidence'],[['BR-001','[Normal flow and expected output]','[Tester, result, evidence link]'],['BR-001','[Missing input or rejection]','[Tester, result, evidence link]'],['[ID]','[Restricted user and integration failure]','[Tester, result, evidence link]']])
fields(d,[('UAT plan','Environment, roles, safe test data, dates and entry criteria'),('Exit criteria','Required coverage, acceptable defects and approval authority'),('Adoption and benefits','Training, communication, support and benefit review date')])
h(d,'Open questions and risks')
table(d,['ID and type','Issue and impact','Owner and due date'],[['Q-001','[Question, decision needed and affected requirement]','[Name and date]'],['R-001','[Risk, likelihood, impact and mitigation]','[Name and date]']])
h(d,'Approval and change control')
p(d,'Approval confirms the business requirements baseline. Record later changes with the change ID, reason, affected requirements, cost and schedule impact, decision maker and date. Revalidate affected acceptance criteria and work instructions.')
table(d,['Approver role','Decision and conditions','Name and date'],[['Business owner','[Approve / Rework; conditions]','[Name and date]'],['Sponsor','[Approve / Rework; conditions]','[Name and date]'],['System or delivery owner','[Feasibility and dependencies reviewed]','[Name and date]']])
d.save(OUT/'Probo_Medical_Business_Requirements.docx')

d=make('Process and Work Instruction Template','Use this template to document how an approved business process is performed consistently. The process owner defines the workflow and controls; staff use the work instructions to complete tasks and verify the result.')
fields(d,[('Process name and reference','Name and SOP reference'),('Process owner and author','Names and roles'),('Version and status','Version; Draft / In review / Approved'),('Effective date and review date','Dates confirmed by the process owner'),('Business areas and sites','Teams and locations covered'),('Related requirements','Approved BR IDs and source document'),('Audience','Roles expected to follow this procedure')])
h(d,'Purpose and boundaries')
fields(d,[('Purpose','Outcome this procedure delivers'),('Start and end','Trigger and verifiable completion condition'),('Included and excluded work','Activities and explicit exclusions')])
h(d,'Completion guidance')
p(d,'Replace all bracketed prompts with verified instructions. Duplicate the detailed task record for each task. Validate steps with someone who performs the work, include exception paths, and obtain the process owner’s approval before the effective date. Examples are illustrative and are not approved operating procedures.')
page(d,'Roles prerequisites and controls')
table(d,['Role','Responsibility','Backup or escalation'],[['Process owner','Approve procedure, controls and revisions.','[Named delegate]'],['Task performer','Complete steps and retain required evidence.','[Supervisor or team]'],['Approver or reviewer','Review evidence and accept or reject the outcome.','[Named alternate]'],['System support','Investigate technical failures and restore service.','[Support route]']])
h(d,'Before starting')
fields(d,[('Access','Application, environment and required role or permission'),('Training','Required training and how completion is verified'),('Inputs','Record identifiers, documents, authorizations and quality checks'),('Tools','Approved systems, equipment or reference materials'),('Operating conditions','Relevant handling, quality or site instructions'),('Stop conditions','When to pause and who must authorize continuation')])
h(d,'Control requirements')
p(d,'Identify who may create, edit, approve and complete the work. State any independent review, duplicate checks, mandatory evidence and record-retention ownership. Reference the approved policy for retention or handling requirements rather than inventing a period or rule.')
h(d,'Terms and references')
fields(d,[('Definitions','Business terms, abbreviations and status meanings'),('Reference documents','Policy or requirement ID, title, version and location')])
page(d,'Process overview and handoffs')
table(d,['Stage and owner','Action and output','Decision or handoff'],[['1  [Intake role]','[Receive trigger; validate inputs]','[Accept or return for correction]'],['2  [Assigned role]','[Perform task; record result]','[Route for review when required]'],['3  [Reviewer]','[Check evidence and decision criteria]','[Approve or return with reason]'],['4  [Completing role]','[Confirm output; notify receiver]','[Close and retain evidence]']])
h(d,'Decision rules')
table(d,['Condition','Required action','Authority'],[['[Required information absent]','[Return to sender; record reason]','[Role]'],['[Approval threshold met]','[Route to defined approver]','[Role]'],['[Exception outside procedure]','[Pause and escalate]','[Role]']])
fields(d,[('Handoff agreement','Sender, receiver, information required and acknowledgement'),('Timing','Expected completion times, working hours and escalation trigger'),('Completion definition','Output, status and evidence that prove the process is complete')])
p(d,'Validate the overview against the detailed task instructions. Each stage should have a named role, a clear output and a path for rejected or incomplete work.')
page(d,'Detailed work instruction record')
p(d,'Copy this page for each task. Write one action per numbered step using the actual screen labels or approved tool names. Screenshots should omit sensitive information and identify the relevant field or control.')
fields(d,[('Task ID and name','WI-001 and task title'),('Performer and trigger','Role and event that starts the task'),('Preconditions','Access, valid inputs and prior approvals')])
table(d,['Step','Action','Expected result and evidence'],[['1','[Open the correct application and locate the record using its unique identifier.]','[Correct environment and record confirmed.]'],['2','[Review required information against the approved checklist.]','[Inputs complete; otherwise follow exception ID.]'],['3','[Perform the task using verified field names and decision rules.]','[Expected values or output recorded.]'],['4','[Save or submit using the appropriate approved action.]','[Success indicator and resulting status confirmed.]'],['5','[Verify the result and complete the handoff.]','[Evidence reference and recipient acknowledgement.]']],[.45,3.55,3])
fields(d,[('Verification','What to check independently before marking the task complete'),('Exception references','Applicable exception IDs and escalation contacts'),('Evidence location','Record, approved repository and naming convention')])
h(d,'Illustrative task boundary')
p(d,'EXAMPLE ONLY — For service request intake, verify the request identifier and agreed required information before assigning work. The owner must confirm the actual application screens, field labels, routing rules and completion status before this becomes a usable instruction.')
page(d,'Exceptions validation and release')
table(d,['Exception','Immediate response','Owner and resolution evidence'],[['Missing or inconsistent input','Pause affected work and request correction.','[Owner; corrected input reference]'],['Duplicate or wrong record','Verify the identifier before proceeding.','[Owner; confirmed record reference]'],['System or integration failure','Record error and time; follow approved support route.','[Support; incident and recovery evidence]'],['Approval rejected','Record reason and return to responsible role.','[Owner; revised submission or closure]']])
h(d,'Procedure validation')
p(d,'A representative user should walk through the normal path, missing-input path, rejection path and a technical failure scenario in the agreed test environment. Record instruction IDs, tester, date, actual result, evidence and corrections. Repeat affected steps after changes.')
fields(d,[('Quality measures','Measure, target, frequency, owner and response to a breach'),('Training and release','Audience, training method, completion evidence and communication'),('Review triggers','Periodic review, policy changes, system changes or incidents')])
table(d,['Approval','Decision','Name and date'],[['Process owner','[Approve / Rework; conditions]','[Name and date]'],['Business representative','[Instructions validated]','[Name and date]'],['Control or system owner if applicable','[Relevant controls and steps verified]','[Name and date]']])
h(d,'Revision history')
table(d,['Version and date','Change and reason','Owner'],[['[Version and date]','[Affected steps, reason and approval reference]','[Name]']])
d.save(OUT/'Probo_Medical_Process_Work_Instructions.docx')
print(str(OUT.resolve()))
