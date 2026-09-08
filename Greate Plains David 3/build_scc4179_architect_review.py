from __future__ import annotations

import shutil
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Pt

from build_architect_review import (
    AMBER,
    BLUE,
    INK,
    LIGHT_AMBER,
    LIGHT_BLUE,
    LIGHT_RED,
    LIGHT_TEAL,
    MUTED,
    NAVY,
    RED,
    TEAL,
    add_body,
    add_figure,
    add_heading,
    arrow,
    draw_wrapped,
    load_font,
    rounded,
    set_run_font,
)
from build_scc3387_architect_review import (
    configure_document,
    normalize_table_margins,
    review_callout,
    review_table,
)


ROOT = Path(__file__).resolve().parent
REFERENCE = ROOT / "SCC-3387-architect-review.docx"
ASSET_DIR = ROOT / "doc-assets-scc4179"
OUTPUT = ROOT / "SCC-4179-architect-review.docx"
ASSET_DIR.mkdir(exist_ok=True)


def clear_document_body(document: Document) -> None:
    body = document._body._element
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)


def create_access_diagram(path: Path) -> None:
    image = Image.new("RGB", (1600, 1000), "#F7F9FC")
    draw = ImageDraw.Draw(image)
    title = load_font(42, bold=True)
    subtitle = load_font(22)
    heading = load_font(24, bold=True)
    body = load_font(18)
    small = load_font(16, bold=True)

    draw.text((70, 50), "SCC-4179 | CPNI least-privilege control path", font=title, fill="#0B2545")
    draw.text(
        (70, 110),
        "Profile and broad permission-set grants must be removed before the approved CPNI permission set is exclusive.",
        font=subtitle,
        fill="#5B6777",
    )

    boxes = [
        (
            (55, 225, 350, 605),
            "Profile exposure",
            ["Merge: 3 profiles", "UAT: 6 profiles", "Read or edit still granted"],
            "#FDECEC",
            "#C53030",
        ),
        (
            (430, 225, 755, 605),
            "Broad permission sets",
            ["Account and Contact - Full Access", "Heroku and Data Hub", "Both retain read access"],
            "#FDECEC",
            "#C53030",
        ),
        (
            (835, 225, 1160, 605),
            "Approved access",
            ["Contact CPNI Read/Write", "Both fields read/edit", "1 Merge + 2 UAT assignees"],
            "#E7F5F1",
            "#16836B",
        ),
        (
            (1240, 225, 1545, 605),
            "Release proof",
            ["Inventory Production", "Named-user UI/API UAT", "Post-deploy FLS read-back"],
            "#FFF6DE",
            "#B7791F",
        ),
    ]

    for box, box_title, lines, fill, outline in boxes:
        rounded(draw, box, fill, outline=outline, radius=22, width=4)
        draw.text((box[0] + 24, box[1] + 24), box_title, font=heading, fill="#0B2545")
        y = box[1] + 98
        for line in lines:
            draw.ellipse((box[0] + 26, y + 7, box[0] + 38, y + 19), fill=outline)
            wrapped = draw_wrapped(
                draw,
                (box[0] + 52, y),
                line,
                body,
                "#1F2937",
                box[2] - box[0] - 78,
                spacing=6,
            )
            y += max(76, 27 * len(wrapped) + 34)

    for start, end, label, color in (
        ((350, 415), (430, 415), "REMOVE", "#C53030"),
        ((755, 415), (835, 415), "RETAIN", "#16836B"),
        ((1160, 415), (1240, 415), "VERIFY", "#B7791F"),
    ):
        arrow(draw, start, end, color=color, width=6)
        draw.text((start[0] + 5, start[1] - 42), label, font=small, fill=color)

    rounded(draw, (70, 685, 1530, 905), "#EAF2FA", outline="#2E74B5", radius=22, width=3)
    draw.text((105, 720), "Architect reading", font=heading, fill="#0B2545")
    decision = (
        "The intended permission set is correctly configured, but access is not exclusive while profile and broad "
        "permission-set grants remain. Promotion requires a complete Production inventory, focused FLS removal, "
        "named-user UAT, and post-deployment permission read-back."
    )
    draw_wrapped(draw, (105, 780), decision, body, "#1F2937", 1360, spacing=8)
    image.save(path)


def add_metadata_table(document: Document) -> None:
    metadata = [
        ("Environments", "GreatPlainsMerge 00DEa00000GkAsHMAV | GreatPlainsUAT 00DEa00000FZlLBMA1"),
        ("Review date", "26 August 2026"),
        ("Evidence", "SOQL read-back | FieldPermissions and PermissionSetAssignment"),
        ("Decision", "Changes Requested - not ready for promotion"),
        ("Boundary", "No records, assignments, metadata, deployment, activation, or Production changes"),
    ]
    table = document.add_table(rows=len(metadata), cols=2)
    table.style = "Table Grid"
    from build_architect_review import LIGHT_BLUE, set_cell_text, set_table_geometry, shade_cell

    set_table_geometry(table, [1900, 7460])
    normalize_table_margins(table)
    for row, (label, value) in zip(table.rows, metadata):
        shade_cell(row.cells[0], LIGHT_BLUE)
        set_cell_text(row.cells[0], label, size=9.5, color=NAVY, bold=True)
        set_cell_text(row.cells[1], value, size=9.5)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def build_document() -> None:
    diagram_path = ASSET_DIR / "scc4179-cpni-control-path.png"
    create_access_diagram(diagram_path)

    shutil.copy2(REFERENCE, OUTPUT)
    document = Document(OUTPUT)
    clear_document_body(document)
    configure_document(document)

    # Page 1: executive decision.
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(2)
    set_run_font(paragraph.add_run("ARCHITECT REVIEW"), size=10, color=BLUE, bold=True)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(
        paragraph.add_run("SCC-4179 | Restrict CPNI Field Access"),
        size=25,
        color=NAVY,
        bold=True,
    )

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(16)
    set_run_font(
        paragraph.add_run("Great Plains sandboxes | Read-only least-privilege readiness assessment"),
        size=13,
        color=MUTED,
    )

    add_metadata_table(document)
    review_callout(
        document,
        "Architect decision",
        "Do not promote yet. The approved CPNI permission set is correct, but profile and broad permission-set grants still expose both fields, so permission-set-only access is not achieved.",
        fill=LIGHT_RED,
        accent=RED,
    )

    add_heading(document, "Acceptance outcome", 1)
    acceptance_rows = [
        ("1", "Named profiles", "Fail", "System Administrator, API Only, and Standard User still grant read and/or edit in both sandboxes."),
        ("2", "Other profiles", "Fail", "UAT also exposes both fields through Product, GPC-BI & Reporting, and GPC-Residential CRC Team."),
        ("3", "Non-target permission sets", "Fail", "Account and Contact - Full Access and Heroku and Data Hub Permissions retain read access."),
        ("4", "Approved permission set", "Pass", "Contact Encrypted CPNI Read/Write Access grants read/edit to both fields in Merge and UAT."),
        ("5", "Production and runtime proof", "Not run", "Production was not inventoried and no named-user UI/API UAT or deployment was performed."),
    ]
    review_table(
        document,
        ["#", "Scope", "Result", "Evidence"],
        acceptance_rows,
        [500, 2150, 1450, 5260],
        status_column=2,
    )

    # Page 2: solution visual.
    document.add_page_break()
    add_heading(document, "Solution and control path", 1)
    add_body(
        document,
        "The solution is declarative FLS: remove both CPNI fields from every profile and non-approved permission set, retain read/edit only in the approved permission set, then prove exclusivity with named users and permission read-back.",
    )
    add_figure(
        document,
        diagram_path,
        "Figure 1. SCC-4179 CPNI least-privilege control path",
        "Diagram showing current profile exposure and broad permission-set exposure being removed, the approved Contact CPNI permission set retained, and Production inventory plus named-user verification required before release.",
        width=6.45,
    )
    review_callout(
        document,
        "Design reading",
        "Green is the intended grant. Red access paths must be removed. Amber is the release evidence gate before promotion.",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )

    # Page 3: environment evidence.
    document.add_page_break()
    add_heading(document, "Current CPNI access evidence", 1)
    add_heading(document, "GreatPlainsMerge", 2)
    merge_rows = [
        ("System Administrator", "Read", "Read / Edit", "Fail"),
        ("System Administrator - API Only", "Read / Edit", "Read / Edit", "Fail"),
        ("Standard User", "Read", "Read", "Fail"),
        ("Account and Contact - Full Access", "Read", "Read", "Fail"),
        ("Heroku and Data Hub Permissions", "Read", "Read", "Fail"),
        ("Contact Encrypted CPNI Read/Write Access", "Read / Edit", "Read / Edit", "Pass - 1 assignment"),
    ]
    review_table(
        document,
        ["Access path", "CPNI Password", "CPNI Pin", "Acceptance"],
        merge_rows,
        [3600, 1800, 1800, 2160],
        status_column=3,
    )

    add_heading(document, "GreatPlainsUAT", 2)
    uat_rows = [
        ("System Administrator + API Only", "Read / Edit", "Read / Edit", "Fail - 2 profiles"),
        ("Standard, Product, GPC-BI, GPC-Residential", "Read", "Read", "Fail - 4 profiles"),
        ("Account and Contact - Full Access", "Read", "Read", "Fail - grant path"),
        ("Heroku and Data Hub Permissions", "Read", "Read", "Fail - 1 assignment"),
        ("Contact Encrypted CPNI Read/Write Access", "Read / Edit", "Read / Edit", "Pass - 2 assignments"),
    ]
    review_table(
        document,
        ["Access path", "CPNI Password", "CPNI Pin", "Acceptance"],
        uat_rows,
        [3600, 1800, 1800, 2160],
        status_column=3,
    )
    review_callout(
        document,
        "Exclusivity blocker",
        "The target permission set works, but it is not the only grant path. Every profile and non-target permission set must lose both read and edit before SCC-4179 can pass.",
        fill=LIGHT_RED,
        accent=RED,
    )

    # Page 4: implementation and evidence boundary.
    document.add_page_break()
    add_heading(document, "Implementation and release boundary", 1)
    component_rows = [
        ("Contact fields", "CPNI_Pin__c; CPNI_Password__c", "Reference only - no schema change"),
        ("Named profiles", "System Administrator; API Only; Standard User", "Remove read and edit for both fields"),
        ("Additional UAT profiles", "Product; GPC-BI & Reporting; GPC-Residential CRC Team", "Remove read access; inventory Production equivalents"),
        ("Broad permission sets", "Account_and_Contact_Full_Access; Heroku_and_Data_Hub_Permissions", "Remove both field grants"),
        ("Approved permission set", "Contact_CPNI_Read_Write_Access", "Preserve read=true and edit=true"),
        ("Code and automation", "Apex, Flow, validation rules, layouts", "No change"),
    ]
    review_table(
        document,
        ["Component type", "Focused scope", "Required action"],
        component_rows,
        [2100, 3900, 3360],
    )

    add_heading(document, "Release sequence", 2)
    sequence_rows = [
        ("Pre-deployment", "Retrieve Production FLS; inventory all grant paths and target-set assignees; assess integrations; run focused check-only validation; obtain Security and release approval."),
        ("Deployment", "Promote only approved FLS removals to UAT first; retain the target permission set; deploy to Production only after named-user UAT and an approved change window."),
        ("Post-deployment", "Read back FieldPermissions and assignments; run negative and positive UI/API tests; review Setup Audit Trail; capture job ID, scope, results, and sign-offs."),
        ("Rollback", "Restore the approved metadata backup only if a critical dependency fails; remove any temporary emergency permission assignment after resolution."),
    ]
    review_table(document, ["Phase", "Required evidence"], sequence_rows, [1900, 7460])
    review_callout(
        document,
        "Production gate",
        "Do not use the sandbox component list as a complete Production package. Production was not inspected; its full FLS and assignment inventory is mandatory before validation or deployment.",
        fill=LIGHT_AMBER,
        accent=AMBER,
    )

    # Page 5: UAT, approvals, and story response.
    document.add_page_break()
    add_heading(document, "UAT and approval gates", 1)
    uat_test_rows = [
        ("Negative UI", "Named users without the target set cannot see either field on View or Edit."),
        ("Negative API", "The same users cannot query or update either field through the approved API test path."),
        ("Positive access", "A named user with only the target set can view and edit both fields using synthetic UAT data."),
        ("Exclusivity", "FieldPermissions read-back shows the target permission set as the only remaining read/edit grant path."),
        ("Regression", "Approved integration and reporting owners confirm expected behavior after profile and broad-set removal."),
    ]
    review_table(document, ["Test", "Pass evidence"], uat_test_rows, [2100, 7260])

    add_heading(document, "Required approval gates", 2)
    gate_rows = [
        ("1", "Inventory Production", "Salesforce Admin / Security", "Complete profile, permission-set, assignment, integration, and reporting impact inventory."),
        ("2", "Approve authorized users", "Data Owner / Security", "Final target-permission-set assignee list and emergency-access process."),
        ("3", "Validate focused package", "Release Manager", "Check-only job ID, exact components, test result, and explicit deployment approval."),
        ("4", "Run named-user UAT", "QA / Business Owner", "UI/API negative and positive tests with record IDs, timestamps, and screenshots."),
        ("5", "Promote and read back", "DevOps / Salesforce Admin", "Deployment receipt, FLS/assignment read-back, audit trail, and sign-offs."),
    ]
    review_table(
        document,
        ["#", "Required action", "Owner", "Evidence needed"],
        gate_rows,
        [500, 2400, 2200, 4260],
    )

    add_heading(document, "Ready-to-paste architect response", 2)
    response = (
        "Architect review completed for SCC-4179 in GreatPlainsMerge and GreatPlainsUAT. Changes Requested: "
        "Contact_CPNI_Read_Write_Access correctly grants read/edit to CPNI Pin and CPNI Password, but permission-set-only "
        "access is not achieved. Merge still grants access through three profiles and two broad permission sets; UAT "
        "grants access through six profiles and two broad permission-set definitions. Production was not inspected. "
        "Complete the Production inventory, remove both field grants from every profile and non-target permission set, "
        "retain the approved permission set, run focused check-only validation and named-user UI/API UAT, then capture "
        "post-deployment FLS and assignment read-back. No records, assignments, metadata, activation, deployment, or "
        "Production changes were made during this review."
    )
    review_callout(document, "Story comment", response, fill=LIGHT_BLUE, accent=BLUE)

    document.core_properties.title = "SCC-4179 Restrict CPNI Field Access Architect Review"
    document.core_properties.subject = "Read-only least-privilege readiness assessment"
    document.core_properties.author = "Solution Architecture Review"
    document.core_properties.keywords = (
        "Salesforce, Great Plains, SCC-4179, CPNI, FLS, permission set, least privilege, architect review"
    )
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
