from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

folder=Path('outputs/business_templates_2026-09-08')
src=folder/'Probo_Medical_Combined_Requirements_and_Fillable_Workbook.docx'
d=Document(src)
before=d.element.xpath('.//w:sdt//w:t/text()')
for p in d.paragraphs:
 if p.text=='PROBO MEDICAL':
  p.runs[0].text='PROBO MEDICAL  |  AutoFast'
seen=set()
for s in d.sections:
 for header in [s.header,s.first_page_header]:
  if id(header._element) in seen:continue
  seen.add(id(header._element))
  for p in header.paragraphs:
   if p.text:
    p.text='PROBO MEDICAL  |  AutoFast                         BUSINESS REQUIREMENTS'
    for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0)
 for footer in [s.footer,s.first_page_footer]:
  if id(footer._element) in seen:continue
  seen.add(id(footer._element))
  for p in footer.paragraphs:
   p.clear()
   p.paragraph_format.tab_stops.add_tab_stop(s.page_width-s.left_margin-s.right_margin,WD_TAB_ALIGNMENT.RIGHT)
   p.add_run('Probo Medical  |  AutoFast  |  Requirements workbook').font.size=Pt(9)
   p.add_run('\t')
   x=OxmlElement('w:fldSimple');x.set(qn('w:instr'),'PAGE');p._p.append(x)
d.core_properties.title='Probo Medical and AutoFast Business Requirements and Fillable Workbook'
d.core_properties.author='Probo Medical | AutoFast'
assert d.element.xpath('.//w:sdt//w:t/text()')==before
assert len(d.element.xpath('.//w:sdt'))==171
out=folder/'Probo_Medical_AutoFast_Combined_Requirements_Workbook.docx'
d.save(out)
print(out.resolve())
