from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_architect_review import (
    BLUE,
    INK,
    LIGHT_BLUE,
    LIGHT_GRAY,
    LIGHT_RED,
    LIGHT_TEAL,
    MUTED,
    NAVY,
    RED,
    TEAL,
    add_body,
    add_callout,
    add_figure,
    add_heading,
    add_table,
    arrow,
    configure_header_footer,
    draw_wrapped,
    load_font,
    rgb,
    rounded,
    set_cell_text,
    set_run_font,
    set_table_geometry,
    shade_cell,
)


ROOT = Path(__file__).resolve().parent
ASSET_DIR = ROOT / "doc-assets-scc3380"
OUTPUT = ROOT / "SCC-3380-architect-review.docx"
ASSET_DIR.mkdir(exist_ok=True)


def create_relationship_diagram(path: Path) -> None:
    image = Image.new("RGB", (1600, 1050), "#F7F9FC")
    draw = ImageDraw.Draw(image)
    title = load_font(42, bold=True)
    subtitle = load_font(22)
    heading = load_font(25, bold=True)
    body = load_font(19)
    small = load_font(17)

    draw.text((65, 48), "SCC-3380 | Incident and Trouble Ticket control path", font=title, fill="#0B2545")
    draw.text(
        (67, 108),
        "The source automation is active, but the guardrails, NOC assignment, and UAT prerequisites are incomplete.",
        font=subtitle,
        fill="#5B6777",
    )

    boxes = [
        ((55, 215, 375, 555), "NOC user", [
            "Permission set exists",
            "0 assignments",
            "0 active users identified by NOC title/department",
        ], "#FFF6DE", "#B7791F"),
        ((455, 215, 790, 555), "Incident", [
            "Close confirmation checkbox",
            "Incident_After_Update active/latest",
            "Priority sync and confirmed bulk close",
        ], "#E7F5F1", "#16836B"),
        ((870, 215, 1215, 555), "Case Related Issue", [
            "Before-save type check active",
            "After-save stamps Related Outage",
            "Delete flow clears Related Outage",
        ], "#E7F5F1", "#16836B"),
        ((1295, 215, 1550, 555), "Trouble Ticket", [
            "10 source records",
            "Parent close flow active",
            "Individual case details retained",
        ], "#EAF2FA", "#2E74B5"),
    ]

    for box, box_title, lines, fill, outline in boxes:
        rounded(draw, box, fill, outline=outline, radius=22, width=4)
        draw.text((box[0] + 24, box[1] + 24), box_title, font=heading, fill="#0B2545")
        y = box[1] + 92
        for line in lines:
            draw.ellipse((box[0] + 26, y + 7, box[0] + 38, y + 19), fill=outline)
            draw_wrapped(draw, (box[0] + 52, y), line, body, "#1F2937", box[2] - box[0] - 78, spacing=5)
            y += 72

    arrow(draw, (375, 385), (455, 385), color="#B7791F", width=6)
    arrow(draw, (790, 385), (870, 385), color="#16836B", width=6)
    arrow(draw, (1215, 385), (1295, 385), color="#2E74B5", width=6)
    draw.text((386, 340), "access", font=small, fill="#B7791F")
    draw.text((802, 340), "relates", font=small, fill="#16836B")
    draw.text((1225, 340), "updates", font=small, fill="#2E74B5")

    rounded(draw, (55, 650, 760, 960), "#FDECEC", outline="#C53030", radius=22, width=4)
    draw.text((85, 680), "Source guardrail blocker", font=heading, fill="#0B2545")
    guardrails = [
        "Three story-listed validation rules do not exist.",
        "Self-parent, two-node cycles, and non-Trouble Ticket children are not fully prevented.",
        "Parent and Related Outage updates are not filtered to Trouble Ticket children.",
    ]
    y = 742
    for line in guardrails:
        draw.ellipse((87, y + 7, 99, y + 19), fill="#C53030")
        draw_wrapped(draw, (115, y), line, body, "#1F2937", 590, spacing=5)
        y += 72

    rounded(draw, (840, 650, 1550, 960), "#FDECEC", outline="#C53030", radius=22, width=4)
    draw.text((870, 680), "UAT check-only blocker", font=heading, fill="#0B2545")
    uat_lines = [
        "Job 0AfEa00000bbQnpKAE failed: 2/15 components passed.",
        "UAT lacks the SCC-3384/3386 fields, Trouble Ticket record type, and relationship required by this package.",
        "Profiles, FlexiPage, and the pre-existing Case email action also fail validation.",
    ]
    y = 742
    for line in uat_lines:
        draw.ellipse((872, y + 7, 884, y + 19), fill="#C53030")
        draw_wrapped(draw, (900, y), line, body, "#1F2937", 600, spacing=5)
        y += 72

    image.save(path)


def configure_document(document: Document) -> None:
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        configure_header_footer(section)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 12, 6),
        ("Heading 2", 13, BLUE, 10, 5),
        ("Heading 3", 12, NAVY, 8, 4),
    ):
        style = document.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def build_document() -> None:
    diagram_path = ASSET_DIR / "scc3380-relationship-control-path.png"
    create_relationship_diagram(diagram_path)

    document = Document()
    configure_document(document)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(2)
    set_run_font(paragraph.add_run("ARCHITECT REVIEW"), size=10, color=BLUE, bold=True)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(paragraph.add_run("SCC-3380 | Incident-Case Relationships"), size=25, color=NAVY, bold=True)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(16)
    set_run_font(
        paragraph.add_run("Great Plains Merge to UAT | Read-only retrieval and check-only readiness assessment"),
        size=13,
        color=MUTED,
    )

    metadata = [
        ("Source", "GreatPlainsMerge (Sandbox) | Org 00DEa00000GkAsHMAV"),
        ("Target", "GreatPlainsUAT (Sandbox) | Org 00DEa00000FZlLBMA1"),
        ("Review date", "26 August 2026"),
        ("Retrieval", "Succeeded with 3 missing story items | Job 09SEa00000icZA2MAM"),
        ("Check-only", "Failed | Job 0AfEa00000bbQnpKAE | 2/15 passed"),
        ("Decision", "NO-GO - Changes Requested before promotion"),
        ("Boundary", "No records, deployment, activation, or org metadata were changed"),
    ]
    table = document.add_table(rows=len(metadata), cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [1900, 7460])
    for row, (label, value) in zip(table.rows, metadata):
        shade_cell(row.cells[0], LIGHT_BLUE)
        set_cell_text(row.cells[0], label, size=9.4, color=NAVY, bold=True)
        set_cell_text(row.cells[1], value, size=9.4)
    document.add_paragraph().paragraph_format.space_after = Pt(2)

    add_callout(
        document,
        "Architect decision",
        "Do not promote SCC-3380. The five source flows are active, but three required guardrails are absent, the NOC permission set is unassigned, and the UAT dry run failed 13 of 15 components.",
        fill=LIGHT_RED,
        accent=RED,
    )

    add_heading(document, "Acceptance outcome", 1)
    acceptance_rows = [
        ("1", "Incident/Case linking", "Partial pass", "Case Related Issue layouts and active stamp/clear flows exist in source; UAT lacks Related_Outage__c and other prerequisites."),
        ("2", "Type and circular controls", "Fail", "Only the non-Trouble Ticket Incident-link check exists. The three listed validation rules for parent/type/circular protection are absent."),
        ("3", "Incident field sync", "Partial pass", "Active Incident flow maps standard and custom Priority. Runtime was not executed; UAT validation fails on missing Incident and Case fields."),
        ("4", "Confirmed bulk close", "Partial pass", "Source logic requires Closed transition plus checkbox. Runtime was not executed and target validation failed."),
        ("5", "Parent close and unlink", "Partial fail", "Parent close and unlink logic exist, but child updates are not constrained to Trouble Ticket record type and parent guardrails are absent."),
        ("6", "Reporting", "Fail in UAT", "Source report type exists; UAT has no Trouble_Tickets__r relationship, so the report type cannot validate."),
    ]
    add_table(document, ["#", "Scope", "Result", "Evidence"], acceptance_rows, [500, 2050, 1450, 5360], status_column=2)

    document.add_page_break()
    add_heading(document, "Solution and control path", 1)
    add_body(document, "The intended relationship model is visible in source, but access, guardrails, and target prerequisites break the end-to-end promotion path.")
    add_figure(
        document,
        diagram_path,
        "Figure 1. SCC-3380 Incident and Trouble Ticket control path",
        "Diagram showing an NOC user reaching an Incident, Case Related Issue, and Trouble Ticket. Active source automation is green or blue. Missing validation rules, zero permission-set assignments, and the failed UAT check-only package are red blockers.",
        width=6.45,
    )
    add_callout(document, "Design reading", "Green/blue is configured in source. Amber needs ownership. Red blocks approval.", fill=LIGHT_BLUE, accent=BLUE)

    document.add_page_break()
    add_heading(document, "Detailed source findings", 1)
    add_heading(document, "1. Automation and controls", 2)
    automation_rows = [
        ("SCC_3380_Case_Related_Issue_Before_Save", "Active/latest", "Blocks a non-Trouble Ticket Case from being linked to an Incident."),
        ("SCC_3380_Case_Related_Issue_After_Save", "Active/latest", "On create, stamps Case.Related_Outage__c for an Incident link."),
        ("SCC_3380_Case_Related_Issue_After_Delete", "Active/latest", "Before delete, clears Related Outage only when it matches the removed Incident."),
        ("Incident_After_Update", "Active/latest", "Syncs Priority and performs checkbox-confirmed bulk close through both relationship paths."),
        ("Case_After_Trigger_Flow", "Active/latest", "Closes open child Cases when a Trouble Ticket parent closes; also retains the existing email action."),
    ]
    add_table(document, ["Automation", "State", "Architect evidence"], automation_rows, [3600, 1600, 4160], status_column=1)

    add_heading(document, "2. Missing guardrails", 2)
    guardrail_rows = [
        ("Case.SCC_3380_Prevent_Circular_Parent_Case", "Absent", "No SCC-3380 rule exists to block self-parenting or A-to-B-to-A cycles."),
        ("Case.SCC_3380_Child_Must_Be_Trouble_Ticket", "Absent", "A non-Trouble Ticket can remain a child; the parent-close flow updates all open children."),
        ("CaseRelatedIssue.SCC_3380_Case_Must_Be_Trouble_Ticket", "Absent", "The before-save Flow provides a narrower Incident-link control, not the listed validation rule."),
        ("Record-type filters", "Incomplete", "Incident Related Outage updates and parent-child closure do not filter target Cases by Trouble Ticket record type."),
    ]
    add_table(document, ["Control", "State", "Risk"], guardrail_rows, [3500, 1400, 4460], status_column=1)
    add_callout(
        document,
        "Guardrail blocker",
        "The story's circular and child-type acceptance criteria are not satisfied by the retrieved source package.",
        fill=LIGHT_RED,
        accent=RED,
    )

    document.add_page_break()
    add_heading(document, "Check-only deployment evidence", 1)
    add_heading(document, "GreatPlainsUAT validation", 2)
    validation_rows = [
        ("Job", "0AfEa00000bbQnpKAE", "checkOnly=true; rollbackOnError=true; NoTestRun"),
        ("Package", "15 components / 23 files", "2 component successes; 13 component failures"),
        ("Passed", "2", "Incident.Close_Related_Trouble_Tickets__c and SCC_3380_Case_Related_Issue_Before_Save"),
        ("Tests", "0", "No Apex in scope; no runtime story tests were executed"),
    ]
    add_table(document, ["Metric", "Result", "Evidence"], validation_rows, [1800, 2400, 5160])

    failure_rows = [
        ("Missing SCC prerequisites", "7", "Report type; Case-Trouble Ticket; Case-Case Layout; Incident_After_Update; After Save; After Delete; plus permission-set dependency on Related_Outage__c."),
        ("Incident layout", "1", "Outage Incident Layout references missing Incident.Affected_Area_Region__c."),
        ("Profiles", "3", "Admin/API profiles use unknown ArchiveArticles permission; Standard references missing Case.Trouble_Ticket record type."),
        ("Lightning page", "1", "Incident_Record_Page1 has an unsupported shouldDisplayFiltersOnRight property."),
        ("Existing Case flow", "1", "Case_After_Trigger_Flow emailTemplateName input type does not match the assigned value in UAT validation."),
        ("Story omissions", "3 not packaged", "All three listed SCC-3380 validation rules are missing in the source org."),
    ]
    add_table(document, ["Failure group", "Count", "Exact blocker"], failure_rows, [2300, 1000, 6060], status_column=1)

    add_heading(document, "Read-only runtime evidence", 2)
    runtime_rows = [
        ("Trouble Ticket Cases", "10", "Source org aggregate only; no records changed."),
        ("Case Related Issues", "1", "The observed relationship type is Incident."),
        ("Cases with Related Outage", "1", "Shows the source relationship is in use, but does not prove all acceptance paths."),
        ("Permission-set assignments", "0", "SCC_3380_NOC_Incident_Case_Relationship is not assigned."),
        ("Identified active NOC users", "0", "No active user matched NOC in Title or Department."),
    ]
    add_table(document, ["Evidence", "Count", "Interpretation"], runtime_rows, [2600, 900, 5860])

    document.add_page_break()
    add_heading(document, "Required approval gates", 1)
    gate_rows = [
        ("1", "Implement the three guardrails", "Admin / Architect", "Deploy active controls for self/circular parent chains, Trouble Ticket-only children, and Trouble Ticket-only Incident links."),
        ("2", "Sequence prerequisite stories", "Release manager", "Move SCC-3384/3386 record type, Related Outage, Service Address, Priority, Incident fields, and relationship metadata before SCC-3380."),
        ("3", "Constrain update scope", "Flow owner / Architect", "Filter all bulk updates to Trouble Ticket Cases and decide whether Case Related Issue edits must resync Related Outage."),
        ("4", "Repair deploy artifacts", "Salesforce admin", "Replace broad profiles with focused access, remove unsupported FlexiPage drift, and fix the Case email-action input type."),
        ("5", "Identify the NOC persona", "Product owner / Security", "Provide the named NOC user/profile and assign the focused permission set; prove field, tab, layout, and report access."),
        ("6", "Revalidate and execute UAT", "QA / Release manager", "Require a clean check-only result, then execute the five story tests with IDs/screenshots before any promotion approval."),
    ]
    add_table(document, ["Gate", "Required action", "Owner", "Evidence needed"], gate_rows, [600, 2700, 2100, 3960])

    add_heading(document, "Ready-to-paste architect response", 1)
    response = (
        "Architect review completed for SCC-3380 using GreatPlainsMerge as source and GreatPlainsUAT as the check-only target. Changes Requested / NO-GO: all five retrieved flows are active/latest in source, but the three story-listed validation rules are absent, so circular parent relationships and Trouble Ticket-only child enforcement are not fully implemented. The NOC permission set has zero assignments and no active NOC persona was identified. Check-only job 0AfEa00000bbQnpKAE failed with 13 of 15 components; only the Incident confirmation checkbox and the before-save Case Related Issue flow passed. The failures include missing SCC-3384/3386 fields and record type, profile drift, an unsupported Incident page property, and the existing Case email-action input mismatch. Complete the six gates above and require a clean validation plus the five named UAT tests before promotion. No records, metadata, activation, or deployment changes were made."
    )
    add_callout(document, "Story comment", response, fill=LIGHT_BLUE, accent=BLUE)

    document.core_properties.title = "SCC-3380 Incident-Case Relationships Architect Review"
    document.core_properties.subject = "Read-only retrieval and check-only readiness assessment"
    document.core_properties.author = "Solution Architecture Review"
    document.core_properties.keywords = "Salesforce, Great Plains, SCC-3380, Incident, Case, architect review, dry run"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
