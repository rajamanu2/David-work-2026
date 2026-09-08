from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
import re,json,zipfile

# Reuse the current document's established table and font treatment.
exec(Path('scratch/build_presentation_brd.py').read_text(encoding='utf-8').split("p('PROBO MEDICAL')")[0])
folder=Path('outputs/business_templates_2026-09-08')
src=folder/'Probo_Medical_AutoFast_Combined_Requirements_Workbook.docx'
d=Document(src)
prior_fields=d.element.xpath('.//w:sdt//w:t/text()')
old_children=list(d.element.body)
def field(label,hint,space=16):
 h(label)
 z=d.add_paragraph(hint);z.paragraph_format.space_after=Pt(5)
 z.runs[0].italic=True;z.runs[0].font.size=Pt(10)
 z=d.add_paragraph('[Click here to enter your response]');z.paragraph_format.space_after=Pt(space)
def section(number,title,intro):
 d.add_page_break();z=d.add_paragraph(f'{number:02d}  /  PROJECT AND DELIVERY DETAILS');z.paragraph_format.space_after=Pt(6)
 z.runs[0].font.size=Pt(9);z.runs[0].bold=True
 d.add_heading(title,1);p(intro)

section(12,'Project overview and Business Ask','Complete these project intake fields alongside the earlier workbook. Use the same initiative and requirement IDs throughout so each detail can be traced to the business request.')
p('Project or initiative ID  [Enter reference]\nProject name  [Enter name]\nBusiness sponsor  [Enter name]     Business owner  [Enter name]')
field('Project overview','Summarize the project, its purpose, intended business outcome and the teams or processes affected.',22)
field('Business Ask','State exactly what change, capability or decision the business is requesting and why it is needed.',22)
field('Additional project details','Record relevant background, locations, dependencies, assumptions and links to supporting documents.',22)
field('Business notes','Capture business context, preferences, constraints and points that should guide the delivery team.',22)
p('Related requirement IDs  [Enter references]\nPrepared by and date  [Enter name and date]')

section(13,'Working team and stakeholder questions','Identify the people doing the work and keep business and IT questions visible until there is an agreed answer. Add rows when additional people or questions are identified.')
h('Working team')
table(['Name and team','Role and responsibility','Contact or availability'],[['[Enter name and team]','[Enter delivery responsibility]','[Enter contact or availability]'],['[Enter name and team]','[Enter delivery responsibility]','[Enter contact or availability]'],['[Enter name and team]','[Enter delivery responsibility]','[Enter contact or availability]']],[1.9,3.2,1.8])
h('Stakeholder questions and responses')
table(['Stakeholder and question','Response or decision','Owner and due date'],[['[Enter stakeholder and question]','[Enter answer or Pending]','[Enter owner and date]'],['[Enter stakeholder and question]','[Enter answer or Pending]','[Enter owner and date]']],[2.75,2.4,1.75])
h('IT comments and questions')
table(['Comment or question','Response and impact','Owner and status'],[['[Enter IT comment or question]','[Enter answer and requirement impact]','[Enter owner and status]'],['[Enter IT comment or question]','[Enter answer and requirement impact]','[Enter owner and status]']],[2.6,2.6,1.7])
p('Decision authority  [Enter name or role]\nUnresolved items requiring escalation  [Enter references and next action]')

section(14,'Feature and dependency assessment','Complete this page for each feature or requirement that needs a business and IT review. Duplicate it as needed and link to the detailed requirement record and acceptance tests.')
p('Requirement or feature ID  [Enter reference]\nFeature owner  [Enter name]')
field('Description and feature','Describe the requested capability, expected behavior and the boundary of this feature.',12)
field('User story','As a role, I need a capability so that a business outcome is achieved.',12)
field('Acceptance criteria','Record the measurable conditions for acceptance or link to the approved criteria for this requirement.',12)
h('Impacted application')
table(['Application and environment','Impact and owner'],[['[Enter application and environment]','[Enter change or impact and owner]'],['[Enter application or Not applicable]','[Enter impact and owner or reason]']],[2.8,4.1])
h('Component and Record Dependency')
table(['Dependency type','Item and relationship','Required action and owner'],[['Component','[Enter component and what depends on it]','[Enter sequence, action and owner]'],['Record or data','[Enter object or record type and relationship]','[Enter prerequisite, action and owner]']],[1.2,3,2.7])
p('Use approved record references and document any required setup or migration. Record Not applicable with a reason where no dependency exists.')

section(15,'Priority timeframe and total estimate','Complete the estimate with the delivery owner. Use a consistent effort unit and state the scope covered. Leave figures as Not yet estimated until an owner provides them.')
p('Estimate scope and linked requirements  [Enter project, feature or requirement IDs]\nPriority  [Enter Must, Should, Could or Deferred]\nPriority rationale  [Enter business justification]')
h('Timeframe')
table(['Milestone','Target date or range','Dependency or reason'],[['Requested start','[Enter date or range]','[Enter dependency]'],['Business need by','[Enter date or range]','[Enter reason for deadline]'],['Planned delivery','[Enter date or range]','[Enter assumptions or Not agreed]']],[1.6,2.1,3.2])
h('Estimate breakdown')
table(['Work item','Estimated effort','Owner or basis'],[['Analysis and design','[Enter amount and unit]','[Enter estimator or basis]'],['Build or configuration','[Enter amount and unit]','[Enter estimator or basis]'],['Testing and UAT support','[Enter amount and unit]','[Enter estimator or basis]'],['Release and training','[Enter amount and unit]','[Enter estimator or basis]'],['Contingency if included','[Enter amount and unit]','[Enter rationale]']],[2.3,2.1,2.5])
h('Total Estimated')
p('Total estimated effort  [Enter total and unit]\nEstimated elapsed duration  [Enter duration and unit]\nEstimated cost and currency  [Enter amount and currency or Not applicable]')
p('Estimate basis assumptions and exclusions  [Enter details]\nEstimated by and date  [Enter name and date]\nReview status and approval reference  [Enter status and reference]')
p('Sum the effort rows manually using the same unit. Effort and elapsed duration differ when tasks overlap. Confirm whether contingency is included; do not count it twice.')

# Keep the estimate guidance with its totals without shrinking the text.
estimate_page=False
for para in list(d.paragraphs):
 if para.text=='15  /  PROJECT AND DELIVERY DETAILS':estimate_page=True
 if not estimate_page:continue
 if not para.text and not para._p.xpath('.//w:br | .//w:sectPr'):
  para._p.getparent().remove(para._p)
 elif para.style.name=='Heading 2':
  para.paragraph_format.space_before=Pt(8)
 elif para.style.name=='Normal':
  para.paragraph_format.space_after=Pt(6)

# Only newly added response fields receive new native Word controls.
new_children=[x for x in d.element.body if x not in old_children]
for child in new_children:
 for run in list(child.xpath('.//w:r')):
  if len(run.findall(qn('w:t')))<=1:continue
  parent=run.getparent();pos=parent.index(run)
  for item in run:
   if item.tag==qn('w:rPr'):continue
   nr=OxmlElement('w:r')
   if run.find(qn('w:rPr')) is not None:nr.append(deepcopy(run.find(qn('w:rPr'))))
   nr.append(deepcopy(item));parent.insert(pos,nr);pos+=1
  parent.remove(run)
counter=0
for child in new_children:
 for run in list(child.xpath('.//w:r')):
  ts=run.findall(qn('w:t'))
  if len(ts)!=1 or '[' not in (ts[0].text or ''):continue
  parent=run.getparent();pos=parent.index(run)
  for piece in re.split(r'(\[[^\]]+\])',ts[0].text):
   if not piece:continue
   nr=deepcopy(run);nr.find(qn('w:t')).text=piece;nr.find(qn('w:t')).set(qn('xml:space'),'preserve')
   if piece.startswith('[') and piece.endswith(']'):
    counter+=1;sdt=OxmlElement('w:sdt');pr=OxmlElement('w:sdtPr')
    for key,val in [('alias',piece[1:-1]),('tag',f'PROJECT_DETAIL_{counter:03d}'),('id',str(5000+counter))]:
     el=OxmlElement('w:'+key);el.set(qn('w:val'),val);pr.append(el)
    txt=OxmlElement('w:text');txt.set(qn('w:multiLine'),'1');pr.append(txt);sdt.append(pr)
    content=OxmlElement('w:sdtContent');content.append(nr);sdt.append(content);parent.insert(pos,sdt)
   else:parent.insert(pos,nr)
   pos+=1
  parent.remove(run)
assert d.element.xpath('.//w:sdt//w:t/text()')[:len(prior_fields)]==prior_fields
all_text=' '.join(d.element.xpath('.//w:t/text()'))
required=['Project overview','Business Ask','Working team','Stakeholder questions and responses','Additional project details','Description and feature','User story','Business notes','IT comments and questions','Priority','Timeframe','Acceptance criteria','Impacted application','Component and Record Dependency','Total Estimated']
assert all(term in all_text for term in required)
out=folder/'Probo_Medical_AutoFast_Complete_Requirements_Workbook.docx'
d.core_properties.title='Probo Medical and AutoFast Complete Business Requirements Workbook'
d.save(out)
with zipfile.ZipFile(out) as z:assert z.testzip() is None
summary={'output':str(out.resolve()),'existing_fields_preserved':171,'added_fields':counter,'total_fields':len(d.element.xpath('.//w:sdt')),'requested_labels_checked':required}
Path('scratch/business_templates_qa/complete_workbook_checks.json').write_text(json.dumps(summary,indent=2))
print(json.dumps(summary,indent=2))
