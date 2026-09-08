from __future__ import annotations

import importlib.util
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
BASE_BUILDER = ROOT / "scratch" / "next-four-documents" / "build_story_docs.py"
OUTPUT = (
    ROOT
    / "outputs"
    / "case-00010716"
    / "Case_00010716_Asset_554470_Evaluation_Authentication_Standard_AutoFast.docx"
)


def load_helpers():
    spec = importlib.util.spec_from_file_location("autofast_helpers", BASE_BUILDER)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def replace_header(h, doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    for run in list(p.runs):
        p._p.remove(run._r)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    left = p.add_run("PROBO MEDICAL | SALESFORCE")
    h.set_run_font(left, size=8.5, color=h.MUTED, bold=True)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    right = p.add_run("\tCASE 00010716 | PRODUCTION IMPLEMENTATION")
    h.set_run_font(right, size=8.5, color=h.MUTED, bold=True)


def build() -> Path:
    h = load_helpers()
    doc = Document()
    h.configure_styles(doc)
    h.configure_section(doc, "00010716")
    replace_header(h, doc)

    doc.core_properties.title = "Case 00010716 Asset 554470 Evaluation Authentication Standard AutoFast"
    doc.core_properties.subject = "Production implementation and verification record"
    doc.core_properties.author = "Probo Medical Salesforce Team"
    doc.core_properties.keywords = "Salesforce, AutoFast, Evaluation, Authentication, 00010716, 554470"

    title = doc.add_paragraph("AUTOFAST IMPLEMENTATION RECORD", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle = doc.add_paragraph(
        "Case 00010716 | Asset 554470 Evaluation Authentication Error",
        style="Subtitle",
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT

    h.add_metadata_rows(
        doc,
        [
            ("Salesforce Case", "00010716"),
            ("Business Reference", "554470"),
            ("RMA", "A-219744"),
            ("Affected Evaluation", "E-2026-036262"),
            ("Prepared", "September 3, 2026"),
            ("Environment Status", "DevDO verified | Production deployed"),
            ("Final Status", "DEPLOYED | AUTOMATED VERIFICATION PASSED"),
        ],
    )
    h.add_callout(
        doc,
        "Implementation outcome",
        "The Evaluation authentication failure reported for asset 554470 has been corrected and deployed to Production. The focused Production regression completed with four passing tests and no failures.",
        "green",
    )

    doc.add_paragraph("1. Business Request", style="Heading 1")
    h.add_labeled_paragraph(
        doc,
        "Reported problem",
        "The repair team completed the evaluation for asset 554470 but could not authenticate the Evaluation in Salesforce. The page returned 'Attempt to de-reference a null object' and prevented the user from continuing.",
    )
    h.add_labeled_paragraph(
        doc,
        "Operational impact",
        "The technician could provide evaluation notes, but the system error interrupted the normal Evaluation completion process and required Salesforce support involvement.",
    )
    h.add_labeled_paragraph(
        doc,
        "Requested outcome",
        "Allow the existing Evaluation authentication action to complete against the Evaluation where the user started, without a null-reference exception or an update to another Evaluation.",
    )

    doc.add_paragraph("2. Affected Production Records", style="Heading 1")
    h.add_table(
        doc,
        ["Record type", "Record", "Dry-run finding"],
        [
            ("Product Item / Asset", "554470", "Record exists and is linked to the reported repair activity."),
            ("RMA", "A-219744", "Record exists and has two related Evaluations."),
            ("Evaluation", "E-2026-036262", "Original Evaluation created by Cole Rhode; reported failure."),
            ("Evaluation", "E-2026-036275", "Newer related Evaluation created by Rachel Croft."),
            ("Site user", "Approval Form Site Guest User", "Active under the Approval Form Profile."),
        ],
        [2200, 2400, 4760],
    )
    h.add_callout(
        doc,
        "Why the second Evaluation matters",
        "Because two Evaluations are related to the same RMA, any logic that guesses the most recently modified Evaluation can select the wrong record. The fix removes that ambiguity.",
        "blue",
    )

    h.add_page_break(doc)
    doc.add_paragraph("3. Steps to Reproduce the Original Issue", style="Heading 1")
    h.add_labeled_paragraph(
        doc,
        "Precondition",
        "Use an authorized repair user with access to asset 554470, RMA A-219744, and Evaluation E-2026-036262. These steps document the original Production behavior and are not instructions to alter the historical records.",
    )
    h.add_table(
        doc,
        ["Step", "Action", "Original result"],
        [
            ("1", "Open asset 554470 and confirm the related RMA is A-219744.", "The reported repair context is visible."),
            ("2", "Open Evaluation E-2026-036262 from the RMA.", "The completed evaluation details are available."),
            ("3", "Select the existing Evaluation authentication action.", "Salesforce starts the authentication return process."),
            ("4", "Allow the process to return to Salesforce.", "The page displays 'Attempt to de-reference a null object'."),
            ("5", "Review the Evaluation after the error.", "The user cannot reliably continue the Evaluation process."),
        ],
        [700, 4800, 3860],
    )
    h.add_labeled_paragraph(
        doc,
        "Expected result",
        "Salesforce returns to E-2026-036262, continues the existing Evaluation process, and does not modify E-2026-036275 or any unrelated Evaluation.",
    )

    doc.add_paragraph("4. Root Cause", style="Heading 1")
    h.add_labeled_paragraph(
        doc,
        "Record context was lost",
        "The external authentication return did not reliably retain the Id of the Evaluation where the user began.",
    )
    h.add_labeled_paragraph(
        doc,
        "Unsafe recovery logic",
        "When the Evaluation Id was unavailable, the controller attempted to recover by selecting the most recently modified Evaluation for the current user. That approach could return no usable record or a different Evaluation.",
    )
    h.add_labeled_paragraph(
        doc,
        "Missing defensive checks",
        "The controller assumed that the Evaluation and its related RMA, Product Item, test data, and folder configuration were present. A missing value could therefore be dereferenced instead of producing a readable message.",
    )

    doc.add_paragraph("5. Implemented Solution", style="Heading 1")
    h.add_table(
        doc,
        ["Control", "Implementation"],
        [
            ("Exact-record return", "The originating Evaluation Id is retained through the authentication state and restored when Salesforce receives the return request."),
            ("Type validation", "The returned Id is accepted only when it belongs to the Evaluation object."),
            ("No record guessing", "The latest-modified-Evaluation fallback was removed."),
            ("Null-safe processing", "Missing Evaluation, RMA, Product Item, Product Test, or folder data now stops safely with a clear page message."),
            ("Folder normalization", "Saved folder values are normalized before use so a full URL or query suffix does not corrupt the folder Id."),
        ],
        [2500, 6860],
    )
    h.add_callout(
        doc,
        "No-hardcoding confirmation",
        "No Salesforce record Ids, asset numbers, RMA numbers, user Ids, or fallback business values were added to the implementation. Runtime record context determines the Evaluation to process.",
        "green",
    )

    h.add_page_break(doc)
    doc.add_paragraph("6. Deployment Scope", style="Heading 1")
    h.add_table(
        doc,
        ["Component type", "API name", "Purpose"],
        [
            ("Apex Class", "cAuthURIForEval", "Retains the originating Evaluation context."),
            ("Apex Class", "cGoogleAppAuthenticationWithSalesforce", "Restores the exact Evaluation and applies defensive validation."),
            ("Apex Test", "GoogleAuthTestClass", "Covers correct-record return and missing/invalid context."),
        ],
        [1800, 4300, 3260],
    )

    doc.add_paragraph("7. Validation and Deployment Evidence", style="Heading 1")
    h.add_table(
        doc,
        ["Gate", "Job / Run", "Result"],
        [
            ("DevDO focused dry run", "0AfiK0000000o85SAA", "PASSED - 3/3 components; 4/4 tests; 0 errors"),
            ("Production check-only", "0AfjR0000000w41SAA", "PASSED - 3/3 components; 4/4 tests; 0 errors"),
            ("Production deployment", "0AfjR0000000w5dSAA", "SUCCEEDED - 3/3 components; 0 component errors"),
            ("Post-deployment regression", "05mjR00000009MfQAI", "PASSED - 4/4 methods; 0 failures"),
            ("Production read-back", "September 3, 2026 18:15:48 UTC", "All three classes updated by David Okolo"),
        ],
        [2750, 3000, 3610],
        status_column=2,
    )
    h.add_labeled_paragraph(
        doc,
        "Focused coverage",
        "The validated package recorded 301 of 397 executable lines covered for the controller (75.8%) and 15 of 17 executable lines covered for the authentication URI helper (88.2%).",
    )
    h.add_labeled_paragraph(
        doc,
        "Test method evidence",
        "case00010716RejectsMissingOrInvalidCallbackState, mytest, mytest2, and mytest3 all passed in the post-deployment Production run.",
    )
    h.add_callout(
        doc,
        "Execution boundary",
        "Validation and deployment were completed through Salesforce metadata and automated Apex tests. No browser sign-in, consent screen, external Drive access, or manual authentication session was performed as part of this release.",
        "blue",
    )

    doc.add_paragraph("8. Acceptance Criteria", style="Heading 1")
    h.add_table(
        doc,
        ["Acceptance criterion", "Evidence", "Status"],
        [
            ("The originating Evaluation is retained across the authentication return.", "Focused Apex regression", "PASS"),
            ("A competing Evaluation is not selected or updated.", "Correct-record regression assertion", "PASS"),
            ("Missing or invalid Evaluation context fails safely.", "Case-specific negative-path test", "PASS"),
            ("No Salesforce record Ids or business values are hard-coded.", "Source review", "PASS"),
            ("The package validates and deploys cleanly in Production.", "Check-only and deployment jobs", "PASS"),
            ("Business user confirms the normal Evaluation workflow.", "Optional operational confirmation", "RECOMMENDED"),
        ],
        [4300, 3300, 1760],
        status_column=2,
    )

    h.add_page_break(doc)
    doc.add_paragraph("9. Rollback and Support Notes", style="Heading 1")
    h.add_labeled_paragraph(
        doc,
        "Rollback scope",
        "If an unexpected regression is confirmed, redeploy the previously retained Production versions of the two controller classes and their test class through the approved release process. Do not change Production data as part of code rollback.",
    )
    h.add_labeled_paragraph(
        doc,
        "Support monitoring",
        "If a user reports another authentication error, capture the Evaluation name, RMA, asset number, timestamp, and user. Enable a short, time-bounded debug trace for the Approval Form Site Guest User only when troubleshooting evidence is required, then disable it after capture.",
    )
    h.add_labeled_paragraph(
        doc,
        "Case closure",
        "The implementation and automated verification are complete. The support owner may close case 00010716 after the business confirms there is no remaining operational issue for asset 554470.",
    )
    h.add_callout(
        doc,
        "Final determination",
        "DEPLOYED AND VERIFIED. No additional development is currently required. Retain this document with the case and deployment evidence.",
        "green",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
