from __future__ import annotations

import copy
import difflib
import os
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\LIKKI\Documents\ChatGPT\David 2 Probo medical")
REFERENCE = Path(
    r"C:\Users\LIKKI\.codex\plugins\cache\openai-curated-remote\openai-templates"
    r"\0.1.1\skills\artifact-template-strategy-memorandum\assets\reference.docx"
)
OLD_DIR = ROOT / "scratch" / "case-00010716-uat" / "force-app" / "main" / "default" / "classes"
NEW_DIR = ROOT / "scratch" / "case-00010716-uat-after" / "force-app" / "main" / "default" / "classes"
OUTPUT = ROOT / "outputs" / "case-00010716" / "Case_00010716_Apex_Code_Changes_Old_vs_New.docx"
WORKING = ROOT / "scratch" / "case-00010716-code-diff-docx" / "working.docx"

NAVY = "112075"
BLUE = "7A86DE"
TEXT = "202124"
MUTED = "60646C"
OLD_FILL = "FDECEC"
NEW_FILL = "EAF6EC"
OLD_TEXT = "9C0006"
NEW_TEXT = "006100"
HEADER_FILL = "E8EBF8"
LIGHT_BLUE = "F2F4FC"
WHITE = "FFFFFF"


CLASS_SPECS = [
    ("cAuthURIForEval", 11, 1),
    ("cGoogleAppAuthenticationWithSalesforce", 153, 134),
    ("GoogleAuthTestClass", 92, 7),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_inches: list[float], indent_dxa: int = 120) -> None:
    table.autofit = False
    widths_dxa = [int(round(width * 1440)) for width in widths_inches]
    total_dxa = sum(widths_dxa)
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_inches[index])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name: str, size: float, color: str = TEXT, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def add_body_paragraph(doc: DocumentObject, text: str = "", style: str | None = None, after=6, before=0):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    if text:
        run = p.add_run(text)
        set_run_font(run, "Helvetica Neue", 10, TEXT)
    return p


def bullet_num_id(doc: DocumentObject) -> int:
    cached = getattr(doc, "_case10716_bullet_num_id", None)
    if cached is not None:
        return cached
    numbering = doc.part.numbering_part.element
    bullet_abstract_ids = []
    for abstract in numbering.findall(qn("w:abstractNum")):
        if abstract.find(f".//{qn('w:numFmt')}[@{qn('w:val')}='bullet']") is not None:
            bullet_abstract_ids.append(abstract.get(qn("w:abstractNumId")))
    if not bullet_abstract_ids:
        raise RuntimeError("The retained template does not contain a bullet numbering definition.")
    chosen_abstract = bullet_abstract_ids[0]
    for num in numbering.findall(qn("w:num")):
        abstract_id = num.find(qn("w:abstractNumId"))
        if abstract_id is not None and abstract_id.get(qn("w:val")) == chosen_abstract:
            num_id = int(num.get(qn("w:numId")))
            setattr(doc, "_case10716_bullet_num_id", num_id)
            return num_id
    existing = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = (max(existing) + 1) if existing else 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), chosen_abstract)
    num.append(abstract_ref)
    numbering.append(num)
    setattr(doc, "_case10716_bullet_num_id", num_id)
    return num_id


def add_bullet(doc: DocumentObject, text: str) -> None:
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Pt(16)
    p.paragraph_format.first_line_indent = Pt(-10)
    bullet = p.add_run("• ")
    set_run_font(bullet, "Helvetica Neue", 10, NAVY, True)
    run = p.add_run(text)
    set_run_font(run, "Helvetica Neue", 10, TEXT)


def clear_document_body(doc: DocumentObject) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def format_heading(paragraph, level: int) -> None:
    paragraph.style = f"Heading {level}"
    for run in paragraph.runs:
        set_run_font(run, "Helvetica Neue", 15 if level == 1 else 11, NAVY, True)


def add_heading(doc: DocumentObject, text: str, level: int = 1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    format_heading(p, level)
    return p


def add_callout(doc: DocumentObject, label: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.2], indent_dxa=170)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    set_cell_margins(cell, 150, 170, 150, 170)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    lead = p.add_run(label + " ")
    set_run_font(lead, "Helvetica Neue", 10, NAVY, True)
    body = p.add_run(text)
    set_run_font(body, "Helvetica Neue", 10, TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_overview_table(doc: DocumentObject) -> None:
    table = doc.add_table(rows=1, cols=4)
    headers = ["APEX CLASS", "OLD UAT", "CURRENT UAT", "PURPOSE"]
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, "Helvetica Neue", 8.5, WHITE, True)
    set_repeat_table_header(table.rows[0])

    purposes = {
        "cAuthURIForEval": "Carry the Evaluation ID through Google OAuth.",
        "cGoogleAppAuthenticationWithSalesforce": "Restore the correct record, clean the folder ID, and fail safely.",
        "GoogleAuthTestClass": "Prove the callback updates only the intended Evaluation.",
    }
    for name, adds, deletes in CLASS_SPECS:
        cells = table.add_row().cells
        values = [name, f"-{deletes} lines", f"+{adds} lines", purposes[name]]
        for idx, value in enumerate(values):
            set_cell_margins(cells[idx])
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if idx == 1:
                set_cell_shading(cells[idx], OLD_FILL)
            elif idx == 2:
                set_cell_shading(cells[idx], NEW_FILL)
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in (1, 2) else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(value)
            set_run_font(run, "Helvetica Neue", 8.5, TEXT, idx == 0)
    set_table_geometry(table, [2.05, 0.9, 0.95, 2.3], indent_dxa=120)


def grouped_hunks(old_lines: list[str], new_lines: list[str], context: int = 2):
    matcher = difflib.SequenceMatcher(None, old_lines, new_lines, autojunk=False)
    return list(matcher.get_grouped_opcodes(context))


def align_hunk(old_lines: list[str], new_lines: list[str], group):
    rows = []
    for tag, i1, i2, j1, j2 in group:
        old_chunk = old_lines[i1:i2]
        new_chunk = new_lines[j1:j2]
        if tag == "equal":
            for offset, (old_text, new_text) in enumerate(zip(old_chunk, new_chunk)):
                rows.append((i1 + offset + 1, old_text, False, j1 + offset + 1, new_text, False))
        elif tag == "replace":
            count = max(len(old_chunk), len(new_chunk))
            for offset in range(count):
                old_number = i1 + offset + 1 if offset < len(old_chunk) else None
                new_number = j1 + offset + 1 if offset < len(new_chunk) else None
                old_text = old_chunk[offset] if offset < len(old_chunk) else ""
                new_text = new_chunk[offset] if offset < len(new_chunk) else ""
                rows.append((old_number, old_text, old_number is not None, new_number, new_text, new_number is not None))
        elif tag == "delete":
            for offset, old_text in enumerate(old_chunk):
                rows.append((i1 + offset + 1, old_text, True, None, "", False))
        elif tag == "insert":
            for offset, new_text in enumerate(new_chunk):
                rows.append((None, "", False, j1 + offset + 1, new_text, True))
    return rows


def explanation_for(class_name: str, index: int, rows) -> tuple[str, str]:
    old_text = "\n".join(row[1] for row in rows)
    new_text = "\n".join(row[4] for row in rows)
    combined = old_text + "\n" + new_text

    if class_name == "cAuthURIForEval":
        if "this(Clientkey, redirect_uri, null)" in combined:
            return "Accept the Evaluation ID", "We kept the original constructor for compatibility and added a second constructor that can carry the exact Evaluation ID."
        if "String stateValue = 'url='" in combined:
            return "Build a return state with the record ID", "The OAuth state now contains both the return page and the exact Evaluation ID, creating a reliable claim ticket for the Google round trip."
        return "Use the encoded state in the Google URL", "Before leaving Salesforce for Google, the code URL-encodes the claim ticket and sends it in the OAuth request. Google returns it unchanged, so Salesforce knows exactly which record to resume."

    if class_name == "cGoogleAppAuthenticationWithSalesforce":
        rules = [
            ("private final ProductItem__c pitem", "Allow safe early exit", "The Product Item variable is no longer forced to be assigned before the constructor can stop. That lets validation errors end the request cleanly without creating another failure."),
            ("new cAuthURIForEval(key, redirect_uri, itemid)", "Send the exact Evaluation into Google authentication", "The Authenticate action now passes the current Evaluation ID into the OAuth URL builder. This is the hand-off that keeps the record identity intact during the round trip."),
            ("if (String.isBlank(itemid))", "Restore the record from callback state", "When Google returns without the Evaluation in the normal page parameter, the controller recovers it from the OAuth claim ticket and stops safely if neither source is valid."),
            ("getEvaluationIdFromState", "Recover and validate the Evaluation", "The callback decodes the returned OAuth state, confirms the value is a real Evaluation ID, and shows a readable message instead of guessing or crashing."),
            ("List<Evaluation_Natalie__c> evaluations", "Query the exact Evaluation safely", "The controller queries only the validated Evaluation ID and checks the result before using it, preventing a missing record from causing a query exception."),
            ("Product_Item_ID__c))", "Validate the Product Item reference", "The controller checks that the Evaluation has a Product Item and that its stored ID is valid before running any Google Drive logic."),
            ("most recently modified evaluation created by user", "Remove the unsafe 'latest record' guess", "The old fallback could choose whichever Evaluation the user touched most recently. That whole branch was removed, so the callback cannot silently switch to another Evaluation."),
            ("associatedRMAs.isEmpty()", "Check that the related RMA exists", "The controller now verifies that the Evaluation has a usable related RMA before continuing, and returns a readable message when it does not."),
            ("productItems.isEmpty()", "Load and normalize the Product Item folder", "The controller safely loads the Product Item, checks that it exists, and converts either a full Google Drive URL or a raw folder value into one clean folder ID."),
            ("No Evaluation is available to save", "Stop an invalid save safely", "If the callback did not load a valid Evaluation, Save now stops with a clear message instead of dereferencing a null record."),
            ("Product Item does not have a valid Google Drive folder", "Require a usable folder ID", "The file lookup now stops early and explains the problem when the Product Item has no valid Drive folder."),
            ("mappedFileName", "Ignore empty Drive file names", "Google responses can contain missing or incomplete items. The controller now skips blank names instead of adding bad map entries."),
            ("f.length() < 6", "Protect dated-folder parsing", "The old code always took the first six characters. The new guard skips short or blank names and avoids substring errors."),
            ("No valid image folder was found", "Explain when no image folder matches", "If no valid dated image folder can be selected, the user receives a clear message and the callout is not attempted with a blank folder ID."),
            ("nestedMappedFileName", "Clean nested-folder file results", "The same blank-name protection is applied to files found inside the selected image folder."),
            ("return testFileMap.clone", "Make file mocks deterministic", "During tests, the controller returns the prepared file map directly. That lets the test exercise filename matching without making a real Google call."),
        ]
        for needle, title, explanation in rules:
            if needle in combined:
                return title, explanation
        return f"Controller safety change {index}", "This change keeps the Google callback tied to valid Salesforce data and prevents incomplete Google Drive responses from breaking the page."

    rules = [
        ("Image__c='https://example.invalid", "Add a safe test folder URL", "The unit test now includes a fake folder URL with a query string so folder-ID cleanup is tested without using a real Google Drive folder."),
        ("competingEvaluation", "Prove only the intended Evaluation changes", "The test creates a target Evaluation and a second control Evaluation. It rebuilds the callback state, runs the controller, and verifies the target changes while the control record remains untouched."),
        ("normalizeDriveFolderId('0Bzmq", "Verify callback state and folder cleanup", "The test confirms the callback restores the expected Evaluation and that a full Google Drive URL is reduced to the correct folder ID before file lookup."),
        ("case00010716RejectsMissingOrInvalidCallbackState", "Test missing and invalid inputs", "The new negative test confirms that missing IDs, invalid callback state, wrong object IDs, and missing folders produce safe results and readable messages."),
    ]
    for needle, title, explanation in rules:
        if needle in combined:
            return title, explanation
    return f"Test coverage change {index}", "This test change checks the new callback and folder-handling behavior without relying on a live Google account."


def add_code_line(cell, line_number, text, changed, side: str) -> None:
    p = cell.paragraphs[0] if len(cell.paragraphs) == 1 and not cell.paragraphs[0].text else cell.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    number = f"{line_number:>4}" if line_number is not None else "    "
    marker = "-" if side == "old" and changed else "+" if side == "new" and changed else " "
    run = p.add_run(f"{number} {marker} {text}")
    color = OLD_TEXT if side == "old" and changed else NEW_TEXT if side == "new" and changed else TEXT
    set_run_font(run, "Consolas", 7.5, color, changed)


def add_diff_table(doc: DocumentObject, rows) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.allow_autofit = False
    set_table_geometry(table, [4.75, 4.75], indent_dxa=110)
    headers = (("OLD UAT CODE", OLD_FILL, OLD_TEXT), ("NEW UAT CODE", NEW_FILL, NEW_TEXT))
    for idx, (label, fill, color) in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, fill)
        set_cell_margins(cell, 90, 110, 90, 110)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(label)
        set_run_font(run, "Helvetica Neue", 8.5, color, True)
    set_repeat_table_header(table.rows[0])

    for old_number, old_text, old_changed, new_number, new_text, new_changed in rows:
        old_cell, new_cell = table.add_row().cells
        set_cell_shading(old_cell, "FFF8F8")
        set_cell_shading(new_cell, "F7FCF8")
        for cell in (old_cell, new_cell):
            set_cell_margins(cell, 0, 110, 0, 110)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        add_code_line(old_cell, old_number, old_text, old_changed, "old")
        add_code_line(new_cell, new_number, new_text, new_changed, "new")


def add_human_explanation(doc: DocumentObject, explanation: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    label = p.add_run("In plain English: ")
    set_run_font(label, "Helvetica Neue", 9.5, NAVY, True)
    body = p.add_run(explanation)
    set_run_font(body, "Helvetica Neue", 9.5, TEXT)


def replace_xml_text(path: Path, replacements: dict[str, str]) -> None:
    temp = path.with_suffix(".patched.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename in {"word/header1.xml", "word/footer1.xml", "docProps/core.xml"}:
                text = data.decode("utf-8")
                for old, new in replacements.items():
                    text = text.replace(old, new)
                data = text.encode("utf-8")
            target.writestr(item, data)
    os.replace(temp, path)


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REFERENCE, WORKING)
    doc = Document(WORKING)
    clear_document_body(doc)

    # Use the same review header/footer on every page. The retained template
    # defines separate even-page parts, which otherwise hide the review label
    # and clip the page-number text on landscape continuation pages.
    even_odd_headers = doc.settings._element.find(qn("w:evenAndOddHeaders"))
    if even_odd_headers is not None:
        doc.settings._element.remove(even_odd_headers)

    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Helvetica Neue"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Helvetica Neue")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Helvetica Neue")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("APEX CODE CHANGE REVIEW")
    set_run_font(run, "Helvetica Neue", 30, NAVY, True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(70)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Google Drive Photo Authentication Fix | Story 00010716")
    set_run_font(run, "Helvetica Neue", 15, BLUE, True)

    metadata = doc.add_table(rows=6, cols=2)
    metadata_values = [
        ("Environment", "Probo Medical UAT"),
        ("Scope", "Three Apex classes"),
        ("Deployment", "0AfjH0000000T6rSAE"),
        ("Verification", "3/3 components; 4/4 tests passed"),
        ("Prepared for", "Technical and business review"),
        ("Status", "Current deployed UAT code"),
    ]
    for row, (label, value) in zip(metadata.rows, metadata_values):
        set_cell_margins(row.cells[0], 85, 110, 85, 110)
        set_cell_margins(row.cells[1], 85, 110, 85, 110)
        row.cells[0].text = ""
        row.cells[1].text = ""
        label_run = row.cells[0].paragraphs[0].add_run(label)
        set_run_font(label_run, "Helvetica Neue", 9, NAVY, True)
        value_run = row.cells[1].paragraphs[0].add_run(value)
        set_run_font(value_run, "Helvetica Neue", 9, TEXT)
    set_table_geometry(metadata, [1.55, 4.65], indent_dxa=110)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(55)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Code-only comparison. No Google Drive folder setup is included.")
    set_run_font(run, "Helvetica Neue", 9, MUTED)

    doc.add_page_break()
    add_heading(doc, "1. What was wrong", 1)
    add_body_paragraph(
        doc,
        "Google authentication temporarily sends the user away from Salesforce. The old code did not reliably carry the Evaluation ID through that trip. When the user came back, the controller could fall back to the most recently modified Evaluation, which meant the wrong record could be selected.",
    )
    add_body_paragraph(
        doc,
        "The controller also treated a full Google Drive URL as if it were already a clean folder ID. Missing records, short filenames, or incomplete Drive responses could cause query, substring, or null-reference failures.",
    )

    add_heading(doc, "2. How we fixed it", 1)
    add_callout(
        doc,
        "Simple explanation:",
        "We gave the Google round trip a claim ticket. The claim ticket is the exact Salesforce Evaluation ID. When Google sends the user back, Salesforce reads that ticket, validates it, and resumes the correct Evaluation instead of guessing.",
    )
    for item in (
        "Carry the exact Evaluation ID inside the OAuth state parameter.",
        "Restore and validate that ID after Google returns to Salesforce.",
        "Remove the fallback that selected the user's most recently modified Evaluation.",
        "Convert a complete Drive folder URL into a clean folder ID.",
        "Stop safely with readable messages when required records or folder data are missing.",
        "Strengthen the test class with a target Evaluation and a control Evaluation.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "3. Apex scope", 1)
    add_body_paragraph(
        doc,
        "Exactly three Apex class files changed. Their three .cls-meta.xml files did not change, and no Visualforce page was changed in this deployment. These changes fix callback record selection, folder-ID normalization, defensive error handling, and automated coverage. They do not create a new Cap field and do not change the pre-existing image-to-field assignment logic. This document intentionally covers code changes only.",
    )
    add_overview_table(doc)

    code_section = doc.add_section(WD_SECTION.NEW_PAGE)
    code_section.orientation = WD_ORIENT.LANDSCAPE
    code_section.page_width = Inches(11)
    code_section.page_height = Inches(8.5)
    code_section.top_margin = Inches(0.6)
    code_section.bottom_margin = Inches(0.6)
    code_section.left_margin = Inches(0.6)
    code_section.right_margin = Inches(0.6)
    code_section.header_distance = Inches(0.3)
    code_section.footer_distance = Inches(0.3)
    code_section.header.is_linked_to_previous = True
    code_section.footer.is_linked_to_previous = True
    page_number_type = code_section._sectPr.find(qn("w:pgNumType"))
    if page_number_type is not None and qn("w:start") in page_number_type.attrib:
        del page_number_type.attrib[qn("w:start")]

    add_heading(doc, "4. Side-by-side Apex diff", 1)
    add_body_paragraph(
        doc,
        "Red-tinted lines are old UAT code that was removed or replaced. Green-tinted lines are the current UAT code that was added or replaced. Uncolored lines provide nearby context.",
        after=8,
    )

    for class_number, (class_name, adds, deletes) in enumerate(CLASS_SPECS, start=1):
        if class_number > 1:
            doc.add_page_break()
        add_heading(doc, f"4.{class_number} {class_name}", 1)
        add_body_paragraph(doc, f"Verified change size: +{adds} lines / -{deletes} lines.", after=6)

        old_lines = (OLD_DIR / f"{class_name}.cls").read_text(encoding="utf-8").splitlines()
        new_lines = (NEW_DIR / f"{class_name}.cls").read_text(encoding="utf-8").splitlines()
        hunks = grouped_hunks(old_lines, new_lines, context=2)
        for hunk_index, group in enumerate(hunks, start=1):
            rows = align_hunk(old_lines, new_lines, group)
            title, explanation = explanation_for(class_name, hunk_index, rows)
            add_heading(doc, f"Change {hunk_index}: {title}", 2)
            add_diff_table(doc, rows)
            add_human_explanation(doc, explanation)

    doc.add_page_break()
    add_heading(doc, "5. Verification", 1)
    for item in (
        "Old side: retrieved from ProboUAT before deployment.",
        "New side: retrieved again from ProboUAT after deployment.",
        "Deployment 0AfjH0000000T6rSAE succeeded with 3/3 components.",
        "GoogleAuthTestClass completed 4/4 tests with no failures.",
        "All changed Apex hunks are included in the side-by-side section.",
        "No unchanged OAuth credential lines are reproduced in this document.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "6. Bottom line", 1)
    add_body_paragraph(
        doc,
        "Before the fix, the Google callback could return without knowing which Evaluation started the process and could choose the wrong recent record. After the fix, the exact Evaluation travels through OAuth, is validated on return, and is the only record the save test allows to change. Folder links and incomplete Drive results are also handled more safely.",
    )

    doc.core_properties.title = "Case 00010716 Apex Code Changes - Old vs New"
    doc.core_properties.subject = "Google Drive photo authentication callback fix"
    doc.core_properties.author = "Probo Medical Technical Review"
    doc.core_properties.keywords = "Salesforce, Apex, OAuth, Google Drive, UAT, 00010716"
    doc.settings.update_fields_on_open = True
    doc.save(OUTPUT)

    replace_xml_text(
        OUTPUT,
        {
            "Strategy Memo": "Apex Change Review",
            "[Confidentiality]": "Internal - Code Review",
            "[Name]": "Probo Medical Technical Review",
        },
    )

    print(f"WROTE {OUTPUT}")
    print(f"SIZE {OUTPUT.stat().st_size}")


if __name__ == "__main__":
    build()
