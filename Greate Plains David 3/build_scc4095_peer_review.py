from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_architect_review import (
    INK,
    LIGHT_BLUE,
    MUTED,
    NAVY,
    add_body,
    add_heading,
    add_page_field,
    add_table,
    rgb,
    set_cell_margins,
    set_cell_text,
    set_run_font,
    set_table_geometry,
    shade_cell,
)


ROOT = Path(__file__).resolve().parent
REFERENCE = ROOT / "SCC-3384-peer-review.docx"
OUTPUT = ROOT / "SCC-4095-peer-review.docx"


def clear_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_light_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9D9D9")


def normalize_table(table) -> None:
    set_light_borders(table)
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=85, start=120, bottom=85, end=120)


def review_table(document, headers, rows, widths, status_column=None):
    table = add_table(document, headers, rows, widths, status_column=status_column)
    normalize_table(table)
    return table


def configure_footer(section) -> None:
    for header in (section.header, section.even_page_header, section.first_page_header):
        paragraph = header.paragraphs[0]
        paragraph.text = ""
        paragraph.paragraph_format.space_after = Pt(0)
    for footer in (section.footer, section.even_page_footer, section.first_page_footer):
        paragraph = footer.paragraphs[0]
        paragraph.text = ""
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph.paragraph_format.space_before = Pt(0)
        run = paragraph.add_run("GreatPlainsDevA  |  Page ")
        set_run_font(run, size=9, color=MUTED)
        add_page_field(paragraph)


def configure_document(document: Document) -> None:
    document.settings.odd_and_even_pages_header_footer = False
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        section.different_first_page_header_footer = False
        configure_footer(section)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = document.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = rgb("000000")
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.keep_with_next = True
    p_pr = title._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)

    for name, size, before, after in (
        ("Heading 1", 16, 12, 6),
        ("Heading 2", 13, 10, 5),
        ("Heading 3", 12, 8, 4),
    ):
        style = document.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb("000000")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_decision(document: Document, label: str, message: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.10
    lead = paragraph.add_run(f"{label}  ")
    set_run_font(lead, size=11, color=NAVY, bold=True)
    body = paragraph.add_run(message)
    set_run_font(body, size=11, color=INK)


def add_page_heading(document: Document, text: str) -> None:
    paragraph = add_heading(document, text, 1)
    paragraph.paragraph_format.page_break_before = True


def add_metadata(document: Document) -> None:
    rows = [
        ("Environment", "GreatPlainsDevA sandbox | Great Plains Communications | Org 00DEa00000Fc086MAB"),
        ("Review date", "5 September 2026"),
        ("Focused retrieval", "Succeeded | DevA job 09SEa00000jCq7RMAS"),
        ("Check only", "Succeeded | Job 0AfEa00000bn9GAKAY"),
        ("Decision", "Approved with Conditions - development complete; QA evidence required"),
        ("Boundary", "No deployment, activation, persistent record creation, or metadata change"),
    ]
    table = document.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [1900, 7460])
    normalize_table(table)
    for row, (label, value) in zip(table.rows, rows):
        shade_cell(row.cells[0], LIGHT_BLUE)
        set_cell_text(row.cells[0], label, size=9.5, color=NAVY, bold=True)
        set_cell_text(row.cells[1], value, size=9.3)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def build_document() -> None:
    document = Document(REFERENCE)
    clear_body(document)
    configure_document(document)

    eyebrow = document.add_paragraph()
    eyebrow.paragraph_format.space_before = Pt(8)
    eyebrow.paragraph_format.space_after = Pt(2)
    set_run_font(eyebrow.add_run("PEER REVIEW"), size=10, color="000000", bold=True)

    title = document.add_paragraph(style="Title")
    title.add_run("SCC 4095 Gantt Label Peer Review")

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    set_run_font(
        subtitle.add_run("Great Plains DevA Sandbox Check only implementation assessment"),
        size=13,
        color="000000",
    )

    add_metadata(document)
    add_decision(
        document,
        "Decision",
        "The requested Gantt Label behavior is already implemented in DevA and the focused check-only validation succeeds. No additional functional coding is required. Approval remains conditional on completing the story documentation and a named-persona UI save test.",
    )

    add_heading(document, "Acceptance outcome", 1)
    rows = [
        ("1", "Gantt Label on new screen", "Pass", "RecurringAbsence.page binds an input field to ResourceAbsence.FSL__GanttLabel__c in the Absence Details section."),
        ("2", "Value copied to every recurrence", "Pass", "The controller clones the configured ResourceAbsence for each occurrence, so the populated Gantt Label is preserved."),
        ("3", "Automated persistence test", "Pass", "The test sets PTO, creates three recurring absences, queries the saved records, and asserts PTO on each one."),
        ("4", "Focused deployment validation", "Pass", "The check-only job validates all three components with no component or test failures."),
        ("5", "Controller coverage", "Pass", "281 of 349 executable lines are covered, producing 80.52 percent coverage."),
        ("6", "Functional source parity", "Pass", "The Visualforce page, controller, test class, and page metadata are identical in current DevA and Merge retrievals."),
        ("7", "Story definition and UI evidence", "Conditional", "The story has no description or acceptance criteria, and no record-creating named-persona UI test was performed during this dry run."),
    ]
    review_table(document, ["#", "Scope", "Result", "Evidence"], rows, [500, 2200, 1350, 5310], status_column=2)

    add_page_heading(document, "Development scope")
    scope_rows = [
        ("ApexPage", "RecurringAbsence", "Page plus metadata XML", "Displays the existing managed Gantt Label field on the New Recurring Absence screen."),
        ("ApexClass", "RecurringAbsenceController", "Controller plus metadata XML", "Generates recurring records by cloning the populated ResourceAbsence; no separate field assignment is required."),
        ("ApexClass", "RecurringAbsenceControllerTest", "Test class plus metadata XML", "Proves Gantt Label persistence across all generated records and covers recurrence behavior."),
    ]
    review_table(document, ["Type", "Component", "Physical files", "Purpose"], scope_rows, [1200, 3350, 2150, 2660])

    add_heading(document, "Items not required for this story", 2)
    nonscope_rows = [
        ("New custom field", "Not required", "FSL__GanttLabel__c already exists as the managed Gantt Label Text 255 field."),
        ("Page layout", "Not required", "The feature is implemented on the custom Visualforce page named RecurringAbsence."),
        ("Permission set change", "Not indicated", "CRC and principal Field Service permission sets grant read and edit access in DevA and Merge."),
        ("Controller field mapping", "Not required", "The existing clone operation carries the field value into each new ResourceAbsence."),
    ]
    review_table(document, ["Potential item", "Scope reading", "Evidence"], nonscope_rows, [2300, 1800, 5260], status_column=1)

    add_heading(document, "Environment comparison", 2)
    parity_rows = [
        ("Visualforce page source", "Identical", "The Gantt Label input is present in both current DevA and Merge."),
        ("Controller source", "Identical", "No functional source delta was found."),
        ("Test source", "Identical", "The Gantt Label persistence assertion exists in both environments."),
        ("Page metadata", "Identical", "Both use API version 60 and FSL package version 262.0."),
        ("Apex metadata", "Different", "DevA controller and test metadata use API versions 64 and 67; Merge uses 60 for both. Do not treat these version-only differences as SCC-4095 scope without an explicit decision."),
    ]
    review_table(document, ["Item", "Result", "Peer review reading"], parity_rows, [2600, 1600, 5160], status_column=1)

    add_page_heading(document, "Validation evidence")
    validation_rows = [
        ("Target identity", "Pass", "Great Plains Communications sandbox 00DEa00000Fc086MAB on USA20S."),
        ("Check-only status", "Pass", "Succeeded with checkOnly true; job 0AfEa00000bn9GAKAY."),
        ("Components", "Pass", "3 of 3 succeeded with 0 component errors; all were reported unchanged in DevA."),
        ("Specified tests", "Pass", "11 tests ran with 0 failures."),
        ("Controller coverage", "Pass", "281 covered and 68 uncovered executable lines; 80.52 percent."),
        ("Field metadata", "Pass", "FSL__GanttLabel__c is present as Gantt Label, Text 255."),
        ("Field access", "Pass", "CRC Access Permission Set Group and the principal Field Service permission sets grant read and edit access."),
    ]
    review_table(document, ["Evidence", "Result", "Observed outcome"], validation_rows, [2500, 1450, 5410], status_column=1)

    add_heading(document, "Review boundaries", 2)
    boundary_rows = [
        ("UI save test", "Not performed", "Saving the screen would create Resource Absence records. The dry run used code, metadata, security, and test evidence only."),
        ("Named QA persona", "Not verified", "The story does not define a required profile or permission-set assignment for the final UI test."),
        ("Merge deployment", "Not performed", "Merge was inspected read-only for source and field-access parity; no package was deployed or check-only transmitted to Merge."),
        ("Story completeness", "Open", "Description, business value, acceptance criteria, pre-deployment, and post-deployment instructions are absent or recorded as None."),
        ("Instruction boundary", "Preserved", "The document's request to coordinate with named colleagues was treated as ticket context and was not actioned."),
    ]
    review_table(document, ["Boundary", "Status", "Evidence"], boundary_rows, [2300, 1700, 5360], status_column=1)

    add_decision(
        document,
        "Peer review reading",
        "Technical validation supports moving the existing implementation to controlled QA confirmation. It does not support claiming complete business acceptance until the missing criteria and record-creating UI evidence are supplied.",
    )

    add_page_heading(document, "Required completion gates")
    gates = [
        ("1", "Write acceptance criteria", "Product Owner", "Define field visibility, allowed personas, required value persistence, recurrence types, and expected success behavior."),
        ("2", "Confirm branch provenance", "Developer", "Verify the three component sources in the SCC-4095 branch match the reviewed DevA implementation."),
        ("3", "Decide API versions", "Technical Lead", "Exclude the controller and test metadata version differences unless the release intentionally upgrades them."),
        ("4", "Run named-persona UI test", "QA", "Create a recurring absence with a Gantt Label and confirm every generated record contains the label."),
        ("5", "Complete technical inventory", "Developer", "List the ApexPage, controller, test class, and all three companion metadata files in the story."),
        ("6", "Retest in Merge or UAT", "QA and Release Manager", "Confirm the field renders for the affected persona after the approved package is promoted through the pipeline."),
    ]
    review_table(document, ["#", "Required action", "Owner", "Evidence needed"], gates, [550, 2650, 2000, 4160])

    add_heading(document, "Ready to paste peer review response", 1)
    response = (
        "Peer review completed in GreatPlainsDevA. Approved with Conditions. The New Recurring Absence Visualforce page contains the Gantt Label field, the controller's clone logic preserves that value on every generated Resource Absence, and the focused test asserts the value on all three created records. Check-only job 0AfEa00000bn9GAKAY succeeded for all three components with 11 tests, no failures, and 80.52 percent coverage for RecurringAbsenceController. Current DevA and Merge functional source is identical; only controller and test metadata API versions differ. No additional functional development, new field, page layout, or permission-set change is indicated. Before final approval, add explicit acceptance criteria and the component inventory, decide whether the API version differences are intentional, and attach a named-persona UI save test showing that every generated record retains the label. No deployment, activation, persistent record creation, or org metadata change was made."
    )
    add_body(document, response, after=4)

    document.core_properties.title = "SCC 4095 Gantt Label Peer Review"
    document.core_properties.subject = "Check only implementation assessment in GreatPlainsDevA"
    document.core_properties.author = "Peer Review"
    document.core_properties.keywords = "Salesforce, Great Plains, SCC-4095, Resource Absence, Gantt Label, peer review"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
