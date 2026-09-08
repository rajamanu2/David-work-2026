from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


SOURCE = Path(r"C:\Users\LIKKI\Downloads\SCC-4179.docx")
OUTPUT = Path(r"C:\Users\LIKKI\Documents\ChatGPT\Greate Plains David 3\SCC-4179-completed.docx")


def set_text_preserve_format(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        first = paragraph.runs[0]
        first.text = text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def insert_after(paragraph: Paragraph, text: str = "") -> Paragraph:
    new_xml = OxmlElement("w:p")
    paragraph._p.addnext(new_xml)
    new_paragraph = Paragraph(new_xml, paragraph._parent)
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def create_numbering(document: Document, kind: str = "decimal") -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), kind)
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1." if kind == "decimal" else "•")
    level.append(level_text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    p_pr.append(indent)
    level.append(p_pr)
    abstract.append(level)

    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        first_num.addprevious(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph: Paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    else:
        for child in list(num_pr):
            num_pr.remove(child)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)


def paragraph_exact(document: Document, text: str, occurrence: int = 0) -> Paragraph:
    matches = [p for p in document.paragraphs if p.text.strip() == text]
    if occurrence >= len(matches):
        raise ValueError(f"Paragraph not found: {text!r} occurrence {occurrence}")
    return matches[occurrence]


def answer_after(document: Document, section_name: str, occurrence: int = 0) -> Paragraph:
    paragraphs = document.paragraphs
    matches = [i for i, p in enumerate(paragraphs) if p.text.strip() == section_name]
    if occurrence >= len(matches):
        raise ValueError(f"Section not found: {section_name!r} occurrence {occurrence}")
    start = matches[occurrence]
    for paragraph in paragraphs[start + 1 :]:
        if paragraph.text.strip() == "Top of Form":
            continue
        return paragraph
    raise ValueError(f"Answer not found after {section_name!r}")


def replace_with_numbered_items(
    paragraph: Paragraph, items: tuple[str, ...], document: Document
) -> Paragraph:
    num_id = create_numbering(document, "decimal")
    set_text_preserve_format(paragraph, items[0])
    apply_numbering(paragraph, num_id)
    anchor = paragraph
    for item in items[1:]:
        anchor = insert_after(anchor, item)
        apply_numbering(anchor, num_id)
    return anchor


document = Document(SOURCE)

# Correct source-export artifacts and make the core requirement testable.
set_text_preserve_format(
    paragraph_exact(document, "Per review with Matt Evrates… System Admin, System Admin API, and Standard User profiles should no longer have Field Level Security access to the two CPNI fields: CPNI Pin and CPNI Password ."),
    "Per review with Matt Evrates, the System Administrator, System Administrator - API Only, and Standard User profiles should no longer have Field-Level Security access to CPNI Pin or CPNI Password.",
)
set_text_preserve_format(
    paragraph_exact(document, "Remove FLS (read/ and or write) on both CPNI fields for: System Admin, System Admin API, Standard User profiles (Admin to remove from other profiles if this was configured for other profiles in UAT/Prod)."),
    "Remove read and edit FLS for both fields from the three named profiles and from every other profile or permission set that grants access in UAT or Production.",
)
set_text_preserve_format(document.paragraphs[9], "")

acceptance = paragraph_exact(
    document,
    'CPNI FLS is removed from the three profiles listed above + any SF admin has configured. Only users with the "Contact Encrypted CPNI Read/Write Access" permission set can view/edit the CPNI fields.',
)
set_text_preserve_format(
    acceptance,
    'Pass only when Contact.CPNI_Pin__c and Contact.CPNI_Password__c are unreadable and uneditable through every profile and every non-approved permission set, and read/edit access remains only in "Contact Encrypted CPNI Read/Write Access." Validate both UI and API behavior with named users.',
)

# Preserve the approved declarative design while recording the actual readiness result.
solution = paragraph_exact(
    document,
    "Remove FLS on CPNI fields, delegate access to these fields via permission set only.",
)
set_text_preserve_format(
    solution,
    "Use declarative FLS only: remove both CPNI field permissions from all profiles and all non-approved permission sets; retain read/edit access only in Contact Encrypted CPNI Read/Write Access. No field schema, Apex, Flow, or page-layout change is required.",
)
evidence = insert_after(
    solution,
    "Read-only verification on 26 August 2026: changes are still required before promotion. GreatPlainsMerge has CPNI access through the three named profiles plus Account and Contact - Full Access and Heroku and Data Hub Permissions. GreatPlainsUAT also has profile access through Product, GPC-BI & Reporting, and GPC-Residential CRC Team. The intended permission set correctly grants read/edit in both sandboxes.",
)
evidence.runs[0].bold = True
status = paragraph_exact(document, "Approved")
set_text_preserve_format(status, "Approved design; implementation changes required before deployment")

# Complete the empty deployment-disposition area before the testing section.
deployment_overview = paragraph_exact(document, "Deployment", 0)
anchor = deployment_overview
for item in (
    "Disposition: not ready for promotion. No Salesforce changes were made during this review.",
    "GreatPlainsMerge (00DEa00000GkAsHMAV): the target permission set has 1 assignment; both named fields are still exposed through 3 profiles and 2 additional permission sets.",
    "GreatPlainsUAT (00DEa00000FZlLBMA1): the target permission set has 2 assignments; both named fields are still exposed through 6 profiles and 2 additional permission-set definitions.",
    "Production was not inspected or changed. Its complete profile, permission-set, and assignment inventory is a mandatory pre-deployment gate.",
):
    anchor = insert_after(anchor, item)
    apply_numbering(anchor, create_numbering(document, "bullet"))

# Replace the run-on testing paragraph with reproducible, synthetic-data UAT tests.
testing = next(
    p for p in document.paragraphs if p.text.strip().startswith("Test Case 1: Confirm profile users")
)
set_text_preserve_format(testing, "Test case 1 - Negative access without the target permission set")
testing.runs[0].bold = True
negative_steps = (
    "Use an approved UAT user on each affected profile and confirm Contact Encrypted CPNI Read/Write Access is not assigned.",
    "Open the same synthetic UAT Contact in the record page and Edit action; confirm CPNI Pin and CPNI Password are absent.",
    "Run an approved API/SOQL check as that user; confirm neither field can be queried or updated.",
    "Repeat for every profile and non-approved permission set identified by the final FLS inventory.",
)
anchor = testing
negative_num = create_numbering(document, "decimal")
for step in negative_steps:
    anchor = insert_after(anchor, step)
    apply_numbering(anchor, negative_num)

anchor = insert_after(anchor, "Test case 2 - Positive access through the target permission set")
anchor.runs[0].bold = True
positive_num = create_numbering(document, "decimal")
for step in (
    "Use a named UAT user assigned only Contact Encrypted CPNI Read/Write Access for these fields.",
    "Open the same synthetic UAT Contact; confirm both fields are visible and editable.",
    "Enter approved synthetic values, save, reopen the record, and confirm the values persist and render with the expected encrypted-field masking.",
    "Verify UI and approved API behavior, then remove any temporary test assignment or data according to the UAT cleanup plan.",
):
    anchor = insert_after(anchor, step)
    apply_numbering(anchor, positive_num)

anchor = insert_after(anchor, "Test case 3 - Exclusivity and regression")
anchor.runs[0].bold = True
regression_num = create_numbering(document, "decimal")
for step in (
    "Query all FieldPermissions rows for both fields and prove the target permission set is the only remaining read/edit grant path.",
    "Regression-test approved integrations and reporting users that previously depended on profile or broad permission-set access.",
    "Capture user, profile, permission-set assignments, record ID, timestamps, expected result, actual result, and screenshots in the SCC-4179 evidence.",
):
    anchor = insert_after(anchor, step)
    apply_numbering(anchor, regression_num)
runtime = insert_after(
    anchor,
    "Runtime status: pending. No user assignment, Contact record, metadata, deployment, or activation was changed during this read-only completion review.",
)
runtime.runs[0].bold = True

# Fill the remaining release-runbook sections.
pre = answer_after(document, "Pre-Deployment")
replace_with_numbered_items(
    pre,
    (
        "Retrieve and archive the current Production metadata for both Contact fields, every profile, and every permission set that grants either field.",
        "Inventory current target-permission-set assignees and obtain Security/Data Owner approval for the final authorized-user list.",
        "Build a focused package that removes read/edit from every unauthorized access path while preserving read/edit in Contact_CPNI_Read_Write_Access.",
        "Assess integrations, reports, list views, and automated users that may rely on the two fields; obtain owners and regression tests before removal.",
        "Run a check-only validation against the target org, record the job ID and exact components, and obtain explicit deployment approval.",
    ),
    document,
)

deploy = answer_after(document, "Deployment", 1)
replace_with_numbered_items(
    deploy,
    (
        "Deploy only the approved, focused profile and permission-set FLS changes; do not modify field definitions, layouts, Apex, Flow, or unrelated permissions.",
        "Retain read=true and edit=true only for Contact.CPNI_Pin__c and Contact.CPNI_Password__c in Contact_CPNI_Read_Write_Access.",
        "Remove both field-permission entries, or set read/edit false as supported, in every profile and non-approved permission set found by the inventory.",
        "Deploy to GreatPlainsUAT first. Promote to Production only after named-user UAT, integration regression, Security approval, and an approved change window.",
        "Record the deployment job ID, final component count, test level/result, start/end time, and any omissions. Do not auto-assign the permission set to new users.",
    ),
    document,
)

post = answer_after(document, "Post-Deployment")
replace_with_numbered_items(
    post,
    (
        "Read back FieldPermissions for both fields and confirm no profile or non-approved permission set grants read or edit.",
        "Confirm the target permission set still grants read/edit and its assignee list exactly matches the approved list.",
        "Execute the negative, positive, exclusivity, and integration regression tests above with named users.",
        "Review Setup Audit Trail and deployment results, attach evidence to SCC-4179, and obtain Business, Security, and QA sign-off.",
        "If a critical dependency fails, roll back with the approved metadata backup and remove any emergency assignment after the incident is resolved.",
    ),
    document,
)

set_text_preserve_format(
    answer_after(document, "DevOps Assignee"),
    "Salesforce Release Manager / Salesforce Administrator; record the named owner before the change window.",
)

# Complete technical-documentation fields while keeping the solution declarative.
section_answers = {
    "Libraries/Packages Used": "None. Standard Salesforce profile, permission-set, and Field-Level Security metadata only.",
    "APIs": "No runtime API is introduced. Use Metadata API for the approved deployment and SOQL read-back of FieldPermissions and PermissionSetAssignment for verification.",
    "Triggers or Automations": "None. No Apex trigger, Flow, validation rule, or automation change is required.",
    "New or Updated Objects/Fields": "No object or field schema change. Existing fields referenced: Contact.CPNI_Pin__c and Contact.CPNI_Password__c.",
    "Settings Modifications": "Field-Level Security only: remove both field grants from every profile and non-approved permission set; preserve read/edit only in Contact_CPNI_Read_Write_Access.",
    "Error Logging": "No application logging change. Release evidence must include validation/deployment results, Setup Audit Trail, FLS read-back, and named-user UAT results.",
    "Temporal Changes": "Permanent least-privilege control. Any emergency access must use the approved permission set, be time-bounded by process, and be removed after the approved window.",
}
for section_name, answer in section_answers.items():
    set_text_preserve_format(answer_after(document, section_name), answer)

components = answer_after(document, "Components")
set_text_preserve_format(
    components,
    "New: none. Update: Profiles - System Administrator, System Administrator - API Only, Standard User, plus every additional Production/UAT profile found by inventory (UAT currently includes Product, GPC-BI & Reporting, and GPC-Residential CRC Team). Permission Sets - preserve Contact_CPNI_Read_Write_Access; remove both field grants from Account_and_Contact_Full_Access, Heroku_and_Data_Hub_Permissions, and any other non-approved grant path found in Production.",
)

# Remove hidden form-export residue after the final answer; it otherwise creates
# a blank trailing page after the added release documentation.
final_answer = answer_after(document, "Temporal Changes")
node = final_answer._p.getnext()
while node is not None and node.tag != qn("w:sectPr"):
    next_node = node.getnext()
    node.getparent().remove(node)
    node = next_node

document.core_properties.title = "SCC-4179 Restrict CPNI Field Access - Completed Review"
document.core_properties.subject = "Completed deployment plan, technical documentation, and verification gates"
document.core_properties.keywords = (
    "Salesforce, SCC-4179, CPNI, Field-Level Security, permission set, GreatPlainsMerge, GreatPlainsUAT"
)
document.save(OUTPUT)
print(OUTPUT)
