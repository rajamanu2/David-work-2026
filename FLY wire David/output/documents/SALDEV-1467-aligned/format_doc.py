from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import re

root=Path(__file__).parent
src=Document(r'C:/Users/LIKKI/Downloads/SALDEV-1467 ARCHITECT re-review.docx')
texts=[p.text for p in src.paragraphs]
doc=Document()
s=doc.sections[0]
s.top_margin=s.bottom_margin=Inches(.7)
s.left_margin=s.right_margin=Inches(.75)
for name in ['Normal','Title','Heading 1','Heading 2','List Bullet']:
 st=doc.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
 st.font.size=Pt(11)
 st.paragraph_format.space_after=Pt(7)
 st.paragraph_format.line_spacing=1.06
doc.styles['Title'].font.size=Pt(23)
doc.styles['Title'].font.bold=True
for name in ['Heading 1','Heading 2']:
 doc.styles[name].font.size=Pt(14 if name=='Heading 1' else 11)
 doc.styles[name].font.bold=True
 doc.styles[name].paragraph_format.space_before=Pt(10)
 doc.styles[name].paragraph_format.keep_with_next=True

def clean(t):
 t=t.strip().replace('`','').replace('\ufffd','-')
 return re.sub(r'\.(?=[A-Z])','. ',t)

def para(t,style=None):
 p=doc.add_paragraph(style=style)
 t=clean(t)
 if ':' in t and len(t.split(':',1)[0])<85:
  a,b=t.split(':',1);p.add_run(a+':').bold=True;p.add_run(b)
 else:p.add_run(t)
 p.paragraph_format.widow_control=True
 return p

doc.add_paragraph('SALDEV-1467 Architect Re-review','Title')
para(texts[2])
doc.add_heading('Verified environment and validation',1)
for i in range(6,12):
 p=para(texts[i][2:],'List Bullet')
 p.paragraph_format.left_indent=Inches(.16)
 p.paragraph_format.first_line_indent=Inches(-.16)

table=doc.add_table(rows=1,cols=3)
table.alignment=WD_TABLE_ALIGNMENT.CENTER
table.autofit=False
widths=[3.9,1.8,1.3]
headers=['Component','Covered / total lines','Coverage']
for j,w in enumerate(widths):table.columns[j].width=Inches(w)
for j,t in enumerate(headers):table.rows[0].cells[j].text=t
for i in range(15,21):
 vals=[v.strip() for v in texts[i].strip('|').split('|')]
 for c,t in zip(table.add_row().cells,vals):c.text=t
for ri,row in enumerate(table.rows):
 pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
 if ri==0:pr.append(OxmlElement('w:tblHeader'))
 for j,c in enumerate(row.cells):
  c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  cp=c._tc.get_or_add_tcPr()
  mar=OxmlElement('w:tcMar')
  for side in ['top','left','bottom','right']:
   e=OxmlElement('w:'+side);e.set(qn('w:w'),'100');e.set(qn('w:type'),'dxa');mar.append(e)
  cp.append(mar)
  borders=OxmlElement('w:tcBorders')
  for side in ['top','left','bottom','right']:
   e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
  cp.append(borders)
  sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E5EBF1' if ri==0 else 'FFFFFF');cp.append(sh)
  for p in c.paragraphs:
   p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
   p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
   for r in p.runs:r.font.size=Pt(10);r.bold=ri==0

doc.add_heading('Reassessment of the developer comment',1).paragraph_format.page_break_before=True
for i in range(24,30):
 t=clean(texts[i]);label,body=t.split(':',1)
 doc.add_heading(label,2)
 body=body.strip()
 para(body[0].upper()+body[1:])
doc.add_heading('What to do next',1).paragraph_format.page_break_before=True
for i in range(33,37):
 p=para(texts[i]);p.paragraph_format.left_indent=Inches(.2);p.paragraph_format.first_line_indent=Inches(-.2)
doc.add_heading('Suggested Jira comment',1)
para(texts[40])
for el in list(doc.styles.element.iter(qn('w:pBdr'))):
 el.getparent().remove(el)
for el in list(doc.element.iter(qn('w:pBdr'))):
 el.getparent().remove(el)
doc.save(root/'SALDEV-1467_Architect_Re-review_Formatted.docx')
print(root/'SALDEV-1467_Architect_Re-review_Formatted.docx')

