import json
import subprocess
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

WORKBOOK = Path(r"C:\Users\LIKKI\Downloads\9-2 Copy of Copy of EPIQ-CPC-Probo System List.xlsx")
SF = r"C:\Program Files (x86)\sf\bin\sf.cmd"
SERVICE_CONTRACT_IDS = [
    "810Ro00000XRhAjIAL",
    "810Ro00000XRS0QIAX",
    "810Ro00000XRdlSIAT",
    "810Ro00000XR5LcIAL",
    "810Ro00000XRuHdIAL",
    "810Ro00000XRRxCIAX",
]
OPPORTUNITY_IDS = [
    "006Ro00000iKOb7IAG",
    "006Ro00000iL6QrIAK",
    "006Ro00000iL5eTIAS",
    "006Ro00000iL2NGIA0",
    "006Ro00000iKrbVIAS",
    "006Ro00000iLF5xIAG",
]


def norm(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip().upper()


wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
ws = wb["MainList"]
rows = [
    row
    for row in ws.iter_rows(min_row=2, max_col=11, values_only=True)
    if any(value is not None and str(value).strip() for value in row)
]
system_ids = {norm(row[0]) for row in rows if norm(row[0])}
serials = {norm(row[6]) for row in rows if norm(row[6])}
descriptions = Counter(norm(row[2]) for row in rows if norm(row[2]))

serial_quality = {
    "blank": sum(1 for row in rows if not norm(row[6])),
    "zero": sum(1 for row in rows if norm(row[6]) == "0"),
    "needed": sum(1 for row in rows if norm(row[6]) == "NEEDED"),
}
duplicate_system_ids = sum(count - 1 for count in Counter(norm(row[0]) for row in rows if norm(row[0])).values() if count > 1)
duplicate_serials = sum(count - 1 for count in Counter(norm(row[6]) for row in rows if norm(row[6]) not in {"", "0", "NEEDED"}).values() if count > 1)

sc_values = ",".join(f"'{v}'" for v in SERVICE_CONTRACT_IDS)
opp_values = ",".join(f"'{v}'" for v in OPPORTUNITY_IDS)
query = (
    "SELECT Id,WorkOrderNumber,Opportunity__r.Name,Asset_Serial_Number__c "
    f"FROM WorkOrder WHERE ServiceContractId IN ({sc_values}) OR Opportunity__c IN ({opp_values})"
)
proc = subprocess.run(
    [SF, "data", "query", "--target-org", "ProboMedical", "--query", query, "--json"],
    capture_output=True,
    text=True,
    encoding="utf-8",
)
if proc.returncode != 0:
    print(proc.stderr or proc.stdout)
    raise SystemExit(proc.returncode)
payload = json.loads(proc.stdout)
records = payload["result"]["records"]

summary = Counter()
by_opportunity = defaultdict(Counter)
unmatched = []
for record in records:
    serial = norm(record.get("Asset_Serial_Number__c"))
    opportunity = (record.get("Opportunity__r") or {}).get("Name") or "(blank)"
    summary["total"] += 1
    by_opportunity[opportunity]["total"] += 1
    if not serial:
        category = "blank"
    elif serial in {"N/A", "NA", "UNKNOWN", "NEEDED", "0"}:
        category = "placeholder"
    elif serial in serials:
        category = "matched_mainlist_serial"
    elif serial in system_ids:
        category = "matched_mainlist_system_id"
    else:
        category = "unmatched"
        unmatched.append({"workOrder": record.get("WorkOrderNumber"), "serial": serial, "opportunity": opportunity})
    summary[category] += 1
    by_opportunity[opportunity][category] += 1

result = {
    "workbook": {
        "mainListDataRows": len(rows),
        "modelCounts": dict(descriptions),
        "serialQuality": serial_quality,
        "duplicateSystemIdRowsBeyondFirst": duplicate_system_ids,
        "duplicateValidSerialRowsBeyondFirst": duplicate_serials,
    },
    "productionWorkOrders": dict(summary),
    "byOpportunity": {name: dict(counts) for name, counts in sorted(by_opportunity.items())},
    "unmatchedSample": unmatched[:20],
}
sys.stdout.reconfigure(encoding="utf-8")
print(json.dumps(result, indent=2))
