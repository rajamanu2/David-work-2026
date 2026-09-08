from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
import zipfile,re

OUT=Path(__file__).parent
SRC=OUT.parent/'SALDEV-1420-test-document-review'
d=Document(); s=d.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=s.bottom_margin=Inches(.65);s.left_margin=s.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Normal'].paragraph_format.line_spacing=1.06
d.styles['Title'].font.size=Pt(25)
d.styles['Heading 1'].font.size=Pt(18)
d.styles['Heading 2'].font.size=Pt(13)
h=s.header.paragraphs[0];h.text='SALDEV-1420  |  Approval email testing';h.style='Caption'
f=s.footer.paragraphs[0];f.text='Partial sandbox  •  Test review and reproduction guide     |     '
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
d.core_properties.author='Salesforce QA';d.core_properties.title='SALDEV 1420 Test Review and Reproduction Guide'
d.core_properties.subject='Approval email template review and business test steps'
def p(t,style=None): return d.add_paragraph(t,style)
def head(t): d.add_heading(t,1)
def sub(t): d.add_heading(t,2)
def page(t): d.add_page_break();head(t)
def steps(items):
 for i,t in enumerate(items,1): p(f'{i}. {t}')
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 for i,w in enumerate(widths): t.columns[i].width=Inches(w)
 for c,x in zip(t.rows[0].cells,headers): c.text=x
 rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
 for row in rows:
  for c,x in zip(t.add_row().cells,row): c.text=str(x)
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  for i,c in enumerate(row.cells):
   c.width=Inches(widths[i]);tc=c._tc.get_or_add_tcPr()
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tc.append(borders)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EDF3' if ri==0 else 'FFFFFF');tc.append(shade)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(5);para.paragraph_format.space_before=Pt(5)
    for run in para.runs:run.font.size=Pt(10);run.bold=ri==0
 return t
def shot(n,box,width,caption):
 path=SRC/f'image-{n:02}.png';iw,ih=Image.open(path).size
 x1,y1,x2,y2=box
 para=p('');shape=para.add_run().add_picture(str(path),width=Inches(width),height=Inches(width*(y2-y1)/(x2-x1)))
 fill=shape._inline.graphic.graphicData.pic.blipFill
 crop=OxmlElement('a:srcRect')
 for k,v in dict(l=x1/iw,t=y1/ih,r=(iw-x2)/iw,b=(ih-y2)/ih).items():crop.set(k,str(round(v*100000)))
 fill.insert(1,crop);shape._inline.docPr.set('descr',caption)
 p(caption,'Caption')

d.add_heading('SALDEV 1420 Test Review and Reproduction Guide',0)
p('Approval email templates in Partial sandbox','Subtitle')
p('Review date: 7 September 2026  |  Status: business validation pending')
p('This document reviews the supplied test evidence and gives reproducible checks for Quote Approval, Legal, Credit Risk, and Revenue emails. The evidence is useful as a reference, but does not establish complete business acceptance. Fresh results and the Quote Approval description decision are still required.')
sub('Review findings')
for t in [
'Use start date ascending for group order. Use longest subscription term only when start dates match, followed by the Quote Line Editor sequence. The August 20 term-only sorting note is outdated.',
'The final six-scenario section in the supplied test execution document has template headings but no attached results. Earlier screenshots do not complete that section.',
'The September Legal examples need matching source dates and line-order evidence. Their visible term values alone cannot prove chronological order.',
'Quote Approval has an unresolved description difference: the old reference combines Description and Rate Details; the recorded September 7 Partial review found Description alone. The story owner must confirm the intended behavior.',
'Renewal coverage, representative user access, and final business sign-off are not demonstrated by the supplied evidence.']:
 p(t,'List Bullet')
sub('How to use the screenshots')
p('The four email-table excerpts on pages 4–7 come from Story Test Execution Document for SALDEV-1420. They are historical reference screenshots, not results from a new execution. Their older Group and date banners are not the final header acceptance baseline. Confirm the agreed Term presentation with the story owner if date display remains disputed.')
sub('Evidence basis')
p('Sources: the supplied SALDEV-1420 story, August 20 design document, three September Legal examples, and Story Test Execution Document for SALDEV-1420. Review observations are recorded as of September 7; inspect the current quote and settings when testing.')

page('Test setup and preview steps')
steps(['Sign in to FlywirePartial. Confirm the address contains flywire--partial.sandbox. Record the tester and test date. Use a representative CPQ Sales user for the business access check.',
'Open the retained SALDEV-1420 Test Opp quote using the link below. Open Edit Lines and capture its group start dates, subscription terms, product order, and Multiple Ship-to Accounts value before testing.',
'In Setup, open Classic Email Templates and select the applicable template below. Use the available preview action with the correct related Quote or Opportunity approval record. Generate a non-sending preview; do not select Send.',
'If preview requires an approval record that is unavailable, record the missing prerequisite as Blocked. A static template screen without merged record data is not a successful output test.',
'Compare the merged table against the Quote Line Editor. Save a screenshot of each template and source data, including the record ID, expected result, actual result, and status.'])
p('Retained quote ID: a2NhG000003xBvNUAU')
p('https://flywire--partial.sandbox.lightning.force.com/lightning/r/SBQQ__Quote__c/a2NhG000003xBvNUAU/view')
table(['Template label','Template API name'],[
['Request - Quote Approval Request - AQS','AQS_Quote_Approval_Request'],['Legal - Approval Request','Legal_Approval_Request'],['Credit Risk - Approval Request','Credit_Risk_Approval_Request'],['Revenue - Approval Request','Revenue_Approval_Request']],[3.7,3.4])
sub('Retained quote expected result')
p('At the recorded review, both groups were non-ramped, started 1 October 2026, and had 72-month terms. Recheck these inputs because other testers may edit the record.')
table(['Header','Expected product sequence'],[['Term 1: 72 months','eStore → eStore (Usage) → eStore - Implementation Fee'],['Term 2: 72 months','Domestic Payments → Cross Border Payments']],[1.7,5.4])
p('Each expected line appears once. Bundle children remain beneath their parent. If current inputs differ, document the difference and establish the expected result from the saved QLE before evaluating the output.')

page('Group ordering reproduction')
p('Use an approved sandbox test copy for input changes. Preserve the retained story quote. Set the quote to 72 months beginning 1 October 2026, then prepare the groups below through the normal CPQ workflow.')
table(['Group','Ramping','Start date','Months','Position'],[['G1','Yes','2026-10-01','24','2'],['G2','Yes','2028-10-01','12','3'],['G3','Yes','2029-10-01','36','4'],['G4','No','2026-10-01','72','1']],[.7,1.1,2.1,1.5,1.7])
steps(['Add distinctive products to identify each group. Save and calculate the quote. Reopen Edit Lines and verify the persisted dates, terms, and sequence.',
'Preview each of the four templates using the matching approval context. Capture both source inputs and the complete output across all groups.',
'Expect G4 → G1 → G2 → G3, with terms 72 → 24 → 12 → 36 months. The later 36-month group must not move ahead of the earlier 12-month group.'])
sub('Additional cases for every template')
table(['Case','Inputs and expected result'],[
['Three ramped groups','Successive start dates; terms 12, 6, 18. Expect 12 → 6 → 18.'],
['Equal dates and terms','Put a Z-named group before an A-named group. Preserve QLE sequence.'],
['Equal dates with different terms','Longer term first; retain QLE order for remaining ties.'],
['Two non-ramped and two ramped','Full-term groups on the first date precede shorter same-date groups; later dates follow chronologically.'],
['Later non-ramped group','Earlier ramped start still comes first. Record Blocked if CPQ does not allow the inputs.'],
['No groups','One flat table; no empty Group 1 or Term header. Preserve QLE line order.']],[2.1,5.0])
p('Execution status: Pending for each case and each template until fresh evidence is recorded. An unsupported data setup must be marked Blocked with the exact validation message.')

page('Quote Approval email template')
p('Template: Request - Quote Approval Request - AQS')
shot(7,(319,462,1256,858),7.05,'Figure 1. Historical Quote Approval table excerpt, August 17. The excerpt shows column layout and sample lines; it is not a complete group or a fresh test result.')
sub('Steps to reproduce and verify')
steps(['Follow the preview steps on page 2 using AQS_Quote_Approval_Request and the correct related approval record.',
'Verify Billing Frequency is column 5 and Subscription Type is column 6. Compare Product, Applied Discount, Price, Quantity, Description, Quote Line Type, and conditional Ship-To values with the source quote.',
'Use a line with distinct Description and Rate Details text. Record whether the preview displays one or both values in Description. Capture the source values and the rendered cell.',
'Repeat the mixed-group test on page 3 and the ungrouped test. Check that each line appears once under the correct group.'])
sub('Open decision')
p('The reference mapping combines SBQQ__Description__c and Rate_Details__c. The September 7 review recorded SBQQ__Description__c alone in Partial. The explicit Rate Details acceptance wording targets Legal. Obtain the story owner’s decision for Quote Approval and record it before assigning Pass to this description check.')
p('Expected column correction: Billing Frequency 5; Subscription Type 6. Do not copy the duplicate position 6 from the old test document.')

page('Legal email template')
p('Template: Legal - Approval Request')
shot(15,(319,462,1025,1090),5.15,'Figure 2. Historical Legal table excerpt, August 17. Description contains sample external description and rate text. Group and date banner reflects the older presentation.')
sub('Steps to reproduce and verify')
steps(['Preview Legal_Approval_Request against the appropriate Legal approval record using the same saved quote inputs.',
'Test four lines: External Description plus Rate Details; description only; rates only; both blank. Expect one Description column, each populated value once, readable separation, and no NULL text or dangling separators.',
'Verify Product, Price, Quantity, Billing Frequency, descriptions, group order, and bundle hierarchy against QLE. Have the Legal reviewer confirm the text is suitable for the contract workflow.'])

page('Credit Risk email template')
p('Template: Credit Risk - Approval Request')
shot(11,(319,457,1025,1087),5.15,'Figure 3. Historical Credit Risk table excerpt, August 17. This is separate template evidence even though its table context is shared with Legal.')
sub('Steps to reproduce and verify')
steps(['Preview Credit_Risk_Approval_Request using the matching Credit Risk approval record. Capture this template separately from Legal.',
'Compare Product, Price, Quantity, Billing Frequency, and Description with the source values. Repeat the populated and blank Description and Rate Details cases.',
'Repeat grouped and ungrouped cases. Check Ship-To with the flag set to Yes and No, independent of whether groups exist. Confirm the intended approver can access the rendered output.'])

page('Revenue email template')
p('Template: Revenue - Approval Request')
shot(19,(265,416,855,775),6.1,'Figure 4. Historical Revenue table excerpt, August 18. Sample Description includes rate text. The rows shown are an excerpt and do not prove full quote ordering.')
sub('Steps to reproduce and verify')
steps(['Preview Revenue_Approval_Request using the matching Revenue approval record and saved quote inputs.',
'Compare Product, Applied Discount, Price, Quantity, Billing Frequency, and Description against the source quote. Verify Description combines SBQQ__Description__c and Rate_Details__c as recorded in the reference mapping.',
'Test populated, partially blank, and blank description inputs. Expect each value once, readable spacing, and no placeholder text.',
'Repeat the mixed-group, equal-date, and ungrouped cases. Verify each parent is followed by its bundle options and that standalone product order matches QLE.'])
sub('Expected result')
p('All eligible lines appear once in the expected group and sequence. Group order follows start date first and longest term only for equal dates. Prices, discounts, quantities, and frequencies match the saved source record. Record any mismatch with the quote ID and a screenshot of the source and output.')

page('Regression checks and test results')
sub('Steps for remaining coverage')
steps(['Bundle order: on an approved test copy, include a bundle and two standalone products in nonalphabetical order. Include nested options if supported. Save and calculate, then check every template for missing, duplicated, or misplaced lines.',
'Ship-To: repeat grouped and ungrouped cases with Multiple Ship-to Accounts set to Yes and No. Expect the conditional Ship-To column for Yes and its absence for No. Grouping alone does not determine visibility.',
'Amendment: use an existing supported amendment test record. Inspect calculated Quote_Line_Type__c values; lines equal to Amendment must be excluded. Reconcile eligible lines with intended Opportunity Products. Do not type values into this calculated field.',
'Renewal: use a supported renewal test record. Verify eligible lines remain visible and that prices, quantities, descriptions, group order, and bundle order match the saved quote. Record missing test data as Blocked.',
'Access: verify output as the representative CPQ Sales user and intended approver. If access fails, record the user and error. An administrator can inspect CPQ Sales Permissions and EmailTableController class access without changing settings.'])
sub('Execution record to complete for each case')
table(['Field','Record'],[['Case and template','_______________________________'],['Tester and date','_______________________________'],['Quote and related approval ID','_______________________________'],['Source QLE and output evidence','_______________________________'],['Expected and actual result','_______________________________'],['Status and issue reference','Pending / Pass / Fail / Blocked: ______________']],[2.5,4.6])
sub('Acceptance before closure')
p('Complete the required scenarios across all four templates, including new, amendment, and renewal behavior. Resolve the Quote Approval mapping decision and any disputed header presentation. Obtain Legal acceptance and the business owner’s sign-off using fresh screenshots and explicit results.')
p('Business reviewer: __________________    Review date: ______________')
out=OUT/'SALDEV-1420-Test-Review-and-Reproduction-Guide.docx';d.save(out)
with zipfile.ZipFile(out) as z:
 text=' '.join(z.read(n).decode('utf-8',errors='ignore') for n in z.namelist() if n.endswith('.xml'))
 assert not re.search(r'\b(?:AI|Codex|ChatGPT|OpenAI|GPT|deployment|deploy|deployed)\b',text,re.I)
print(out)
