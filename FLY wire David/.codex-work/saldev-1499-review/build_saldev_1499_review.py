from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\LIKKI\Documents\ChatGPT\david")
TASK = ROOT / ".codex-work" / "saldev-1499-review"
REFERENCE = Path(r"C:\Users\LIKKI\Downloads\SALDEV-1413-architect-review.docx")
OUTPUT = ROOT / "flywire-docgen" / "deliverables" / "SALDEV-1499_Architect_Review.docx"
DIAGRAM = TASK / "saldev-1499-control-path.png"
REFERENCE_SHA256 = "9C1A9F936367232B5DD8F9C08C98A6C801259257B220C84E76CE3CFA3937C13F"

NAVY = "0B2545"
BLUE = "2E75B6"
GRAY = "5E6B7E"
TEXT = "263445"
LIGHT_GRAY = "F1F3F6"
RED = "D13232"
PALE_RED = "FCEAEA"
AMBER = "B97816"
PALE_AMBER = "FFF4D9"
GREEN = "168A73"
PALE_GREEN = "E5F3EF"
PALE_BLUE = "EAF2FB"
WHITE = "FFFFFF"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest().upper()


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, size: float | None = None, bold: bool | None = None,
                 color: str | None = None, italic: bool | None = None):
    run.font.name = "Arial"
    if run._element.get_or_add_rPr().rFonts is None:
        run._element.get_or_add_rPr().append(OxmlElement("w:rFonts"))
    rfonts = run._element.get_or_add_rPr().rFonts
    rfonts.set(qn("w:ascii"), "Arial")
    rfonts.set(qn("w:hAnsi"), "Arial")
    rfonts.set(qn("w:eastAsia"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = rgb(color)
    if italic is not None:
        run.italic = italic


def set_cell_fill(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=72, start=120, bottom=72, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="000000", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_table_geometry(table, widths_inches: list[float]):
    widths = [int(round(v * 1440)) for v in widths_inches]
    total = sum(widths)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[i])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_table_borders(table)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def format_cell(cell, text: str, *, bold=False, color=TEXT, size=9.4,
                align=WD_ALIGN_PARAGRAPH.LEFT, fill: str | None = None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    if fill:
        set_cell_fill(cell, fill)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[float],
              status_col: int | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    repeat_header(table.rows[0])
    for i, header in enumerate(headers):
        format_cell(table.rows[0].cells[i], header, bold=True, color=NAVY, size=9.6, fill=LIGHT_GRAY)
    for row_values in rows:
        row = table.add_row()
        for i, value in enumerate(row_values):
            fill = None
            color = TEXT
            bold = False
            align = WD_ALIGN_PARAGRAPH.LEFT
            if status_col is not None and i == status_col:
                bold = True
                align = WD_ALIGN_PARAGRAPH.CENTER
                normalized = value.lower()
                if normalized == "pass" or normalized == "confirmed":
                    fill, color = PALE_GREEN, GREEN
                elif normalized in {"partial", "risk"}:
                    fill, color = PALE_AMBER, AMBER
                elif normalized in {"fail", "blocker"}:
                    fill, color = PALE_RED, RED
            format_cell(row.cells[i], value, bold=bold, color=color, size=9.0, align=align, fill=fill)
    set_table_geometry(table, widths)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(1)
    return table


def add_callout(doc, label: str, body: str, kind="info", body_size=11.0):
    color, fill = {
        "fail": (RED, PALE_RED),
        "risk": (AMBER, PALE_AMBER),
        "pass": (GREEN, PALE_GREEN),
        "info": (BLUE, PALE_BLUE),
    }[kind]
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_fill(cell, fill)
    set_cell_margins(cell, top=110, start=140, bottom=110, end=140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(label.upper())
    set_run_font(run, size=9.8, bold=True, color=color)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.0
    run2 = p2.add_run(body)
    set_run_font(run2, size=body_size, bold=True, color=NAVY)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.space_before = Pt(1)
    return table


def add_heading(doc, text: str, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = False
    run = p.add_run(text)
    set_run_font(run, size=16.8 if level == 1 else 13.2, bold=True, color=BLUE)
    return p


def add_body(doc, text: str, *, bold=False, italic=False, color=TEXT,
             align=WD_ALIGN_PARAGRAPH.LEFT, after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    set_run_font(r, size=10.7, bold=bold, italic=italic, color=color)
    return p


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.7)
    normal.font.color.rgb = rgb(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for level, size in ((1, 16.8), (2, 13.2)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(BLUE)
        style.paragraph_format.space_before = Pt(10 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True


def make_diagram(path: Path):
    w, h = 1500, 830
    image = Image.new("RGB", (w, h), "#F6F8FB")
    draw = ImageDraw.Draw(image)
    arial = r"C:\Windows\Fonts\arial.ttf"
    arial_bold = r"C:\Windows\Fonts\arialbd.ttf"
    title_font = ImageFont.truetype(arial_bold, 35)
    subtitle_font = ImageFont.truetype(arial, 19)
    box_title = ImageFont.truetype(arial_bold, 23)
    body_font = ImageFont.truetype(arial, 18)
    label_font = ImageFont.truetype(arial_bold, 15)
    note_title = ImageFont.truetype(arial_bold, 22)
    note_body = ImageFont.truetype(arial, 18)

    draw.text((65, 48), "SALDEV-1499 | safe correction and promotion path", font=title_font, fill="#0B2545")
    draw.text((65, 94), "Keep the before-save design, correct sequencing and trigger flow, then prove behavior at scale.", font=subtitle_font, fill="#5E6B7E")

    boxes = [
        (55, 175, 340, 525, "Incident mechanism", ["After-save Flow", "updates the same line", "Automation recurses", "SOQL reaches 101"], "#D13232", "#FCEAEA"),
        (410, 175, 695, 525, "Trigger regression", ["Early return on", "ordinary updates", "Average calculation", "can be skipped"], "#D13232", "#FCEAEA"),
        (765, 175, 1050, 525, "Ordering risk", ["Background copy runs", "before price-change", "Formula depends on", "changed values"], "#B97816", "#FFF4D9"),
        (1120, 175, 1405, 525, "Release proof", ["Correct both paths", "Assert field values", "Run 200-line test", "Complete business UAT"], "#168A73", "#E5F3EF"),
    ]
    for x1, y1, x2, y2, title, lines, edge, fill in boxes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=20, fill=fill, outline=edge, width=5)
        draw.text((x1 + 22, y1 + 28), title, font=box_title, fill="#0B2545")
        y = y1 + 105
        for line in lines:
            draw.ellipse((x1 + 23, y + 8, x1 + 35, y + 20), fill=edge)
            draw.text((x1 + 50, y), line, font=body_font, fill="#263445")
            y += 48

    arrows = [(340, 410, "REMOVE"), (695, 765, "CORRECT"), (1050, 1120, "PROVE")]
    for start, end, label in arrows:
        y = 350
        draw.line((start + 8, y, end - 15, y), fill="#A36A12", width=6)
        draw.polygon([(end - 15, y - 12), (end + 3, y), (end - 15, y + 12)], fill="#A36A12")
        tw = draw.textlength(label, font=label_font)
        draw.text(((start + end - tw) / 2, y - 35), label, font=label_font, fill="#A36A12")

    draw.rounded_rectangle((70, 610, 1430, 770), radius=20, fill="#EAF2FB", outline="#2E75B6", width=4)
    draw.text((105, 642), "Architect reading", font=note_title, fill="#0B2545")
    note = ("The deactivated Flow removes the direct recursion source. Promotion still requires correcting the live trigger's "
            "early return, sequencing formula-dependent assignments after price-change logic, and replacing coverage-only tests "
            "with value assertions, bulk proof, and stepped/non-stepped UAT.")
    words = note.split()
    lines, current = [], ""
    for word in words:
        candidate = (current + " " + word).strip()
        if draw.textlength(candidate, font=note_body) > 1270:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    y = 687
    for line in lines[:3]:
        draw.text((105, y), line, font=note_body, fill="#263445")
        y += 30
    image.save(path)


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def build():
    if sha256(REFERENCE) != REFERENCE_SHA256:
        raise RuntimeError("Reference DOCX hash changed; fresh template distillation required.")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    TASK.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    clear_body(doc)
    configure_styles(doc)
    doc.core_properties.title = "SALDEV-1499 Architect Review"
    doc.core_properties.subject = "Quote Line Type background field update Flow deactivation and Apex replacement review"
    doc.core_properties.author = "Flywire Salesforce Architecture Review"
    doc.core_properties.keywords = "SALDEV-1499, Salesforce, CPQ, Flow, Apex, Quote Line"

    # PAGE 1 - executive decision
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("ARCHITECT REVIEW")
    set_run_font(r, size=11.2, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 0.95
    r = p.add_run("SALDEV-1499 | Quote Line Type\nFlow Deactivation")
    set_run_font(r, size=25.5, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("Flywire Partial sandbox | Read-only implementation and promotion-readiness assessment")
    set_run_font(r, size=13.2, color=GRAY)

    add_table(doc, ["Environment", "FlywirePartial | Org 00DhG0000000jOX | Sandbox"], [
        ["Review date", "28 August 2026"],
        ["Evidence", "Live FlowDefinition/Flow/Apex read-back | unit-test results | coverage | source inspection"],
        ["Decision", "Changes Requested - not ready for promotion"],
        ["Boundary", "No code edits, deployment, activation, data changes, or Production changes"],
    ], [1.32, 5.18])

    add_callout(doc, "Architect decision",
                "Do not promote yet. The legacy Flow is genuinely inactive and the before-save Apex direction removes its extra DML cycle, but the deployed trigger can skip existing rollup logic, formula-dependent fields are copied in the wrong sequence, and the tests do not prove the required field outcomes.",
                "fail")
    add_heading(doc, "Acceptance outcome", 1)
    add_table(doc, ["#", "Scope", "Result", "Evidence"], [
        ["1", "Flow deactivation", "Pass", "FlowDefinition.ActiveVersionId is null; latest v13 is Obsolete."],
        ["2", "Root-cause direction", "Pass", "Background fields moved to before-save memory assignment; no Flow self-update remains."],
        ["3", "Trigger regression safety", "Fail", "An after-update early return can bypass calculateAveragePrice()."],
        ["4", "Functional sequencing", "Fail", "Quote Line Type is copied before price-change logic updates a field used by its formula."],
        ["5", "Automated proof", "Partial", "Eight tests pass, but SALDEV-1499 assertions do not verify the output fields or 200-line limits."],
    ], [0.34, 1.48, 1.0, 3.68], status_col=2)

    # PAGE 2 - process path
    add_page_break(doc)
    add_heading(doc, "Incident mechanism and control path", 1)
    add_body(doc, "The incident was raised after an after-save record-triggered Flow updated the same Quote Line that initiated the transaction. That second DML operation re-entered Quote Line automation and the transaction exhausted the 100-query synchronous limit. Moving simple assignments to a before-save trigger is the correct architectural direction because it modifies Trigger.new without a second update.")
    make_diagram(DIAGRAM)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    inline_shape = p.add_run().add_picture(str(DIAGRAM), width=Inches(6.45))
    inline_shape._inline.docPr.set("descr", "Four-stage SALDEV-1499 correction path from the Flow incident through trigger and sequencing fixes to release proof")
    inline_shape._inline.docPr.set("title", "SALDEV-1499 safe correction and promotion path")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Figure 1. SALDEV-1499 safe correction and promotion path")
    set_run_font(r, size=10.2, italic=True, color=GRAY)
    add_callout(doc, "Design reading",
                "Preserve the before-save, bulk-collected query pattern. Correct the control-flow and sequencing defects surgically, then prove field values and governor headroom. Do not reactivate the Flow as a workaround.",
                "info")

    # PAGE 3 - live evidence
    add_page_break(doc)
    add_heading(doc, "Current SALDEV-1499 implementation evidence", 1)
    add_heading(doc, "Flow and Apex state", 2)
    add_table(doc, ["Evidence area", "Observed state", "Result", "Architect reading"], [
        ["Target org", "FlywirePartial; sandbox; USA1148S", "Pass", "Verified against Organization Id 00DhG0000000jOXUAY."],
        ["Flow definition", "ActiveVersionId = null", "Pass", "The Flow is genuinely deactivated."],
        ["Latest Flow version", "v13; API 65; Obsolete", "Pass", "No active version remains."],
        ["QuoteLineTrigger", "Active; API 67", "Partial", "Before-save replacement exists, but after-update control flow can exit early."],
        ["Trigger handler", "Active; API 67", "Partial", "Bulk collection exists; background-field order and behavior require correction."],
        ["Failed Quote Line", "a1zhG000000MBi7QAG not found", "Partial", "The exact incident record cannot be replayed or field-read back."],
    ], [1.42, 1.52, 0.86, 2.70], status_col=2)

    add_heading(doc, "Automated test and coverage evidence", 2)
    add_table(doc, ["Evidence area", "Observed state", "Result", "Architect reading"], [
        ["Latest focused run", "8 test methods passed", "Pass", "Latest results recorded 27 Aug 2026 around 21:15 UTC."],
        ["Trigger coverage", "23 covered / 10 uncovered", "Partial", "69.7% aggregate coverage; uncovered branches include control paths."],
        ["Handler coverage", "245 covered / 24 uncovered", "Pass", "91.1% aggregate coverage, but percentage does not prove behavior."],
        ["SALDEV-1499 assertions", "Record existence / non-empty lists", "Fail", "Expected background, mismatch, fallback, and price-change values are not asserted."],
        ["Bulk governor proof", "No 200-line limits test", "Fail", "The original SOQL-101 acceptance concern is not directly proven."],
    ], [1.42, 1.52, 0.86, 2.70], status_col=2)
    add_callout(doc, "Evidence boundary",
                "The linked design document is not accessible to the connected Google account. Requirement alignment for fallback and adjusted-rate behavior therefore remains conditional until the team shares the document or records the decisions in Jira.",
                "risk")

    # PAGE 4 - findings
    add_page_break(doc)
    add_heading(doc, "Detailed code-review findings", 1)
    add_table(doc, ["Priority", "Finding", "Why it matters", "Required correction"], [
        ["P1", "After-update early return", "When OLI synchronization is unnecessary, the trigger returns before calculateAveragePrice(). Ordinary updates can stop recalculating existing rollups.", "Replace the trigger-level return with a guarded OLI-sync block so later after-save logic always reaches its own recursion guard."],
        ["P1", "Formula copy occurs too early", "processBackgroundFields() copies Quote_Line_Type__c before handlePriceChangeOnAmendmentClones() updates Price_Change__c. The live formula depends on Price_Change__c and Group_Allow_Product_Ramping__c.", "Compute price-change inputs first, then derive/copy Quote Line Type from final values, or calculate the classification explicitly in one method."],
        ["P1", "Tests do not assert outcomes", "Passing tests can still leave Quote_Line_Type_background__c, Rate_Details_Changed__c, or Default_Rate_Details__c wrong.", "Re-query records and assert exact true/false/text values for insert, update, ramp, renewal, amendment, cancellation, and blank cases."],
        ["P2", "Behavior exceeds the legacy Flow", "Apex adds default-rate fallback and adjusted-rate override comparisons beyond the Flow's direct Rate_Details versus Default_Rate_Details check.", "Confirm these rules in the design/Jira acceptance criteria and test each branch; otherwise narrow the implementation."],
        ["P2", "Local source is stale", "The checkout still shows Flow v13 Active and an older after-only trigger. It does not match the live sandbox implementation.", "Retrieve the exact live components into a focused review package before producing the promotion artifact or diff."],
    ], [0.82, 1.30, 2.05, 2.33])
    add_callout(doc, "Promotion blocker",
                "Do not approve promotion until the trigger return and formula-ordering defects are corrected, explicit field-value assertions pass, and a bulk transaction demonstrates that the replacement stays below governor limits.",
                "fail")

    # PAGE 5 - equivalence and focused correction scope
    add_page_break(doc)
    add_heading(doc, "Functional equivalence and implementation boundary", 1)
    add_heading(doc, "Legacy-to-Apex behavior that must remain equivalent", 2)
    add_table(doc, ["Behavior", "Legacy Flow", "Replacement expectation"], [
        ["Quote Line Type background", "Copy formula value after save when blank or different", "Persist the classification calculated from the transaction's final price-change/ramping state."],
        ["Rate mismatch", "Rate_Details__c != Default_Rate_Details__c", "Match approved null/blank semantics and document any fallback/override expansion."],
        ["Record update", "Explicit DML on the initiating record", "Before-save memory assignment only; no self-DML."],
        ["Existing rollups", "Separate trigger behavior", "Continue calculateAveragePrice() on all intended insert/update/delete/undelete paths."],
    ], [1.68, 2.18, 2.64])
    add_heading(doc, "Focused correction scope", 2)
    add_table(doc, ["Component type", "Focused scope", "Required action"], [
        ["Quote Line trigger", "Before-save call order", "Run price-change classification before the background copy, or combine both deterministically."],
        ["Quote Line trigger", "After-update OLI-sync branch", "Remove the trigger-level early return and preserve subsequent rollup execution."],
        ["Trigger handler", "processBackgroundFields", "Define null/blank, fallback, adjusted-rate, and formula-copy behavior explicitly."],
        ["Apex tests", "QuoteLineTriggerHandleTest", "Assert exact outputs and prior rollup behavior; avoid pass-by-record-existence assertions."],
        ["Bulk proof", "200 mixed Quote Lines", "Capture query/DML counts and verify all classifications, mismatch flags, and rollups."],
        ["Focused package", "Trigger, handler, test, inactive Flow state", "Retrieve current live metadata and validate only the reviewed components plus required dependencies."],
    ], [1.38, 2.42, 2.70])

    # PAGE 6 - release evidence
    add_page_break(doc)
    add_heading(doc, "Release sequence and production gate", 1)
    add_table(doc, ["Phase", "Required evidence"], [
        ["Requirement confirmation", "Share the design document or copy its approved decisions into Jira: formula timing, blank handling, fallback source, adjusted-rate rules, and stepped-up behavior."],
        ["Focused correction", "Surgically correct trigger order/control flow and tests; keep the Flow deactivated and preserve unrelated CPQ automation."],
        ["Sandbox validation", "Run the focused Apex class plus relevant Quote Line regression tests; record test-run ID, coverage, query/DML headroom, and field-level results."],
        ["Business UAT", "Create stepped and non-stepped quotes; save/recalculate; confirm no SOQL-101 error and verify Quote Line Type, mismatch flags, and existing averages."],
        ["Promotion/read-back", "After approval only: validate the production-target package, deploy, verify Flow ActiveVersionId remains null, read back Apex versions, and repeat smoke tests."],
    ], [1.45, 5.05])
    add_heading(doc, "Validation evidence to capture", 2)
    add_table(doc, ["Evidence", "Required detail"], [
        ["Automated test receipt", "Test-run ID, named methods, pass/fail totals, runtime, and exact trigger/handler coverage."],
        ["Governor headroom", "Queries and DML statements before/after the 200-line transaction, with no SOQL-101 exception."],
        ["Record-level read-back", "Quote and Quote Line IDs plus expected/actual background type, mismatch flag, default rate details, and average fields."],
        ["Metadata state", "FlowDefinition.ActiveVersionId, latest Flow status, Apex API/status/modified date, and reviewed package component count."],
    ], [1.70, 4.80])
    add_callout(doc, "Production gate",
                "This review covered FlywirePartial only. Production was not inventoried or changed. Re-retrieve current Production metadata and run focused check-only validation before any deployment decision.",
                "risk")

    # PAGE 7 - UAT and team response
    add_page_break(doc)
    add_heading(doc, "UAT and approval gates", 1)
    add_table(doc, ["Test", "Pass evidence"], [
        ["Ordinary Quote Line update", "An update with no OLI-sync requirement still runs calculateAveragePrice() and preserves expected average placement."],
        ["Stepped-up renewal", "A ramp-enabled renewal saves without error and background type equals the final Quote_Line_Type__c classification."],
        ["Amendment and cancellation", "Price Change, Cancellation, and Cancellation Price Change outcomes match final quantity and price-change inputs."],
        ["Rate details equal/different", "Equal values set Rate_Details_Changed__c false; approved mismatch/fallback/override cases set it true."],
        ["Blank and null handling", "Blank formula/default/source/subscription cases produce documented values without stale background data or exceptions."],
        ["Bulk and limits", "200 mixed Quote Lines complete in one transaction with correct fields, rollups, query headroom, and no SOQL-101 error."],
        ["Non-stepped regression", "Existing ordinary quoting and average-price behavior match the approved sandbox baseline."],
    ], [1.58, 4.92])

    add_heading(doc, "Required approval gates", 2)
    add_table(doc, ["#", "Required action", "Owner", "Evidence needed"], [
        ["1", "Confirm expanded rate-detail rules", "Product Owner / Architect", "Accessible design decision or Jira acceptance detail."],
        ["2", "Correct trigger control flow/order", "Development Lead", "Reviewed focused diff with no unrelated rewrite."],
        ["3", "Strengthen automated tests", "Developer / QA", "Exact field assertions, regression results, and 200-line limits evidence."],
        ["4", "Complete business UAT", "CPQ Business Owner", "Quote IDs, scenarios, expected/actual values, and no-error confirmation."],
        ["5", "Approve promotion", "Architect / Release Manager", "Check-only job ID, component/test counts, exclusions, and sign-offs."],
    ], [0.34, 1.64, 1.48, 3.04])

    add_heading(doc, "Ready-to-paste architect response", 2)
    add_callout(doc, "Story comment",
                "Architect review completed for SALDEV-1499 in FlywirePartial. Changes Requested. The Quote Line Type background Flow is correctly deactivated (ActiveVersionId is null; v13 is Obsolete), and the before-save Apex direction removes the Flow's self-update. However, the deployed after-update trigger can return before calculateAveragePrice(), and processBackgroundFields() copies the formula-backed Quote Line Type before price-change logic updates fields used by that formula. Eight focused tests pass, but the SALDEV-1499 tests do not assert the resulting background/rate-detail fields and there is no 200-line governor-limit proof. Correct the trigger control flow and sequencing, add exact field-value and bulk tests, share the detailed design decision, and complete stepped/non-stepped UAT before promotion. No code, metadata, activation, deployment, data, or Production changes were made during this review.",
                "info", body_size=9.25)

    doc.save(OUTPUT)
    if sha256(REFERENCE) != REFERENCE_SHA256:
        raise RuntimeError("Reference DOCX was modified unexpectedly.")
    print(OUTPUT)


if __name__ == "__main__":
    build()
